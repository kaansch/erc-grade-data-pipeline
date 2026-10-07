"""Synthetic checks for Turkish section letters and scan coverage."""
import argparse
import contextlib
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook

path=Path(__file__).resolve().parents[1]/'scripts/04_extra_source_search/discover_sources.py'
spec=importlib.util.spec_from_file_location('discovery',path)
scanner=importlib.util.module_from_spec(spec);spec.loader.exec_module(scanner)

def candidate_rows(path):
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        return list(book.active.values)
    finally:
        book.close()

class DiscoveryTests(unittest.TestCase):
    def test_running_without_arguments_prompts_for_paths(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'extra files';root.mkdir()
            report=Path(tmp)/'missing sections.xlsx';report.touch()
            output=Path(tmp)/'scan results'
            entered=[f'"{root}"',f'"{report}"',f'"{output}"']
            with patch('builtins.input',side_effect=entered),patch.object(scanner,'scan',return_value=0) as run,contextlib.redirect_stdout(io.StringIO()) as printed:
                self.assertEqual(scanner.main([]),0)
            args=run.call_args.args[0]
            self.assertEqual(args.root,[root])
            self.assertEqual(args.missing_report,report)
            self.assertEqual(args.output_dir,output)
            self.assertIn(str(output.resolve()),printed.getvalue())

    def test_sections_stay_distinct(self):
        self.assertEqual(scanner.section_letters('3/I 3/İ 3/ı 3/i'),['I','İ'])

    def test_turkish_school_and_heading(self):
        self.assertEqual(scanner.normalize('IŞIK İLKOKULU'),'isik ilkokulu')
        found=scanner.evidence('3. SINIF 2. DÖNEM 3/İ TÜRKÇE MATEMATİK')
        self.assertEqual(found['grades'],['3']);self.assertEqual(found['terms'],['2'])
        self.assertEqual(found['sections'],['İ'])

    def test_content_duplicates_errors_and_source_preservation(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'sources';root.mkdir()
            raw='Işık İlkokulu,3. Sınıf 2. Dönem,3/İ,Matematik\nSynthetic Student,3,3,3\n'
            (root/'a.csv').write_text(raw,encoding='utf-8')
            (root/'b.csv').write_text(raw,encoding='utf-8')
            (root/'broken.pdf').write_bytes(b'not a PDF')
            before={p.name:scanner.sha256(p) for p in root.iterdir()}
            args=argparse.Namespace(root=[root],output_dir=Path(tmp)/'result',school=['Işık İlkokulu'],missing_report=None,inventory_only=False,max_pdf_pages=0,max_excel_rows=0,text_encoding='utf-8-sig')
            with contextlib.redirect_stdout(io.StringIO()):result=scanner.scan(args)
            self.assertEqual(result,2)
            manifest=json.loads((args.output_dir/'manifest.json').read_text(encoding='utf-8'))
            self.assertEqual(manifest['file_instances'],3);self.assertEqual(manifest['candidate_rows'],2);self.assertEqual(manifest['errors'],1)
            self.assertEqual(before,{p.name:scanner.sha256(p) for p in root.iterdir()})
            rows=candidate_rows(args.output_dir/'candidates.xlsx')
            self.assertEqual(rows[0],('File','Location','Why candidate'))
            self.assertEqual(len(rows),3)
            self.assertEqual({Path(row[0]).name for row in rows[1:]},{'a.csv','b.csv'})

    def test_output_inside_sources_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            args=argparse.Namespace(root=[Path(tmp)],output_dir=Path(tmp)/'out')
            with self.assertRaises(ValueError):scanner.scan(args)

    def test_missing_report_rows_are_linked_as_evidence_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'sources';root.mkdir()
            (root/'Işık School.csv').write_text('Işık İlkokulu,3. Sınıf 2. Dönem,3/İ,Matematik\n',encoding='utf-8')
            entry={'report_sheet':'3. SINIF 2. DÖNEM','report_row':7,'district':'Örnek İlçe',
                   'school':'Işık İlkokulu','section':'İ','grade':'3','term':'2',
                   'keyword':scanner.school_keyword('Işık İlkokulu')}
            args=argparse.Namespace(root=[root],output_dir=Path(tmp)/'result',school=[],
                                    missing_report=Path(tmp)/'missing.xlsx',inventory_only=False,
                                    max_pdf_pages=0,max_excel_rows=0,text_encoding='utf-8-sig')
            with patch.object(scanner,'missing_entries',return_value=[entry]),contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(scanner.scan(args),0)
            rows=candidate_rows(args.output_dir/'candidates.xlsx')
            self.assertEqual(len(rows),2)
            self.assertEqual(rows[1][1],None)
            self.assertIn('isik',rows[1][2])
            self.assertIn('Işık İlkokulu',rows[1][2])

    def test_excel_candidates_use_exact_rows_even_with_inflated_dimension(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'sources';root.mkdir()
            source=root/'grades.xlsx'
            book=Workbook();sheet=book.active
            sheet['B2']='Atatürk İlkokulu'
            sheet['B3']='Atatürk İlkokulu'
            sheet['B7']='Atatürk İlkokulu'
            book.save(source);book.close()
            with ZipFile(source) as archive:
                contents={name:archive.read(name) for name in archive.namelist()}
            sheet_xml=contents['xl/worksheets/sheet1.xml']
            contents['xl/worksheets/sheet1.xml']=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:XFD1048576"',sheet_xml,count=1)
            with ZipFile(source,'w',ZIP_DEFLATED) as archive:
                for name,data in contents.items(): archive.writestr(name,data)
            args=argparse.Namespace(root=[root],output_dir=Path(tmp)/'result',school=['Atatürk İlkokulu'],missing_report=None,inventory_only=False,max_pdf_pages=0,max_excel_rows=0,text_encoding='utf-8-sig')
            with contextlib.redirect_stdout(io.StringIO()):self.assertEqual(scanner.scan(args),0)
            rows=candidate_rows(args.output_dir/'candidates.xlsx')
            self.assertEqual(len(rows),2)
            self.assertEqual(rows[1][1],'rows 2-3; row 7')

    def test_multiple_sheets_show_sheet_and_duplicate_paths_are_rechecked(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)/'sources';root.mkdir()
            book=Workbook();book.active.title='First'
            second=book.create_sheet('Second');second['A5']='Atatürk İlkokulu'
            book.save(root/'grades.xlsx');book.close()
            (root/'A.csv').write_text('nothing related',encoding='utf-8')
            (root/'Atatürk.csv').write_text('nothing related',encoding='utf-8')
            args=argparse.Namespace(root=[root],output_dir=Path(tmp)/'result',school=['Atatürk İlkokulu'],missing_report=None,inventory_only=False,max_pdf_pages=0,max_excel_rows=0,text_encoding='utf-8-sig')
            with contextlib.redirect_stdout(io.StringIO()):self.assertEqual(scanner.scan(args),0)
            rows=candidate_rows(args.output_dir/'candidates.xlsx')
            self.assertEqual(len(rows),3)
            self.assertEqual(rows[1][1],None)
            self.assertEqual(Path(rows[1][0]).name,'Atatürk.csv')
            self.assertIn('file path',rows[1][2])
            self.assertEqual(rows[2][1],'Second, row 5')

if __name__=='__main__':unittest.main()
