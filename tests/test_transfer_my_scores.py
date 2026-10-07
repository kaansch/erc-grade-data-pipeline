"""Check score transfer without using the original grade workbooks."""

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts/05_erc_roster_matching"))
from transfer_my_scores import (
    BASE_MISMATCH_HEADERS,
    SCORED_MISMATCH_HEADERS,
    transfer_my_scores,
)


class TransferMyScoresTests(unittest.TestCase):
    def make_report(self, path, scored=False, school="School A"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Mismatched Students"
        headers = SCORED_MISMATCH_HEADERS if scored else BASE_MISMATCH_HEADERS
        sheet.append(headers)
        values = {
            "Source Sheet": "Term 1",
            "Source District": "District A",
            "Source School": "School A",
            "Source Section": "4A",
            "Source Student Name": "Student A",
            "ERC District": "District A",
            "ERC School": school,
            "Best Candidate in Source Section": "Student A",
            "Source-Section ERC Section": "4A",
            "Best Candidate Across All Sections": "Student A",
            "All-Sections ERC Section": "4A",
            "Reason Not Matched": "Review",
            "District Similarity": 1,
            "School Similarity": 1,
            "Source-Section Similarity": 1,
            "All-Sections Similarity": 1,
            "My score": 3,
        }
        sheet.append([values[header] for header in headers])
        workbook.save(path)
        workbook.close()

    def test_transfers_score_to_a_new_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_path, new_path, output_path = (root / name for name in ("old.xlsx", "new.xlsx", "output.xlsx"))
            self.make_report(old_path, scored=True)
            self.make_report(new_path)
            original_old, original_new = old_path.read_bytes(), new_path.read_bytes()

            stats = transfer_my_scores(old_path, new_path, output_path)
            result = load_workbook(output_path)
            sheet = result["Mismatched Students"]
            self.assertEqual(sheet["G1"].value, "My score")
            self.assertEqual(sheet["G2"].value, 3)
            self.assertEqual(stats.transferred, 1)
            self.assertTrue(stats.inserted_score_column)
            self.assertEqual(old_path.read_bytes(), original_old)
            self.assertEqual(new_path.read_bytes(), original_new)
            result.close()

    def test_skips_a_changed_erc_match(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_path, new_path, output_path = (root / name for name in ("old.xlsx", "new.xlsx", "output.xlsx"))
            self.make_report(old_path, scored=True)
            self.make_report(new_path, school="School B")

            stats = transfer_my_scores(old_path, new_path, output_path)
            result = load_workbook(output_path)
            self.assertIsNone(result["Mismatched Students"]["G2"].value)
            self.assertEqual(stats.changed_context, 1)
            result.close()


if __name__ == "__main__":
    unittest.main()
