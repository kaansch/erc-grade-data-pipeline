"""
Merge multiple Excel workbooks in one folder into the first workbook.

Usage:
    python excel_merge.py

    The script will ask you to paste the full folder path in the terminal.

Notes:
    - The first Excel file is chosen after sorting file names alphabetically.
    - Temporary Excel files whose names start with "~$" are ignored.
    - Row 1 is treated as the header row and is copied only from the
      destination workbook.
    - Source data starts at row 2.
    - Values are appended after the destination sheet's last row containing data.
    - This script supports .xlsx and .xlsm files through openpyxl.

Install dependency if needed:
    pip install openpyxl
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Iterable, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}


def has_value(value: object) -> bool:
    """Return True when a cell value should count as data."""
    return value is not None and value != ""


def row_has_data(values: Iterable[object]) -> bool:
    """Return True when at least one value in a row contains data."""
    return any(has_value(value) for value in values)


def find_excel_files(folder: Path) -> List[Path]:
    """Find supported Excel files in the folder, ignoring Excel temp files."""
    files = [
        path
        for path in folder.iterdir()
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and not path.name.startswith("~$")
    ]
    return sorted(files, key=lambda path: path.name.lower())


def get_worksheet(workbook, sheet_name: Optional[str]) -> Worksheet:
    """Get the requested sheet or the first sheet when none is provided."""
    if sheet_name is None:
        return workbook.worksheets[0]

    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' was not found in {workbook.sheetnames}.")

    return workbook[sheet_name]


def find_last_data_row(sheet: Worksheet) -> int:
    """Find the last row that contains at least one non-empty cell value."""
    for row_number in range(sheet.max_row, 0, -1):
        row_values = (cell.value for cell in sheet[row_number])
        if row_has_data(row_values):
            return row_number
    return 0


def ensure_blank_row(sheet: Worksheet, row_number: int) -> None:
    """Stop before writing if the target row already contains data."""
    row_values = (cell.value for cell in sheet[row_number])
    if row_has_data(row_values):
        raise RuntimeError(
            f"Refusing to overwrite row {row_number} in sheet '{sheet.title}'."
        )


def load_excel_workbook(path: Path):
    """Load a workbook, preserving VBA content when the destination is .xlsm."""
    return load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")


def append_source_rows(
    destination_sheet: Worksheet,
    source_sheet: Worksheet,
    next_blank_row: int,
    include_blank_rows: bool,
) -> Tuple[int, int]:
    """Append source rows from row 2 onward into the destination sheet."""
    last_source_row = find_last_data_row(source_sheet)

    if last_source_row <= 1:
        return next_blank_row, 0

    copied_rows = 0

    for row_values in source_sheet.iter_rows(
        min_row=2,
        max_row=last_source_row,
        values_only=True,
    ):
        if not include_blank_rows and not row_has_data(row_values):
            continue

        ensure_blank_row(destination_sheet, next_blank_row)

        for column_number, value in enumerate(row_values, start=1):
            destination_sheet.cell(
                row=next_blank_row,
                column=column_number,
                value=value,
            )

        copied_rows += 1
        next_blank_row += 1

    return next_blank_row, copied_rows


def merge_excel_files(
    folder: Path,
    sheet_name: Optional[str] = None,
    include_blank_rows: bool = False,
) -> Tuple[Path, int, int]:
    """Merge all supported Excel files in a folder into the first sorted file."""
    if not folder.exists():
        raise FileNotFoundError(f"Folder does not exist: {folder}")

    if not folder.is_dir():
        raise NotADirectoryError(f"Path is not a folder: {folder}")

    excel_files = find_excel_files(folder)

    if len(excel_files) < 2:
        raise ValueError("At least two .xlsx or .xlsm files are required.")

    destination_path = excel_files[0]
    source_paths = excel_files[1:]

    destination_workbook = load_excel_workbook(destination_path)
    destination_sheet = get_worksheet(destination_workbook, sheet_name)

    last_destination_row = find_last_data_row(destination_sheet)
    if last_destination_row == 0:
        raise ValueError(
            f"Destination sheet '{destination_sheet.title}' has no header row."
        )

    next_blank_row = last_destination_row + 1
    total_copied_rows = 0
    processed_source_files = 0

    for source_path in source_paths:
        source_workbook = load_excel_workbook(source_path)
        source_sheet = get_worksheet(source_workbook, sheet_name)

        next_blank_row, copied_rows = append_source_rows(
            destination_sheet=destination_sheet,
            source_sheet=source_sheet,
            next_blank_row=next_blank_row,
            include_blank_rows=include_blank_rows,
        )

        if copied_rows > 0:
            processed_source_files += 1
            total_copied_rows += copied_rows

        source_workbook.close()

    destination_workbook.save(destination_path)
    destination_workbook.close()

    return destination_path, processed_source_files, total_copied_rows


def clean_path_text(path_text: str) -> str:
    """Remove surrounding whitespace and quotes from a pasted folder path."""
    return path_text.strip().strip('"').strip("'")


def ask_for_folder_path() -> Path:
    """Ask the user for the folder path when no path is given in the command."""
    while True:
        folder_text = input("Paste the full folder path that contains the Excel files: ")
        folder_text = clean_path_text(folder_text)

        if folder_text:
            return Path(folder_text)

        print("Folder path cannot be empty. Please paste the full folder path.")


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Merge Excel files in one folder into the first sorted file."
    )
    parser.add_argument(
        "folder",
        nargs="?",
        type=Path,
        help="Optional folder path containing the Excel files.",
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Sheet name to merge. Defaults to the first sheet in each workbook.",
    )
    parser.add_argument(
        "--include-blank-rows",
        action="store_true",
        help="Copy completely blank rows that appear between source data rows.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_arguments()
    folder = args.folder if args.folder is not None else ask_for_folder_path()

    try:
        destination_path, processed_files, copied_rows = merge_excel_files(
            folder=folder,
            sheet_name=args.sheet,
            include_blank_rows=args.include_blank_rows,
        )
    except Exception as error:
        print(f"Error: {error}")
        sys.exit(1)

    print(f"Destination file: {destination_path}")
    print(f"Source files with copied rows: {processed_files}")
    print(f"Rows appended: {copied_rows}")


if __name__ == "__main__":
    main()
