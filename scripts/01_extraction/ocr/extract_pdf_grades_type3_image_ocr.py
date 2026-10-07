# -*- coding: utf-8 -*-
r"""
Convert image/scanned type-3 grade PDFs into Excel using OCR.

This is the OCR version of extract_pdf_grades_type3_selectable.py for the
sideways image-based Sultangazi/Basaksehir-style teacher grade sheets.

Important table logic:
- One PDF contains one class section ("şube").
- Pages are course pages for that same section.
- The pages are image-based and usually rotated sideways inside the PDF.
- The grade is taken from the "1. DÖNEM PUANI" column for each student.
- "SERBEST ETKİNLİKLER" is detected but skipped.

Install Python packages in the environment used by VS Code:

    pip install pymupdf pillow pytesseract numpy openpyxl

You must also install the Tesseract OCR program itself. On Windows, if Tesseract
is not on PATH, pass its full path with --tesseract-cmd.

Run with a PDF or folder:

    python extract_pdf_grades_type3_image_ocr.py "C:\path\to\pdf_or_folder"
"""

from __future__ import annotations

import argparse
import csv
import io
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

try:
    import fitz  # PyMuPDF
    import numpy as np
    import pytesseract
    from PIL import Image, ImageDraw, ImageEnhance, ImageOps
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    missing_name = getattr(exc, "name", "required package")
    raise SystemExit(
        f"Missing Python package: {missing_name}\n"
        "Install the requirements with:\n\n"
        "    pip install pymupdf pillow pytesseract numpy openpyxl\n"
    ) from exc


METADATA_COLUMNS = ["İlçe", "Okul", "Sınıf", "Şube", "Okul No", "Ad Soyad"]

KNOWN_COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "HAYAT BİLGİSİ",
    "YABANCI DİL",
    "GÖRSEL SANATLAR",
    "MÜZİK",
    "FEN BİLİMLERİ",
    "BEDEN EĞİTİMİ VE OYUN",
    "SERBEST ETKİNLİKLER",
]

# The attached example uses exactly this order. OCR still tries to read the
# title first; this is only a fallback when OCR cannot read a course heading.
PAGE_COURSE_FALLBACK = [
    "TÜRKÇE",
    "MATEMATİK",
    "HAYAT BİLGİSİ",
    "YABANCI DİL",
    "SERBEST ETKİNLİKLER",
    "FEN BİLİMLERİ",
    "BEDEN EĞİTİMİ VE OYUN",
]

COURSE_SYNONYMS = {
    "YABANCI DİL": ["YABANCI DİL", "İNGİLİZCE", "INGILIZCE"],
    "BEDEN EĞİTİMİ VE OYUN": [
        "BEDEN EĞİTİMİ VE OYUN",
        "BEDEN EĞİTİMİ",
        "BEDEN EGITIMI",
    ],
}

SUMMARY_KEYWORDS = [
    "TOPLAM ÖĞRENCİ SAYISI",
    "BAŞARILI ÖĞRENCİ SAYISI",
    "BAŞARISIZ ÖĞRENCİ SAYISI",
    "BAŞARI YÜZDESİ",
    "NOT ORTALAMASI",
    "ZAYIFI OLAN",
]

TURKISH_TO_ASCII = str.maketrans(
    {
        "ç": "c",
        "Ç": "C",
        "ğ": "g",
        "Ğ": "G",
        "ı": "i",
        "I": "I",
        "İ": "I",
        "ö": "o",
        "Ö": "O",
        "ş": "s",
        "Ş": "S",
        "ü": "u",
        "Ü": "U",
    }
)

try:
    RESAMPLE_LANCZOS = Image.Resampling.LANCZOS
except AttributeError:  # Pillow < 9.1
    RESAMPLE_LANCZOS = Image.LANCZOS


@dataclass
class GradeGrid:
    x_lines: list[int]
    y_lines: list[int]
    row_intervals: list[tuple[int, int]]
    school_no_col: int
    name_col: int
    grade_col: int
    table_bbox: tuple[int, int, int, int]


@dataclass
class IdentitySlot:
    interval_index: int
    school_no: str
    name: str


@dataclass
class ParsedPdf:
    pdf_path: Path
    courses: list[str]
    header: dict[str, str]
    rows: list[dict[str, object]]


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", "\n").replace("\xa0", " ")
    text = unicodedata.normalize("NFC", text)
    return re.sub(r"\s+", " ", text).strip()


def fold_for_match(value: object) -> str:
    text = clean_text(value).translate(TURKISH_TO_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def has_letter(value: str) -> bool:
    return any(char.isalpha() for char in value)


def is_excluded_course(course: str) -> bool:
    folded = fold_for_match(course)
    return "SERBEST" in folded and "ETKINLIK" in folded


def normalize_school_no(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"-?\s*\d+", text)
    if not match:
        return ""
    return re.sub(r"\s+", "", match.group(0))


def normalize_name(value: object) -> str:
    text = clean_text(value)
    text = text.replace("|", " ").replace("_", " ")
    text = re.sub(r"[^A-Za-zÇĞİÖŞÜçğıöşü' .-]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text.upper()


def normalize_grade(value: object) -> int | str:
    text = clean_text(value)
    if not text:
        return ""

    compact = text.strip().replace(",", ".")
    match = re.search(r"[123]", compact)
    if match:
        return int(match.group(0))

    folded = fold_for_match(compact)
    # Common OCR confusions for single digit grades.
    if folded in {"I", "L"}:
        return 1
    if folded in {"Z"}:
        return 2

    return ""


def canonical_course_from_text(text: str) -> str:
    folded = fold_for_match(text)
    if not folded:
        return ""

    for course in KNOWN_COURSES:
        choices = COURSE_SYNONYMS.get(course, [course])
        for choice in choices:
            choice_folded = fold_for_match(choice)
            if choice_folded and choice_folded in folded:
                return course

    compact_folded = folded.replace(" ", "")
    for course in KNOWN_COURSES:
        choices = COURSE_SYNONYMS.get(course, [course])
        for choice in choices:
            choice_folded = fold_for_match(choice).replace(" ", "")
            if choice_folded and choice_folded in compact_folded:
                return course

    return ""


def line_has_summary_text(value: str) -> bool:
    folded = fold_for_match(value)
    return any(fold_for_match(keyword) in folded for keyword in SUMMARY_KEYWORDS)


def resolve_input_pdfs(raw_path: str | Path) -> tuple[list[Path], bool, Path]:
    path = Path(raw_path).expanduser()

    if path.is_dir():
        pdf_files = sorted(item for item in path.glob("*.pdf") if item.is_file())
        if not pdf_files:
            raise FileNotFoundError(f"No PDF files were found in folder: {path}")
        return pdf_files, True, path

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(f"The input file is not a PDF: {path}")
        return [path], False, path.parent

    if path.suffix.lower() != ".pdf":
        pdf_version = path.with_suffix(".pdf")
        if pdf_version.is_file():
            return [pdf_version], False, pdf_version.parent

    raise FileNotFoundError(
        f"PDF file or folder was not found: {path}\n"
        "Tip: pass one PDF file path or a folder containing PDF files."
    )


def configure_tesseract(tesseract_cmd: str | None) -> None:
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
        return

    common_windows_path = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    if common_windows_path.exists():
        pytesseract.pytesseract.tesseract_cmd = str(common_windows_path)


def verify_tesseract_available() -> None:
    try:
        pytesseract.get_tesseract_version()
    except Exception as error:
        raise RuntimeError(
            "Tesseract OCR is not available. Install Tesseract OCR, then rerun. "
            "If it is already installed, pass --tesseract-cmd with the full "
            "path to tesseract.exe."
        ) from error


def resolve_ocr_language(requested_lang: str) -> str:
    try:
        available_languages = set(pytesseract.get_languages(config=""))
    except Exception:
        return requested_lang

    requested_parts = [part.strip() for part in requested_lang.split("+") if part.strip()]
    if not requested_parts:
        return requested_lang

    available_parts = [part for part in requested_parts if part in available_languages]
    missing_parts = [part for part in requested_parts if part not in available_languages]

    if available_parts and missing_parts:
        resolved_lang = "+".join(available_parts)
        print(
            "Warning: Tesseract language data not found for "
            f"{'+'.join(missing_parts)}; using {resolved_lang}.",
            file=sys.stderr,
        )
        return resolved_lang

    if available_parts:
        return requested_lang

    if "eng" in available_languages:
        print(
            f"Warning: Tesseract language data not found for {requested_lang}; using eng.",
            file=sys.stderr,
        )
        return "eng"

    return requested_lang


def render_pdf_page(page, dpi: int) -> Image.Image:
    zoom = dpi / 72
    pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    return Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB")


def orient_page_image(image: Image.Image) -> Image.Image:
    # The sample pages are landscape forms rotated inside portrait PDF pages.
    # Rotating -90 degrees makes the title readable and the table upright.
    if image.height > image.width:
        return image.rotate(-90, expand=True, fillcolor="white")
    return image


def threshold_dark_pixels(image: Image.Image, threshold: int = 100) -> np.ndarray:
    grayscale = ImageOps.grayscale(image)
    return np.array(grayscale) < threshold


def group_positions(values: Iterable[int], max_gap: int = 2) -> list[tuple[int, int, int]]:
    positions = [int(value) for value in values]
    if not positions:
        return []

    groups: list[tuple[int, int, int]] = []
    start = previous = positions[0]
    for value in positions[1:]:
        if value - previous <= max_gap:
            previous = value
            continue
        groups.append((start, previous, (start + previous) // 2))
        start = previous = value
    groups.append((start, previous, (start + previous) // 2))
    return groups


def detect_grid(image: Image.Image) -> GradeGrid:
    dark = threshold_dark_pixels(image)
    height, width = dark.shape
    row_counts = dark.sum(axis=1)
    col_counts = dark.sum(axis=0)

    row_groups = group_positions(np.where(row_counts > width * 0.35)[0])
    if len(row_groups) < 10:
        row_groups = group_positions(np.where(row_counts > width * 0.25)[0])
    y_lines = [group[2] for group in row_groups]

    if len(y_lines) < 10:
        raise ValueError("Could not detect enough horizontal grid lines.")

    table_top = y_lines[0]
    table_bottom = y_lines[-1]
    table_height = table_bottom - table_top

    col_groups = group_positions(np.where(col_counts > height * 0.10)[0])
    x_lines = [
        group[2]
        for group in col_groups
        if col_counts[group[2]] > table_height * 0.55
    ]

    if len(x_lines) < 12:
        raise ValueError("Could not detect enough vertical grid lines.")

    intervals = list(zip(x_lines, x_lines[1:]))
    widths = [right - left for left, right in intervals]

    # The student name column is the widest early column, immediately after
    # S.No and school-number columns.
    early_limit = max(5, min(len(intervals), len(intervals) // 3))
    name_col = max(range(early_limit), key=lambda index: widths[index])
    if name_col < 1:
        raise ValueError("Could not locate school-number and student-name columns.")
    school_no_col = name_col - 1

    min_row_gap = max(8, int(height * 0.008))
    max_row_gap = max(35, int(height * 0.04))
    row_intervals = [
        (top, bottom)
        for top, bottom in zip(y_lines, y_lines[1:])
        if min_row_gap <= bottom - top <= max_row_gap
    ]

    if len(row_intervals) < 5:
        raise ValueError("Could not locate student row intervals.")

    # The grade column is the filled numeric column after the name column.
    candidate_scores: list[tuple[int, int, int]] = []
    for column_index in range(name_col + 1, len(intervals)):
        left, right = intervals[column_index]
        if right - left < 10:
            continue

        score = 0
        nonempty_rows = 0
        for top, bottom in row_intervals:
            y_margin = max(3, int((bottom - top) * 0.18))
            x_margin = max(3, int((right - left) * 0.10))
            crop = dark[top + y_margin : bottom - y_margin, left + x_margin : right - x_margin]
            dark_count = int(crop.sum())
            score += dark_count
            if dark_count > 8:
                nonempty_rows += 1

        candidate_scores.append((nonempty_rows, score, column_index))

    if not candidate_scores:
        raise ValueError("Could not locate the 1st-term grade column.")

    candidate_scores.sort(reverse=True)
    grade_col = candidate_scores[0][2]

    return GradeGrid(
        x_lines=x_lines,
        y_lines=y_lines,
        row_intervals=row_intervals,
        school_no_col=school_no_col,
        name_col=name_col,
        grade_col=grade_col,
        table_bbox=(x_lines[0], table_top, x_lines[-1], table_bottom),
    )


def crop_table_cell(
    image: Image.Image,
    grid: GradeGrid,
    column_index: int,
    row_interval: tuple[int, int],
    pad_x: int = 4,
    pad_y: int = 3,
) -> Image.Image:
    left = grid.x_lines[column_index] + pad_x
    right = grid.x_lines[column_index + 1] - pad_x
    top = row_interval[0] + pad_y
    bottom = row_interval[1] - pad_y
    return image.crop((max(0, left), max(0, top), max(left + 1, right), max(top + 1, bottom)))


def enhance_for_ocr(cell: Image.Image, scale: float = 3.0, threshold: int = 190) -> Image.Image:
    grayscale = ImageOps.grayscale(cell)
    grayscale = ImageEnhance.Contrast(grayscale).enhance(2.0)

    if scale != 1:
        new_size = (
            max(1, int(grayscale.width * scale)),
            max(1, int(grayscale.height * scale)),
        )
        grayscale = grayscale.resize(new_size, RESAMPLE_LANCZOS)

    array = np.array(grayscale)
    binary = np.where(array < threshold, 0, 255).astype("uint8")
    result = Image.fromarray(binary, mode="L")
    return ImageOps.expand(result, border=12, fill=255)


def ocr_image(
    image: Image.Image,
    lang: str,
    psm: int,
    whitelist: str | None = None,
    preserve_spaces: bool = False,
) -> str:
    config_parts = [f"--psm {psm}", "--oem 3"]
    if whitelist:
        config_parts.append(f"-c tessedit_char_whitelist={whitelist}")
    if preserve_spaces:
        config_parts.append("-c preserve_interword_spaces=1")
    config = " ".join(config_parts)

    try:
        return pytesseract.image_to_string(image, lang=lang, config=config)
    except pytesseract.TesseractError as error:
        if "Failed loading language" in str(error) and lang != "eng":
            return pytesseract.image_to_string(image, lang="eng", config=config)
        raise


def ocr_header(image: Image.Image, grid: GradeGrid, lang: str) -> str:
    left, table_top, right, _bottom = grid.table_bbox
    header_top = max(0, int(table_top - image.height * 0.18))
    header = image.crop((left, header_top, right, max(header_top + 1, table_top - 5)))
    header = enhance_for_ocr(header, scale=2.0, threshold=200)
    return ocr_image(header, lang=lang, psm=6, preserve_spaces=True)


def ocr_school_no(cell: Image.Image, lang: str) -> str:
    candidates: list[str] = []
    for psm in (7, 8, 10):
        text = ocr_image(
            enhance_for_ocr(cell, scale=4.0, threshold=195),
            lang=lang,
            psm=psm,
            whitelist="-0123456789",
        )
        school_no = normalize_school_no(text)
        if school_no:
            candidates.append(school_no)
    return candidates[0] if candidates else ""


def ocr_student_name(cell: Image.Image, lang: str) -> str:
    text = ocr_image(
        enhance_for_ocr(cell, scale=3.2, threshold=200),
        lang=lang,
        psm=7,
        preserve_spaces=True,
    )
    return normalize_name(text)


def ocr_grade(cell: Image.Image, lang: str) -> int | str:
    candidates: list[int] = []
    for psm in (7, 8, 10):
        text = ocr_image(
            enhance_for_ocr(cell, scale=5.0, threshold=205),
            lang=lang,
            psm=psm,
            whitelist="123,0.OoIlZz",
        )
        grade = normalize_grade(text)
        if grade != "":
            candidates.append(int(grade))

    if not candidates:
        return ""

    # Prefer the most repeated answer; this is useful when psm modes disagree.
    return max(sorted(set(candidates)), key=candidates.count)


def extract_district_and_school(lines: list[str]) -> tuple[str, str]:
    for line in lines[:40]:
        folded = fold_for_match(line)
        if "/" not in line:
            continue
        if "MUDURLUGU" not in folded and "MUDURLUG" not in folded and "OKUL" not in folded:
            continue

        district_part, school_part = line.split("/", 1)
        district = clean_text(district_part)
        school_part = clean_text(school_part)
        school = re.split(
            r"\s+M[üu]d[üu]rl[üu][ğg][üu]\b",
            school_part,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0].strip()
        if school == school_part:
            school = re.sub(r"\s*MUDURLUGU?.*$", "", fold_for_match(school_part)).title()
        return district, school

    return "", ""


def extract_class_and_section(text: str) -> tuple[str, str]:
    raw_match = re.search(
        r"(?P<class>\d+)\.?\s*S[ıiİI]n[ıiİI]f\s*/\s*(?P<section>[A-ZÇĞİÖŞÜ0-9]+)\s*Şubesi",
        text,
        flags=re.IGNORECASE,
    )
    if raw_match:
        return clean_text(raw_match.group("class")), clean_text(raw_match.group("section")).upper()

    folded = fold_for_match(text)
    folded_match = re.search(
        r"\b(?P<class>\d+)\s*SINIF\s*/?\s*(?P<section>[A-Z0-9]+)\s*SUBESI\b",
        folded,
    )
    if folded_match:
        return folded_match.group("class"), folded_match.group("section")

    return "", ""


def extract_header_info(header_texts: list[str]) -> dict[str, str]:
    full_text = "\n".join(header_texts)
    lines = [clean_text(line) for line in full_text.splitlines() if clean_text(line)]
    district, school = extract_district_and_school(lines)
    class_no, section = extract_class_and_section(" ".join(lines[:80]))

    return {
        "district": district,
        "school": school,
        "class_no": class_no,
        "section": section,
    }


def course_for_page(header_text: str, page_index: int) -> str:
    course = canonical_course_from_text(header_text)
    if course:
        return course

    if 0 <= page_index < len(PAGE_COURSE_FALLBACK):
        return PAGE_COURSE_FALLBACK[page_index]

    return ""


def valid_student_identity(school_no: str, name: str) -> bool:
    if not school_no or not name or not has_letter(name):
        return False
    folded_name = fold_for_match(name)
    if any(marker in folded_name for marker in ("OGRENCI", "SOYAD", "OKUL", "PUANI", "DERSLER")):
        return False
    if line_has_summary_text(name):
        return False
    return True


def extract_page_rows_from_grid(
    image: Image.Image,
    grid: GradeGrid,
    header: dict[str, str],
    course: str,
    lang: str,
    identity_slots: list[IdentitySlot] | None,
) -> tuple[list[dict[str, object]], list[IdentitySlot]]:
    if not course or is_excluded_course(course):
        return [], identity_slots or []

    rows: list[dict[str, object]] = []
    new_identity_slots: list[IdentitySlot] = []

    if identity_slots:
        for slot in identity_slots:
            if slot.interval_index >= len(grid.row_intervals):
                continue
            row_interval = grid.row_intervals[slot.interval_index]
            grade = ocr_grade(crop_table_cell(image, grid, grid.grade_col, row_interval), lang)
            rows.append(
                {
                    "İlçe": header["district"],
                    "Okul": header["school"],
                    "Sınıf": header["class_no"],
                    "Şube": header["section"],
                    "Okul No": slot.school_no,
                    "Ad Soyad": slot.name,
                    course: grade,
                }
            )
        return rows, identity_slots

    for interval_index, row_interval in enumerate(grid.row_intervals):
        school_no = ocr_school_no(
            crop_table_cell(image, grid, grid.school_no_col, row_interval),
            lang,
        )
        name = ocr_student_name(
            crop_table_cell(image, grid, grid.name_col, row_interval, pad_x=5, pad_y=3),
            lang,
        )

        if not valid_student_identity(school_no, name):
            continue

        grade = ocr_grade(crop_table_cell(image, grid, grid.grade_col, row_interval), lang)

        slot = IdentitySlot(interval_index=interval_index, school_no=school_no, name=name)
        new_identity_slots.append(slot)
        rows.append(
            {
                "İlçe": header["district"],
                "Okul": header["school"],
                "Sınıf": header["class_no"],
                "Şube": header["section"],
                "Okul No": school_no,
                "Ad Soyad": name,
                course: grade,
            }
        )

    return rows, new_identity_slots


def draw_debug_grid(image: Image.Image, grid: GradeGrid) -> Image.Image:
    debug = image.copy()
    draw = ImageDraw.Draw(debug)
    for x in grid.x_lines:
        draw.line((x, grid.table_bbox[1], x, grid.table_bbox[3]), fill="red", width=2)
    for y in grid.y_lines:
        draw.line((grid.table_bbox[0], y, grid.table_bbox[2], y), fill="blue", width=2)

    for column, color in (
        (grid.school_no_col, "orange"),
        (grid.name_col, "green"),
        (grid.grade_col, "purple"),
    ):
        left = grid.x_lines[column]
        right = grid.x_lines[column + 1]
        draw.rectangle((left, grid.table_bbox[1], right, grid.table_bbox[3]), outline=color, width=5)

    return debug


def save_debug_files(
    debug_dir: Path,
    pdf_path: Path,
    page_number: int,
    image: Image.Image,
    grid: GradeGrid | None,
    header_text: str,
    course: str,
    rows: list[dict[str, object]],
) -> None:
    debug_dir.mkdir(parents=True, exist_ok=True)
    safe_stem = safe_file_part(pdf_path.stem)
    prefix = debug_dir / f"{safe_stem}_page_{page_number:03d}"

    debug_text = (
        f"PDF: {pdf_path}\n"
        f"Page: {page_number}\n"
        f"Course: {course}\n"
        f"Rows parsed on page: {len(rows)}\n"
        "\n--- HEADER OCR TEXT ---\n"
        f"{header_text}\n"
    )
    (prefix.with_suffix(".txt")).write_text(debug_text, encoding="utf-8")

    if grid is not None:
        draw_debug_grid(image, grid).save(f"{prefix}_grid.png")
    else:
        image.save(f"{prefix}_rendered.png")


def extract_pdf(
    pdf_path: Path,
    dpi: int,
    lang: str,
    errors: list[dict[str, str]],
    debug_dir: Path | None,
) -> ParsedPdf:
    with fitz.open(str(pdf_path)) as document:
        page_images: list[Image.Image] = []
        page_grids: list[GradeGrid] = []
        header_texts: list[str] = []

        for page_index in range(document.page_count):
            page = document.load_page(page_index)
            image = orient_page_image(render_pdf_page(page, dpi=dpi))
            grid = detect_grid(image)
            header_text = ocr_header(image, grid, lang=lang)

            page_images.append(image)
            page_grids.append(grid)
            header_texts.append(header_text)

        header = extract_header_info(header_texts)
        rows: list[dict[str, object]] = []
        course_order: list[str] = []
        identity_slots: list[IdentitySlot] | None = None

        for page_index, (image, grid, header_text) in enumerate(
            zip(page_images, page_grids, header_texts),
            start=0,
        ):
            course = course_for_page(header_text, page_index)
            page_rows, maybe_identity_slots = extract_page_rows_from_grid(
                image=image,
                grid=grid,
                header=header,
                course=course,
                lang=lang,
                identity_slots=identity_slots,
            )

            if identity_slots is None and maybe_identity_slots:
                identity_slots = maybe_identity_slots

            if course and not is_excluded_course(course) and course not in course_order:
                course_order.append(course)
            rows.extend(page_rows)

            if not page_rows and not is_excluded_course(course):
                errors.append(
                    {
                        "pdf": str(pdf_path),
                        "page": str(page_index + 1),
                        "reason": "No student rows parsed from detected grid.",
                        "raw_text": header_text[:500],
                    }
                )

            if debug_dir is not None:
                save_debug_files(
                    debug_dir=debug_dir,
                    pdf_path=pdf_path,
                    page_number=page_index + 1,
                    image=image,
                    grid=grid,
                    header_text=header_text,
                    course=course,
                    rows=page_rows,
                )

        if not rows:
            raise ValueError(f"No student rows could be extracted from {pdf_path.name}.")

        return ParsedPdf(pdf_path=pdf_path, courses=course_order, header=header, rows=rows)


def combine_parsed_pdfs(
    parsed_pdfs: list[ParsedPdf],
) -> tuple[list[dict[str, object]], list[str]]:
    students: dict[str, dict[str, object]] = {}
    student_order: list[str] = []
    course_order: list[str] = []

    for parsed in parsed_pdfs:
        for course in parsed.courses:
            if course and not is_excluded_course(course) and course not in course_order:
                course_order.append(course)

        for row in parsed.rows:
            school_no = str(row.get("Okul No", ""))
            if not school_no:
                continue

            student_key = "\x1f".join(
                [
                    str(row.get("İlçe", "")),
                    str(row.get("Okul", "")),
                    str(row.get("Sınıf", "")),
                    str(row.get("Şube", "")),
                    school_no,
                ]
            )

            if student_key not in students:
                students[student_key] = {column: row.get(column, "") for column in METADATA_COLUMNS}
                student_order.append(student_key)
            else:
                for column in METADATA_COLUMNS:
                    if not students[student_key].get(column) and row.get(column):
                        students[student_key][column] = row[column]

            for column, value in row.items():
                if column in METADATA_COLUMNS or is_excluded_course(column):
                    continue
                if column not in course_order:
                    course_order.append(column)
                students[student_key][column] = value

    final_rows: list[dict[str, object]] = []
    for student_key in student_order:
        record = students[student_key]
        for course in course_order:
            record.setdefault(course, "")
        final_rows.append(record)

    return final_rows, course_order


def save_excel(rows: list[dict[str, object]], course_columns: list[str], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Notlar"

    output_columns = [*METADATA_COLUMNS, *course_columns]
    worksheet.append(output_columns)
    for row in rows:
        worksheet.append([row.get(column, "") for column in output_columns])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    school_no_column = output_columns.index("Okul No") + 1
    school_no_letter = get_column_letter(school_no_column)
    for row_number in range(2, worksheet.max_row + 1):
        worksheet[f"{school_no_letter}{row_number}"].number_format = "@"

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header_value = column_cells[0].value
        max_length = max(len(str(cell.value or "")) for cell in column_cells)

        if header_value == "Ad Soyad":
            width = min(max(max_length + 2, 20), 42)
        elif header_value in {"İlçe", "Okul"}:
            width = min(max(max_length + 2, 14), 36)
        elif header_value in course_columns:
            width = min(max(max_length + 2, 12), 28)
        else:
            width = min(max(max_length + 2, 10), 18)

        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 42
    workbook.save(output_path)


def safe_file_part(value: str) -> str:
    value = clean_text(value)
    value = value.translate(TURKISH_TO_ASCII)
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" ._")
    return value or "grades"


def build_output_path(
    pdf_paths: list[Path],
    input_is_folder: bool,
    input_base_dir: Path,
    output_dir: Path | None,
) -> Path:
    if output_dir is None:
        output_dir = input_base_dir / "type3_ocr_output"

    if input_is_folder:
        return output_dir / "combined_grades.xlsx"

    return output_dir / f"{safe_file_part(pdf_paths[0].stem)}.xlsx"


def write_error_log(errors: list[dict[str, str]], output_dir: Path) -> Path | None:
    if not errors:
        return None

    output_dir.mkdir(parents=True, exist_ok=True)
    log_path = output_dir / "type3_ocr_extraction_errors.csv"
    fieldnames = ["pdf", "page", "reason", "raw_text"]

    with log_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for error in errors:
            writer.writerow({field: error.get(field, "") for field in fieldnames})

    return log_path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert type-3 image/scanned grade PDFs into Excel using grid-based OCR."
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=None,
        help="PDF file path or folder path. If omitted, the script asks for it.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Optional folder where the Excel file and error log will be saved.",
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=350,
        help="Rendering DPI. Try 300, 350, or 400. Default: 350.",
    )
    parser.add_argument(
        "--ocr-lang",
        default="tur+eng",
        help="Tesseract language code. Default: tur+eng.",
    )
    parser.add_argument(
        "--tesseract-cmd",
        default=None,
        help="Full path to tesseract.exe if it is not on PATH.",
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="Save OCR header text and grid-overlay images for troubleshooting.",
    )
    return parser.parse_args(argv)


def get_input_path(args: argparse.Namespace) -> str:
    if args.input_path:
        return args.input_path

    user_input = input("Enter a PDF file path or a folder path containing PDF files: ").strip()
    if not user_input:
        raise ValueError("No input path was entered.")
    return user_input.strip('"')


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    configure_tesseract(args.tesseract_cmd)
    verify_tesseract_available()
    ocr_lang = resolve_ocr_language(args.ocr_lang)

    input_path = get_input_path(args)
    pdf_paths, input_is_folder, input_base_dir = resolve_input_pdfs(input_path)
    output_dir = Path(args.output_dir).expanduser() if args.output_dir else None
    output_path = build_output_path(pdf_paths, input_is_folder, input_base_dir, output_dir)
    log_output_dir = output_path.parent
    debug_dir = log_output_dir / "_type3_ocr_debug" if args.debug else None

    parsed_pdfs: list[ParsedPdf] = []
    errors: list[dict[str, str]] = []

    print(f"PDF files found: {len(pdf_paths)}")
    print(f"OCR language: {ocr_lang}")
    print(f"Render DPI: {args.dpi}")

    for index, pdf_path in enumerate(pdf_paths, start=1):
        try:
            parsed = extract_pdf(
                pdf_path=pdf_path,
                dpi=args.dpi,
                lang=ocr_lang,
                errors=errors,
                debug_dir=debug_dir,
            )
            parsed_pdfs.append(parsed)
            course_text = ", ".join(parsed.courses) if parsed.courses else "no included course detected"
            print(
                f"[{index}/{len(pdf_paths)}] OK: {pdf_path.name} "
                f"-> courses={course_text}, rows={len(parsed.rows)}"
            )
        except Exception as error:
            errors.append(
                {
                    "pdf": str(pdf_path),
                    "page": "",
                    "reason": str(error),
                    "raw_text": "",
                }
            )
            print(f"[{index}/{len(pdf_paths)}] ERROR: {pdf_path.name}: {error}", file=sys.stderr)

    if not parsed_pdfs:
        log_path = write_error_log(errors, log_output_dir)
        if log_path:
            print(f"Error log saved: {log_path}")
        raise RuntimeError("No PDFs could be converted.")

    rows, course_columns = combine_parsed_pdfs(parsed_pdfs)
    if not rows:
        raise RuntimeError("No student rows were extracted from the processed PDFs.")

    save_excel(rows, course_columns, output_path)
    log_path = write_error_log(errors, log_output_dir)

    print(f"Converted PDFs: {len(parsed_pdfs)}")
    print(f"Created: {output_path}")
    if log_path:
        print(f"Error log saved: {log_path}")
    if debug_dir:
        print(f"Debug files saved: {debug_dir}")

    return 1 if errors else 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
