# -*- coding: utf-8 -*-
r"""
OCR-only extractor for the type-2 one-page grade table format.

Use this script when the type-2 PDF is scanned/photographed, or when the PDF
has a broken/unreliable selectable-text layer. The script renders each PDF page
as a high-resolution image, detects the visual table grid, OCRs each needed
cell separately, and writes one Excel file per PDF.

This version preserves blank grades by table cell position. It never assigns
grades by counting visible grade numbers in a row.

Install Python packages:

    pip install pymupdf pillow pytesseract opencv-python numpy openpyxl

You must also install the Tesseract OCR program itself. On Windows, install
Tesseract and Turkish language data if possible. If Tesseract is installed in
a non-standard location, pass its path with --tesseract-cmd.

Run with a path:

    python extract_pdf_grades_type2_image_ocr.py "C:\path\to\pdf_or_folder"

Or run without a path and the script will ask you to enter one:

    python extract_pdf_grades_type2_image_ocr.py

Useful options:

    --dpi 350
    --ocr-lang tur+eng
    --tesseract-cmd "C:\Program Files\Tesseract-OCR\tesseract.exe"
    --debug-images
"""

from __future__ import annotations

import argparse
import csv
import io
import math
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

try:
    import cv2
    import fitz  # PyMuPDF
    import numpy as np
    import pytesseract
    from PIL import Image, ImageDraw, ImageEnhance, ImageOps
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    missing_name = getattr(exc, "name", "required package")
    raise SystemExit(
        f"Missing Python package: {missing_name}\n"
        "Install the requirements with:\n\n"
        "    pip install pymupdf pillow pytesseract opencv-python numpy openpyxl\n"
    ) from exc


OUTPUT_COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "HAYAT BİLGİSİ",
    "GÖRSEL SANATLAR",
    "YABANCI DİL",
    "MÜZİK",
    "FEN BİLİMLERİ",
    "BEDEN EĞİTİMİ VE OYUN",
]

PDF_COURSES = [
    *OUTPUT_COURSES,
    "SERBEST ETKİNLİKLER",
]

OUTPUT_COLUMNS = [
    "İlçe",
    "Okul",
    "Dönem",
    "Sınıf",
    "Şube",
    "Okul No",
    "Ad Soyad",
    *OUTPUT_COURSES,
]

SUMMARY_KEYWORDS = [
    "TOPLAM ÖĞRENCİ SAYISI",
    "BAŞARILI ÖĞRENCİ SAYISI",
    "BAŞARISIZ ÖĞRENCİ SAYISI",
    "BAŞARI YÜZDESİ",
    "NOT ORTALAMASI",
    "ZAYIFI OLAN",
]

SUMMARY_TOKEN_GROUPS = [
    ("TOPLAM", "OGRENCI"),
    ("BASARILI", "OGRENCI"),
    ("BASARISIZ", "OGRENCI"),
    ("BASARI", "YUZDE"),
    ("NOT", "ORTALAMA"),
    ("ZAYIFI", "OLAN"),
]

TURKISH_TO_ASCII = str.maketrans(
    {
        "ç": "c",
        "Ç": "C",
        "ğ": "g",
        "Ğ": "G",
        "ı": "i",
        "I": "I",
        "İ": "I",
        "ö": "o",
        "Ö": "O",
        "ş": "s",
        "Ş": "S",
        "ü": "u",
        "Ü": "U",
    }
)

RESAMPLE_LANCZOS = getattr(getattr(Image, "Resampling", Image), "LANCZOS")


@dataclass
class TableGrid:
    horizontal_lines: list[int]
    vertical_lines: list[int]


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def fold_for_match(value: object) -> str:
    text = clean_text(value).translate(TURKISH_TO_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def has_letter(value: str) -> bool:
    return any(char.isalpha() for char in value)


def line_has_summary_text(value: str) -> bool:
    folded = fold_for_match(value)
    if any(fold_for_match(keyword) in folded for keyword in SUMMARY_KEYWORDS):
        return True
    return any(all(token in folded for token in tokens) for tokens in SUMMARY_TOKEN_GROUPS)


def normalize_school_no(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"\d+", text)
    if not match:
        return ""
    return re.sub(r"\s+", "", match.group(0))


def normalize_name(value: object) -> str:
    text = clean_text(value)
    text = text.replace("|", " ").replace("_", " ")
    text = re.sub(r"[^0-9A-Za-zÇĞİÖŞÜçğıöşüÂâÎîÛû\s.'-]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text.upper()


def normalize_grade(value: object) -> str:
    text = clean_text(value)
    if not text:
        return ""

    match = re.search(r"\d+(?:[,.]\d+)?", text)
    if not match:
        return ""

    grade = match.group(0).replace(",", ".")
    return grade


def normalize_course_grade(value: object) -> str:
    grade = normalize_grade(value)
    if not grade:
        return ""

    if re.fullmatch(r"[123]\d+", grade):
        return grade[0]

    if re.fullmatch(r"[123]0+", grade):
        return grade[0]

    decimal_match = re.fullmatch(r"([123])\.0+", grade)
    if decimal_match:
        return decimal_match.group(1)

    return grade


def resolve_input_pdfs(raw_path: str | Path) -> list[Path]:
    path = Path(raw_path).expanduser()

    if path.is_dir():
        pdf_files = sorted(item for item in path.glob("*.pdf") if item.is_file())
        if not pdf_files:
            raise FileNotFoundError(f"No PDF files were found in folder: {path}")
        return pdf_files

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(f"The input file is not a PDF: {path}")
        return [path]

    if path.suffix.lower() != ".pdf":
        pdf_version = path.with_suffix(".pdf")
        if pdf_version.is_file():
            return [pdf_version]

    raise FileNotFoundError(
        f"PDF file or folder was not found: {path}\n"
        "Tip: pass one PDF file path or a folder containing PDF files."
    )


def configure_tesseract(tesseract_cmd: str | None) -> None:
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
        return

    common_windows_path = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    if common_windows_path.exists():
        pytesseract.pytesseract.tesseract_cmd = str(common_windows_path)


def verify_tesseract_available() -> None:
    try:
        pytesseract.get_tesseract_version()
    except Exception as error:
        raise RuntimeError(
            "Tesseract OCR is not available. Install Tesseract OCR, then rerun. "
            "If it is already installed, pass --tesseract-cmd with the full "
            "path to tesseract.exe."
        ) from error


def resolve_ocr_language(requested_lang: str) -> str:
    try:
        available_languages = set(pytesseract.get_languages(config=""))
    except Exception:
        return requested_lang

    requested_parts = [part.strip() for part in requested_lang.split("+") if part.strip()]
    if not requested_parts:
        return requested_lang

    available_parts = [part for part in requested_parts if part in available_languages]
    missing_parts = [part for part in requested_parts if part not in available_languages]

    if available_parts and missing_parts:
        resolved_lang = "+".join(available_parts)
        print(
            "Warning: Tesseract language data not found for "
            f"{'+'.join(missing_parts)}; using {resolved_lang}.",
            file=sys.stderr,
        )
        return resolved_lang

    if available_parts:
        return requested_lang

    if "eng" in available_languages:
        print(
            f"Warning: Tesseract language data not found for {requested_lang}; using eng.",
            file=sys.stderr,
        )
        return "eng"

    return requested_lang


def render_pdf_page(page, dpi: int) -> Image.Image:
    zoom = dpi / 72
    pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    return Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB")


def preprocess_for_ocr(image: Image.Image) -> Image.Image:
    grayscale = ImageOps.grayscale(image)
    grayscale = ImageEnhance.Contrast(grayscale).enhance(1.8)
    array = np.array(grayscale)
    threshold = cv2.adaptiveThreshold(
        array,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY,
        35,
        11,
    )
    return Image.fromarray(threshold)


def pil_to_gray_array(image: Image.Image) -> np.ndarray:
    if image.mode != "L":
        image = ImageOps.grayscale(image)
    return np.array(image)


def enhance_cell_for_ocr(cell: Image.Image, scale: float = 2.0) -> Image.Image:
    grayscale = ImageOps.grayscale(cell)
    grayscale = ImageEnhance.Contrast(grayscale).enhance(2.0)

    if scale != 1:
        new_size = (
            max(1, int(grayscale.width * scale)),
            max(1, int(grayscale.height * scale)),
        )
        grayscale = grayscale.resize(new_size, RESAMPLE_LANCZOS)

    array = np.array(grayscale)
    threshold = cv2.threshold(array, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
    return Image.fromarray(threshold)


def ocr_image(
    image: Image.Image,
    lang: str,
    psm: int,
    whitelist: str | None = None,
    preserve_spaces: bool = False,
) -> str:
    config_parts = [f"--psm {psm}", "--oem 3"]
    if whitelist:
        config_parts.append(f"-c tessedit_char_whitelist={whitelist}")
    if preserve_spaces:
        config_parts.append("-c preserve_interword_spaces=1")

    config = " ".join(config_parts)
    try:
        return pytesseract.image_to_string(image, lang=lang, config=config)
    except pytesseract.TesseractError as error:
        if "Failed loading language" in str(error) and lang != "eng":
            return pytesseract.image_to_string(image, lang="eng", config=config)
        raise


def extract_header_info(header_text: str) -> dict[str, str]:
    lines = [clean_text(line) for line in header_text.splitlines() if clean_text(line)]

    district = ""
    school = ""
    period = ""
    class_no = ""
    section = ""

    for line in lines:
        folded = fold_for_match(line)
        if "/" not in line:
            continue
        if "MUDURLUGU" not in folded and "MUDURLUG" not in folded and "OKUL" not in folded:
            continue

        district_part, school_part = line.split("/", 1)
        district = clean_text(district_part)
        school_part = clean_text(school_part)
        school = re.split(
            r"\s+M[üu]d[üu]rl[üu][ğg][üu]\b",
            school_part,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0].strip()
        if school == school_part:
            folded_school = fold_for_match(school_part)
            school = re.sub(r"\s*MUDURLUGU?.*$", "", folded_school).title()
        break

    for line in lines:
        if "DONEM" in fold_for_match(line):
            period = line
            break

    folded_header = fold_for_match(header_text)
    class_match = re.search(
        r"\b(?P<class>\d+)\s*SINIF\s*(?:/)?\s*(?P<section>[A-Z0-9]+)\s*SUBESI\b",
        folded_header,
    )
    if class_match:
        class_no = class_match.group("class")
        section = class_match.group("section")

    return {
        "district": district,
        "school": school,
        "period": period,
        "class_no": class_no,
        "section": section,
    }


def merge_positions(indices: np.ndarray, max_gap: int = 6) -> list[int]:
    if len(indices) == 0:
        return []

    groups: list[list[int]] = []
    current = [int(indices[0])]
    for index in map(int, indices[1:]):
        if index - current[-1] <= max_gap:
            current.append(index)
        else:
            groups.append(current)
            current = [index]
    groups.append(current)

    return [int(round(sum(group) / len(group))) for group in groups]


def detect_table_lines(processed_image: Image.Image) -> tuple[list[int], list[int]]:
    gray = pil_to_gray_array(processed_image)
    binary_inverse = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]

    height, width = binary_inverse.shape
    horizontal_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (max(width // 30, 50), 3))
    vertical_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (3, max(height // 35, 45)))

    horizontal = cv2.morphologyEx(binary_inverse, cv2.MORPH_OPEN, horizontal_kernel)
    vertical = cv2.morphologyEx(binary_inverse, cv2.MORPH_OPEN, vertical_kernel)

    horizontal_projection = np.count_nonzero(horizontal, axis=1)
    vertical_projection = np.count_nonzero(vertical, axis=0)

    y_indices = np.where(horizontal_projection > width * 0.18)[0]
    x_indices = np.where(vertical_projection > height * 0.08)[0]

    horizontal_lines = merge_positions(y_indices, max_gap=8)
    vertical_lines = merge_positions(x_indices, max_gap=8)

    horizontal_lines = [y for y in horizontal_lines if 0 <= y < height]
    vertical_lines = [x for x in vertical_lines if 0 <= x < width]
    return horizontal_lines, vertical_lines


def choose_vertical_lines(vertical_lines: list[int], width: int) -> list[int]:
    lines = sorted(x for x in vertical_lines if width * 0.005 <= x <= width * 0.995)
    expected_boundaries = len(PDF_COURSES) + 4  # row no, okul no, name, 9 courses -> 13 lines

    if len(lines) < expected_boundaries:
        raise ValueError(
            f"table vertical lines could not be detected: found {len(lines)}, "
            f"need at least {expected_boundaries}"
        )

    if len(lines) == expected_boundaries:
        return lines

    layout_candidates: list[tuple[float, list[int]]] = []
    span_candidates: list[tuple[int, list[int]]] = []
    for start in range(0, len(lines) - expected_boundaries + 1):
        subset = lines[start : start + expected_boundaries]
        column_widths = [right - left for left, right in zip(subset, subset[1:])]
        span = subset[-1] - subset[0]
        if span >= width * 0.55:
            span_candidates.append((span, subset))

        if len(column_widths) < 5:
            continue

        row_no_width = column_widths[0]
        school_no_width = column_widths[1]
        name_width = column_widths[2]
        course_widths = column_widths[3:]
        median_course_width = float(np.median(course_widths))

        if median_course_width <= 0:
            continue

        name_ratio = name_width / median_course_width
        school_ratio = school_no_width / median_course_width
        row_no_ratio = row_no_width / median_course_width
        course_uniformity = float(np.std(course_widths) / median_course_width)

        # The correct type-2 layout starts with row number, school number, and
        # a wide student-name column. Extra blank grade-like columns may appear
        # to the right, so "widest span" alone can start one column too late.
        if (
            name_ratio >= 3.0
            and 0.25 <= row_no_ratio <= 1.8
            and 1.0 <= school_ratio <= 4.5
            and course_uniformity <= 0.45
        ):
            score = (
                (name_ratio * 100)
                + ((1.0 - min(course_uniformity, 1.0)) * 50)
                - abs(school_ratio - 2.2) * 8
                - abs(row_no_ratio - 0.8) * 8
                - start
            )
            layout_candidates.append((score, subset))

    if layout_candidates:
        return max(layout_candidates, key=lambda item: item[0])[1]

    if span_candidates:
        return max(span_candidates, key=lambda item: item[0])[1]

    return lines[:expected_boundaries]


def choose_horizontal_lines(horizontal_lines: list[int], height: int) -> list[int]:
    lines = sorted(y for y in horizontal_lines if height * 0.05 <= y <= height * 0.98)

    if len(lines) < 4:
        raise ValueError(
            f"table horizontal lines could not be detected: found {len(lines)}, need at least 4"
        )

    # If there are decorative/header lines above the student table, keep the lower cluster
    # that contains row-sized gaps.
    for start in range(len(lines)):
        candidate = lines[start:]
        if len(candidate) < 4:
            break
        gaps = [b - a for a, b in zip(candidate, candidate[1:])]
        normal_gaps = [gap for gap in gaps if 8 <= gap <= height * 0.08]
        if len(normal_gaps) >= 3:
            return candidate

    return lines


def detect_table_grid(processed_image: Image.Image) -> TableGrid:
    horizontal_lines, vertical_lines = detect_table_lines(processed_image)
    width, height = processed_image.size

    return TableGrid(
        horizontal_lines=choose_horizontal_lines(horizontal_lines, height),
        vertical_lines=choose_vertical_lines(vertical_lines, width),
    )


def estimate_page_skew_angle(image: Image.Image) -> float:
    """Estimate the visual tilt of long horizontal table lines in degrees."""
    grayscale = np.array(ImageOps.grayscale(image))
    height, width = grayscale.shape

    max_width = 2200
    if width > max_width:
        scale = max_width / width
        grayscale = cv2.resize(
            grayscale,
            (max_width, max(1, int(height * scale))),
            interpolation=cv2.INTER_AREA,
        )
        height, width = grayscale.shape

    blurred = cv2.GaussianBlur(grayscale, (5, 5), 0)
    binary_inverse = cv2.threshold(
        blurred,
        0,
        255,
        cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU,
    )[1]

    min_line_length = max(int(width * 0.16), 250)
    max_line_gap = max(int(width * 0.015), 20)
    threshold = max(int(width * 0.05), 80)
    lines = cv2.HoughLinesP(
        binary_inverse,
        rho=1,
        theta=np.pi / 180,
        threshold=threshold,
        minLineLength=min_line_length,
        maxLineGap=max_line_gap,
    )

    if lines is None:
        return 0.0

    angles: list[float] = []
    for line in lines:
        x1, y1, x2, y2 = line[0]
        length = math.hypot(float(x2 - x1), float(y2 - y1))
        if length < min_line_length:
            continue

        angle = math.degrees(math.atan2(float(y2 - y1), float(x2 - x1)))
        if angle <= -90:
            angle += 180
        elif angle > 90:
            angle -= 180

        if abs(angle) <= 15:
            angles.append(angle)

    if len(angles) < 3:
        return 0.0

    median_angle = float(np.median(angles))
    if abs(median_angle) < 0.20 or abs(median_angle) > 8.0:
        return 0.0
    return median_angle


def rotate_image(image: Image.Image, angle: float) -> Image.Image:
    """Rotate without cutting page corners, filling the new background white."""
    array = np.array(image.convert("RGB"))
    height, width = array.shape[:2]
    center = (width / 2, height / 2)
    matrix = cv2.getRotationMatrix2D(center, angle, 1.0)

    cos_value = abs(matrix[0, 0])
    sin_value = abs(matrix[0, 1])
    new_width = int((height * sin_value) + (width * cos_value))
    new_height = int((height * cos_value) + (width * sin_value))

    matrix[0, 2] += (new_width / 2) - center[0]
    matrix[1, 2] += (new_height / 2) - center[1]

    rotated = cv2.warpAffine(
        array,
        matrix,
        (new_width, new_height),
        flags=cv2.INTER_CUBIC,
        borderMode=cv2.BORDER_CONSTANT,
        borderValue=(255, 255, 255),
    )
    return Image.fromarray(rotated)


def score_grid_candidate(image: Image.Image) -> tuple[int, int, int]:
    processed_image = preprocess_for_ocr(image)
    horizontal_lines, vertical_lines = detect_table_lines(processed_image)
    width, height = processed_image.size

    horizontal_count = len([y for y in horizontal_lines if height * 0.05 <= y <= height * 0.98])
    vertical_count = len([x for x in vertical_lines if width * 0.005 <= x <= width * 0.995])
    score = (horizontal_count * 20) + (vertical_count * 8)

    try:
        choose_horizontal_lines(horizontal_lines, height)
        score += 500
    except ValueError:
        pass

    try:
        choose_vertical_lines(vertical_lines, width)
        score += 500
    except ValueError:
        pass

    return score, horizontal_count, vertical_count


def deskew_page_image(image: Image.Image) -> tuple[Image.Image, float]:
    angle = estimate_page_skew_angle(image)
    if not angle:
        return image, 0.0

    candidates: list[tuple[float, Image.Image]] = [(0.0, image)]
    for candidate_angle in (angle, -angle):
        if all(abs(candidate_angle - existing_angle) > 0.05 for existing_angle, _ in candidates):
            candidates.append((candidate_angle, rotate_image(image, candidate_angle)))

    scored_candidates: list[tuple[int, int, int, float, Image.Image]] = []
    for candidate_angle, candidate_image in candidates:
        score, horizontal_count, vertical_count = score_grid_candidate(candidate_image)
        scored_candidates.append(
            (score, horizontal_count, vertical_count, candidate_angle, candidate_image)
        )

    best_score, _horizontal_count, _vertical_count, best_angle, best_image = max(
        scored_candidates,
        key=lambda item: item[:3],
    )
    original_score = scored_candidates[0][0]

    if best_angle and best_score > original_score:
        return best_image, best_angle
    return image, 0.0


def crop_cell(
    image: Image.Image,
    left: int,
    top: int,
    right: int,
    bottom: int,
    pad_x: int = 6,
    pad_y: int = 3,
) -> Image.Image:
    width, height = image.size
    cell_width = max(1, right - left)
    cell_height = max(1, bottom - top)
    safe_pad_x = min(pad_x, max(1, cell_width // 8))
    safe_pad_y = min(pad_y, max(1, cell_height // 6))

    left = max(0, left + safe_pad_x)
    top = max(0, top + safe_pad_y)
    right = min(width, right - safe_pad_x)
    bottom = min(height, bottom - safe_pad_y)

    if right <= left or bottom <= top:
        return image.crop((0, 0, 1, 1))
    return image.crop((left, top, right, bottom))


def crop_cell_asymmetric(
    image: Image.Image,
    left: int,
    top: int,
    right: int,
    bottom: int,
    pad_left: int = 1,
    pad_top: int = 1,
    pad_right: int = 10,
    pad_bottom: int = 1,
) -> Image.Image:
    width, height = image.size
    left = max(0, left + pad_left)
    top = max(0, top + pad_top)
    right = min(width, right - pad_right)
    bottom = min(height, bottom - pad_bottom)

    if right <= left or bottom <= top:
        return image.crop((0, 0, 1, 1))
    return image.crop((left, top, right, bottom))


def ocr_school_no(cell: Image.Image, lang: str) -> str:
    candidates: list[str] = []
    for psm, scale in ((7, 2.7), (7, 4.0), (8, 4.0), (10, 4.0), (13, 4.0)):
        text = ocr_image(
            enhance_cell_for_ocr(cell, scale=scale),
            lang=lang,
            psm=psm,
            whitelist="-0123456789",
        )
        school_no = normalize_school_no(text)
        if school_no:
            candidates.append(school_no)

    if not candidates:
        return ""

    return max(candidates, key=len)


def ocr_student_name(cell: Image.Image, lang: str) -> str:
    text = ocr_image(
        enhance_cell_for_ocr(cell, scale=2.2),
        lang=lang,
        psm=7,
        preserve_spaces=True,
    )
    return normalize_name(text)


def extract_grade_shape_features(cell: Image.Image) -> dict[str, float] | None:
    grayscale = np.array(ImageOps.grayscale(cell))
    if grayscale.size == 0:
        return None

    scaled = cv2.resize(
        grayscale,
        (max(1, grayscale.shape[1] * 4), max(1, grayscale.shape[0] * 4)),
        interpolation=cv2.INTER_CUBIC,
    )
    blurred = cv2.GaussianBlur(scaled, (3, 3), 0)
    binary_inverse = cv2.threshold(
        blurred,
        0,
        255,
        cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU,
    )[1]

    height, width = binary_inverse.shape
    horizontal_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (max(width // 3, 30), 3),
    )
    vertical_kernel = cv2.getStructuringElement(
        cv2.MORPH_RECT,
        (3, max(height // 3, 20)),
    )
    horizontal_lines = cv2.morphologyEx(binary_inverse, cv2.MORPH_OPEN, horizontal_kernel)
    vertical_lines = cv2.morphologyEx(binary_inverse, cv2.MORPH_OPEN, vertical_kernel)
    without_lines = cv2.subtract(binary_inverse, cv2.bitwise_or(horizontal_lines, vertical_lines))
    without_lines = cv2.morphologyEx(
        without_lines,
        cv2.MORPH_CLOSE,
        cv2.getStructuringElement(cv2.MORPH_RECT, (3, 3)),
    )

    component_count, labels, stats, centroids = cv2.connectedComponentsWithStats(
        without_lines,
        8,
    )
    candidates: list[tuple[int, int, int, int, int, int]] = []

    for component_index in range(1, component_count):
        x, y, component_width, component_height, area = stats[component_index]
        center_x, center_y = centroids[component_index]

        if area < max(20, height * width * 0.003):
            continue
        if component_height < height * 0.18 or component_width < width * 0.03:
            continue
        if center_x < width * 0.12 or center_x > width * 0.88:
            continue
        if center_y < height * 0.12 or center_y > height * 0.88:
            continue

        candidates.append(
            (
                int(area),
                int(x),
                int(y),
                int(component_width),
                int(component_height),
                component_index,
            )
        )

    if not candidates:
        return None

    area, x, y, component_width, component_height, component_index = max(
        candidates,
        key=lambda item: item[0],
    )
    digit_mask = (labels[y : y + component_height, x : x + component_width] == component_index)
    digit_height, digit_width = digit_mask.shape

    if digit_width == 0 or digit_height == 0:
        return None

    lower_half = digit_mask[(digit_height * 2) // 3 :, :]
    upper_third = digit_mask[: max(1, digit_height // 3), :]
    lower_left = int(lower_half[:, : max(1, digit_width // 2)].sum())
    lower_right = int(lower_half[:, max(1, digit_width // 2) :].sum())
    upper_left = int(upper_third[:, : max(1, digit_width // 2)].sum())
    upper_right = int(upper_third[:, max(1, digit_width // 2) :].sum())

    return {
        "area": float(area),
        "width_ratio": component_width / max(width, 1),
        "height_ratio": component_height / max(height, 1),
        "aspect_ratio": component_width / max(component_height, 1),
        "lower_left_ratio": lower_left / max(lower_right, 1),
        "upper_left_ratio": upper_left / max(upper_right, 1),
    }


def classify_grade_by_shape(cell: Image.Image) -> str:
    features = extract_grade_shape_features(cell)
    if features is None:
        return ""

    width_ratio = features["width_ratio"]
    height_ratio = features["height_ratio"]
    area = features["area"]
    lower_left_ratio = features["lower_left_ratio"]

    # A true "1" is a narrow central mark. Misread 3s are usually much wider.
    if width_ratio <= 0.11 and area < 1300:
        return "1"

    # A "2" has a strong lower-left stroke; 3s are mostly right-heavy below.
    if lower_left_ratio >= 1.20 and height_ratio >= 0.38 and area >= 1500:
        return "2"

    return "3"


def ocr_grade(cell: Image.Image, lang: str) -> str:
    text = ocr_image(
        enhance_cell_for_ocr(cell, scale=2.8),
        lang=lang,
        psm=10,
        whitelist="123",
    )
    ocr_grade_value = normalize_course_grade(text)
    shape_grade = classify_grade_by_shape(cell)

    if not ocr_grade_value:
        return shape_grade

    if ocr_grade_value == "1" and shape_grade in {"2", "3"}:
        return shape_grade

    return ocr_grade_value


def build_output_row(
    header: dict[str, str],
    school_no: str,
    name: str,
    grades: list[str],
) -> dict[str, str]:
    row = {
        "İlçe": header["district"],
        "Okul": header["school"],
        "Dönem": header["period"],
        "Sınıf": header["class_no"],
        "Şube": header["section"],
        "Okul No": school_no,
        "Ad Soyad": name,
    }

    for index, course in enumerate(OUTPUT_COURSES):
        row[course] = grades[index] if index < len(grades) else ""

    return row


def extract_rows_from_grid(
    page_image: Image.Image,
    grid: TableGrid,
    header: dict[str, str],
    lang: str,
    pdf_path: Path,
    errors: list[dict[str, str]],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    horizontal_lines = grid.horizontal_lines
    vertical_lines = grid.vertical_lines

    # Row 0 is the course/header row; student rows begin after it.
    student_row_bands = list(zip(horizontal_lines[1:-1], horizontal_lines[2:]))

    for row_number, (row_top, row_bottom) in enumerate(student_row_bands, start=1):
        if row_bottom - row_top < 8:
            continue

        serial_cell = crop_cell(page_image, vertical_lines[0], row_top, vertical_lines[1], row_bottom)
        school_no_cell = crop_cell_asymmetric(
            page_image,
            vertical_lines[1],
            row_top,
            vertical_lines[2],
            row_bottom,
            pad_left=1,
            pad_top=1,
            pad_right=10,
            pad_bottom=1,
        )
        name_cell = crop_cell(page_image, vertical_lines[2], row_top, vertical_lines[3], row_bottom)

        serial_text = ocr_image(
            enhance_cell_for_ocr(serial_cell, scale=2.2),
            lang=lang,
            psm=10,
            whitelist="0123456789",
        )
        serial = normalize_grade(serial_text)
        school_no = ocr_school_no(school_no_cell, lang)
        name = ocr_student_name(name_cell, lang)

        if not serial and not school_no and not name:
            continue
        if line_has_summary_text(name):
            continue
        if not school_no and not name:
            continue

        if not school_no or not name:
            errors.append(
                {
                    "pdf": str(pdf_path),
                    "row": str(row_number),
                    "reason": "missing school number or student name",
                    "raw_text": f"serial={serial_text!r}; school_no={school_no!r}; name={name!r}",
                }
            )
            continue

        grades: list[str] = []
        for course_index in range(len(OUTPUT_COURSES)):
            left = vertical_lines[3 + course_index]
            right = vertical_lines[4 + course_index]
            grade_cell = crop_cell(page_image, left, row_top, right, row_bottom, pad_x=1, pad_y=1)
            grades.append(ocr_grade(grade_cell, lang))

        rows.append(build_output_row(header, school_no, name, grades))

    return rows


def safe_file_part(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" ._")
    return value or "bilinmeyen"


def save_debug_images(
    debug_dir: Path,
    pdf_path: Path,
    page_image: Image.Image,
    processed_image: Image.Image,
    grid: TableGrid | None,
    raw_page_image: Image.Image | None = None,
) -> None:
    debug_dir.mkdir(parents=True, exist_ok=True)
    safe_stem = safe_file_part(pdf_path.stem)
    if raw_page_image is not None:
        raw_page_image.save(debug_dir / f"{safe_stem}_raw_render.png")
    page_image.save(debug_dir / f"{safe_stem}_render.png")
    processed_image.save(debug_dir / f"{safe_stem}_processed.png")

    if grid is None:
        return

    overlay = page_image.copy()
    draw = ImageDraw.Draw(overlay)
    width, height = overlay.size

    for y in grid.horizontal_lines:
        draw.line((0, y, width, y), fill=(255, 0, 0), width=3)
    for x in grid.vertical_lines:
        draw.line((x, 0, x, height), fill=(0, 0, 255), width=3)

    overlay.save(debug_dir / f"{safe_stem}_grid_overlay.png")


def extract_pdf(
    pdf_path: Path,
    dpi: int,
    lang: str,
    errors: list[dict[str, str]],
    debug_dir: Path | None,
) -> list[dict[str, str]]:
    with fitz.open(str(pdf_path)) as document:
        if document.page_count != 1:
            raise ValueError(
                f"Expected exactly 1 page, but found {document.page_count} pages in {pdf_path.name}."
            )

        page = document.load_page(0)
        raw_page_image = render_pdf_page(page, dpi=dpi)
        page_image, _deskew_angle = deskew_page_image(raw_page_image)
        processed_image = preprocess_for_ocr(page_image)

        grid: TableGrid | None = None
        try:
            grid = detect_table_grid(processed_image)
        finally:
            if debug_dir:
                save_debug_images(
                    debug_dir,
                    pdf_path,
                    page_image,
                    processed_image,
                    grid,
                    raw_page_image=raw_page_image,
                )

        header_height = int(page_image.height * 0.35)
        header_crop = page_image.crop((0, 0, page_image.width, header_height))
        header_ocr_image = enhance_cell_for_ocr(header_crop, scale=1.4)
        header_text = ocr_image(header_ocr_image, lang=lang, psm=6, preserve_spaces=True)
        header = extract_header_info(header_text)

        rows = extract_rows_from_grid(
            page_image=page_image,
            grid=grid,
            header=header,
            lang=lang,
            pdf_path=pdf_path,
            errors=errors,
        )

        if not rows:
            raise ValueError(f"No student rows could be extracted from {pdf_path.name}.")

        return rows


def save_excel(rows: list[dict[str, str]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Notlar"

    worksheet.append(OUTPUT_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column, "") for column in OUTPUT_COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    okul_no_column = OUTPUT_COLUMNS.index("Okul No") + 1
    okul_no_letter = get_column_letter(okul_no_column)
    for row_number in range(2, worksheet.max_row + 1):
        worksheet[f"{okul_no_letter}{row_number}"].number_format = "@"

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        header_value = column_cells[0].value

        if header_value == "Ad Soyad":
            width = min(max(max_length + 2, 18), 38)
        elif header_value in {"İlçe", "Okul"}:
            width = min(max(max_length + 2, 14), 34)
        elif header_value in OUTPUT_COURSES:
            width = min(max(max_length + 2, 12), 26)
        else:
            width = min(max(max_length + 2, 10), 18)

        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 42
    workbook.save(output_path)


def build_output_path(pdf_path: Path, output_dir: Path | None) -> Path:
    folder = output_dir if output_dir is not None else pdf_path.parent
    return folder / f"{pdf_path.stem}.xlsx"


def write_error_log(errors: list[dict[str, str]], output_dir: Path) -> Path | None:
    if not errors:
        return None

    output_dir.mkdir(parents=True, exist_ok=True)
    log_path = output_dir / "type2_ocr_extraction_errors.csv"
    fieldnames = ["pdf", "row", "reason", "raw_text"]

    with log_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for error in errors:
            writer.writerow({field: error.get(field, "") for field in fieldnames})

    return log_path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert type-2 image/scanned grade-table PDFs into Excel using OCR."
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=None,
        help="PDF file path or folder path. If omitted, the script asks for it.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Optional folder where Excel files and error logs will be saved.",
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=350,
        help="Rendering DPI. Try 300, 350, or 400. Default: 350.",
    )
    parser.add_argument(
        "--ocr-lang",
        default="tur+eng",
        help="Tesseract language code. Default: tur+eng.",
    )
    parser.add_argument(
        "--tesseract-cmd",
        default=None,
        help="Full path to tesseract.exe if it is not on PATH.",
    )
    parser.add_argument(
        "--debug-images",
        action="store_true",
        help="Save rendered, processed, and grid-overlay images for troubleshooting.",
    )
    return parser.parse_args(argv)


def get_input_path(args: argparse.Namespace) -> str:
    if args.input_path:
        return args.input_path

    user_input = input("Enter a PDF file path or a folder path containing PDF files: ").strip()
    if not user_input:
        raise ValueError("No input path was entered.")
    return user_input.strip('"')


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    configure_tesseract(args.tesseract_cmd)
    verify_tesseract_available()
    ocr_lang = resolve_ocr_language(args.ocr_lang)

    input_path = get_input_path(args)
    pdf_paths = resolve_input_pdfs(input_path)
    output_dir = Path(args.output_dir).expanduser() if args.output_dir else None
    log_output_dir = output_dir if output_dir else pdf_paths[0].parent
    debug_dir = (log_output_dir / "_type2_ocr_debug_images") if args.debug_images else None

    created_files: list[Path] = []
    errors: list[dict[str, str]] = []
    failures: list[tuple[Path, Exception]] = []

    print(f"PDF files found: {len(pdf_paths)}")
    print(f"OCR language: {ocr_lang}")
    print(f"Render DPI: {args.dpi}")

    for index, pdf_path in enumerate(pdf_paths, start=1):
        try:
            rows = extract_pdf(
                pdf_path=pdf_path,
                dpi=args.dpi,
                lang=ocr_lang,
                errors=errors,
                debug_dir=debug_dir,
            )
            output_path = build_output_path(pdf_path, output_dir)
            save_excel(rows, output_path)
            created_files.append(output_path)
            print(f"[{index}/{len(pdf_paths)}] OK: {pdf_path.name} -> {output_path}")
        except Exception as error:
            failures.append((pdf_path, error))
            errors.append(
                {
                    "pdf": str(pdf_path),
                    "row": "",
                    "reason": "pdf processing failed",
                    "raw_text": str(error),
                }
            )
            print(f"[{index}/{len(pdf_paths)}] ERROR: {pdf_path.name}: {error}", file=sys.stderr)

    log_path = write_error_log(errors, log_output_dir)

    print(f"Converted PDFs: {len(created_files)}")
    for output_path in created_files:
        print(f"Created: {output_path}")

    if log_path:
        print(f"Error log saved: {log_path}")

    if failures:
        print("\nFiles with errors:", file=sys.stderr)
        for pdf_path, error in failures:
            print(f"  - {pdf_path}: {error}", file=sys.stderr)
        return 1

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
