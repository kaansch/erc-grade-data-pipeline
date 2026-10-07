from __future__ import annotations

import math
import re
import sys
import unicodedata
from dataclasses import dataclass
from numbers import Number
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.cell_range import CellRange
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError as exc:
    raise SystemExit(
        "This program requires openpyxl. Install it in the VS Code terminal with:\n"
        "python -m pip install openpyxl"
    ) from exc


COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "FEN BİLİMLERİ",
    "SOSYAL BİLGİLER",
    "YABANCI DİL",
    "DİN KÜLTÜRÜ VE AHLAK BİLGİSİ",
    "GÖRSEL SANATLAR",
    "MÜZİK",
    "BEDEN EĞİTİMİ VE OYUN",
    "TRAFİK GÜVENLİĞİ",
    "İNSAN HAKLARI, YURTTAŞLIK VE DEMOKRASİ",
]

COURSE_FIRST_COLUMN = 3  # C
COURSE_LAST_COLUMN = 13  # M
OVERALL_GRADE_COLUMN = 15  # O
EXPECTED_STUDENT_COUNT = 99
LABEL_SEARCH_LAST_COLUMN = 40

NAME_LABEL = "ADI VE SOYADI"
WEEKLY_LABEL = "HAFTALIK DERS SAATI"
FIRST_TERM_LABEL = "1. DONEM PUANI"
SECOND_TERM_LABEL = "2. DONEM PUANI"
COURSE_YEAR_END_LABEL = "YIL SONU PUANI"
WEIGHTED_LABEL = "AGIRLIKLI PUANI"


@dataclass(frozen=True)
class NameAnchor:
    row: int
    column: int
    name: str | None


@dataclass(frozen=True)
class LabelAnchor:
    row: int
    column: int
    merged_range: CellRange | None

    @property
    def min_row(self) -> int:
        return self.merged_range.min_row if self.merged_range else self.row

    @property
    def max_row(self) -> int:
        return self.merged_range.max_row if self.merged_range else self.row


class MergedCellResolver:
    """Returns the top-left value and source coordinate for merged cells."""

    def __init__(self, worksheet: Worksheet) -> None:
        self.worksheet = worksheet
        self.ranges = [CellRange(str(item)) for item in worksheet.merged_cells.ranges]
        self._ranges_by_row: dict[int, list[CellRange]] = {}

        for merged_range in self.ranges:
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                self._ranges_by_row.setdefault(row, []).append(merged_range)

    def range_at(self, row: int, column: int) -> CellRange | None:
        for merged_range in self._ranges_by_row.get(row, []):
            if merged_range.min_col <= column <= merged_range.max_col:
                return merged_range
        return None

    def get(self, row: int, column: int) -> tuple[Any, tuple[int, int], CellRange | None]:
        merged_range = self.range_at(row, column)
        if merged_range is None:
            return self.worksheet.cell(row, column).value, (row, column), None

        source = (merged_range.min_row, merged_range.min_col)
        return self.worksheet.cell(*source).value, source, merged_range


def normalize_text(value: Any) -> str:
    if value is None:
        return ""

    text = str(value).replace("\u00a0", " ")
    text = text.translate(
        str.maketrans(
            {
                "ı": "i",
                "İ": "I",
                "ş": "s",
                "Ş": "S",
                "ğ": "g",
                "Ğ": "G",
                "ü": "u",
                "Ü": "U",
                "ö": "o",
                "Ö": "O",
                "ç": "c",
                "Ç": "C",
            }
        )
    )
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"\s+", " ", text).strip().upper()


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip()


def clean_grade(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, bool):
        return value

    if isinstance(value, Number):
        numeric_value = float(value)
        if not math.isfinite(numeric_value):
            return None
        if numeric_value.is_integer():
            return int(numeric_value)
        return numeric_value

    text = clean_text(value)
    if not text:
        return None

    numeric_text = text.replace(" ", "")
    if re.fullmatch(r"[+-]?\d+(?:[.,]\d*)?", numeric_text):
        normalized_numeric_text = numeric_text.rstrip(".,")
        numeric_value = float(normalized_numeric_text.replace(",", "."))
        if numeric_value.is_integer():
            return int(numeric_value)
        return numeric_value

    return text


def find_unused_output_path(input_path: Path) -> Path:
    candidate = input_path.with_name(f"{input_path.stem}_extracted.xlsx")
    counter = 2
    while candidate.exists():
        candidate = input_path.with_name(f"{input_path.stem}_extracted_{counter}.xlsx")
        counter += 1
    return candidate


def validate_input_path(raw_path: str) -> Path:
    cleaned_path = raw_path.strip().strip('"').strip("'")
    if not cleaned_path:
        raise ValueError("No input path was entered.")

    input_path = Path(cleaned_path).expanduser()
    if not input_path.exists():
        raise ValueError(f"The path does not exist: {input_path}")
    if not input_path.is_file():
        raise ValueError(f"The path is not a file: {input_path}")
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("The input file must have the .xlsx extension.")
    return input_path


def iter_relevant_cells(
    worksheet: Worksheet,
    start_row: int = 1,
    end_row: int | None = None,
) -> Iterable[tuple[int, int, Any]]:
    final_row = worksheet.max_row if end_row is None else min(end_row, worksheet.max_row)
    final_column = min(max(worksheet.max_column, OVERALL_GRADE_COLUMN), LABEL_SEARCH_LAST_COLUMN)

    for row in range(max(1, start_row), final_row + 1):
        for column in range(1, final_column + 1):
            value = worksheet.cell(row, column).value
            if value is not None:
                yield row, column, value


def extract_name_from_anchor(
    worksheet: Worksheet,
    resolver: MergedCellResolver,
    row: int,
    column: int,
    raw_value: Any,
) -> str | None:
    raw_text = clean_text(raw_value)

    for separator in (":", "："):
        if separator in raw_text:
            after_separator = clean_text(raw_text.split(separator, 1)[1])
            if after_separator:
                return after_separator

    source_seen = resolver.get(row, column)[1]
    final_column = min(max(worksheet.max_column, OVERALL_GRADE_COLUMN), LABEL_SEARCH_LAST_COLUMN)
    for candidate_column in range(column + 1, final_column + 1):
        candidate_value, candidate_source, _ = resolver.get(row, candidate_column)
        if candidate_source == source_seen:
            continue

        candidate_text = clean_text(candidate_value)
        candidate_normalized = normalize_text(candidate_value)
        if not candidate_text:
            continue
        if candidate_normalized in {
            NAME_LABEL,
            "BABA ADI",
            "ANA ADI",
            "DOGUM YERI",
            "UYRUGU",
            "OKUL NUMARASI",
        }:
            continue
        return candidate_text.lstrip(":： ") or None

    return None


def find_name_anchors(
    worksheet: Worksheet,
    resolver: MergedCellResolver,
) -> list[NameAnchor]:
    anchors: list[NameAnchor] = []

    for row, column, value in iter_relevant_cells(worksheet):
        if NAME_LABEL not in normalize_text(value):
            continue

        name = extract_name_from_anchor(worksheet, resolver, row, column, value)
        anchors.append(NameAnchor(row=row, column=column, name=name))

    anchors.sort(key=lambda item: (item.row, item.column))

    # A merged or duplicated label must still represent only one student.
    unique_anchors: list[NameAnchor] = []
    seen_rows: set[int] = set()
    for anchor in anchors:
        if anchor.row in seen_rows:
            continue
        seen_rows.add(anchor.row)
        unique_anchors.append(anchor)
    return unique_anchors


def find_label_rows(
    worksheet: Worksheet,
    label: str,
    start_row: int,
    end_row: int,
) -> list[int]:
    rows: set[int] = set()
    for row, _column, value in iter_relevant_cells(worksheet, start_row, end_row):
        if label in normalize_text(value):
            rows.add(row)
    return sorted(rows)


def closest_preceding_weekly_row(weekly_rows: list[int], name_row: int) -> int | None:
    preceding_rows = [row for row in weekly_rows if row <= name_row]
    return preceding_rows[-1] if preceding_rows else None


def next_weekly_row(weekly_rows: list[int], current_weekly_row: int) -> int | None:
    return next((row for row in weekly_rows if row > current_weekly_row), None)


def find_label_anchors(
    worksheet: Worksheet,
    resolver: MergedCellResolver,
    label: str,
    start_row: int,
    end_row: int,
) -> list[LabelAnchor]:
    anchors: list[LabelAnchor] = []
    seen_sources: set[tuple[int, int]] = set()

    for row, column, value in iter_relevant_cells(worksheet, start_row, end_row):
        if label not in normalize_text(value):
            continue

        _raw_value, source, merged_range = resolver.get(row, column)
        if source in seen_sources:
            continue
        seen_sources.add(source)
        anchors.append(
            LabelAnchor(
                row=source[0],
                column=source[1],
                merged_range=merged_range,
            )
        )

    return sorted(anchors, key=lambda item: (item.row, item.column))


def first_anchor_at_or_after(
    anchors: list[LabelAnchor],
    minimum_row: int,
) -> LabelAnchor | None:
    return next((anchor for anchor in anchors if anchor.row >= minimum_row), None)


def row_structure_score(
    worksheet: Worksheet,
    resolver: MergedCellResolver,
    row: int,
) -> tuple[int, int, int]:
    vertical_merges = 0
    styled_positions = 0
    populated_positions = 0

    for column in range(COURSE_FIRST_COLUMN, COURSE_LAST_COLUMN + 1):
        raw_value, source, merged_range = resolver.get(row, column)
        if merged_range is not None and merged_range.max_row > merged_range.min_row:
            vertical_merges += 1

        source_cell = worksheet.cell(source[0], source[1])
        current_cell = worksheet.cell(row, column)
        if source_cell.has_style or current_cell.has_style:
            styled_positions += 1
        if clean_grade(raw_value) is not None:
            populated_positions += 1

    return vertical_merges, styled_positions, populated_positions


def row_directly_right_of_label(
    worksheet: Worksheet,
    resolver: MergedCellResolver,
    anchor: LabelAnchor,
) -> int:
    candidate_rows = range(anchor.min_row, anchor.max_row + 1)
    return max(
        candidate_rows,
        key=lambda row: (row_structure_score(worksheet, resolver, row), -row),
    )


def course_band_bottom_row(
    resolver: MergedCellResolver,
    band_row: int,
) -> int:
    bottom_row_counts: dict[int, int] = {}

    for column in range(COURSE_FIRST_COLUMN, COURSE_LAST_COLUMN + 1):
        _raw_value, source, merged_range = resolver.get(band_row, column)
        if merged_range is not None and source[1] == column:
            bottom_row = merged_range.max_row
        else:
            bottom_row = band_row
        bottom_row_counts[bottom_row] = bottom_row_counts.get(bottom_row, 0) + 1

    # Course cells normally share one height. The most common bottom edge is the
    # reliable boundary when an individual cell has an unusual merge.
    return max(
        bottom_row_counts,
        key=lambda row: (bottom_row_counts[row], row),
    )


def read_course_grade_band(
    resolver: MergedCellResolver,
    row: int,
    student_name: str,
    band_name: str,
    warnings: list[str],
) -> list[Any]:
    values: list[Any] = []
    misplaced_columns: list[str] = []

    for course, column in zip(
        COURSES,
        range(COURSE_FIRST_COLUMN, COURSE_LAST_COLUMN + 1),
    ):
        raw_value, source, _merged_range = resolver.get(row, column)
        cleaned_value = clean_grade(raw_value)

        # Grade cells may be merged vertically, but never across course columns.
        if cleaned_value is not None and source[1] != column:
            values.append(None)
            misplaced_columns.append(course)
            continue

        values.append(cleaned_value)

    if misplaced_columns:
        warnings.append(
            f"{student_name}: {band_name} value(s) appeared outside their expected "
            f"course column(s): {', '.join(misplaced_columns)}."
        )

    return values


def same_course_positions(
    resolver: MergedCellResolver,
    first_row: int,
    second_row: int,
) -> bool:
    for column in range(COURSE_FIRST_COLUMN, COURSE_LAST_COLUMN + 1):
        _first_value, first_source, _first_range = resolver.get(first_row, column)
        _second_value, second_source, _second_range = resolver.get(second_row, column)
        if first_source != second_source:
            return False
    return True


def extract_student_grade_bands(
    worksheet: Worksheet,
    resolver: MergedCellResolver,
    weekly_row: int,
    area_end_row: int,
    student_name: str,
    warnings: list[str],
) -> tuple[list[Any], list[Any], list[Any]]:
    empty_band = [None] * len(COURSES)

    first_anchors = find_label_anchors(
        worksheet, resolver, FIRST_TERM_LABEL, weekly_row, area_end_row
    )
    second_anchors = find_label_anchors(
        worksheet, resolver, SECOND_TERM_LABEL, weekly_row, area_end_row
    )
    year_end_anchors = find_label_anchors(
        worksheet, resolver, COURSE_YEAR_END_LABEL, weekly_row, area_end_row
    )
    weighted_anchors = find_label_anchors(
        worksheet, resolver, WEIGHTED_LABEL, weekly_row, area_end_row
    )

    first_anchor = first_anchor_at_or_after(first_anchors, weekly_row)
    second_anchor = first_anchor_at_or_after(
        second_anchors,
        (first_anchor.row + 1) if first_anchor else weekly_row,
    )
    year_end_anchor = first_anchor_at_or_after(
        year_end_anchors,
        (second_anchor.row + 1) if second_anchor else weekly_row,
    )
    weighted_anchor = first_anchor_at_or_after(
        weighted_anchors,
        (year_end_anchor.row + 1) if year_end_anchor else weekly_row,
    )

    missing_labels = [
        label
        for label, anchor in (
            ("1. Dönem Puanı", first_anchor),
            ("2. Dönem Puanı", second_anchor),
            ("Yıl Sonu Puanı course band", year_end_anchor),
            ("Ağırlıklı Puanı", weighted_anchor),
        )
        if anchor is None
    ]
    if missing_labels:
        warnings.append(
            f"{student_name}: required grade-region label(s) could not be identified: "
            f"{', '.join(missing_labels)}."
        )

    if first_anchor is not None:
        first_row = row_directly_right_of_label(
            worksheet,
            resolver,
            first_anchor,
        )
        first_values = read_course_grade_band(
            resolver,
            first_row,
            student_name,
            "1. Dönem Puanı",
            warnings,
        )
    else:
        first_row = None
        first_values = empty_band.copy()

    if second_anchor is not None and first_row is not None:
        second_row = course_band_bottom_row(resolver, first_row) + 1
        if second_row >= second_anchor.min_row or second_row > area_end_row:
            warnings.append(
                f"{student_name}: 2. Dönem Puanı grade region could not be "
                "identified directly underneath the 1. Dönem region."
            )
            second_values = empty_band.copy()
        elif same_course_positions(resolver, first_row, second_row):
            warnings.append(
                f"{student_name}: 2. Dönem Puanı resolved to the same merged "
                "course positions as 1. Dönem Puanı."
            )
            second_values = empty_band.copy()
        else:
            second_values = read_course_grade_band(
                resolver,
                second_row,
                student_name,
                "2. Dönem Puanı",
                warnings,
            )
    else:
        second_values = empty_band.copy()

    if year_end_anchor is not None:
        year_end_row = row_directly_right_of_label(
            worksheet,
            resolver,
            year_end_anchor,
        )
        year_end_values = read_course_grade_band(
            resolver,
            year_end_row,
            student_name,
            "Yıl Sonu Puanı",
            warnings,
        )
    else:
        year_end_values = empty_band.copy()

    return first_values, second_values, year_end_values


def extract_toplam_not(
    resolver: MergedCellResolver,
    weekly_row: int,
    area_end_row: int,
    student_name: str,
    warnings: list[str],
) -> Any:
    candidates: list[tuple[tuple[int, int, int], Any]] = []

    for merged_range in resolver.ranges:
        if not (merged_range.min_col <= OVERALL_GRADE_COLUMN <= merged_range.max_col):
            continue
        if merged_range.max_row <= merged_range.min_row:
            continue
        if merged_range.min_row > area_end_row:
            continue

        raw_value = resolver.worksheet.cell(
            merged_range.min_row,
            merged_range.min_col,
        ).value
        cleaned_value = clean_grade(raw_value)
        if not isinstance(cleaned_value, Number) or isinstance(cleaned_value, bool):
            continue

        if merged_range.min_row == weekly_row - 1:
            priority = (0, 0, merged_range.min_row)
        elif merged_range.min_row < weekly_row <= merged_range.max_row:
            priority = (1, weekly_row - merged_range.min_row, merged_range.min_row)
        elif merged_range.max_row == weekly_row - 1:
            priority = (2, 0, merged_range.min_row)
        else:
            continue

        candidates.append((priority, cleaned_value))

    if not candidates:
        warnings.append(f"{student_name}: Toplam Not was not found in column O.")
        return None

    return min(candidates, key=lambda item: item[0])[1]


def build_headers() -> list[str]:
    return (
        ["Ad Soyad"]
        + [f"{course} 1. Dönem Puanı" for course in COURSES]
        + [f"{course} 2. Dönem Puanı" for course in COURSES]
        + [f"{course} Yıl Sonu Puanı" for course in COURSES]
        + ["Toplam Not"]
    )


def create_output_workbook(records: list[list[Any]], output_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Notlar"
    worksheet.sheet_view.showGridLines = False

    headers = build_headers()
    worksheet.append(headers)
    for record in records:
        worksheet.append(record)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_border = Border(bottom=Side(style="medium", color="163A5C"))
    body_border = Border(bottom=Side(style="hair", color="D9E2F3"))

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for index, cell in enumerate(row, start=1):
            cell.border = body_border
            if index == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, int) and not isinstance(cell.value, bool):
                    cell.number_format = "0"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.####"

    worksheet.row_dimensions[1].height = 48
    worksheet.column_dimensions["A"].width = 30
    for column in range(2, len(headers) + 1):
        course_index = (column - 2) % len(COURSES)
        course_name = COURSES[course_index] if column < len(headers) else ""
        width = 24
        if len(course_name) > 24:
            width = 31
        worksheet.column_dimensions[get_column_letter(column)].width = width

    worksheet.freeze_panes = "B2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, worksheet.max_row)}"

    workbook.save(output_path)
    workbook.close()


def extract_records(worksheet: Worksheet) -> tuple[list[list[Any]], list[str]]:
    warnings: list[str] = []
    records: list[list[Any]] = []
    resolver = MergedCellResolver(worksheet)
    anchors = find_name_anchors(worksheet, resolver)
    weekly_rows = find_label_rows(worksheet, WEEKLY_LABEL, 1, worksheet.max_row)

    if not anchors:
        warnings.append('No cells containing "Adı ve Soyadı" were found.')
        return records, warnings

    used_weekly_rows: set[int] = set()
    for index, anchor in enumerate(anchors, start=1):
        student_name = anchor.name or f"[Name unreadable at row {anchor.row}]"
        if anchor.name is None:
            warnings.append(f"Student {index}: name could not be read from row {anchor.row}.")

        weekly_row = closest_preceding_weekly_row(weekly_rows, anchor.row)

        if weekly_row is None:
            warnings.append(
                f"{student_name}: closest preceding Haftalık Ders Saati was not found "
                "inside this student's area."
            )
            first_values = [None] * len(COURSES)
            second_values = [None] * len(COURSES)
            year_end_values = [None] * len(COURSES)
            overall_value = None
        else:
            following_weekly_row = next_weekly_row(weekly_rows, weekly_row)
            area_end_row = (
                following_weekly_row - 1
                if following_weekly_row is not None
                else worksheet.max_row
            )

            if weekly_row in used_weekly_rows:
                warnings.append(
                    f"{student_name}: more than one Adı ve Soyadı anchor was associated "
                    f"with the Haftalık Ders Saati row {weekly_row}."
                )
            used_weekly_rows.add(weekly_row)

            first_values, second_values, year_end_values = extract_student_grade_bands(
                worksheet,
                resolver,
                weekly_row,
                area_end_row,
                student_name,
                warnings,
            )
            overall_value = extract_toplam_not(
                resolver,
                weekly_row,
                area_end_row,
                student_name,
                warnings,
            )

        records.append(
            [student_name]
            + first_values
            + second_values
            + year_end_values
            + [overall_value]
        )

    if len(records) != EXPECTED_STUDENT_COUNT:
        warnings.append(
            f"Expected {EXPECTED_STUDENT_COUNT} students, but extracted {len(records)}."
        )

    return records, warnings


def get_raw_input_path() -> str:
    if len(sys.argv) > 2:
        raise ValueError("Use either no argument or one .xlsx input path.")
    if len(sys.argv) == 2:
        return sys.argv[1]
    return input("Enter the path of the input .xlsx file: ")


def main() -> int:
    try:
        input_path = validate_input_path(get_raw_input_path())
    except (EOFError, KeyboardInterrupt):
        print("\nOperation cancelled.")
        return 1
    except ValueError as exc:
        print(f"ERROR: {exc}")
        return 1

    output_path = find_unused_output_path(input_path)
    input_workbook = None

    try:
        # data_only=True reads cached formula results and never changes the source workbook.
        input_workbook = load_workbook(input_path, read_only=False, data_only=True)
        input_worksheet = input_workbook.worksheets[0]
        records, warnings = extract_records(input_worksheet)
        create_output_workbook(records, output_path)
    except Exception as exc:
        print(f"ERROR: Extraction failed: {exc}")
        return 1
    finally:
        if input_workbook is not None:
            input_workbook.close()

    print(f"Students extracted: {len(records)}")
    print(f"Warnings: {len(warnings)}")
    for warning in warnings:
        print(f"WARNING: {warning}")
    print(f"Output workbook: {output_path}")
    print("The original workbook was opened for reading and was not modified.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
