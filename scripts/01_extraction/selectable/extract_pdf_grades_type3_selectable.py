# -*- coding: utf-8 -*-
r"""
Convert the third selectable-text PDF table type into Excel.

Use this script for the Basaksehir-style table where each PDF contains one
class section and multiple course tables/sections. For each student row, the
grade is read as the single non-empty numeric value after the student-name
column. Empty grade rows remain empty in Excel.

Install requirements:

    pip install pdfplumber openpyxl

Run with a path:

    python extract_pdf_grades_type3_selectable.py "C:\path\to\pdf_or_folder"

Or run without a path and the script will ask you to enter one:

    python extract_pdf_grades_type3_selectable.py
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    missing_name = getattr(exc, "name", "required package")
    raise SystemExit(
        f"Missing Python package: {missing_name}\n"
        "Install the requirements with:\n\n"
        "    pip install pdfplumber openpyxl\n"
    ) from exc


METADATA_COLUMNS = ["İlçe", "Okul", "Sınıf", "Şube", "Okul No", "Ad Soyad"]

KNOWN_COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "HAYAT BİLGİSİ",
    "YABANCI DİL",
    "GÖRSEL SANATLAR",
    "MÜZİK",
    "FEN BİLİMLERİ",
    "BEDEN EĞİTİMİ VE OYUN",
    "SERBEST ETKİNLİKLER",
]

SUMMARY_KEYWORDS = [
    "TOPLAM ÖĞRENCİ SAYISI",
    "BAŞARILI ÖĞRENCİ SAYISI",
    "BAŞARISIZ ÖĞRENCİ SAYISI",
    "BAŞARI YÜZDESİ",
    "NOT ORTALAMASI",
    "ZAYIFI OLAN",
]

SUMMARY_TOKEN_GROUPS = [
    ("TOPLAM", "OGRENCI"),
    ("BASARILI", "OGRENCI"),
    ("BASARISIZ", "OGRENCI"),
    ("BASARI", "YUZDE"),
    ("NOT", "ORTALAMA"),
    ("ZAYIFI", "OLAN"),
]

TABLE_SETTINGS_VARIANTS = [
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "lines",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 3,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "text",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
]

TURKISH_TO_ASCII = str.maketrans(
    {
        "ç": "c",
        "Ç": "C",
        "ğ": "g",
        "Ğ": "G",
        "ı": "i",
        "I": "I",
        "İ": "I",
        "ö": "o",
        "Ö": "O",
        "ş": "s",
        "Ş": "S",
        "ü": "u",
        "Ü": "U",
    }
)


@dataclass
class ParsedPdf:
    pdf_path: Path
    courses: list[str]
    header: dict[str, str]
    rows: list[dict[str, str]]


@dataclass
class ExtractedTable:
    rows: list[list[object]]
    bbox: tuple[float, float, float, float] | None = None


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def fold_for_match(value: object) -> str:
    text = clean_text(value).translate(TURKISH_TO_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def has_letter(value: str) -> bool:
    return any(char.isalpha() for char in value)


def value_from_row(row: list[object], index: int | None) -> str:
    if index is None or index < 0 or index >= len(row):
        return ""
    return clean_text(row[index])


def row_is_empty(row: Iterable[object]) -> bool:
    return not any(clean_text(cell) for cell in row)


def row_text(row: Iterable[object]) -> str:
    return " ".join(clean_text(cell) for cell in row if clean_text(cell))


def line_has_summary_text(value: str) -> bool:
    folded = fold_for_match(value)
    if any(fold_for_match(keyword) in folded for keyword in SUMMARY_KEYWORDS):
        return True
    return any(all(token in folded for token in tokens) for tokens in SUMMARY_TOKEN_GROUPS)


def row_has_summary_text(row: Iterable[object]) -> bool:
    return line_has_summary_text(row_text(row))


def normalize_school_no(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"-?\s*\d+", text)
    if not match:
        return ""
    return re.sub(r"\s+", "", match.group(0))


def normalize_name(value: object) -> str:
    text = clean_text(value)
    text = text.replace("|", " ").replace("_", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text.upper()


def normalize_grade(value: object) -> str:
    text = clean_text(value)
    if not text:
        return ""

    match = re.search(r"\d+(?:[,.]\d+)?", text)
    if not match:
        return ""

    return match.group(0).replace(",", ".")


def is_excluded_course(course: str) -> bool:
    folded = fold_for_match(course)
    return "SERBEST" in folded and "ETKINLIK" in folded


def detect_course_name_from_text(text: str) -> str:
    folded = fold_for_match(text)
    if not folded:
        return ""

    for course in KNOWN_COURSES:
        if fold_for_match(course) in folded:
            return course

    # Fallback for unexpected course names written as "... DERSI".
    for raw_line in text.splitlines():
        line = clean_text(raw_line)
        folded_line = fold_for_match(line)
        if "DERSI" not in folded_line:
            continue

        before_dersi = re.split(r"\s+DERS[İI]\b", line, maxsplit=1, flags=re.IGNORECASE)[0]
        before_dersi = re.sub(
            r"^.*?\bDERS\s+YILI\b",
            "",
            before_dersi,
            flags=re.IGNORECASE,
        )
        before_dersi = re.sub(
            r"^\s*\d{4}\s*[-/]\s*\d{4}\s*",
            "",
            before_dersi,
        )
        course = normalize_name(before_dersi)
        if course:
            return course

    return ""


def resolve_input_pdfs(raw_path: str | Path) -> tuple[list[Path], bool, Path]:
    path = Path(raw_path).expanduser()

    if path.is_dir():
        pdf_files = sorted(item for item in path.glob("*.pdf") if item.is_file())
        if not pdf_files:
            raise FileNotFoundError(f"No PDF files were found in folder: {path}")
        return pdf_files, True, path

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(f"The input file is not a PDF: {path}")
        return [path], False, path.parent

    if path.suffix.lower() != ".pdf":
        pdf_version = path.with_suffix(".pdf")
        if pdf_version.is_file():
            return [pdf_version], False, pdf_version.parent

    raise FileNotFoundError(
        f"PDF file or folder was not found: {path}\n"
        "Tip: pass one PDF file path or a folder containing PDF files."
    )


def extract_district_and_school(lines: list[str]) -> tuple[str, str]:
    for line in lines[:30]:
        folded = fold_for_match(line)
        if "/" not in line:
            continue
        if "MUDURLUGU" not in folded and "MUDURLUG" not in folded:
            continue

        district_part, school_part = line.split("/", 1)
        district = clean_text(district_part)
        school_part = clean_text(school_part)
        school = re.split(
            r"\s+M[üu]d[üu]rl[üu][ğg][üu]\b",
            school_part,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0].strip()
        if school == school_part:
            school = re.sub(
                r"\s*MUDURLUGU?.*$",
                "",
                fold_for_match(school_part),
                flags=re.IGNORECASE,
            ).title()
        return district, school

    return "", ""


def extract_course_name(lines: list[str], pdf_path: Path) -> str:
    for line in lines[:40]:
        folded = fold_for_match(line)
        if "DERSI" not in folded:
            continue

        for course in KNOWN_COURSES:
            if fold_for_match(course) in folded:
                return course

        raw_before_dersi = re.split(r"\s+DERS[İI]\b", line, maxsplit=1, flags=re.IGNORECASE)[0]
        raw_before_dersi = re.sub(
            r"^.*?\bDERS\s+YILI\b",
            "",
            raw_before_dersi,
            flags=re.IGNORECASE,
        )
        course = normalize_name(raw_before_dersi)
        if course:
            return course

    stem = clean_text(pdf_path.stem).upper()
    return stem or "BILINMEYEN DERS"


def extract_class_and_section(text: str) -> tuple[str, str]:
    raw_match = re.search(
        r"(?P<class>\d+)\.\s*S[ıiİI]n[ıiİI]f\s*/\s*(?P<section>.+?)\s*Şubesi",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if raw_match:
        return clean_text(raw_match.group("class")), clean_text(raw_match.group("section"))

    folded = fold_for_match(text)
    folded_match = re.search(
        r"\b(?P<class>\d+)\s*SINIF\s*/?\s*(?P<section>[A-Z0-9]+)\s*SUBESI\b",
        folded,
    )
    if folded_match:
        return folded_match.group("class"), folded_match.group("section")

    return "", ""


def extract_header_info(page_text: str, pdf_path: Path) -> dict[str, str]:
    lines = [clean_text(line) for line in page_text.splitlines() if clean_text(line)]
    district, school = extract_district_and_school(lines)
    class_no, section = extract_class_and_section(" ".join(lines[:50]))

    return {
        "district": district,
        "school": school,
        "class_no": class_no,
        "section": section,
        "course": "",
    }


def nonempty_cells(row: list[object]) -> list[tuple[int, str]]:
    return [(index, clean_text(cell)) for index, cell in enumerate(row) if clean_text(cell)]


def looks_like_student_row(row: list[object]) -> bool:
    if row_is_empty(row) or row_has_summary_text(row):
        return False

    cells = nonempty_cells(row)
    if len(cells) < 3:
        return False

    for position in range(0, len(cells) - 1):
        school_index, school_value = cells[position]
        name_index, name_value = cells[position + 1]

        if school_index > 3:
            break
        if not normalize_school_no(school_value):
            continue
        if not has_letter(name_value):
            continue
        if "OGRENCI" in fold_for_match(name_value) or "SOYAD" in fold_for_match(name_value):
            continue
        if name_index - school_index > 2:
            continue

        return True

    return False


def detect_student_start_row(table: list[list[object]]) -> int:
    for row_index, row in enumerate(table):
        if looks_like_student_row(row):
            return row_index
    return 0


def joined_header_by_column(table: list[list[object]], student_start_row: int) -> list[str]:
    max_columns = max((len(row) for row in table), default=0)
    header_rows = table[:student_start_row]
    column_headers: list[str] = []

    for column_index in range(max_columns):
        parts = [
            clean_text(row[column_index])
            for row in header_rows
            if column_index < len(row) and clean_text(row[column_index])
        ]
        column_headers.append(" ".join(parts))

    return column_headers


def detect_columns_from_header(
    table: list[list[object]],
    student_start_row: int,
) -> dict[str, int]:
    column_headers = joined_header_by_column(table, student_start_row)
    columns: dict[str, int] = {}

    for index, header in enumerate(column_headers):
        folded = fold_for_match(header)
        tokens = set(folded.split())

        if "OKUL" in tokens and "NO" in tokens:
            columns.setdefault("school_no", index)

        if "OGRENCI" in folded or ("AD" in tokens and ("SOYAD" in tokens or "SOYADI" in tokens)):
            columns.setdefault("name", index)

        if {"1", "DONEM", "PUANI"}.issubset(tokens):
            columns.setdefault("grade", index)

    return columns


def detect_identity_columns_from_rows(table: list[list[object]]) -> tuple[int, int] | None:
    pair_counts: Counter[tuple[int, int]] = Counter()

    for row in table:
        if row_is_empty(row) or row_has_summary_text(row):
            continue

        cells = nonempty_cells(row)
        if len(cells) < 2:
            continue

        for position in range(0, len(cells) - 1):
            school_index, school_value = cells[position]
            name_index, name_value = cells[position + 1]

            if school_index > 3:
                break
            if not normalize_school_no(school_value):
                continue
            if not has_letter(name_value):
                continue
            if "OGRENCI" in fold_for_match(name_value) or "SOYAD" in fold_for_match(name_value):
                continue

            pair_counts[(school_index, name_index)] += 1
            break

    if not pair_counts:
        return None

    return pair_counts.most_common(1)[0][0]


def choose_columns(table: list[list[object]], student_start_row: int) -> dict[str, int] | None:
    columns = detect_columns_from_header(table, student_start_row)

    identity_columns = detect_identity_columns_from_rows(table[student_start_row:])
    if identity_columns is not None:
        columns.setdefault("school_no", identity_columns[0])
        columns.setdefault("name", identity_columns[1])

    if {"school_no", "name"}.issubset(columns):
        return columns

    return None


def grade_from_student_row(row: list[object], name_column: int) -> str:
    numeric_values: list[str] = []

    for cell in row[name_column + 1 :]:
        grade = normalize_grade(cell)
        if grade:
            numeric_values.append(grade)

    if not numeric_values:
        return ""

    return numeric_values[0]


def parse_rows_from_table(
    table: list[list[object]],
    header: dict[str, str],
    course: str,
) -> list[dict[str, str]]:
    if not table:
        return []
    if not course or is_excluded_course(course):
        return []

    student_start_row = detect_student_start_row(table)
    columns = choose_columns(table, student_start_row)
    if columns is None:
        return []

    rows: list[dict[str, str]] = []
    for row in table[student_start_row:]:
        if row_is_empty(row) or row_has_summary_text(row):
            continue

        school_no = normalize_school_no(value_from_row(row, columns["school_no"]))
        name = normalize_name(value_from_row(row, columns["name"]))

        if not school_no or not name or not has_letter(name):
            continue
        folded_name = fold_for_match(name)
        if "OGRENCI" in folded_name or "SOYAD" in folded_name:
            continue

        rows.append(
            {
                "İlçe": header["district"],
                "Okul": header["school"],
                "Sınıf": header["class_no"],
                "Şube": header["section"],
                "Okul No": school_no,
                "Ad Soyad": name,
                course: grade_from_student_row(row, columns["name"]),
            }
        )

    return rows


def parse_rows_from_text_fallback(
    page_text: str,
    header: dict[str, str],
    course: str,
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    if not course or is_excluded_course(course):
        return rows

    for raw_line in page_text.splitlines():
        line = clean_text(raw_line)
        if not line or line_has_summary_text(line):
            continue

        match = re.match(r"^(?P<serial>\d+)\s+(?P<school_no>-?\d+)\s+(?P<rest>.+)$", line)
        if not match:
            continue

        school_no = normalize_school_no(match.group("school_no"))
        rest_tokens = clean_text(match.group("rest")).split()
        if not school_no or not rest_tokens:
            continue

        numeric_positions = [
            index for index, token in enumerate(rest_tokens) if normalize_grade(token)
        ]
        if not numeric_positions:
            name = " ".join(rest_tokens)
            grade = ""
        else:
            first_numeric = numeric_positions[0]
            name = " ".join(rest_tokens[:first_numeric])
            grade = normalize_grade(rest_tokens[first_numeric])

        name = normalize_name(name)
        if not name or not has_letter(name):
            continue

        rows.append(
            {
                "İlçe": header["district"],
                "Okul": header["school"],
                "Sınıf": header["class_no"],
                "Şube": header["section"],
                "Okul No": school_no,
                "Ad Soyad": name,
                course: grade,
            }
        )

    return rows


def parse_rows_from_text_multi_course(
    page_text: str,
    header: dict[str, str],
    previous_course: str = "",
) -> tuple[list[dict[str, str]], str]:
    rows: list[dict[str, str]] = []
    current_course = previous_course

    for raw_line in page_text.splitlines():
        line = clean_text(raw_line)
        if not line:
            continue

        detected_course = detect_course_name_from_text(line)
        if detected_course:
            current_course = detected_course
            continue

        if not current_course or is_excluded_course(current_course):
            continue

        rows.extend(parse_rows_from_text_fallback(line, header, current_course))

    return rows, current_course


def extract_tables_from_page(page) -> list[ExtractedTable]:
    for settings in TABLE_SETTINGS_VARIANTS:
        try:
            found_tables = page.find_tables(table_settings=settings)
        except Exception:
            found_tables = []

        extracted_tables: list[ExtractedTable] = []
        for table in found_tables:
            try:
                rows = table.extract()
            except Exception:
                rows = []
            if rows:
                extracted_tables.append(ExtractedTable(rows=rows, bbox=table.bbox))

        if extracted_tables:
            return extracted_tables

    try:
        raw_tables = page.extract_tables()
    except Exception:
        raw_tables = []

    return [ExtractedTable(rows=table, bbox=None) for table in raw_tables if table]


def text_near_table(page, extracted_table: ExtractedTable) -> str:
    if extracted_table.bbox is None:
        return page.extract_text() or ""

    x0, y0, x1, _y1 = extracted_table.bbox
    top = max(0, y0 - 150)
    bottom = min(page.height, y0 + 5)

    try:
        nearby_text = page.crop((0, top, page.width, bottom)).extract_text() or ""
    except Exception:
        nearby_text = ""

    if nearby_text:
        return nearby_text

    return page.extract_text() or ""


def course_for_table(page, extracted_table: ExtractedTable, previous_course: str = "") -> str:
    student_start_row = detect_student_start_row(extracted_table.rows)
    table_header_text = "\n".join(row_text(row) for row in extracted_table.rows[:student_start_row])

    course = detect_course_name_from_text(table_header_text)
    if course:
        return course

    nearby_text = text_near_table(page, extracted_table)
    nearby_lines = [clean_text(line) for line in nearby_text.splitlines() if clean_text(line)]
    for line in reversed(nearby_lines):
        course = detect_course_name_from_text(line)
        if course:
            return course

    return previous_course


def extract_pdf(pdf_path: Path) -> ParsedPdf:
    with pdfplumber.open(str(pdf_path)) as pdf:
        page_texts = [page.extract_text() or "" for page in pdf.pages]
        full_text = "\n".join(page_texts)
        header = extract_header_info(full_text, pdf_path)

        rows: list[dict[str, str]] = []
        course_order: list[str] = []
        previous_course = ""

        for page in pdf.pages:
            page_text = page.extract_text() or ""
            extracted_tables = extract_tables_from_page(page)
            page_row_count_before = len(rows)

            for extracted_table in extracted_tables:
                course = course_for_table(page, extracted_table, previous_course)
                if course:
                    previous_course = course

                if not course or is_excluded_course(course):
                    continue

                parsed_rows = parse_rows_from_table(extracted_table.rows, header, course)
                if parsed_rows:
                    rows.extend(parsed_rows)
                    if course not in course_order:
                        course_order.append(course)

            if not extracted_tables or len(rows) == page_row_count_before:
                parsed_rows, previous_course = parse_rows_from_text_multi_course(
                    page_text,
                    header,
                    previous_course=previous_course,
                )
                rows.extend(parsed_rows)
                for row in parsed_rows:
                    for column in row:
                        if column not in METADATA_COLUMNS and column not in course_order:
                            course_order.append(column)

        if not rows:
            fallback_rows: list[dict[str, str]] = []
            previous_course = ""
            for page_text in page_texts:
                parsed_rows, previous_course = parse_rows_from_text_multi_course(
                    page_text,
                    header,
                    previous_course=previous_course,
                )
                fallback_rows.extend(parsed_rows)
                for row in parsed_rows:
                    for column in row:
                        if column not in METADATA_COLUMNS and column not in course_order:
                            course_order.append(column)
            rows = fallback_rows

        if not rows:
            raise ValueError(f"No student rows could be extracted from {pdf_path.name}.")

        return ParsedPdf(pdf_path=pdf_path, courses=course_order, header=header, rows=rows)


def combine_parsed_pdfs(parsed_pdfs: list[ParsedPdf]) -> tuple[list[dict[str, str]], list[str]]:
    students: dict[str, dict[str, str]] = {}
    student_order: list[str] = []
    course_order: list[str] = []

    for parsed in parsed_pdfs:
        for course in parsed.courses:
            if course and not is_excluded_course(course) and course not in course_order:
                course_order.append(course)

        for row in parsed.rows:
            school_no = row.get("Okul No", "")
            if not school_no:
                continue

            student_key = "\x1f".join(
                [
                    row.get("İlçe", ""),
                    row.get("Okul", ""),
                    row.get("Sınıf", ""),
                    row.get("Şube", ""),
                    school_no,
                ]
            )

            if student_key not in students:
                students[student_key] = {column: row.get(column, "") for column in METADATA_COLUMNS}
                student_order.append(student_key)
            else:
                for column in METADATA_COLUMNS:
                    if not students[student_key].get(column) and row.get(column):
                        students[student_key][column] = row[column]

            for column, value in row.items():
                if column in METADATA_COLUMNS or is_excluded_course(column):
                    continue
                if column not in course_order:
                    course_order.append(column)
                students[student_key][column] = value

    final_rows: list[dict[str, str]] = []
    for student_key in student_order:
        record = students[student_key]
        for course in course_order:
            record.setdefault(course, "")
        final_rows.append(record)

    return final_rows, course_order


def save_excel(rows: list[dict[str, str]], course_columns: list[str], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Notlar"

    output_columns = [*METADATA_COLUMNS, *course_columns]
    worksheet.append(output_columns)
    for row in rows:
        worksheet.append([row.get(column, "") for column in output_columns])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    school_no_column = output_columns.index("Okul No") + 1
    school_no_letter = get_column_letter(school_no_column)
    for row_number in range(2, worksheet.max_row + 1):
        worksheet[f"{school_no_letter}{row_number}"].number_format = "@"

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header_value = column_cells[0].value
        max_length = max(len(str(cell.value or "")) for cell in column_cells)

        if header_value == "Ad Soyad":
            width = min(max(max_length + 2, 20), 42)
        elif header_value in {"İlçe", "Okul"}:
            width = min(max(max_length + 2, 14), 36)
        elif header_value in course_columns:
            width = min(max(max_length + 2, 12), 28)
        else:
            width = min(max(max_length + 2, 10), 18)

        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 42
    workbook.save(output_path)


def safe_file_part(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" ._")
    return value or "grades"


def build_output_path(
    pdf_paths: list[Path],
    input_is_folder: bool,
    input_base_dir: Path,
    output_dir: Path | None,
) -> Path:
    if output_dir is None:
        output_dir = input_base_dir / "type3_selectable_output"

    if input_is_folder:
        return output_dir / "combined_grades.xlsx"

    return output_dir / f"{safe_file_part(pdf_paths[0].stem)}.xlsx"


def write_error_log(errors: list[dict[str, str]], output_dir: Path) -> Path | None:
    if not errors:
        return None

    output_dir.mkdir(parents=True, exist_ok=True)
    log_path = output_dir / "type3_selectable_extraction_errors.csv"
    fieldnames = ["pdf", "reason"]

    with log_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for error in errors:
            writer.writerow({field: error.get(field, "") for field in fieldnames})

    return log_path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert type-3 selectable-text grade PDFs into Excel."
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=None,
        help="PDF file path or folder path. If omitted, the script asks for it.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Optional folder where the Excel file and error log will be saved.",
    )
    return parser.parse_args(argv)


def get_input_path(args: argparse.Namespace) -> str:
    if args.input_path:
        return args.input_path

    user_input = input("Enter a PDF file path or a folder path containing PDF files: ").strip()
    if not user_input:
        raise ValueError("No input path was entered.")
    return user_input.strip('"')


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    input_path = get_input_path(args)
    pdf_paths, input_is_folder, input_base_dir = resolve_input_pdfs(input_path)
    output_dir = Path(args.output_dir).expanduser() if args.output_dir else None
    output_path = build_output_path(pdf_paths, input_is_folder, input_base_dir, output_dir)

    parsed_pdfs: list[ParsedPdf] = []
    errors: list[dict[str, str]] = []

    print(f"PDF files found: {len(pdf_paths)}")
    for index, pdf_path in enumerate(pdf_paths, start=1):
        try:
            parsed = extract_pdf(pdf_path)
            parsed_pdfs.append(parsed)
            course_text = ", ".join(parsed.courses) if parsed.courses else "no included course detected"
            print(
                f"[{index}/{len(pdf_paths)}] OK: {pdf_path.name} "
                f"-> courses={course_text}, rows={len(parsed.rows)}"
            )
        except Exception as error:
            errors.append({"pdf": str(pdf_path), "reason": str(error)})
            print(f"[{index}/{len(pdf_paths)}] ERROR: {pdf_path.name}: {error}", file=sys.stderr)

    if not parsed_pdfs:
        log_path = write_error_log(errors, output_path.parent)
        if log_path:
            print(f"Error log saved: {log_path}")
        raise RuntimeError("No PDFs could be converted.")

    rows, course_columns = combine_parsed_pdfs(parsed_pdfs)
    if not rows:
        raise RuntimeError("No student rows were extracted from the processed PDFs.")

    save_excel(rows, course_columns, output_path)
    log_path = write_error_log(errors, output_path.parent)

    print(f"Converted PDFs: {len(parsed_pdfs)}")
    print(f"Created: {output_path}")
    if log_path:
        print(f"Error log saved: {log_path}")

    return 1 if errors else 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
