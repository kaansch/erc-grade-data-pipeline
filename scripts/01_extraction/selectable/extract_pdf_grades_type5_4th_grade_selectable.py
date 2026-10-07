# -*- coding: utf-8 -*-
r"""
Convert the fourth-grade selectable-text report-card PDF type into Excel files.

Use this script when each PDF has many pages and each page belongs to one
student. The script reads the student's identity information and the course
table, then creates one Excel file for each detected class/section:

    4_A_Şubesi_1_Dönem.xlsx

Only these course-grade columns are exported:

    TÜRKÇE
    MATEMATİK
    FEN BİLİMLERİ
    SOSYAL BİLGİLER
    YABANCI DİL
    DİN KÜLTÜRÜ VE AHLAK BİLGİSİ
    GÖRSEL SANATLAR
    MÜZİK
    BEDEN EĞİTİMİ VE OYUN
    TRAFİK GÜVENLİĞİ
    İNSAN HAKLARI, YURTTAŞLIK VE DEMOKRASİ

Only "1. DÖNEM NOTU" is exported. "2. DÖNEM NOTU" is ignored if it exists.
"SERBEST ETKİNLİKLER" is ignored if it appears.

Install requirements:

    pip install pdfplumber openpyxl

Run with a PDF path:

    python extract_pdf_grades_type5_4th_grade_selectable.py "C:\path\to\file.pdf"

Or run without a path and the script will ask you to enter one:

    python extract_pdf_grades_type5_4th_grade_selectable.py
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from pathlib import Path
from typing import Iterable

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    missing_name = getattr(exc, "name", "required package")
    raise SystemExit(
        f"Missing Python package: {missing_name}\n"
        "Install the requirements with:\n\n"
        "    pip install pdfplumber openpyxl\n"
    ) from exc


OUTPUT_COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "FEN BİLİMLERİ",
    "SOSYAL BİLGİLER",
    "YABANCI DİL",
    "DİN KÜLTÜRÜ VE AHLAK BİLGİSİ",
    "GÖRSEL SANATLAR",
    "MÜZİK",
    "BEDEN EĞİTİMİ VE OYUN",
    "TRAFİK GÜVENLİĞİ",
    "İNSAN HAKLARI, YURTTAŞLIK VE DEMOKRASİ",
]

EXCLUDED_COURSES = ["SERBEST ETKİNLİKLER"]
ALL_COURSES = [*OUTPUT_COURSES, *EXCLUDED_COURSES]

OUTPUT_COLUMNS = [
    "Okul",
    "Ders Yılı",
    "Sınıf",
    "Şube",
    "Okul No",
    "Ad Soyad",
    *OUTPUT_COURSES,
]

METADATA_KEYS = ["Ad Soyad", "Okul", "Sınıf", "Şube", "Ders Yılı", "Okul No"]
LEFT_PAGE_RATIO = 0.52

TABLE_SETTINGS_VARIANTS = [
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "lines",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 3,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "text",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
]

TURKISH_TO_ASCII = str.maketrans(
    {
        "ç": "c",
        "Ç": "C",
        "ğ": "g",
        "Ğ": "G",
        "ı": "i",
        "İ": "I",
        "ö": "o",
        "Ö": "O",
        "ş": "s",
        "Ş": "S",
        "ü": "u",
        "Ü": "U",
    }
)

LABEL_PATTERNS = {
    "Ad Soyad": [("ADI", "SOYADI"), ("AD", "SOYAD")],
    "Okul": [("OKULU",)],
    "Sınıf": [("SINIFI",)],
    "Şube": [("SUBESI",)],
    "Ders Yılı": [("DERS", "YILI")],
    "Okul No": [("OKUL", "NUMARASI"), ("OKUL", "NO")],
}

ALL_LABEL_PATTERNS = [
    pattern
    for key in METADATA_KEYS
    for pattern in LABEL_PATTERNS.get(key, [])
]


def repair_mojibake(value: str) -> str:
    """Repair common UTF-8 text shown as Latin-1, if the PDF exposes it that way."""
    if not isinstance(value, str):
        return value

    suspicious = any(char in value for char in ("Ã", "Ä", "Å", "Â"))
    suspicious = suspicious or any("\x80" <= char <= "\x9f" for char in value)
    if not suspicious:
        return value

    try:
        return value.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return value


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", "\n").replace("\xa0", " ")
    text = repair_mojibake(text)
    return re.sub(r"\s+", " ", text).strip()


def fold_for_match(value: object) -> str:
    text = clean_text(value).translate(TURKISH_TO_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


COURSE_FOLDS = {
    fold_for_match(course): course
    for course in sorted(ALL_COURSES, key=lambda item: len(fold_for_match(item)), reverse=True)
}


def row_text(row: Iterable[object]) -> str:
    return " ".join(clean_text(cell) for cell in row if clean_text(cell))


def row_is_empty(row: Iterable[object]) -> bool:
    return not any(clean_text(cell) for cell in row)


def token_parts(line: str) -> tuple[list[str], list[str]]:
    raw_tokens = re.findall(r"\S+", clean_text(line))
    folded_tokens = [fold_for_match(token) for token in raw_tokens]
    return raw_tokens, folded_tokens


def pattern_matches_at(tokens: list[str], start: int, pattern: tuple[str, ...]) -> bool:
    if start + len(pattern) > len(tokens):
        return False
    return tokens[start : start + len(pattern)] == list(pattern)


def find_label_start(tokens: list[str], start: int = 0) -> int | None:
    for index in range(start, len(tokens)):
        for pattern in ALL_LABEL_PATTERNS:
            if pattern_matches_at(tokens, index, pattern):
                return index
    return None


def detect_metadata_label_key(value: object) -> str:
    raw_tokens, folded_tokens = token_parts(clean_text(value))
    if not raw_tokens:
        return ""

    for key in ("Okul No", "Ad Soyad", "Ders Yılı", "Sınıf", "Şube", "Okul"):
        for pattern in LABEL_PATTERNS[key]:
            for index in range(len(folded_tokens)):
                if pattern_matches_at(folded_tokens, index, pattern):
                    return key
    return ""


def extract_labeled_value_from_line(line: str, key: str) -> str:
    raw_tokens, folded_tokens = token_parts(line)
    if not raw_tokens:
        return ""

    for pattern in LABEL_PATTERNS[key]:
        for index in range(len(folded_tokens)):
            if not pattern_matches_at(folded_tokens, index, pattern):
                continue

            value_start = index + len(pattern)
            while value_start < len(raw_tokens) and folded_tokens[value_start] == "":
                value_start += 1

            next_label = find_label_start(folded_tokens, value_start)
            value_end = next_label if next_label is not None else len(raw_tokens)
            value = " ".join(raw_tokens[value_start:value_end])
            return trim_metadata_value(key, value)

    return ""


def trim_at_right_side_text(value: str) -> str:
    text = clean_text(value)
    stop_patterns = [
        r"[\"“”]",
        r"\bÖğretmenler\b",
        r"\bMustafa\s+Kemal\b",
        r"\bŞube\s+Rehber\b",
        r"\bDavranışlar\b",
        r"\bGelişim\s+Düzeyleri\b",
        r"\bSınıf\s+Geçme\b",
        r"\bSonuç\b",
        r"\bSosyal\s+Etkinlikler\b",
        r"\bİmzalar\b",
        r"\bOkul\s+Müdürü\b",
        r"\bDersler\b",
        r"\bBaşarı\s+Durumu\b",
    ]

    for pattern in stop_patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            text = text[: match.start()]

    return clean_text(text)


def trim_metadata_value(key: str, value: str) -> str:
    value = clean_text(value).lstrip(":：")
    value = trim_at_right_side_text(value)

    if key == "Ad Soyad":
        match = re.search(
            r"(.+?)(?:\s+OKULU\b|\s+SINIFI\b|\s+DERS\s+YILI\b|$)",
            value,
            flags=re.IGNORECASE,
        )
        if match:
            value = match.group(1)

    if key == "Okul":
        match = re.search(
            r"(.+?\b(?:İLKOKULU|ORTAOKULU|LİSESİ|OKULU)\b)",
            value,
            flags=re.IGNORECASE,
        )
        if match:
            value = match.group(1)

    if key == "Sınıf":
        match = re.search(r"\b(\d+)\s*\.?\s*S[ıiİI]n[ıiİI]f\b", value, flags=re.IGNORECASE)
        if match:
            value = match.group(1)
        else:
            match = re.search(r"\b(\d+)\b", value)
            if match:
                value = match.group(1)

    if key == "Şube":
        match = re.search(
            r"\b([A-ZÇĞİÖŞÜ0-9]+)\s*Şubesi\b",
            value,
            flags=re.IGNORECASE,
        )
        if match:
            value = match.group(1)
        else:
            folded_tokens = fold_for_match(value).split()
            for token in folded_tokens:
                if token in {"A", "B", "C", "D", "E", "F", "G", "H"}:
                    value = token
                    break

    if key == "Ders Yılı":
        match = re.search(r"\b(20\d{2}\s*[-/]\s*20\d{2})\b", value)
        if match:
            value = match.group(1).replace(" ", "")

    if key == "Okul No":
        value = normalize_school_no(value)

    return clean_text(value)


def normalize_school_no(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"-?\s*\d+", text)
    if not match:
        return ""
    return re.sub(r"\s+", "", match.group(0))


def normalize_name(value: object) -> str:
    text = clean_text(value)
    text = re.sub(r"\s+", " ", text.replace("|", " ").replace("_", " ")).strip()
    return text.upper()


def normalize_school(value: object) -> str:
    return clean_text(value).upper()


def normalize_class(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"\d+", text)
    if match:
        return match.group(0)
    return text


def normalize_section(value: object) -> str:
    text = clean_text(value)
    folded = fold_for_match(text)

    raw_match = re.search(
        r"\b([A-ZÇĞİÖŞÜ0-9]+)\s*Şubesi\b",
        text,
        flags=re.IGNORECASE,
    )
    if raw_match:
        return raw_match.group(1).upper()

    match = re.search(r"\b([A-Z0-9]+)\s+SUBESI\b", folded)
    if match:
        return match.group(1)

    tokens = [token for token in folded.split() if token not in {"SUBE", "SUBESI"}]
    for token in tokens:
        if token in {"A", "B", "C", "D", "E", "F", "G", "H"}:
            return token

    if tokens and len(tokens[0]) <= 3:
        return tokens[0]

    return ""


def normalize_grade(value: object) -> str:
    text = clean_text(value)
    if not text:
        return ""

    match = re.search(r"\b\d+(?:[,.]\d+)?\b", text)
    if not match:
        return ""

    return match.group(0).replace(",", ".")


def extract_numbers(value: object) -> list[str]:
    text = clean_text(value)
    return [match.group(0).replace(",", ".") for match in re.finditer(r"\b\d+(?:[,.]\d+)?\b", text)]


def safe_crop_text(page, bbox: tuple[float, float, float, float]) -> str:
    try:
        cropped = page.crop(bbox)
        return cropped.extract_text(x_tolerance=2, y_tolerance=3) or ""
    except Exception:
        return ""


def header_text_candidates(page, page_text: str) -> list[str]:
    candidates: list[str] = []
    if page.width and page.height:
        candidates.append(safe_crop_text(page, (0, 0, page.width * LEFT_PAGE_RATIO, page.height * 0.20)))
        candidates.append(safe_crop_text(page, (0, 0, page.width * LEFT_PAGE_RATIO, page.height * 0.28)))
        candidates.append(safe_crop_text(page, (0, 0, page.width * LEFT_PAGE_RATIO, page.height)))
    candidates.append(page_text)
    return [candidate for candidate in candidates if clean_text(candidate)]


def extract_tables_from_area(page) -> list[list[list[object]]]:
    areas = []
    if page.width and page.height:
        areas.append(page.crop((0, 0, page.width * LEFT_PAGE_RATIO, page.height)))
        areas.append(page.crop((0, 0, page.width * 0.56, page.height)))
    else:
        areas.append(page)

    tables: list[list[list[object]]] = []
    seen_signatures: set[str] = set()

    for area in areas:
        for settings in TABLE_SETTINGS_VARIANTS:
            try:
                found_tables = area.extract_tables(table_settings=settings)
            except Exception:
                found_tables = []

            for table in found_tables:
                signature = "\n".join(row_text(row) for row in table if row_text(row))
                if signature and signature not in seen_signatures:
                    seen_signatures.add(signature)
                    tables.append(table)

    if tables:
        return tables

    try:
        return page.extract_tables() or []
    except Exception:
        return []


def value_to_right(row: list[object], label_index: int) -> str:
    values: list[str] = []

    for index in range(label_index + 1, len(row)):
        cell = clean_text(row[index])
        if not cell:
            continue
        if detect_metadata_label_key(cell):
            break
        values.append(cell)

    return clean_text(" ".join(values))


def update_metadata_from_table(metadata: dict[str, str], table: list[list[object]]) -> None:
    for row_index, row in enumerate(table[:20]):
        for column_index, cell in enumerate(row):
            key = detect_metadata_label_key(cell)
            if not key or metadata.get(key):
                continue

            value = extract_labeled_value_from_line(clean_text(cell), key)
            if not value:
                value = value_to_right(row, column_index)

            if not value and row_index + 1 < len(table):
                next_row_values = [
                    clean_text(next_cell)
                    for next_cell in table[row_index + 1]
                    if clean_text(next_cell) and not detect_metadata_label_key(next_cell)
                ]
                value = clean_text(" ".join(next_row_values))

            value = trim_metadata_value(key, value)
            if value:
                metadata[key] = value


def update_metadata_from_text(metadata: dict[str, str], text: str) -> None:
    for line in text.splitlines():
        line = clean_text(line)
        if not line:
            continue

        for key in METADATA_KEYS:
            if metadata.get(key):
                continue
            value = extract_labeled_value_from_line(line, key)
            if value:
                metadata[key] = value


def finalize_metadata(metadata: dict[str, str]) -> dict[str, str]:
    finalized = {key: clean_text(metadata.get(key, "")) for key in METADATA_KEYS}

    finalized["Ad Soyad"] = normalize_name(finalized["Ad Soyad"])
    finalized["Okul"] = normalize_school(finalized["Okul"])
    finalized["Sınıf"] = normalize_class(finalized["Sınıf"])
    finalized["Şube"] = normalize_section(finalized["Şube"])
    finalized["Okul No"] = normalize_school_no(finalized["Okul No"])

    return finalized


def extract_metadata(page, page_text: str, tables: list[list[list[object]]]) -> dict[str, str]:
    metadata = {key: "" for key in METADATA_KEYS}

    for text in header_text_candidates(page, page_text):
        update_metadata_from_text(metadata, text)

    for table in tables:
        update_metadata_from_table(metadata, table)

    return finalize_metadata(metadata)


def is_excluded_course(course: str) -> bool:
    folded = fold_for_match(course)
    return "SERBEST" in folded and "ETKINLIK" in folded


def detect_course_name_from_text(text: str) -> str:
    folded = fold_for_match(text)
    if not folded:
        return ""

    for folded_course, course in COURSE_FOLDS.items():
        if folded_course in folded:
            return course

    return ""


def detect_course_name_from_row(row: list[object]) -> str:
    course = detect_course_name_from_text(row_text(row))
    return course


def find_course_cell_index(row: list[object], course: str) -> int | None:
    folded_course = fold_for_match(course)
    if not folded_course:
        return None

    for index, cell in enumerate(row):
        if folded_course in fold_for_match(cell):
            return index

    for start in range(len(row)):
        combined = " ".join(fold_for_match(cell) for cell in row[start : start + 4])
        if folded_course in combined:
            return start

    return None


def grade_pair_from_text(row_value: str, course: str) -> tuple[str, str]:
    folded_row = fold_for_match(row_value)
    folded_course = fold_for_match(course)
    if folded_course and folded_course in folded_row:
        # Use the full row because the folded/raw positions do not map one-to-one.
        # The course names do not contain numbers, so this keeps the numeric order.
        numbers = extract_numbers(row_value)
    else:
        numbers = extract_numbers(row_value)

    if len(numbers) >= 3:
        return numbers[1], numbers[2]
    if len(numbers) == 2:
        return numbers[1], ""
    return "", ""


def grade_pair_from_table_row(row: list[object], course: str) -> tuple[str, str]:
    course_index = find_course_cell_index(row, course)
    if course_index is None:
        return grade_pair_from_text(row_text(row), course)

    after_course = [clean_text(cell) for cell in row[course_index + 1 :]]

    # The first useful numeric cell after the course is "Haftalık Ders Saati".
    while after_course and not normalize_grade(after_course[0]):
        after_course.pop(0)

    if len(after_course) >= 3:
        first_term = normalize_grade(after_course[1])
        second_term = normalize_grade(after_course[2])
        return first_term, second_term

    return grade_pair_from_text(row_text(row), course)


def empty_course_grade_map() -> dict[str, dict[str, str]]:
    return {course: {"1": "", "2": ""} for course in OUTPUT_COURSES}


def update_course_grade(
    grades: dict[str, dict[str, str]],
    course: str,
    first_term: str,
    second_term: str,
) -> None:
    if course not in grades:
        return

    if first_term or not grades[course]["1"]:
        grades[course]["1"] = first_term
    if second_term or not grades[course]["2"]:
        grades[course]["2"] = second_term


def extract_course_grades_from_tables(tables: list[list[list[object]]]) -> dict[str, dict[str, str]]:
    grades = empty_course_grade_map()

    for table in tables:
        for row in table:
            if row_is_empty(row):
                continue

            course = detect_course_name_from_row(row)
            if not course or is_excluded_course(course):
                continue

            first_term, second_term = grade_pair_from_table_row(row, course)
            update_course_grade(grades, course, first_term, second_term)

    return grades


def left_page_text(page, page_text: str) -> str:
    if page.width and page.height:
        text = safe_crop_text(page, (0, 0, page.width * LEFT_PAGE_RATIO, page.height))
        if clean_text(text):
            return text
    return page_text


def extract_course_grades_from_text(text: str) -> dict[str, dict[str, str]]:
    grades = empty_course_grade_map()

    for raw_line in text.splitlines():
        line = clean_text(raw_line)
        if not line:
            continue

        course = detect_course_name_from_text(line)
        if not course or is_excluded_course(course):
            continue

        first_term, second_term = grade_pair_from_text(line, course)
        update_course_grade(grades, course, first_term, second_term)

    return grades


def merge_grade_maps(
    primary: dict[str, dict[str, str]],
    fallback: dict[str, dict[str, str]],
) -> dict[str, dict[str, str]]:
    merged = empty_course_grade_map()

    for course in OUTPUT_COURSES:
        for term_key in ("1", "2"):
            merged[course][term_key] = primary[course][term_key] or fallback[course][term_key]

    return merged


def has_any_course_grade(grades: dict[str, dict[str, str]]) -> bool:
    return any(term_grades["1"] or term_grades["2"] for term_grades in grades.values())


def missing_metadata_fields(metadata: dict[str, str]) -> list[str]:
    return [key for key in METADATA_KEYS if not metadata.get(key)]


def build_student_row(
    metadata: dict[str, str],
    grades: dict[str, dict[str, str]],
) -> dict[str, str]:
    row = {
        "Okul": metadata.get("Okul", ""),
        "Ders Yılı": metadata.get("Ders Yılı", ""),
        "Sınıf": metadata.get("Sınıf", ""),
        "Şube": metadata.get("Şube", ""),
        "Okul No": metadata.get("Okul No", ""),
        "Ad Soyad": metadata.get("Ad Soyad", ""),
    }

    for course in OUTPUT_COURSES:
        row[course] = grades[course]["1"]

    return row


def extract_student_page(page, page_number: int) -> tuple[dict[str, str] | None, list[dict[str, str]]]:
    warnings: list[dict[str, str]] = []
    page_text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
    tables = extract_tables_from_area(page)

    metadata = extract_metadata(page, page_text, tables)
    table_grades = extract_course_grades_from_tables(tables)
    text_grades = extract_course_grades_from_text(left_page_text(page, page_text))
    grades = merge_grade_maps(table_grades, text_grades)

    if not metadata.get("Ad Soyad") and not metadata.get("Okul No") and not has_any_course_grade(grades):
        warnings.append(
            {
                "page": str(page_number),
                "student": "",
                "reason": "No student metadata or course table could be extracted; page skipped.",
            }
        )
        return None, warnings

    missing_fields = missing_metadata_fields(metadata)
    if missing_fields:
        warnings.append(
            {
                "page": str(page_number),
                "student": metadata.get("Ad Soyad", ""),
                "reason": "Missing metadata fields: " + ", ".join(missing_fields),
            }
        )

    if not has_any_course_grade(grades):
        warnings.append(
            {
                "page": str(page_number),
                "student": metadata.get("Ad Soyad", ""),
                "reason": "No course grades were detected on this page.",
            }
        )

    return build_student_row(metadata, grades), warnings


def resolve_pdf_path(raw_path: str | Path) -> Path:
    path = Path(raw_path).expanduser()

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(f"The input file is not a PDF: {path}")
        return path

    if path.suffix.lower() != ".pdf":
        pdf_version = path.with_suffix(".pdf")
        if pdf_version.is_file():
            return pdf_version

    raise FileNotFoundError(f"PDF file was not found: {path}")


def section_key(row: dict[str, str]) -> tuple[str, str]:
    class_no = normalize_class(row.get("Sınıf", "")) or "Bilinmeyen"
    section = normalize_section(row.get("Şube", "")) or "Bilinmeyen"
    return class_no, section


def group_rows_by_section(rows: list[dict[str, str]]) -> dict[tuple[str, str], list[dict[str, str]]]:
    grouped: dict[tuple[str, str], list[dict[str, str]]] = {}

    for row in rows:
        grouped.setdefault(section_key(row), []).append(row)

    return grouped


def save_excel(rows: list[dict[str, str]], output_path: Path, term_label: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = term_label

    worksheet.append(OUTPUT_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column, "") for column in OUTPUT_COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    school_no_column = OUTPUT_COLUMNS.index("Okul No") + 1
    school_no_letter = get_column_letter(school_no_column)
    for row_number in range(2, worksheet.max_row + 1):
        worksheet[f"{school_no_letter}{row_number}"].number_format = "@"

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header_value = column_cells[0].value
        max_length = max(len(str(cell.value or "")) for cell in column_cells)

        if header_value == "Ad Soyad":
            width = min(max(max_length + 2, 20), 42)
        elif header_value == "Okul":
            width = min(max(max_length + 2, 18), 46)
        elif header_value in OUTPUT_COURSES:
            width = min(max(max_length + 2, 12), 28)
        else:
            width = min(max(max_length + 2, 10), 18)

        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 42
    workbook.save(output_path)


def safe_file_part(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" ._")
    return value or "Bilinmeyen"


def output_path_for_group(
    output_dir: Path,
    class_no: str,
    section: str,
) -> Path:
    class_part = safe_file_part(class_no)
    section_part = safe_file_part(section)
    return output_dir / f"{class_part}_{section_part}_Şubesi_1_Dönem.xlsx"


def write_warning_log(warnings: list[dict[str, str]], output_dir: Path) -> Path | None:
    if not warnings:
        return None

    output_dir.mkdir(parents=True, exist_ok=True)
    log_path = output_dir / "type5_4th_grade_selectable_extraction_warnings.csv"
    fieldnames = ["page", "student", "reason"]

    with log_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for warning in warnings:
            writer.writerow({field: warning.get(field, "") for field in fieldnames})

    return log_path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert fourth-grade selectable report-card PDFs into one-term Excel files by section."
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=None,
        help="PDF file path. If omitted, the script asks for it.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Optional folder where Excel files and warnings will be saved.",
    )
    return parser.parse_args(argv)


def get_input_path(args: argparse.Namespace) -> str:
    if args.input_path:
        return args.input_path

    user_input = input("Enter the PDF file path: ").strip()
    if not user_input:
        raise ValueError("No input path was entered.")
    return user_input.strip('"')


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    pdf_path = resolve_pdf_path(get_input_path(args))
    output_dir = (
        Path(args.output_dir).expanduser()
        if args.output_dir
        else pdf_path.parent / "type5_4th_grade_selectable_output"
    )

    student_rows: list[dict[str, str]] = []
    warnings: list[dict[str, str]] = []

    total_pages = 0
    with pdfplumber.open(str(pdf_path)) as pdf:
        total_pages = len(pdf.pages)
        for page_index, page in enumerate(pdf.pages, start=1):
            row, page_warnings = extract_student_page(page, page_index)
            warnings.extend(page_warnings)

            if row is not None:
                student_rows.append(row)

            print(f"Processed page {page_index}/{total_pages}")

    if not student_rows:
        warning_log = write_warning_log(warnings, output_dir)
        if warning_log:
            print(f"Warning log saved: {warning_log}")
        raise RuntimeError("No student pages could be extracted.")

    grouped = group_rows_by_section(student_rows)
    created_files: list[Path] = []

    for (class_no, section), rows in grouped.items():
        output_path = output_path_for_group(output_dir, class_no, section)
        save_excel(rows, output_path, "1. Dönem")
        created_files.append(output_path)

    warning_log = write_warning_log(warnings, output_dir)
    found_sections = [
        f"{class_no}/{section}"
        for class_no, section in grouped.keys()
    ]

    print(f"Pages processed: {total_pages}")
    print(f"Student rows extracted: {len(student_rows)} from {pdf_path.name}")
    print("Şubeler found: " + (", ".join(found_sections) if found_sections else "none"))

    for output_path in created_files:
        print(f"Created: {output_path}")

    if warning_log:
        print(f"Warning log saved: {warning_log}")

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
