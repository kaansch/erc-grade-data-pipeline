# -*- coding: utf-8 -*-
r"""
Convert the fourth selectable-text PDF table type into Excel.

Use this script for the Beylikduzu-style table where each PDF page contains
up to three student blocks. Each student block contains all course grades.
Only the "1. Donem Puani" row is extracted. The "2. Donem Puani" row and
Serbest Etkinlikler are ignored.

Grade conversion:
    G      -> 1
    I / İ  -> 2
    C.I. / Ç.İ. -> 3

Install requirements:

    pip install pdfplumber openpyxl

Run with a path:

    python extract_pdf_grades_type4_selectable.py "C:\path\to\pdf_or_folder"

Or run without a path and the script will ask you to enter one:

    python extract_pdf_grades_type4_selectable.py
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from pathlib import Path
from typing import Iterable

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    missing_name = getattr(exc, "name", "required package")
    raise SystemExit(
        f"Missing Python package: {missing_name}\n"
        "Install the requirements with:\n\n"
        "    pip install pdfplumber openpyxl\n"
    ) from exc


OUTPUT_COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "HAYAT BİLGİSİ",
    "GÖRSEL SANATLAR",
    "YABANCI DİL",
    "MÜZİK",
    "FEN BİLİMLERİ",
    "BEDEN EĞİTİMİ VE OYUN",
]

ALL_COURSES = [*OUTPUT_COURSES, "SERBEST ETKİNLİKLER"]

OUTPUT_COLUMNS = [
    "İl",
    "İlçe",
    "Öğretim Yılı",
    "Okul",
    "Sınıf",
    "Şube",
    "Okul No",
    "Ad Soyad",
    *OUTPUT_COURSES,
]

TABLE_SETTINGS_VARIANTS = [
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "lines",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 3,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "text",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
]

TURKISH_TO_ASCII = str.maketrans(
    {
        "ç": "c",
        "Ç": "C",
        "ğ": "g",
        "Ğ": "G",
        "ı": "i",
        "I": "I",
        "İ": "I",
        "ö": "o",
        "Ö": "O",
        "ş": "s",
        "Ş": "S",
        "ü": "u",
        "Ü": "U",
    }
)

NEXT_LABEL_RE = re.compile(
    r"\s+(?:İlçe|Ilce|Öğretim\s+Yılı|Ogretim\s+Yili|Okulu?|Okulun\s+Adı|"
    r"Sınıf\s*/\s*Şube|Sinif\s*/\s*Sube|Sınıf|Sinif|Şube|Sube|Sayfa|"
    r"Okul\s+Numarası|Okul\s+Numarasi|Okul\s+No|"
    r"Adı\s+ve\s+Soyadı|Adi\s+ve\s+Soyadi|Ad\s+Soyad|"
    r"Ana\s+Adı|Ana\s+Adi|Baba\s+Adı|Baba\s+Adi|"
    r"Doğum\s+Tarihi|Dogum\s+Tarihi|Doğum\s+Yeri|Dogum\s+Yeri|"
    r"T\.?\s*C\.?\s*Kimlik\s+Numarası|T\.?\s*C\.?\s*Kimlik\s+Numarasi|"
    r"Uyruğu|Uyrugu|Yabancı\s+Dil|Yabanci\s+Dil)\s*[:：]",
    flags=re.IGNORECASE,
)


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def fold_for_match(value: object) -> str:
    text = clean_text(value).translate(TURKISH_TO_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def row_text(row: Iterable[object]) -> str:
    return " ".join(clean_text(cell) for cell in row if clean_text(cell))


def normalize_school_no(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"-?\s*\d+", text)
    if not match:
        return ""
    return re.sub(r"\s+", "", match.group(0))


def normalize_name(value: object) -> str:
    text = clean_text(value)
    text = text.replace("|", " ").replace("_", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text.upper()


def truncate_at_next_label(value: str) -> str:
    match = NEXT_LABEL_RE.search(value)
    if match:
        value = value[: match.start()]
    return clean_text(value)


def labeled_value_from_text(text: str, label_patterns: list[str]) -> str:
    for line in text.splitlines():
        line = clean_text(line)
        if not line:
            continue

        for label_pattern in label_patterns:
            pattern = rf"{label_pattern}\s*[:：]\s*(.+)"
            match = re.search(pattern, line, flags=re.IGNORECASE)
            if match:
                return truncate_at_next_label(match.group(1))

    return ""


def extract_header_info(page_text: str) -> dict[str, str]:
    province = labeled_value_from_text(page_text, [r"\bİl\b", r"\bIl\b"])
    district = labeled_value_from_text(page_text, [r"İlçe", r"Ilce"])
    school_year = labeled_value_from_text(
        page_text,
        [r"Öğretim\s+Yılı", r"Ogretim\s+Yili", r"Eğitim\s+Öğretim\s+Yılı"],
    )
    school = labeled_value_from_text(
        page_text,
        [r"\bOkulu\b", r"\bOkul\b", r"Okulun\s+Adı", r"Okul\s+Adı"],
    )
    class_section = labeled_value_from_text(
        page_text,
        [r"Sınıf\s*/\s*Şube", r"Sinif\s*/\s*Sube"],
    )
    class_no = labeled_value_from_text(page_text, [r"\bSınıf\b", r"\bSinif\b"])
    section = labeled_value_from_text(page_text, [r"\bŞube\b", r"\bSube\b"])

    if class_section:
        class_match = re.search(
            r"(?P<class>\d+)\.\s*S[ıiİI]n[ıiİI]f\s*/\s*(?P<section>.+?)\s*Şubesi",
            class_section,
            flags=re.IGNORECASE,
        )
        if class_match:
            class_no = clean_text(class_match.group("class"))
            section = clean_text(class_match.group("section")).upper()

    if not school_year:
        year_match = re.search(r"\b(20\d{2}\s*[-/]\s*20\d{2})\b", page_text)
        if year_match:
            school_year = clean_text(year_match.group(1).replace(" ", ""))

    if not class_no or not section:
        class_match = re.search(
            r"(?P<class>\d+)\.\s*S[ıiİI]n[ıiİI]f\s*/?\s*(?P<section>[A-ZÇĞİÖŞÜ0-9]+)",
            page_text,
            flags=re.IGNORECASE,
        )
        if class_match:
            class_no = class_no or clean_text(class_match.group("class"))
            section = section or clean_text(class_match.group("section")).upper()

    return {
        "İl": province,
        "İlçe": district,
        "Öğretim Yılı": school_year,
        "Okul": school,
        "Sınıf": class_no,
        "Şube": section,
    }


def resolve_input_pdfs(raw_path: str | Path) -> tuple[list[Path], Path]:
    path = Path(raw_path).expanduser()

    if path.is_dir():
        pdf_files = sorted(item for item in path.glob("*.pdf") if item.is_file())
        if not pdf_files:
            raise FileNotFoundError(f"No PDF files were found in folder: {path}")
        return pdf_files, path

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(f"The input file is not a PDF: {path}")
        return [path], path.parent

    if path.suffix.lower() != ".pdf":
        pdf_version = path.with_suffix(".pdf")
        if pdf_version.is_file():
            return [pdf_version], pdf_version.parent

    raise FileNotFoundError(
        f"PDF file or folder was not found: {path}\n"
        "Tip: pass one PDF file path or a folder containing PDF files."
    )


def row_is_empty(row: Iterable[object]) -> bool:
    return not any(clean_text(cell) for cell in row)


def value_from_row(row: list[object], index: int | None) -> str:
    if index is None or index < 0 or index >= len(row):
        return ""
    return clean_text(row[index])


def grade_symbol_to_number(value: object) -> str:
    text = clean_text(value)
    if not text:
        return ""

    folded = fold_for_match(text)
    compact = folded.replace(" ", "")

    if compact in {"1", "2", "3"}:
        return compact

    if "GELISTIRILMELI" in folded or compact == "G":
        return "1"

    if "COK" in folded or compact in {"CI", "CII", "C"}:
        return "3"

    if compact in {"I", "II", "IYI"} or ("IYI" in folded and "COK" not in folded):
        return "2"

    return ""


def grade_symbols_from_text(value: str) -> list[str]:
    text = clean_text(value)
    if not text:
        return []

    matches = re.finditer(
        r"Ç\s*\.?\s*İ\s*\.?|C\s*\.?\s*I\s*\.?|G|İ|I|[123]|P",
        text,
        flags=re.IGNORECASE,
    )

    return [clean_text(match.group(0)) for match in matches]


def is_first_term_grade_row(row_or_text: object) -> bool:
    folded = fold_for_match(row_or_text)
    tokens = set(folded.split())
    return "1" in tokens and "DONEM" in tokens and "PUANI" in tokens


def is_second_term_grade_row(row_or_text: object) -> bool:
    folded = fold_for_match(row_or_text)
    tokens = set(folded.split())
    return "2" in tokens and "DONEM" in tokens and "PUANI" in tokens


def extract_grade_values_from_table_row(row: list[object]) -> list[str]:
    label_index: int | None = None
    for index, cell in enumerate(row):
        if is_first_term_grade_row(cell):
            label_index = index
            break

    if label_index is None:
        return []

    grade_cells = [clean_text(cell) for cell in row[label_index + 1 :]]
    if not any(grade_cells):
        # Some PDF table extractors collapse the whole row into one cell.
        collapsed = clean_text(row_text(row))
        label_removed = re.sub(
            r"1\s*\.?\s*D[öo]nem\s+Puan[ıi]",
            "",
            collapsed,
            count=1,
            flags=re.IGNORECASE,
        )
        symbols = grade_symbols_from_text(label_removed)
        if symbols and fold_for_match(symbols[0]) == "P":
            symbols = symbols[1:]
        return [grade_symbol_to_number(symbol) for symbol in symbols[:8]]

    p_index = None
    for index, cell in enumerate(grade_cells):
        if fold_for_match(cell) == "P":
            p_index = index
            break

    if p_index is not None:
        grade_cells = grade_cells[p_index + 1 :]

    converted = [grade_symbol_to_number(cell) for cell in grade_cells[:8]]
    while len(converted) < len(OUTPUT_COURSES):
        converted.append("")

    return converted[: len(OUTPUT_COURSES)]


def extract_grade_values_from_grade_only_row(row: list[object]) -> list[str]:
    cells = [clean_text(cell) for cell in row]
    if not any(cells):
        return []

    p_index = None
    for index, cell in enumerate(cells):
        if fold_for_match(cell) == "P":
            p_index = index
            break

    if p_index is not None:
        grade_cells = cells[p_index + 1 : p_index + 1 + len(OUTPUT_COURSES)]
        converted = [grade_symbol_to_number(cell) for cell in grade_cells]
        while len(converted) < len(OUTPUT_COURSES):
            converted.append("")
        return converted[: len(OUTPUT_COURSES)]

    symbols = grade_symbols_from_text(row_text(row))
    if symbols and fold_for_match(symbols[0]) == "P":
        symbols = symbols[1:]

    converted = [grade_symbol_to_number(symbol) for symbol in symbols[: len(OUTPUT_COURSES)]]
    while len(converted) < len(OUTPUT_COURSES):
        converted.append("")

    return converted[: len(OUTPUT_COURSES)] if any(converted) else []


def has_any_grade(grades: list[str]) -> bool:
    return any(clean_text(grade) for grade in grades)


def extract_labeled_value_from_block(
    block_rows: list[list[object]],
    label_tokens: tuple[str, ...],
    value_kind: str,
) -> str:
    for row_index, row in enumerate(block_rows):
        for column_index, cell in enumerate(row):
            cell_text = clean_text(cell)
            folded = fold_for_match(cell_text)
            if not all(token in folded for token in label_tokens):
                continue

            if ":" in cell_text:
                after_colon = truncate_at_next_label(cell_text.split(":", 1)[1])
                if value_kind == "school_no":
                    school_no = normalize_school_no(after_colon)
                    if school_no:
                        return school_no
                elif after_colon:
                    return normalize_name(after_colon)

            for next_column in range(column_index + 1, len(row)):
                candidate = clean_text(row[next_column])
                if not candidate:
                    continue
                if value_kind == "school_no":
                    school_no = normalize_school_no(candidate)
                    if school_no:
                        return school_no
                elif not any(token in fold_for_match(candidate) for token in label_tokens):
                    return normalize_name(candidate)

            if row_index + 1 < len(block_rows):
                for candidate_cell in block_rows[row_index + 1]:
                    candidate = clean_text(candidate_cell)
                    if not candidate:
                        continue
                    if value_kind == "school_no":
                        school_no = normalize_school_no(candidate)
                        if school_no:
                            return school_no
                    elif not any(token in fold_for_match(candidate) for token in label_tokens):
                        return normalize_name(candidate)

    return ""


def extract_student_info_from_text(block_text: str) -> tuple[str, str]:
    school_no = ""
    name = ""

    school_match = re.search(
        r"Okul\s+Numaras[ıi]\s*[:：]?\s*(-?\s*\d+)",
        block_text,
        flags=re.IGNORECASE,
    )
    if school_match:
        school_no = normalize_school_no(school_match.group(1))

    name_match = re.search(
        r"Ad[ıi]\s+ve\s+Soyad[ıi]\s*[:：]?\s*(.+?)(?:\n|$)",
        block_text,
        flags=re.IGNORECASE,
    )
    if name_match:
        name = normalize_name(truncate_at_next_label(name_match.group(1)))

    return school_no, name


def extract_student_info_from_table_block(block_rows: list[list[object]]) -> tuple[str, str]:
    block_text = "\n".join(row_text(row) for row in block_rows)
    school_no, name = extract_student_info_from_text(block_text)

    if not school_no:
        school_no = extract_labeled_value_from_block(
            block_rows,
            ("OKUL", "NUMARASI"),
            value_kind="school_no",
        )
    if not school_no:
        school_no = extract_labeled_value_from_block(
            block_rows,
            ("OKUL", "NO"),
            value_kind="school_no",
        )

    if not name:
        name = extract_labeled_value_from_block(
            block_rows,
            ("ADI", "SOYADI"),
            value_kind="name",
        )
    if not name:
        name = extract_labeled_value_from_block(
            block_rows,
            ("AD", "SOYAD"),
            value_kind="name",
        )

    return school_no, name


def split_table_into_student_blocks(table: list[list[object]]) -> list[list[list[object]]]:
    start_indices: list[int] = []
    for index, row in enumerate(table):
        folded = fold_for_match(row_text(row))
        if "HAFTALIK" in folded and "DERS" in folded and "SAATI" in folded:
            start_indices.append(index)

    if not start_indices:
        for index, row in enumerate(table):
            folded = fold_for_match(row_text(row))
            if is_first_term_grade_row(folded):
                start_indices.append(index)

    if not start_indices:
        for index, row in enumerate(table):
            folded = fold_for_match(row_text(row))
            if "OKUL" in folded and ("NUMARASI" in folded or "NO" in folded):
                start_indices.append(index)

    if not start_indices:
        return [table]

    blocks: list[list[list[object]]] = []
    for position, start_index in enumerate(start_indices):
        end_index = start_indices[position + 1] if position + 1 < len(start_indices) else len(table)
        block = table[start_index:end_index]
        if block:
            blocks.append(block)

    return blocks


def parse_student_block(
    block_rows: list[list[object]],
    header: dict[str, str],
) -> dict[str, str] | None:
    school_no, name = extract_student_info_from_table_block(block_rows)
    if not school_no and not name:
        return None

    first_term_index: int | None = None
    for index, row in enumerate(block_rows):
        if is_first_term_grade_row(row_text(row)):
            first_term_index = index
            break
        if is_second_term_grade_row(row_text(row)):
            continue

    if first_term_index is None:
        return None

    grades = extract_grade_values_from_table_row(block_rows[first_term_index])
    if not has_any_grade(grades):
        for next_row in block_rows[first_term_index + 1 :]:
            if is_second_term_grade_row(row_text(next_row)):
                break
            grades = extract_grade_values_from_grade_only_row(next_row)
            if has_any_grade(grades):
                break

    output_row = {
        **header,
        "Okul No": school_no,
        "Ad Soyad": name,
    }
    for index, course in enumerate(OUTPUT_COURSES):
        output_row[course] = grades[index] if index < len(grades) else ""

    return output_row


def extract_tables_from_page(page) -> list[list[list[object]]]:
    all_tables: list[list[list[object]]] = []
    seen_signatures: set[str] = set()

    for settings in TABLE_SETTINGS_VARIANTS:
        try:
            tables = page.extract_tables(table_settings=settings)
        except Exception:
            tables = []

        for table in tables:
            signature = "\n".join(row_text(row) for row in table)
            if signature and signature not in seen_signatures:
                seen_signatures.add(signature)
                all_tables.append(table)

    if all_tables:
        return all_tables

    try:
        return page.extract_tables()
    except Exception:
        return []


def parse_rows_from_tables(page, header: dict[str, str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for table in extract_tables_from_page(page):
        for block_rows in split_table_into_student_blocks(table):
            row = parse_student_block(block_rows, header)
            if row is not None:
                rows.append(row)
    return rows


def split_text_into_student_blocks(page_text: str) -> list[str]:
    matches = list(
        re.finditer(
            r"Haftalık\s+Ders\s+Saati|Haftalik\s+Ders\s+Saati",
            page_text,
            flags=re.IGNORECASE,
        )
    )
    if not matches:
        matches = list(
            re.finditer(
                r"Okul\s+Numaras[ıi]|Okul\s+No",
                page_text,
                flags=re.IGNORECASE,
            )
        )
    if not matches:
        return []

    blocks: list[str] = []
    for index, match in enumerate(matches):
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(page_text)
        block = page_text[start:end].strip()
        if block:
            blocks.append(block)

    return blocks


def parse_grade_values_from_text_block(block_text: str) -> list[str]:
    lines = [clean_text(line) for line in block_text.splitlines() if clean_text(line)]

    for line_index, line in enumerate(lines):
        if not is_first_term_grade_row(line):
            continue

        line_without_label = re.sub(
            r"1\s*\.?\s*D[öo]nem\s+Puan[ıi]",
            "",
            line,
            count=1,
            flags=re.IGNORECASE,
        )
        symbols = grade_symbols_from_text(line_without_label)
        if symbols and fold_for_match(symbols[0]) == "P":
            symbols = symbols[1:]

        grades = [grade_symbol_to_number(symbol) for symbol in symbols[:8]]
        while len(grades) < len(OUTPUT_COURSES):
            grades.append("")

        if has_any_grade(grades):
            return grades[: len(OUTPUT_COURSES)]

        for next_line in lines[line_index + 1 :]:
            if is_second_term_grade_row(next_line):
                break
            symbols = grade_symbols_from_text(next_line)
            if symbols and fold_for_match(symbols[0]) == "P":
                symbols = symbols[1:]
            grades = [grade_symbol_to_number(symbol) for symbol in symbols[:8]]
            while len(grades) < len(OUTPUT_COURSES):
                grades.append("")
            if has_any_grade(grades):
                return grades[: len(OUTPUT_COURSES)]

    return []


def parse_rows_from_text(page_text: str, header: dict[str, str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    for block_text in split_text_into_student_blocks(page_text):
        school_no, name = extract_student_info_from_text(block_text)
        if not school_no and not name:
            continue

        grades = parse_grade_values_from_text_block(block_text)
        if not grades:
            continue

        output_row = {
            **header,
            "Okul No": school_no,
            "Ad Soyad": name,
        }
        for index, course in enumerate(OUTPUT_COURSES):
            output_row[course] = grades[index] if index < len(grades) else ""
        rows.append(output_row)

    return rows


def extract_pdf(pdf_path: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            header = extract_header_info(page_text)

            page_rows = parse_rows_from_tables(page, header)
            if not page_rows:
                page_rows = parse_rows_from_text(page_text, header)

            rows.extend(page_rows)

    if not rows:
        raise ValueError(f"No student rows could be extracted from {pdf_path.name}.")

    return rows


def save_excel(rows: list[dict[str, str]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Notlar"

    worksheet.append(OUTPUT_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column, "") for column in OUTPUT_COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    school_no_column = OUTPUT_COLUMNS.index("Okul No") + 1
    school_no_letter = get_column_letter(school_no_column)
    for row_number in range(2, worksheet.max_row + 1):
        worksheet[f"{school_no_letter}{row_number}"].number_format = "@"

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header_value = column_cells[0].value
        max_length = max(len(str(cell.value or "")) for cell in column_cells)

        if header_value == "Ad Soyad":
            width = min(max(max_length + 2, 20), 42)
        elif header_value in {"İl", "İlçe", "Okul"}:
            width = min(max(max_length + 2, 12), 36)
        elif header_value in OUTPUT_COURSES:
            width = min(max(max_length + 2, 12), 28)
        else:
            width = min(max(max_length + 2, 10), 18)

        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 42
    workbook.save(output_path)


def safe_file_part(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" ._")
    return value or "grades"


def build_output_path(pdf_path: Path, output_base_dir: Path) -> Path:
    return output_base_dir / f"{safe_file_part(pdf_path.stem)}.xlsx"


def write_error_log(errors: list[dict[str, str]], output_dir: Path) -> Path | None:
    if not errors:
        return None

    output_dir.mkdir(parents=True, exist_ok=True)
    log_path = output_dir / "type4_selectable_extraction_errors.csv"
    fieldnames = ["pdf", "reason"]

    with log_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for error in errors:
            writer.writerow({field: error.get(field, "") for field in fieldnames})

    return log_path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert type-4 selectable student-block grade PDFs into Excel."
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=None,
        help="PDF file path or folder path. If omitted, the script asks for it.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Optional folder where Excel files and error logs will be saved.",
    )
    return parser.parse_args(argv)


def get_input_path(args: argparse.Namespace) -> str:
    if args.input_path:
        return args.input_path

    user_input = input("Enter a PDF file path or a folder path containing PDF files: ").strip()
    if not user_input:
        raise ValueError("No input path was entered.")
    return user_input.strip('"')


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    input_path = get_input_path(args)
    pdf_paths, input_base_dir = resolve_input_pdfs(input_path)
    output_base_dir = (
        Path(args.output_dir).expanduser()
        if args.output_dir
        else input_base_dir / "type4_selectable_output"
    )

    created_files: list[Path] = []
    errors: list[dict[str, str]] = []

    print(f"PDF files found: {len(pdf_paths)}")
    for index, pdf_path in enumerate(pdf_paths, start=1):
        try:
            rows = extract_pdf(pdf_path)
            output_path = build_output_path(pdf_path, output_base_dir)
            save_excel(rows, output_path)
            created_files.append(output_path)
            print(f"[{index}/{len(pdf_paths)}] OK: {pdf_path.name} -> {output_path}")
        except Exception as error:
            errors.append({"pdf": str(pdf_path), "reason": str(error)})
            print(f"[{index}/{len(pdf_paths)}] ERROR: {pdf_path.name}: {error}", file=sys.stderr)

    log_path = write_error_log(errors, output_base_dir)

    print(f"Converted PDFs: {len(created_files)}")
    for output_path in created_files:
        print(f"Created: {output_path}")

    if log_path:
        print(f"Error log saved: {log_path}")

    return 1 if errors else 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
