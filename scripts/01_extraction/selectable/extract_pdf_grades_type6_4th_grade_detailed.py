# -*- coding: utf-8 -*-
r"""
Convert the 4th-grade detailed selectable PDF table type into Excel files.

Expected PDF structure:
    - Each PDF belongs to one class section.
    - Each PDF has exactly 11 course pages.
    - Each page contains one course table.
    - Course page order is fixed.
    - Each student row has detailed 1st-term grades, detailed 2nd-term
      grades, and a year-end score.

For each PDF, this script creates:
    PDF_NAME_CLASS_SECTION_1_donem.xlsx
    PDF_NAME_CLASS_SECTION_2_donem.xlsx
    PDF_NAME_CLASS_SECTION_yil_sonu.xlsx

Install requirements:

    pip install pdfplumber openpyxl

Run with a PDF path or a folder path:

    python extract_pdf_grades_type6_4th_grade_detailed.py "C:\path\to\pdf_or_folder"

Or run without a path and the script will ask you to enter one.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from pathlib import Path
from typing import Iterable

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    missing_name = getattr(exc, "name", "required package")
    raise SystemExit(
        f"Missing Python package: {missing_name}\n"
        "Install the requirements with:\n\n"
        "    pip install pdfplumber openpyxl\n"
    ) from exc


COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "SOSYAL BİLGİLER",
    "YABANCI DİL",
    "DİN KÜLTÜRÜ VE AHLAK BİLGİSİ",
    "GÖRSEL SANATLAR",
    "MÜZİK",
    "TRAFİK GÜVENLİĞİ",
    "FEN BİLİMLERİ",
    "İNSAN HAKLARI, YURTTAŞLIK VE DEMOKRASİ",
    "BEDEN EĞİTİMİ VE OYUN",
]

BASE_COLUMNS = [
    "İlçe",
    "Okul",
    "Ders Yılı",
    "Sınıf",
    "Şube",
    "Okul No",
    "Ad Soyad",
]

TERM_DETAIL_SUFFIXES = [
    "Sınav 1",
    "Sınav 2",
    "Sınav 3",
    "Sınav 4",
    "Katılım 1",
    "Katılım 2",
    "Katılım 3",
    "Katılım 4",
    "Katılım 5",
]

FIRST_TERM_SUFFIXES = [*TERM_DETAIL_SUFFIXES, "1. Dönem Puanı"]
SECOND_TERM_SUFFIXES = [*TERM_DETAIL_SUFFIXES, "2. Dönem Puanı"]
YEAR_END_SUFFIX = "Yıl Sonu Puanı"

TABLE_SETTINGS_VARIANTS = [
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "lines",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 3,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "text",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
]

TURKISH_TO_ASCII = str.maketrans(
    {
        "ç": "c",
        "Ç": "C",
        "ğ": "g",
        "Ğ": "G",
        "ı": "i",
        "İ": "I",
        "ö": "o",
        "Ö": "O",
        "ş": "s",
        "Ş": "S",
        "ü": "u",
        "Ü": "U",
    }
)


def repair_mojibake(value: str) -> str:
    if not isinstance(value, str):
        return value

    suspicious = any(char in value for char in ("Ã", "Ä", "Å", "Â"))
    suspicious = suspicious or any("\x80" <= char <= "\x9f" for char in value)
    if not suspicious:
        return value

    try:
        return value.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return value


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", "\n").replace("\xa0", " ")
    text = repair_mojibake(text)
    return re.sub(r"\s+", " ", text).strip()


def fold_for_match(value: object) -> str:
    text = clean_text(value).translate(TURKISH_TO_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def row_text(row: Iterable[object]) -> str:
    return " ".join(clean_text(cell) for cell in row if clean_text(cell))


def row_is_empty(row: Iterable[object]) -> bool:
    return not any(clean_text(cell) for cell in row)


def normalize_school_no(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"-?\s*\d+", text)
    if not match:
        return ""
    return re.sub(r"\s+", "", match.group(0))


def normalize_name(value: object) -> str:
    text = clean_text(value).replace("|", " ").replace("_", " ")
    return re.sub(r"\s+", " ", text).strip().upper()


def normalize_class(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"\d+", text)
    if match:
        return match.group(0)
    return text


def normalize_section(value: object) -> str:
    text = clean_text(value)
    folded = fold_for_match(text)

    raw_match = re.search(
        r"\b([A-ZÇĞİÖŞÜ0-9]+)\s*Şubesi\b",
        text,
        flags=re.IGNORECASE,
    )
    if raw_match:
        return raw_match.group(1).upper()

    folded_match = re.search(r"\b([A-Z0-9]+)\s+SUBESI\b", folded)
    if folded_match:
        return folded_match.group(1)

    tokens = [token for token in folded.split() if token not in {"SUBE", "SUBESI"}]
    for token in tokens:
        if token in {"A", "B", "C", "D", "E", "F", "G", "H", "I"}:
            return token

    if tokens and len(tokens[0]) <= 3:
        return tokens[0]

    return ""


def normalize_grade(value: object) -> str:
    text = clean_text(value)
    if not text:
        return ""
    return text.replace(",", ".")


def grade_column(course: str, suffix: str) -> str:
    return f"{course} {suffix}"


def term_columns(suffixes: list[str]) -> list[str]:
    columns: list[str] = []
    for course in COURSES:
        for suffix in suffixes:
            columns.append(grade_column(course, suffix))
    return columns


FIRST_TERM_COLUMNS = [*BASE_COLUMNS, *term_columns(FIRST_TERM_SUFFIXES)]
SECOND_TERM_COLUMNS = [*BASE_COLUMNS, *term_columns(SECOND_TERM_SUFFIXES)]
YEAR_END_COLUMNS = [
    *BASE_COLUMNS,
    *[grade_column(course, YEAR_END_SUFFIX) for course in COURSES],
]


def resolve_input_pdfs(raw_path: str | Path) -> tuple[list[Path], Path]:
    path = Path(raw_path).expanduser()

    if path.is_dir():
        pdf_files = sorted(item for item in path.glob("*.pdf") if item.is_file())
        if not pdf_files:
            raise FileNotFoundError(f"No PDF files were found in folder: {path}")
        return pdf_files, path

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(f"The input file is not a PDF: {path}")
        return [path], path.parent

    if path.suffix.lower() != ".pdf":
        pdf_version = path.with_suffix(".pdf")
        if pdf_version.is_file():
            return [pdf_version], pdf_version.parent

    raise FileNotFoundError(
        f"PDF file or folder was not found: {path}\n"
        "Tip: pass one PDF file path or a folder containing PDF files."
    )


def extract_header_info(page_text: str) -> dict[str, str]:
    district = ""
    school = ""
    school_year = ""
    class_no = ""
    section = ""

    lines = [clean_text(line) for line in page_text.splitlines() if clean_text(line)]

    for line in lines[:15]:
        folded_line = fold_for_match(line)
        if "/" not in line:
            continue
        if "MUDURLUGU" not in folded_line and "MUDURLUK" not in folded_line:
            continue

        district_part, school_part = line.split("/", 1)
        district = clean_text(district_part)
        school_part = clean_text(school_part)
        school = re.split(
            r"\s+M[üu]d[üu]rl[üu][ğg][üu]\b",
            school_part,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0].strip()
        break

    header_text = " ".join(lines[:20])

    year_match = re.search(r"\b(20\d{2}\s*[-/]\s*20\d{2})\b", header_text)
    if year_match:
        school_year = year_match.group(1).replace(" ", "")

    raw_section_match = re.search(
        r"(?P<class>\d+)\.\s*S[ıiİI]n[ıiİI]f\s*/\s*"
        r"(?P<section>.+?)\s*Şubesi",
        header_text,
        flags=re.IGNORECASE,
    )
    if raw_section_match:
        class_no = raw_section_match.group("class")
        section = normalize_section(raw_section_match.group("section"))

    if not class_no or not section:
        folded_header = fold_for_match(header_text)
        folded_match = re.search(
            r"\b(?P<class>\d+)\s+SINIF\s*/?\s*(?P<section>[A-Z0-9]+)\s+SUBESI\b",
            folded_header,
        )
        if folded_match:
            class_no = class_no or folded_match.group("class")
            section = section or folded_match.group("section")

    return {
        "İlçe": district,
        "Okul": school,
        "Ders Yılı": school_year,
        "Sınıf": normalize_class(class_no) or "4",
        "Şube": normalize_section(section),
    }


def value_from_row(row: list[object], index: int) -> str:
    if index < 0 or index >= len(row):
        return ""
    return normalize_grade(row[index])


def find_table_columns(table: list[list[object]]) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(table):
        folded_cells = [fold_for_match(cell) for cell in row]
        columns: dict[str, int] = {}

        for col_index, cell in enumerate(folded_cells):
            tokens = set(cell.split())

            if "OKUL" in tokens and "NO" in tokens:
                columns.setdefault("school_no", col_index)

            if "OGRENCININ" in tokens and "ADI" in tokens and "SOYADI" in tokens:
                columns.setdefault("name", col_index)

            if "1" in tokens and "DONEM" in tokens and "PUANI" in tokens:
                columns.setdefault("first_score", col_index)

            if "2" in tokens and "DONEM" in tokens and "PUANI" in tokens:
                columns.setdefault("second_score", col_index)

        if {"school_no", "name", "first_score", "second_score"}.issubset(columns):
            columns.setdefault("year_score", columns["second_score"] + 1)
            return row_index, columns

    return None


def row_looks_like_student(row: list[object], school_no_index: int, name_index: int) -> bool:
    school_no = normalize_school_no(value_from_row(row, school_no_index))
    name = normalize_name(value_from_row(row, name_index))

    if not school_no or not re.fullmatch(r"\d+", school_no):
        return False
    if not name:
        return False

    folded_name = fold_for_match(name)
    return "OGRENCI" not in folded_name and "SOYAD" not in folded_name


def parse_rows_from_table(table: list[list[object]]) -> list[dict[str, object]]:
    column_info = find_table_columns(table)
    if column_info is None:
        return []

    header_row_index, columns = column_info
    school_no_col = columns["school_no"]
    name_col = columns["name"]
    first_score_col = columns["first_score"]
    second_score_col = columns["second_score"]
    year_score_col = columns["year_score"]

    first_exam_cols = [3, 4, 5, 6]
    first_participation_cols = list(range(first_score_col - 5, first_score_col))
    second_exam_cols = list(range(first_score_col + 1, first_score_col + 5))
    second_participation_cols = list(range(second_score_col - 5, second_score_col))

    rows: list[dict[str, object]] = []
    seen_school_numbers: set[str] = set()

    for row in table[header_row_index + 1 :]:
        if row_is_empty(row):
            continue
        if not row_looks_like_student(row, school_no_col, name_col):
            continue

        school_no = normalize_school_no(value_from_row(row, school_no_col))
        if school_no in seen_school_numbers:
            continue
        seen_school_numbers.add(school_no)

        rows.append(
            {
                "school_no": school_no,
                "name": normalize_name(value_from_row(row, name_col)),
                "first_values": [
                    *[value_from_row(row, index) for index in first_exam_cols],
                    *[value_from_row(row, index) for index in first_participation_cols],
                    value_from_row(row, first_score_col),
                ],
                "second_values": [
                    *[value_from_row(row, index) for index in second_exam_cols],
                    *[value_from_row(row, index) for index in second_participation_cols],
                    value_from_row(row, second_score_col),
                ],
                "year_value": value_from_row(row, year_score_col),
            }
        )

    return rows


def parse_rows_from_tables(tables: Iterable[list[list[object]]]) -> list[dict[str, object]]:
    for table in tables:
        if not table:
            continue
        rows = parse_rows_from_table(table)
        if rows:
            return rows
    return []


def split_name_and_numeric_tokens(line: str) -> tuple[str, str, list[str]] | None:
    match = re.match(r"^\s*\d+\s+(?P<school_no>\d+)\s+(?P<rest>.+)$", line)
    if not match:
        return None

    school_no = normalize_school_no(match.group("school_no"))
    tokens = clean_text(match.group("rest")).split()
    first_numeric_index = None
    for index, token in enumerate(tokens):
        if re.fullmatch(r"\d+(?:[,.]\d+)?", token):
            first_numeric_index = index
            break

    if first_numeric_index is None:
        return None

    name = normalize_name(" ".join(tokens[:first_numeric_index]))
    numeric_tokens = [normalize_grade(token) for token in tokens[first_numeric_index:]]
    return school_no, name, numeric_tokens


def parse_rows_from_text(page_text: str) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    seen_school_numbers: set[str] = set()

    for raw_line in page_text.splitlines():
        parsed = split_name_and_numeric_tokens(clean_text(raw_line))
        if parsed is None:
            continue

        school_no, name, values = parsed
        if not school_no or not name or school_no in seen_school_numbers:
            continue

        # Text extraction drops empty cells, so this fallback preserves the final
        # score columns reliably but may not preserve blank exam/participation
        # positions. The table parser is preferred whenever possible.
        year_value = values[-1] if len(values) >= 1 else ""
        second_score = values[-2] if len(values) >= 2 else ""
        first_score = values[-3] if len(values) >= 3 else ""

        first_values = [""] * len(TERM_DETAIL_SUFFIXES) + [first_score]
        second_values = [""] * len(TERM_DETAIL_SUFFIXES) + [second_score]

        rows.append(
            {
                "school_no": school_no,
                "name": name,
                "first_values": first_values,
                "second_values": second_values,
                "year_value": year_value,
            }
        )
        seen_school_numbers.add(school_no)

    return rows


def extract_student_rows(page, page_text: str) -> list[dict[str, object]]:
    for settings in TABLE_SETTINGS_VARIANTS:
        try:
            tables = page.extract_tables(table_settings=settings)
        except Exception:
            tables = []

        rows = parse_rows_from_tables(tables)
        if rows:
            return rows

    try:
        fallback_tables = page.extract_tables() or []
    except Exception:
        fallback_tables = []

    rows = parse_rows_from_tables(fallback_tables)
    if rows:
        return rows

    return parse_rows_from_text(page_text)


def new_record(header: dict[str, str], school_no: str, name: str, columns: list[str]) -> dict[str, str]:
    row = {column: "" for column in columns}
    row.update(
        {
            "İlçe": header.get("İlçe", ""),
            "Okul": header.get("Okul", ""),
            "Ders Yılı": header.get("Ders Yılı", ""),
            "Sınıf": header.get("Sınıf", ""),
            "Şube": header.get("Şube", ""),
            "Okul No": school_no,
            "Ad Soyad": name,
        }
    )
    return row


def ensure_student(
    students: dict[str, dict[str, str]],
    student_order: list[str],
    header: dict[str, str],
    school_no: str,
    name: str,
    columns: list[str],
) -> dict[str, str]:
    if school_no not in students:
        students[school_no] = new_record(header, school_no, name, columns)
        if school_no not in student_order:
            student_order.append(school_no)
    else:
        row = students[school_no]
        for column in BASE_COLUMNS:
            value = header.get(column, "") if column not in {"Okul No", "Ad Soyad"} else ""
            if value and not row.get(column):
                row[column] = value
        if name and not row.get("Ad Soyad"):
            row["Ad Soyad"] = name

    return students[school_no]


def fill_term_values(row: dict[str, str], course: str, suffixes: list[str], values: list[str]) -> None:
    for suffix, value in zip(suffixes, values):
        row[grade_column(course, suffix)] = value


def rows_in_order(students: dict[str, dict[str, str]], student_order: list[str]) -> list[dict[str, str]]:
    return [students[school_no] for school_no in student_order if school_no in students]


def process_pdf(
    pdf_path: Path,
) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]], list[dict[str, str]]]:
    first_students: dict[str, dict[str, str]] = {}
    second_students: dict[str, dict[str, str]] = {}
    year_students: dict[str, dict[str, str]] = {}
    student_order: list[str] = []
    warnings: list[dict[str, str]] = []

    with pdfplumber.open(str(pdf_path)) as pdf:
        if len(pdf.pages) != len(COURSES):
            raise ValueError(
                f"Expected exactly {len(COURSES)} pages, but found {len(pdf.pages)} pages."
            )

        for page_index, page in enumerate(pdf.pages):
            course = COURSES[page_index]
            page_number = page_index + 1
            page_text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            header = extract_header_info(page_text)
            page_rows = extract_student_rows(page, page_text)

            if not page_rows:
                warnings.append(
                    {
                        "pdf": str(pdf_path),
                        "page": str(page_number),
                        "reason": f"No student rows extracted for {course}.",
                    }
                )
                continue

            for source_row in page_rows:
                school_no = str(source_row["school_no"])
                name = str(source_row["name"])

                first_row = ensure_student(
                    first_students,
                    student_order,
                    header,
                    school_no,
                    name,
                    FIRST_TERM_COLUMNS,
                )
                second_row = ensure_student(
                    second_students,
                    student_order,
                    header,
                    school_no,
                    name,
                    SECOND_TERM_COLUMNS,
                )
                year_row = ensure_student(
                    year_students,
                    student_order,
                    header,
                    school_no,
                    name,
                    YEAR_END_COLUMNS,
                )

                fill_term_values(
                    first_row,
                    course,
                    FIRST_TERM_SUFFIXES,
                    list(source_row["first_values"]),
                )
                fill_term_values(
                    second_row,
                    course,
                    SECOND_TERM_SUFFIXES,
                    list(source_row["second_values"]),
                )
                year_row[grade_column(course, YEAR_END_SUFFIX)] = str(source_row["year_value"])

    return (
        rows_in_order(first_students, student_order),
        rows_in_order(second_students, student_order),
        rows_in_order(year_students, student_order),
        warnings,
    )


def save_excel(rows: list[dict[str, str]], columns: list[str], output_path: Path, sheet_title: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_title

    worksheet.append(columns)
    for row in rows:
        worksheet.append([row.get(column, "") for column in columns])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    school_no_column = columns.index("Okul No") + 1
    school_no_letter = get_column_letter(school_no_column)
    for row_number in range(2, worksheet.max_row + 1):
        worksheet[f"{school_no_letter}{row_number}"].number_format = "@"

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header_value = str(column_cells[0].value or "")
        max_length = max(len(str(cell.value or "")) for cell in column_cells)

        if header_value == "Ad Soyad":
            width = min(max(max_length + 2, 20), 42)
        elif header_value in {"İlçe", "Okul"}:
            width = min(max(max_length + 2, 14), 36)
        elif header_value in BASE_COLUMNS:
            width = min(max(max_length + 2, 10), 18)
        else:
            width = min(max(max_length + 2, 12), 26)

        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 44
    workbook.save(output_path)


def safe_file_part(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" ._")
    return value or "bilinmeyen"


def output_paths_for_pdf(pdf_path: Path, output_dir: Path, rows: list[dict[str, str]]) -> tuple[Path, Path, Path]:
    class_no = "sinif"
    section = "sube"

    if rows:
        class_no = safe_file_part(rows[0].get("Sınıf", "") or class_no)
        section = safe_file_part(rows[0].get("Şube", "") or section)

    stem = safe_file_part(pdf_path.stem)
    prefix = f"{stem}_{class_no}_{section}"
    return (
        output_dir / f"{prefix}_1_donem.xlsx",
        output_dir / f"{prefix}_2_donem.xlsx",
        output_dir / f"{prefix}_yil_sonu.xlsx",
    )


def write_warning_log(warnings: list[dict[str, str]], output_dir: Path) -> Path | None:
    if not warnings:
        return None

    output_dir.mkdir(parents=True, exist_ok=True)
    log_path = output_dir / "type6_4th_grade_detailed_warnings.csv"
    fieldnames = ["pdf", "page", "reason"]

    with log_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for warning in warnings:
            writer.writerow({field: warning.get(field, "") for field in fieldnames})

    return log_path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert detailed 4th-grade PDFs into 1st term, 2nd term, and year-end Excel files."
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=None,
        help="PDF file path or folder path. If omitted, the script asks for it.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Optional folder where Excel files and warnings will be saved.",
    )
    return parser.parse_args(argv)


def get_input_path(args: argparse.Namespace) -> str:
    if args.input_path:
        return args.input_path

    user_input = input("Enter a PDF file path or a folder path containing PDFs: ").strip()
    if not user_input:
        raise ValueError("No input path was entered.")
    return user_input.strip('"')


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    pdf_paths, input_base_dir = resolve_input_pdfs(get_input_path(args))
    output_dir = (
        Path(args.output_dir).expanduser()
        if args.output_dir
        else input_base_dir / "type6_4th_grade_detailed_output"
    )

    created_files: list[Path] = []
    all_warnings: list[dict[str, str]] = []
    converted_count = 0

    print(f"PDF files found: {len(pdf_paths)}")
    for pdf_index, pdf_path in enumerate(pdf_paths, start=1):
        try:
            first_rows, second_rows, year_rows, warnings = process_pdf(pdf_path)
            all_warnings.extend(warnings)

            first_path, second_path, year_path = output_paths_for_pdf(
                pdf_path,
                output_dir,
                first_rows,
            )

            save_excel(first_rows, FIRST_TERM_COLUMNS, first_path, "1. Dönem")
            save_excel(second_rows, SECOND_TERM_COLUMNS, second_path, "2. Dönem")
            save_excel(year_rows, YEAR_END_COLUMNS, year_path, "Yıl Sonu")

            created_files.extend([first_path, second_path, year_path])
            converted_count += 1
            print(
                f"[{pdf_index}/{len(pdf_paths)}] OK: {pdf_path.name} "
                f"-> rows={len(first_rows)}"
            )
        except Exception as error:
            all_warnings.append({"pdf": str(pdf_path), "page": "", "reason": str(error)})
            print(f"[{pdf_index}/{len(pdf_paths)}] ERROR: {pdf_path.name}: {error}", file=sys.stderr)

    warning_log = write_warning_log(all_warnings, output_dir)

    if not created_files:
        if warning_log:
            print(f"Warning log saved: {warning_log}")
        raise RuntimeError("No Excel files were created.")

    print(f"PDFs processed successfully: {converted_count}")
    for output_path in created_files:
        print(f"Created: {output_path}")
    if warning_log:
        print(f"Warning log saved: {warning_log}")

    return 0 if converted_count == len(pdf_paths) else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
