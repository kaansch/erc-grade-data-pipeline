# -*- coding: utf-8 -*-
r"""
Convert the original table form into separate Excel files for 1. Donem and
2. Donem.

Expected PDF structure:
    - Each PDF has exactly 8 pages.
    - Each page is one course.
    - Course page order is fixed:
        1. TURKCE
        2. MATEMATIK
        3. HAYAT BILGISI
        4. YABANCI DIL
        5. GORSEL SANATLAR
        6. MUZIK
        7. FEN BILIMLERI
        8. BEDEN EGITIMI VE OYUN
    - Each student row has:
        S.No, Okul No, Ad Soyad, 1. Donem grade, 2. Donem grade

For each PDF, this script creates two Excel files:
    <pdf_name>_<sinif>_<sube>_1_donem.xlsx
    <pdf_name>_<sinif>_<sube>_2_donem.xlsx

Install requirements:

    pip install pdfplumber openpyxl

Run with a PDF path or folder path:

    python extract_pdf_grades_original_2_terms_fixed.py "C:\path\to\file_or_folder"

Or run without a path and the script will ask you to enter one.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from pathlib import Path
from typing import Iterable

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    missing_name = getattr(exc, "name", "required package")
    raise SystemExit(
        f"Missing Python package: {missing_name}\n"
        "Install the requirements with:\n\n"
        "    pip install pdfplumber openpyxl\n"
    ) from exc


COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "HAYAT BİLGİSİ",
    "YABANCI DİL",
    "GÖRSEL SANATLAR",
    "MÜZİK",
    "FEN BİLİMLERİ",
    "BEDEN EĞİTİMİ VE OYUN",
]

OUTPUT_COLUMNS = [
    "İlçe",
    "Okul",
    "Ders Yılı",
    "Sınıf",
    "Şube",
    "Okul No",
    "Ad Soyad",
    *COURSES,
]

GRADE_VALUES = [
    "GELİŞTİRİLMELİ",
    "GELISTIRILMELI",
    "ÇOK İYİ",
    "COK IYI",
    "İYİ",
    "IYI",
    "YETERLİ",
    "YETERLI",
    "ORTA",
    "ZAYIF",
    "PEK İYİ",
    "PEKİYİ",
    "PEK IYI",
    "PEKIYI",
]

TABLE_SETTINGS_VARIANTS = [
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "lines",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 3,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "lines",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
    {
        "vertical_strategy": "text",
        "horizontal_strategy": "text",
        "snap_tolerance": 3,
        "join_tolerance": 3,
        "intersection_tolerance": 4,
        "text_x_tolerance": 2,
        "text_y_tolerance": 3,
    },
]

TURKISH_TO_ASCII = str.maketrans(
    {
        "ç": "c",
        "Ç": "C",
        "ğ": "g",
        "Ğ": "G",
        "ı": "i",
        "İ": "I",
        "ö": "o",
        "Ö": "O",
        "ş": "s",
        "Ş": "S",
        "ü": "u",
        "Ü": "U",
    }
)


def repair_mojibake(value: str) -> str:
    if not isinstance(value, str):
        return value

    suspicious = any(char in value for char in ("Ã", "Ä", "Å", "Â"))
    suspicious = suspicious or any("\x80" <= char <= "\x9f" for char in value)
    if not suspicious:
        return value

    try:
        return value.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return value


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", "\n").replace("\xa0", " ")
    text = repair_mojibake(text)
    return re.sub(r"\s+", " ", text).strip()


def fold_for_match(value: object) -> str:
    text = clean_text(value).translate(TURKISH_TO_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


GRADE_TOKEN_PATTERNS = [
    (fold_for_match(grade).split(), grade)
    for grade in sorted(GRADE_VALUES, key=lambda item: len(fold_for_match(item).split()), reverse=True)
]
VALID_GRADE_FOLDS = {fold_for_match(grade) for grade in GRADE_VALUES}


def row_text(row: Iterable[object]) -> str:
    return " ".join(clean_text(cell) for cell in row if clean_text(cell))


def row_is_empty(row: Iterable[object]) -> bool:
    return not any(clean_text(cell) for cell in row)


def normalize_school_no(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"-?\s*\d+", text)
    if not match:
        return ""
    return re.sub(r"\s+", "", match.group(0))


def normalize_name(value: object) -> str:
    text = clean_text(value).replace("|", " ").replace("_", " ")
    return re.sub(r"\s+", " ", text).strip().upper()


def normalize_grade(value: object) -> str:
    return clean_text(value).upper()


def is_valid_grade(value: object) -> bool:
    text = normalize_grade(value)
    if not text:
        return True
    return fold_for_match(text) in VALID_GRADE_FOLDS


def normalize_class(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"\d+", text)
    if match:
        return match.group(0)
    return text


def normalize_section(value: object) -> str:
    text = clean_text(value)
    folded = fold_for_match(text)

    raw_match = re.search(
        r"\b([A-ZÇĞİÖŞÜ0-9]+)\s*Şubesi\b",
        text,
        flags=re.IGNORECASE,
    )
    if raw_match:
        return raw_match.group(1).upper()

    folded_match = re.search(r"\b([A-Z0-9]+)\s+SUBESI\b", folded)
    if folded_match:
        return folded_match.group(1)

    tokens = [token for token in folded.split() if token not in {"SUBE", "SUBESI"}]
    for token in tokens:
        if token in {"A", "B", "C", "D", "E", "F", "G", "H"}:
            return token

    if tokens and len(tokens[0]) <= 3:
        return tokens[0]

    return ""


def resolve_input_pdfs(raw_path: str | Path) -> tuple[list[Path], Path]:
    path = Path(raw_path).expanduser()

    if path.is_dir():
        pdf_files = sorted(item for item in path.glob("*.pdf") if item.is_file())
        if not pdf_files:
            raise FileNotFoundError(f"No PDF files were found in folder: {path}")
        return pdf_files, path

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(f"The input file is not a PDF: {path}")
        return [path], path.parent

    if path.suffix.lower() != ".pdf":
        pdf_version = path.with_suffix(".pdf")
        if pdf_version.is_file():
            return [pdf_version], pdf_version.parent

    raise FileNotFoundError(
        f"PDF file or folder was not found: {path}\n"
        "Tip: pass one PDF file path or a folder containing PDF files."
    )


def extract_header_info(page_text: str) -> dict[str, str]:
    district = ""
    school = ""
    school_year = ""
    class_no = ""
    section = ""

    lines = [clean_text(line) for line in page_text.splitlines() if clean_text(line)]

    for line in lines[:15]:
        folded_line = fold_for_match(line)
        if "/" not in line:
            continue
        if "MUDURLUGU" not in folded_line and "MUDURLUK" not in folded_line:
            continue

        district_part, school_part = line.split("/", 1)
        district = clean_text(district_part)
        school_part = clean_text(school_part)
        school = re.split(
            r"\s+M[üu]d[üu]rl[üu][ğg][üu]\b",
            school_part,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0].strip()
        break

    header_text = " ".join(lines[:20])

    year_match = re.search(r"\b(20\d{2}\s*[-/]\s*20\d{2})\b", header_text)
    if year_match:
        school_year = year_match.group(1).replace(" ", "")

    raw_section_match = re.search(
        r"(?P<class>\d+)\.\s*S[ıiİI]n[ıiİI]f\s*/\s*"
        r"(?P<section>.+?)\s*Şubesi",
        header_text,
        flags=re.IGNORECASE,
    )
    if raw_section_match:
        class_no = raw_section_match.group("class")
        section = normalize_section(raw_section_match.group("section"))

    if not class_no or not section:
        folded_header = fold_for_match(header_text)
        folded_match = re.search(
            r"\b(?P<class>\d+)\s+SINIF\s*/?\s*(?P<section>[A-Z0-9]+)\s+SUBESI\b",
            folded_header,
        )
        if folded_match:
            class_no = class_no or folded_match.group("class")
            section = section or folded_match.group("section")

    return {
        "district": district,
        "school": school,
        "school_year": school_year,
        "class_no": normalize_class(class_no) or "3",
        "section": normalize_section(section),
    }


def find_table_columns(table: list[list[object]]) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(table):
        folded_cells = [fold_for_match(cell) for cell in row]
        columns: dict[str, int] = {}

        for col_index, cell in enumerate(folded_cells):
            tokens = set(cell.split())

            if "OKUL" in tokens and "NO" in tokens:
                columns.setdefault("school_no", col_index)

            if "OGRENCININ" in tokens and "ADI" in tokens and "SOYADI" in tokens:
                columns.setdefault("name", col_index)

            if (
                "1" in tokens
                and "DONEM" in tokens
                and "DEGERLENDIRME" in tokens
                and "DURUMU" in tokens
            ):
                columns.setdefault("first_grade", col_index)

            if (
                "2" in tokens
                and "DONEM" in tokens
                and "DEGERLENDIRME" in tokens
                and "DURUMU" in tokens
            ):
                columns.setdefault("second_grade", col_index)

        if {"school_no", "name", "first_grade", "second_grade"}.issubset(columns):
            return row_index, columns

    return None


def parse_pdf_tables(tables: Iterable[list[list[object]]]) -> list[dict[str, str]]:
    student_rows: list[dict[str, str]] = []
    seen_school_numbers: set[str] = set()

    for table in tables:
        if not table:
            continue

        column_info = find_table_columns(table)
        if column_info is None:
            widest_row = max((len(row) for row in table), default=0)
            if widest_row < 5:
                continue
            header_row_index = 0
            columns = {
                "school_no": 1,
                "name": 2,
                "first_grade": 3,
                "second_grade": 4,
            }
        else:
            header_row_index, columns = column_info

        max_needed_index = max(columns.values())

        for row in table[header_row_index + 1 :]:
            if len(row) <= max_needed_index or row_is_empty(row):
                continue

            school_no = normalize_school_no(row[columns["school_no"]])
            name = normalize_name(row[columns["name"]])
            first_grade = normalize_grade(row[columns["first_grade"]])
            second_grade = normalize_grade(row[columns["second_grade"]])

            if not school_no or not re.fullmatch(r"\d+", school_no):
                continue
            if not name or "OGRENCI" in fold_for_match(name) or "SOYAD" in fold_for_match(name):
                continue
            if not is_valid_grade(first_grade) or not is_valid_grade(second_grade):
                continue

            if school_no in seen_school_numbers:
                continue
            seen_school_numbers.add(school_no)

            student_rows.append(
                {
                    "school_no": school_no,
                    "name": name,
                    "first_grade": first_grade,
                    "second_grade": second_grade,
                }
            )

    return student_rows


def strip_grade_tokens_from_end(tokens: list[str]) -> tuple[list[str], str]:
    folded_tokens = [fold_for_match(token) for token in tokens]

    for folded_grade_tokens, _canonical_grade in GRADE_TOKEN_PATTERNS:
        if not folded_grade_tokens:
            continue
        token_count = len(folded_grade_tokens)
        if len(tokens) < token_count:
            continue
        if folded_tokens[-token_count:] == folded_grade_tokens:
            grade = normalize_grade(" ".join(tokens[-token_count:]))
            return tokens[:-token_count], grade

    return tokens, ""


def parse_text_rows(page_text: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    for raw_line in page_text.splitlines():
        line = clean_text(raw_line)
        if not line:
            continue

        match = re.match(r"^\s*\d+\s+(?P<school_no>\d+)\s+(?P<rest>.+)$", line)
        if not match:
            continue

        school_no = normalize_school_no(match.group("school_no"))
        rest_tokens = clean_text(match.group("rest")).split()
        if not school_no or not rest_tokens:
            continue

        rest_tokens, second_grade = strip_grade_tokens_from_end(rest_tokens)
        rest_tokens, first_grade = strip_grade_tokens_from_end(rest_tokens)

        # If only one grade was visible in raw text, keep it as the first grade.
        # Table extraction is used first because only it can reliably preserve an
        # empty first-grade cell with a non-empty second-grade cell.
        if first_grade == "" and second_grade != "":
            first_grade = second_grade
            second_grade = ""

        name = normalize_name(" ".join(rest_tokens))
        if not name or "OGRENCI" in fold_for_match(name) or "SOYAD" in fold_for_match(name):
            continue

        rows.append(
            {
                "school_no": school_no,
                "name": name,
                "first_grade": first_grade,
                "second_grade": second_grade,
            }
        )

    return rows


def extract_tables_from_page(page) -> list[list[list[object]]]:
    tables: list[list[list[object]]] = []
    seen_signatures: set[str] = set()

    for settings in TABLE_SETTINGS_VARIANTS:
        try:
            found_tables = page.extract_tables(table_settings=settings)
        except Exception:
            found_tables = []

        for table in found_tables:
            signature = "\n".join(row_text(row) for row in table if row_text(row))
            if signature and signature not in seen_signatures:
                seen_signatures.add(signature)
                tables.append(table)

    if tables:
        return tables

    try:
        return page.extract_tables() or []
    except Exception:
        return []


def extract_student_rows(page, page_text: str) -> list[dict[str, str]]:
    for settings in TABLE_SETTINGS_VARIANTS:
        try:
            tables = page.extract_tables(table_settings=settings)
        except Exception:
            tables = []

        rows = parse_pdf_tables(tables)
        if rows:
            return rows

    try:
        fallback_tables = page.extract_tables() or []
    except Exception:
        fallback_tables = []

    rows = parse_pdf_tables(fallback_tables)
    if rows:
        return rows

    return parse_text_rows(page_text)


def new_student_record(header: dict[str, str], school_no: str, name: str) -> dict[str, str]:
    row = {column: "" for column in OUTPUT_COLUMNS}
    row.update(
        {
            "İlçe": header.get("district", ""),
            "Okul": header.get("school", ""),
            "Ders Yılı": header.get("school_year", ""),
            "Sınıf": header.get("class_no", "") or "3",
            "Şube": header.get("section", ""),
            "Okul No": school_no,
            "Ad Soyad": name,
        }
    )
    return row


def ensure_student(
    students: dict[str, dict[str, str]],
    student_order: list[str],
    header: dict[str, str],
    school_no: str,
    name: str,
) -> dict[str, str]:
    if school_no not in students:
        students[school_no] = new_student_record(header, school_no, name)
        if school_no not in student_order:
            student_order.append(school_no)
    else:
        row = students[school_no]
        for column, value in (
            ("İlçe", header.get("district", "")),
            ("Okul", header.get("school", "")),
            ("Ders Yılı", header.get("school_year", "")),
            ("Sınıf", header.get("class_no", "") or "3"),
            ("Şube", header.get("section", "")),
            ("Ad Soyad", name),
        ):
            if value and not row.get(column):
                row[column] = value

    return students[school_no]


def process_pdf(pdf_path: Path) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]]]:
    first_term_students: dict[str, dict[str, str]] = {}
    second_term_students: dict[str, dict[str, str]] = {}
    student_order: list[str] = []
    warnings: list[dict[str, str]] = []

    with pdfplumber.open(str(pdf_path)) as pdf:
        if len(pdf.pages) != len(COURSES):
            raise ValueError(
                f"Expected exactly {len(COURSES)} pages, but found {len(pdf.pages)} pages."
            )

        for page_index, page in enumerate(pdf.pages):
            course = COURSES[page_index]
            page_number = page_index + 1
            page_text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            header = extract_header_info(page_text)
            page_rows = extract_student_rows(page, page_text)

            if not page_rows:
                warnings.append(
                    {
                        "pdf": str(pdf_path),
                        "page": str(page_number),
                        "reason": f"No student rows extracted for {course}.",
                    }
                )
                continue

            for source_row in page_rows:
                school_no = source_row["school_no"]
                name = source_row["name"]

                first_row = ensure_student(
                    first_term_students,
                    student_order,
                    header,
                    school_no,
                    name,
                )
                second_row = ensure_student(
                    second_term_students,
                    student_order,
                    header,
                    school_no,
                    name,
                )

                first_row[course] = source_row.get("first_grade", "")
                second_row[course] = source_row.get("second_grade", "")

    first_rows = [first_term_students[school_no] for school_no in student_order]
    second_rows = [second_term_students[school_no] for school_no in student_order]
    return first_rows, second_rows, warnings


def save_excel(rows: list[dict[str, str]], output_path: Path, term_label: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = term_label

    worksheet.append(OUTPUT_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column, "") for column in OUTPUT_COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    school_no_column = OUTPUT_COLUMNS.index("Okul No") + 1
    school_no_letter = get_column_letter(school_no_column)
    for row_number in range(2, worksheet.max_row + 1):
        worksheet[f"{school_no_letter}{row_number}"].number_format = "@"

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        header_value = column_cells[0].value
        max_length = max(len(str(cell.value or "")) for cell in column_cells)

        if header_value == "Ad Soyad":
            width = min(max(max_length + 2, 20), 42)
        elif header_value in {"İlçe", "Okul"}:
            width = min(max(max_length + 2, 14), 36)
        elif header_value in COURSES:
            width = min(max(max_length + 2, 14), 30)
        else:
            width = min(max(max_length + 2, 10), 18)

        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 42
    workbook.save(output_path)


def safe_file_part(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" ._")
    return value or "bilinmeyen"


def output_paths_for_pdf(pdf_path: Path, output_dir: Path, rows: list[dict[str, str]]) -> tuple[Path, Path]:
    class_no = "sinif"
    section = "sube"

    if rows:
        class_no = safe_file_part(rows[0].get("Sınıf", "") or class_no)
        section = safe_file_part(rows[0].get("Şube", "") or section)

    stem = safe_file_part(pdf_path.stem)
    first_path = output_dir / f"{stem}_{class_no}_{section}_1_donem.xlsx"
    second_path = output_dir / f"{stem}_{class_no}_{section}_2_donem.xlsx"
    return first_path, second_path


def write_warning_log(warnings: list[dict[str, str]], output_dir: Path) -> Path | None:
    if not warnings:
        return None

    output_dir.mkdir(parents=True, exist_ok=True)
    log_path = output_dir / "original_2_terms_fixed_warnings.csv"
    fieldnames = ["pdf", "page", "reason"]

    with log_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for warning in warnings:
            writer.writerow({field: warning.get(field, "") for field in fieldnames})

    return log_path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert original-form 8-page PDFs into separate 1st/2nd term Excel files."
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=None,
        help="PDF file path or folder path. If omitted, the script asks for it.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Optional folder where Excel files and warnings will be saved.",
    )
    return parser.parse_args(argv)


def get_input_path(args: argparse.Namespace) -> str:
    if args.input_path:
        return args.input_path

    user_input = input("Enter a PDF file path or a folder path containing PDFs: ").strip()
    if not user_input:
        raise ValueError("No input path was entered.")
    return user_input.strip('"')


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    pdf_paths, input_base_dir = resolve_input_pdfs(get_input_path(args))
    output_dir = (
        Path(args.output_dir).expanduser()
        if args.output_dir
        else input_base_dir / "original_2_terms_fixed_output"
    )

    created_files: list[Path] = []
    all_warnings: list[dict[str, str]] = []
    converted_count = 0

    print(f"PDF files found: {len(pdf_paths)}")
    for pdf_index, pdf_path in enumerate(pdf_paths, start=1):
        try:
            first_rows, second_rows, warnings = process_pdf(pdf_path)
            all_warnings.extend(warnings)

            first_output_path, second_output_path = output_paths_for_pdf(
                pdf_path,
                output_dir,
                first_rows,
            )

            save_excel(first_rows, first_output_path, "1. Dönem")
            save_excel(second_rows, second_output_path, "2. Dönem")
            created_files.extend([first_output_path, second_output_path])
            converted_count += 1

            print(
                f"[{pdf_index}/{len(pdf_paths)}] OK: {pdf_path.name} "
                f"-> rows={len(first_rows)}"
            )
        except Exception as error:
            all_warnings.append({"pdf": str(pdf_path), "page": "", "reason": str(error)})
            print(f"[{pdf_index}/{len(pdf_paths)}] ERROR: {pdf_path.name}: {error}", file=sys.stderr)

    warning_log = write_warning_log(all_warnings, output_dir)

    if not created_files:
        if warning_log:
            print(f"Warning log saved: {warning_log}")
        raise RuntimeError("No Excel files were created.")

    print(f"PDFs processed successfully: {converted_count}")
    for output_path in created_files:
        print(f"Created: {output_path}")
    if warning_log:
        print(f"Warning log saved: {warning_log}")

    return 0 if converted_count == len(pdf_paths) else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
