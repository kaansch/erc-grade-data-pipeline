# -*- coding: utf-8 -*-
r"""
Adjustment 2:
Convert separate course PDFs into class Excel files.

Use this version when:
- each PDF file contains only one course,
- each PDF is one or two pages long,
- PDF files are not in course order,
- the course name must be read from the PDF title/header,
- the table structure is the same as in the original code.

The script reads every PDF in a folder, combines the course grades by student,
and saves one Excel file per İlçe + Okul + Sınıf + Şube group.

Before running for the first time, install the required packages in the
Python environment used by VS Code:

    pip install pdfplumber openpyxl

Run on a folder:

    python extract_pdf_grades_adjustment2.py "C:\path\to\folder_with_course_pdfs"

Run on one course PDF only:

    python extract_pdf_grades_adjustment2.py "C:\path\to\course.pdf"
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from pathlib import Path
from typing import Iterable, NamedTuple

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    missing_name = getattr(exc, "name", "required package")
    raise SystemExit(
        f"Missing Python package: {missing_name}\n"
        "Install the requirements with:\n\n"
        "    pip install pdfplumber openpyxl\n"
    ) from exc


DEFAULT_INPUT_PATH = Path(r"C:\ERC-DATA\input")

COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "HAYAT BİLGİSİ",
    "YABANCI DİL",
    "GÖRSEL SANATLAR",
    "MÜZİK",
    "FEN BİLİMLERİ",
    "BEDEN EĞİTİMİ VE OYUN",
]

OUTPUT_COLUMNS = [
    "İlçe",
    "Okul",
    "Sınıf",
    "Şube",
    "Okul No",
    "Ad Soyad",
    *COURSES,
]

GRADE_VALUES = [
    "GELİŞTİRİLMELİ",
    "ÇOK İYİ",
    "İYİ",
    "YETERLİ",
    "ORTA",
    "ZAYIF",
    "PEK İYİ",
    "PEKİYİ",
]

TABLE_SETTINGS = {
    "vertical_strategy": "lines",
    "horizontal_strategy": "lines",
    "snap_tolerance": 3,
    "join_tolerance": 3,
    "intersection_tolerance": 3,
}

TURKISH_TO_ASCII = str.maketrans(
    {
        "ç": "c",
        "Ç": "C",
        "ğ": "g",
        "Ğ": "G",
        "ı": "i",
        "I": "I",
        "İ": "I",
        "ö": "o",
        "Ö": "O",
        "ş": "s",
        "Ş": "S",
        "ü": "u",
        "Ü": "U",
    }
)


class GroupKey(NamedTuple):
    district: str
    school: str
    class_no: str
    section: str


def clean_text(value: object) -> str:
    """Normalize whitespace while preserving Turkish characters."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def fold_for_match(value: object) -> str:
    """Return an ASCII-ish uppercase version for robust Turkish text matching."""
    text = clean_text(value).translate(TURKISH_TO_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


COURSE_BY_FOLDED_NAME = {fold_for_match(course): course for course in COURSES}


def resolve_input_pdfs(raw_path: str | Path) -> list[Path]:
    """Accept a folder, a PDF path, or a PDF path without the .pdf suffix."""
    path = Path(raw_path).expanduser()

    if path.is_dir():
        pdf_files = sorted(item for item in path.glob("*.pdf") if item.is_file())
        if not pdf_files:
            raise FileNotFoundError(f"No PDF files were found in folder: {path}")
        return pdf_files

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(f"The input file is not a PDF: {path}")
        return [path]

    if path.suffix.lower() != ".pdf":
        pdf_version = path.with_suffix(".pdf")
        if pdf_version.is_file():
            return [pdf_version]

    raise FileNotFoundError(
        f"PDF file or folder was not found: {path}\n"
        "Tip: pass a folder containing course PDF files, or pass one PDF file."
    )


def default_output_dir(input_arg: str | Path, pdf_paths: list[Path]) -> Path:
    input_path = Path(input_arg).expanduser()
    if input_path.is_dir():
        return input_path
    if len(pdf_paths) == 1:
        return pdf_paths[0].parent
    return DEFAULT_INPUT_PATH


def extract_header_info(page_text: str) -> dict[str, str]:
    district = ""
    school = ""
    class_no = "3"
    section = ""

    lines = [clean_text(line) for line in page_text.splitlines() if clean_text(line)]

    for line in lines[:12]:
        if "/" not in line:
            continue
        if "MUDURLUGU" not in fold_for_match(line):
            continue

        district_part, school_part = line.split("/", 1)
        district = clean_text(district_part)
        school_part = clean_text(school_part)
        school = re.split(
            r"\s+M[üu]d[üu]rl[üu][ğg][üu]\b",
            school_part,
            maxsplit=1,
            flags=re.IGNORECASE,
        )[0].strip()
        break

    section_pattern = re.compile(
        r"(?P<class>\d+)\.\s*S[ıiİI]n[ıiİI]f\s*/\s*"
        r"(?P<section>.+?)\s*Şubesi",
        flags=re.IGNORECASE,
    )
    section_match = section_pattern.search(page_text)
    if section_match:
        class_no = clean_text(section_match.group("class")) or "3"
        section = clean_text(section_match.group("section"))
    else:
        for line in lines[:12]:
            folded = fold_for_match(line)
            fallback_match = re.search(
                r"\b(?P<class>\d+)\s+SINIF\s+(?P<section>[A-Z0-9]+)\s+SUBESI\b",
                folded,
            )
            if fallback_match:
                class_no = fallback_match.group("class") or "3"
                section = fallback_match.group("section")
                break

    return {
        "district": district,
        "school": school,
        "class_no": class_no,
        "section": section,
    }


def extract_course_name(page_text: str, pdf_name: str = "") -> str:
    """
    Detect the course from the title/header.

    Expected title example:
    2021-2022 DERS YILI TÜRKÇE DERSİ 3. Sınıf / A Şubesi
    """
    lines = [clean_text(line) for line in page_text.splitlines() if clean_text(line)]
    header_text = " ".join(lines[:15])
    folded_header = fold_for_match(header_text)

    title_match = re.search(
        r"\bDERS YILI (?P<course>.+?) DERSI\b",
        folded_header,
    )
    if title_match:
        folded_title_course = clean_text(title_match.group("course"))
        for folded_course, course in COURSE_BY_FOLDED_NAME.items():
            if folded_course == folded_title_course:
                return course
        for folded_course, course in COURSE_BY_FOLDED_NAME.items():
            if folded_course in folded_title_course or folded_title_course in folded_course:
                return course

    for folded_course, course in sorted(
        COURSE_BY_FOLDED_NAME.items(),
        key=lambda item: len(item[0]),
        reverse=True,
    ):
        if folded_course in folded_header:
            return course

    # Last-resort fallback only; normal operation should use the PDF title.
    folded_file_name = fold_for_match(pdf_name)
    for folded_course, course in sorted(
        COURSE_BY_FOLDED_NAME.items(),
        key=lambda item: len(item[0]),
        reverse=True,
    ):
        if folded_course in folded_file_name:
            return course

    return ""


def find_table_columns(table: list[list[object]]) -> tuple[int, dict[str, int]] | None:
    """Find the header row and the important table columns."""
    for row_index, row in enumerate(table):
        folded_cells = [fold_for_match(cell) for cell in row]
        columns: dict[str, int] = {}

        for col_index, cell in enumerate(folded_cells):
            tokens = set(cell.split())

            if "OKUL" in tokens and "NO" in tokens:
                columns.setdefault("school_no", col_index)

            if "OGRENCININ" in tokens and "ADI" in tokens and "SOYADI" in tokens:
                columns.setdefault("name", col_index)

            if (
                "1" in tokens
                and "DONEM" in tokens
                and "DEGERLENDIRME" in tokens
                and "DURUMU" in tokens
            ):
                columns.setdefault("first_term_grade", col_index)

        if {"school_no", "name", "first_term_grade"}.issubset(columns):
            return row_index, columns

    return None


def parse_pdf_tables(tables: Iterable[list[list[object]]]) -> list[dict[str, str]]:
    student_rows: list[dict[str, str]] = []
    seen_rows: set[tuple[str, str]] = set()

    for table in tables:
        if not table:
            continue

        column_info = find_table_columns(table)
        if column_info is None:
            widest_row = max((len(row) for row in table), default=0)
            if widest_row < 4:
                continue
            header_row_index = 0
            columns = {"school_no": 1, "name": 2, "first_term_grade": 3}
        else:
            header_row_index, columns = column_info

        max_needed_index = max(columns.values())

        for row in table[header_row_index + 1 :]:
            if len(row) <= max_needed_index:
                continue

            school_no = clean_text(row[columns["school_no"]])
            name = clean_text(row[columns["name"]])
            grade = clean_text(row[columns["first_term_grade"]])

            if not school_no or not name:
                continue
            if not re.fullmatch(r"\d+", school_no):
                continue
            if "OGRENCI" in fold_for_match(name):
                continue

            row_key = (school_no, name)
            if row_key in seen_rows:
                continue
            seen_rows.add(row_key)

            student_rows.append(
                {
                    "school_no": school_no,
                    "name": name,
                    "grade": grade,
                }
            )

    return student_rows


def parse_text_rows(page_text: str) -> list[dict[str, str]]:
    """Fallback for pages where PDF table detection fails."""
    grade_pattern = "|".join(
        re.escape(grade) for grade in sorted(GRADE_VALUES, key=len, reverse=True)
    )
    row_pattern = re.compile(
        rf"^\s*\d+\s+"
        rf"(?P<school_no>\d+)\s+"
        rf"(?P<name>.+?)\s+"
        rf"(?P<grade>{grade_pattern})"
        rf"(?:\s+(?:{grade_pattern}))?\s*$"
    )

    rows: list[dict[str, str]] = []
    for raw_line in page_text.splitlines():
        line = clean_text(raw_line)
        match = row_pattern.match(line)
        if not match:
            continue
        rows.append(
            {
                "school_no": clean_text(match.group("school_no")),
                "name": clean_text(match.group("name")),
                "grade": clean_text(match.group("grade")),
            }
        )
    return rows


def extract_student_rows(page, page_text: str) -> list[dict[str, str]]:
    try:
        tables = page.extract_tables(table_settings=TABLE_SETTINGS)
    except Exception:
        tables = []

    rows = parse_pdf_tables(tables)
    if rows:
        return rows

    try:
        rows = parse_pdf_tables(page.extract_tables())
    except Exception:
        rows = []
    if rows:
        return rows

    return parse_text_rows(page_text)


def empty_student_record(header: dict[str, str], school_no: str, name: str) -> dict[str, str]:
    record = {column: "" for column in OUTPUT_COLUMNS}
    record.update(
        {
            "İlçe": header["district"],
            "Okul": header["school"],
            "Sınıf": header["class_no"] or "3",
            "Şube": header["section"],
            "Okul No": school_no,
            "Ad Soyad": name,
        }
    )
    return record


def add_page_to_group(
    groups: dict[GroupKey, dict[str, object]],
    header: dict[str, str],
    course: str,
    page_rows: list[dict[str, str]],
) -> None:
    key = GroupKey(
        district=header["district"],
        school=header["school"],
        class_no=header["class_no"] or "3",
        section=header["section"],
    )

    if key not in groups:
        groups[key] = {"students": {}, "order": []}

    students = groups[key]["students"]
    student_order = groups[key]["order"]
    assert isinstance(students, dict)
    assert isinstance(student_order, list)

    for row in page_rows:
        school_no = row["school_no"]

        if school_no not in students:
            students[school_no] = empty_student_record(header, school_no, row["name"])
            student_order.append(school_no)
        else:
            student = students[school_no]
            assert isinstance(student, dict)
            for column, value in (
                ("İlçe", header["district"]),
                ("Okul", header["school"]),
                ("Sınıf", header["class_no"] or "3"),
                ("Şube", header["section"]),
                ("Ad Soyad", row["name"]),
            ):
                if value and not student.get(column):
                    student[column] = value

        student = students[school_no]
        assert isinstance(student, dict)
        student[course] = row["grade"]


def extract_pdf_into_groups(pdf_path: Path, groups: dict[GroupKey, dict[str, object]]) -> int:
    rows_added = 0
    pdf_course = ""

    with pdfplumber.open(str(pdf_path)) as pdf:
        if len(pdf.pages) < 1 or len(pdf.pages) > 2:
            raise ValueError(
                f"Expected 1 or 2 pages, but found {len(pdf.pages)} pages "
                f"in {pdf_path.name}."
            )

        for page_index, page in enumerate(pdf.pages):
            page_number = page_index + 1
            page_text = page.extract_text() or ""
            header = extract_header_info(page_text)

            page_course = extract_course_name(page_text, pdf_path.name)
            if page_course:
                pdf_course = page_course
            elif pdf_course:
                page_course = pdf_course

            if not page_course:
                raise ValueError(
                    f"Could not detect the course name from page {page_number} "
                    f"of {pdf_path.name}."
                )

            page_rows = extract_student_rows(page, page_text)
            if not page_rows:
                raise ValueError(
                    f"No student rows could be extracted from page {page_number} "
                    f"of {pdf_path.name} ({page_course})."
                )

            add_page_to_group(groups, header, page_course, page_rows)
            rows_added += len(page_rows)

    return rows_added


def rows_for_group(group_data: dict[str, object]) -> list[dict[str, str]]:
    students = group_data["students"]
    student_order = group_data["order"]
    assert isinstance(students, dict)
    assert isinstance(student_order, list)
    return [students[school_no] for school_no in student_order]


def save_excel(rows: list[dict[str, str]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Notlar"

    worksheet.append(OUTPUT_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column, "") for column in OUTPUT_COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for row_number in range(2, worksheet.max_row + 1):
        worksheet[f"E{row_number}"].number_format = "@"

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        if column_letter == "F":
            width = min(max(max_length + 2, 18), 36)
        elif column_letter in {"A", "B"}:
            width = min(max(max_length + 2, 14), 34)
        elif column_letter in {"G", "H", "I", "J", "K", "L", "M", "N"}:
            width = min(max(max_length + 2, 14), 28)
        else:
            width = min(max(max_length + 2, 10), 18)
        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 42
    workbook.save(output_path)


def safe_file_part(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" ._")
    return value or "bilinmeyen"


def build_group_output_path(
    key: GroupKey,
    output_dir: Path,
    output_file: str | None,
    single_group: bool,
) -> Path:
    if output_file:
        if not single_group:
            raise ValueError(
                "--output-file can only be used when the input creates one Excel file."
            )
        output_path = Path(output_file).expanduser()
        if not output_path.suffix:
            output_path = output_path.with_suffix(".xlsx")
        if not output_path.is_absolute():
            output_path = output_dir / output_path
        return output_path

    file_name = (
        f"{safe_file_part(key.district)}_"
        f"{safe_file_part(key.school)}_"
        f"{safe_file_part(key.class_no)}"
        f"{safe_file_part(key.section)}_notlar.xlsx"
    )
    return output_dir / file_name


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Extract first-term 3rd grade course evaluations from separate "
            "one-course PDF files."
        )
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=str(DEFAULT_INPUT_PATH),
        help="Folder containing course PDFs, one PDF path, or a PDF path without .pdf.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help=(
            "Folder where Excel files will be saved. If omitted, folder input "
            "saves Excel files in that same folder; one-PDF input saves beside "
            "the PDF."
        ),
    )
    parser.add_argument(
        "--output-file",
        default=None,
        help="Optional Excel file name or full output path. Use only for one output workbook.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    pdf_paths = resolve_input_pdfs(args.input_path)
    output_dir = (
        Path(args.output_dir).expanduser()
        if args.output_dir
        else default_output_dir(args.input_path, pdf_paths)
    )

    groups: dict[GroupKey, dict[str, object]] = {}
    failures: list[tuple[Path, Exception]] = []

    print(f"PDF files found: {len(pdf_paths)}")
    print(f"Excel output folder: {output_dir}")

    for index, pdf_path in enumerate(pdf_paths, start=1):
        try:
            rows_added = extract_pdf_into_groups(pdf_path, groups)
            print(f"[{index}/{len(pdf_paths)}] OK: {pdf_path.name} ({rows_added} row entries)")
        except Exception as error:
            failures.append((pdf_path, error))
            print(f"[{index}/{len(pdf_paths)}] ERROR: {pdf_path.name}: {error}", file=sys.stderr)

    if failures:
        print("\nFiles with errors:", file=sys.stderr)
        for pdf_path, error in failures:
            print(f"  - {pdf_path}: {error}", file=sys.stderr)
        return 1

    if not groups:
        print("No student data was extracted.", file=sys.stderr)
        return 1

    single_group = len(groups) == 1
    for key, group_data in sorted(groups.items(), key=lambda item: item[0]):
        rows = rows_for_group(group_data)
        output_path = build_group_output_path(key, output_dir, args.output_file, single_group)
        save_excel(rows, output_path)
        print(
            "Excel saved: "
            f"{output_path} ({len(rows)} students, "
            f"{key.district} / {key.school} / {key.class_no}-{key.section})"
        )

    print("All course PDFs processed successfully.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
