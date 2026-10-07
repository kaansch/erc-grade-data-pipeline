# -*- coding: utf-8 -*-
r"""
Convert the new one-page selectable-text PDF table format into Excel.

Use this script for the newer table structure where course grades are columns
on a single PDF page. The script can process either one PDF file or every PDF
file directly inside a folder.

Install requirements:

    pip install pdfplumber openpyxl

Run with a path:

    python extract_pdf_grades_type2_selectable.py "C:\path\to\file_or_folder"

Or run without a path and the script will ask you to enter one:

    python extract_pdf_grades_type2_selectable.py
"""

from __future__ import annotations

import argparse
from collections import Counter
import re
import sys
import unicodedata
from pathlib import Path
from typing import Iterable

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    missing_name = getattr(exc, "name", "required package")
    raise SystemExit(
        f"Missing Python package: {missing_name}\n"
        "Install the requirements with:\n\n"
        "    pip install pdfplumber openpyxl\n"
    ) from exc


OUTPUT_COURSES = [
    "TÜRKÇE",
    "MATEMATİK",
    "HAYAT BİLGİSİ",
    "GÖRSEL SANATLAR",
    "YABANCI DİL",
    "MÜZİK",
    "FEN BİLİMLERİ",
    "BEDEN EĞİTİMİ VE OYUN",
]

PDF_COURSES = [
    *OUTPUT_COURSES,
    "SERBEST ETKİNLİKLER",
]

OUTPUT_COLUMNS = [
    "İlçe",
    "Okul",
    "Dönem",
    "Sınıf",
    "Şube",
    "Okul No",
    "Ad Soyad",
    *OUTPUT_COURSES,
]

SUMMARY_KEYWORDS = [
    "TOPLAM ÖĞRENCİ SAYISI",
    "BAŞARILI ÖĞRENCİ SAYISI",
    "BAŞARISIZ ÖĞRENCİ SAYISI",
    "BAŞARI YÜZDESİ",
    "NOT ORTALAMASI",
    "ZAYIFI OLAN",
]

TABLE_SETTINGS = {
    "vertical_strategy": "lines",
    "horizontal_strategy": "lines",
    "snap_tolerance": 3,
    "join_tolerance": 3,
    "intersection_tolerance": 3,
    "text_x_tolerance": 2,
    "text_y_tolerance": 3,
}

TURKISH_TO_ASCII = str.maketrans(
    {
        "ç": "c",
        "Ç": "C",
        "ğ": "g",
        "Ğ": "G",
        "ı": "i",
        "I": "I",
        "İ": "I",
        "ö": "o",
        "Ö": "O",
        "ş": "s",
        "Ş": "S",
        "ü": "u",
        "Ü": "U",
    }
)


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def fold_for_match(value: object) -> str:
    text = clean_text(value).translate(TURKISH_TO_ASCII)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def normalize_school_no(value: object) -> str:
    text = clean_text(value)
    match = re.search(r"-?\d+", text)
    return match.group(0) if match else ""


def is_grade_token(value: str) -> bool:
    return bool(re.fullmatch(r"[0-9]+(?:[,.][0-9]+)?", clean_text(value)))


def resolve_input_pdfs(raw_path: str | Path) -> list[Path]:
    path = Path(raw_path).expanduser()

    if path.is_dir():
        pdf_files = sorted(item for item in path.glob("*.pdf") if item.is_file())
        if not pdf_files:
            raise FileNotFoundError(f"No PDF files were found in folder: {path}")
        return pdf_files

    if path.is_file():
        if path.suffix.lower() != ".pdf":
            raise ValueError(f"The input file is not a PDF: {path}")
        return [path]

    if path.suffix.lower() != ".pdf":
        pdf_version = path.with_suffix(".pdf")
        if pdf_version.is_file():
            return [pdf_version]

    raise FileNotFoundError(
        f"PDF file or folder was not found: {path}\n"
        "Tip: pass one PDF file path or a folder containing PDF files."
    )


def extract_header_info(page_text: str) -> dict[str, str]:
    lines = [clean_text(line) for line in page_text.splitlines() if clean_text(line)]

    district = ""
    school = ""
    period = ""
    class_no = ""
    section = ""

    for line in lines[:20]:
        folded = fold_for_match(line)
        if "/" in line and ("MUDURLUGU" in folded or "MUDURLUG" in folded):
            district_part, school_part = line.split("/", 1)
            district = clean_text(district_part)
            school_part = clean_text(school_part)
            school = re.split(
                r"\s+M[üu]d[üu]rl[üu][ğg][üu]\b",
                school_part,
                maxsplit=1,
                flags=re.IGNORECASE,
            )[0].strip()
            if school == school_part:
                folded_school = fold_for_match(school_part)
                school = re.sub(r"\s*MUDURLUGU?.*$", "", folded_school).title()
            break

    for line in lines[:25]:
        folded = fold_for_match(line)
        if "DONEM" in folded:
            period = line
            break

    joined_header = " ".join(lines[:25])
    folded_header = fold_for_match(joined_header)
    class_match = re.search(
        r"\b(?P<class>\d+)\s*SINIF\s*(?:/)?\s*(?P<section>[A-Z0-9]+)\s*SUBESI\b",
        folded_header,
    )
    if class_match:
        class_no = class_match.group("class")
        section = class_match.group("section")
    else:
        raw_class_match = re.search(
            r"(?P<class>\d+)\.\s*S[ıiİI]n[ıiİI]f\s*/\s*"
            r"(?P<section>.+?)\s*Şubesi",
            joined_header,
            flags=re.IGNORECASE,
        )
        if raw_class_match:
            class_no = clean_text(raw_class_match.group("class"))
            section = clean_text(raw_class_match.group("section"))

    return {
        "district": district,
        "school": school,
        "period": period,
        "class_no": class_no,
        "section": section,
    }


def table_has_summary_text(row: Iterable[object]) -> bool:
    row_folded = fold_for_match(" ".join(clean_text(cell) for cell in row))
    return any(fold_for_match(keyword) in row_folded for keyword in SUMMARY_KEYWORDS)


def line_has_summary_text(line: str) -> bool:
    line_folded = fold_for_match(line)
    return any(fold_for_match(keyword) in line_folded for keyword in SUMMARY_KEYWORDS)


def row_is_empty(row: Iterable[object]) -> bool:
    return not any(clean_text(cell) for cell in row)


def has_letter(value: str) -> bool:
    return any(char.isalpha() for char in value)


def build_output_row(
    header: dict[str, str],
    school_no: str,
    name: str,
    grades: list[str],
) -> dict[str, str]:
    output_row = {
        "İlçe": header["district"],
        "Okul": header["school"],
        "Dönem": header["period"],
        "Sınıf": header["class_no"],
        "Şube": header["section"],
        "Okul No": school_no,
        "Ad Soyad": clean_text(name).upper(),
    }

    grades_to_use = grades[: len(OUTPUT_COURSES)]
    for index, course in enumerate(OUTPUT_COURSES):
        output_row[course] = grades_to_use[index] if index < len(grades_to_use) else ""

    return output_row


def parse_student_line(line: str, header: dict[str, str]) -> dict[str, str] | None:
    line = clean_text(line)
    if not line or line_has_summary_text(line):
        return None

    match = re.match(
        r"^(?P<serial>\d+)\s+(?P<school_no>-?\d+)\s+(?P<rest>.+?)\s*$",
        line,
    )
    if not match:
        return None

    school_no = normalize_school_no(match.group("school_no"))
    rest_tokens = clean_text(match.group("rest")).split()
    if not school_no or not rest_tokens:
        return None

    grades_reversed: list[str] = []
    while rest_tokens and is_grade_token(rest_tokens[-1]) and len(grades_reversed) < len(PDF_COURSES):
        grades_reversed.append(rest_tokens.pop())

    name = clean_text(" ".join(rest_tokens))
    if not name:
        return None

    folded_name = fold_for_match(name)
    if "OGRENCI" in folded_name or "AD SOYAD" in folded_name:
        return None

    grades = list(reversed(grades_reversed))
    return build_output_row(header, school_no, name, grades)


def parse_student_rows_from_text(page_text: str, header: dict[str, str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    for raw_line in page_text.splitlines():
        row = parse_student_line(raw_line, header)
        if row is not None:
            rows.append(row)

    return rows


def nonempty_cells(row: list[object]) -> list[tuple[int, str]]:
    return [(index, clean_text(cell)) for index, cell in enumerate(row) if clean_text(cell)]


def detect_student_identity_columns(table: list[list[object]]) -> tuple[int, int] | None:
    column_pairs: Counter[tuple[int, int]] = Counter()

    for row in table:
        if row_is_empty(row) or table_has_summary_text(row):
            continue

        cells = nonempty_cells(row)
        if len(cells) < 3:
            continue

        first_index, first_value = cells[0]
        if first_index > 1 or not re.fullmatch(r"\d+", first_value):
            continue

        for position in range(1, len(cells) - 1):
            school_index, school_value = cells[position]
            name_index, name_value = cells[position + 1]

            if not re.fullmatch(r"-?\d+", school_value):
                continue
            if is_grade_token(name_value) or not has_letter(name_value):
                continue

            folded_name = fold_for_match(name_value)
            if "OGRENCI" in folded_name or "AD SOYAD" in folded_name:
                continue

            column_pairs[(school_index, name_index)] += 1
            break

    if not column_pairs:
        return None

    return column_pairs.most_common(1)[0][0]


def detect_grade_columns(
    table: list[list[object]],
    school_no_column: int,
    name_column: int,
) -> list[int]:
    grade_counts: Counter[int] = Counter()

    for row in table:
        if row_is_empty(row) or table_has_summary_text(row):
            continue

        school_no = normalize_school_no(value_from_row(row, school_no_column))
        name = clean_text(value_from_row(row, name_column))
        if not school_no or not name or not has_letter(name):
            continue

        for column_index in range(name_column + 1, len(row)):
            value = value_from_row(row, column_index)
            if is_grade_token(value):
                grade_counts[column_index] += 1

    return sorted(grade_counts)


def find_header_row_and_columns(
    table: list[list[object]],
) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(table):
        folded_cells = [fold_for_match(cell) for cell in row]
        columns: dict[str, int] = {}

        for col_index, folded in enumerate(folded_cells):
            tokens = set(folded.split())

            if "OKUL" in tokens and "NO" in tokens:
                columns.setdefault("school_no", col_index)

            if (
                "AD" in tokens
                and "SOYAD" in tokens
                or "SOYADI" in tokens
                or "OGRENCI" in folded
            ):
                columns.setdefault("name", col_index)

            for course in PDF_COURSES:
                folded_course = fold_for_match(course)
                if folded_course in folded:
                    columns.setdefault(course, col_index)

        course_columns_found = sum(1 for course in PDF_COURSES if course in columns)
        if "school_no" in columns and "name" in columns and course_columns_found >= 5:
            return row_index, columns

    return None


def fallback_columns_for_row(row: list[object]) -> dict[str, int] | None:
    nonempty_count = sum(1 for cell in row if clean_text(cell))
    if len(row) >= 12 or nonempty_count >= 12:
        start = 3
        columns = {"school_no": 1, "name": 2}
    elif len(row) >= 11 or nonempty_count >= 11:
        start = 2
        columns = {"school_no": 0, "name": 1}
    else:
        return None

    for offset, course in enumerate(PDF_COURSES):
        columns[course] = start + offset
    return columns


def value_from_row(row: list[object], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return clean_text(row[index])


def parse_student_rows_from_table(
    table: list[list[object]],
    header: dict[str, str],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    identity_columns = detect_student_identity_columns(table)
    if identity_columns is not None:
        school_no_column, name_column = identity_columns
        grade_columns = detect_grade_columns(table, school_no_column, name_column)

        if grade_columns:
            for row in table:
                if row_is_empty(row) or table_has_summary_text(row):
                    continue

                school_no = normalize_school_no(value_from_row(row, school_no_column))
                name = clean_text(value_from_row(row, name_column)).upper()

                if not school_no or not name or not has_letter(name):
                    continue
                if "OGRENCI" in fold_for_match(name) or "AD SOYAD" in fold_for_match(name):
                    continue

                grades = [value_from_row(row, column_index) for column_index in grade_columns]
                rows.append(build_output_row(header, school_no, name, grades))

            if rows:
                return rows

    header_info = find_header_row_and_columns(table)
    if header_info is None:
        header_row_index = 0
        columns = None
    else:
        header_row_index, columns = header_info

    for row in table[header_row_index + 1 :]:
        if row_is_empty(row):
            continue
        if table_has_summary_text(row):
            continue

        compact_row_text = " ".join(clean_text(cell) for cell in row if clean_text(cell))
        compact_row = parse_student_line(compact_row_text, header)
        if compact_row is not None:
            rows.append(compact_row)
            continue

        active_columns = columns or fallback_columns_for_row(row)
        if active_columns is None:
            continue

        school_no = normalize_school_no(value_from_row(row, active_columns.get("school_no")))
        name = clean_text(value_from_row(row, active_columns.get("name"))).upper()

        if not school_no or not name:
            continue
        if "OGRENCI" in fold_for_match(name) or "AD SOYAD" in fold_for_match(name):
            continue

        grades = [value_from_row(row, active_columns.get(course)) for course in PDF_COURSES]
        rows.append(build_output_row(header, school_no, name, grades))

    return rows


def extract_tables(page) -> list[list[list[object]]]:
    try:
        tables = page.extract_tables(table_settings=TABLE_SETTINGS)
    except Exception:
        tables = []

    if tables:
        return tables

    try:
        return page.extract_tables()
    except Exception:
        return []


def extract_pdf(pdf_path: Path) -> list[dict[str, str]]:
    with pdfplumber.open(str(pdf_path)) as pdf:
        if len(pdf.pages) != 1:
            raise ValueError(
                f"Expected exactly 1 page, but found {len(pdf.pages)} pages in {pdf_path.name}."
            )

        page = pdf.pages[0]
        page_text = page.extract_text() or ""
        header = extract_header_info(page_text)

        tables = extract_tables(page)
        all_rows: list[dict[str, str]] = []
        if tables:
            for table in tables:
                all_rows.extend(parse_student_rows_from_table(table, header))

        if all_rows:
            return all_rows

        text_rows = parse_student_rows_from_text(page_text, header)
        if text_rows:
            return text_rows

        raise ValueError(f"No student rows could be extracted from {pdf_path.name}.")



def save_excel(rows: list[dict[str, str]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Notlar"

    worksheet.append(OUTPUT_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column, "") for column in OUTPUT_COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    okul_no_column = OUTPUT_COLUMNS.index("Okul No") + 1
    okul_no_letter = get_column_letter(okul_no_column)
    for row_number in range(2, worksheet.max_row + 1):
        worksheet[f"{okul_no_letter}{row_number}"].number_format = "@"

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        if column_letter == get_column_letter(OUTPUT_COLUMNS.index("Ad Soyad") + 1):
            width = min(max(max_length + 2, 18), 36)
        elif column_cells[0].value in {"İlçe", "Okul"}:
            width = min(max(max_length + 2, 14), 34)
        elif column_cells[0].value in OUTPUT_COURSES:
            width = min(max(max_length + 2, 12), 26)
        else:
            width = min(max(max_length + 2, 10), 18)
        worksheet.column_dimensions[column_letter].width = width

    worksheet.row_dimensions[1].height = 42
    workbook.save(output_path)


def build_output_path(pdf_path: Path, output_dir: Path | None) -> Path:
    folder = output_dir if output_dir is not None else pdf_path.parent
    return folder / f"{pdf_path.stem}.xlsx"


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert the new one-page selectable-text PDF table format to Excel."
    )
    parser.add_argument(
        "input_path",
        nargs="?",
        default=None,
        help="PDF file path or folder path. If omitted, the script asks for it.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Optional folder where Excel files will be saved.",
    )
    return parser.parse_args(argv)


def get_input_path(args: argparse.Namespace) -> str:
    if args.input_path:
        return args.input_path

    user_input = input("Enter a PDF file path or a folder path containing PDF files: ").strip()
    if not user_input:
        raise ValueError("No input path was entered.")
    return user_input.strip('"')


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    input_path = get_input_path(args)
    pdf_paths = resolve_input_pdfs(input_path)
    output_dir = Path(args.output_dir).expanduser() if args.output_dir else None

    created_files: list[Path] = []
    failures: list[tuple[Path, Exception]] = []

    print(f"PDF files found: {len(pdf_paths)}")

    for index, pdf_path in enumerate(pdf_paths, start=1):
        try:
            rows = extract_pdf(pdf_path)
            output_path = build_output_path(pdf_path, output_dir)
            save_excel(rows, output_path)
            created_files.append(output_path)
            print(f"[{index}/{len(pdf_paths)}] OK: {pdf_path.name} -> {output_path}")
        except Exception as error:
            failures.append((pdf_path, error))
            print(f"[{index}/{len(pdf_paths)}] ERROR: {pdf_path.name}: {error}", file=sys.stderr)

    print(f"Converted PDFs: {len(created_files)}")
    for output_path in created_files:
        print(f"Created: {output_path}")

    if failures:
        print("\nFiles with errors:", file=sys.stderr)
        for pdf_path, error in failures:
            print(f"  - {pdf_path}: {error}", file=sys.stderr)
        return 1

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
