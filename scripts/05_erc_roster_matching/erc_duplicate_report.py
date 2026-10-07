"""Create a two-sheet exact-duplicate report from the ERC full list.

The source file is read into memory and is never opened for writing. The report
is always saved under a new timestamped filename beside the source file.

Required source columns:
    v1, v2, v3

Matching rules:
    Section Constraint:         normalized v1 + normalized v2 + normalized v3
    Relaxed Section Constraint: normalized v1 + normalized v3

The student_id column, when present, is removed before matching and is not
included in either output sheet.
"""

from __future__ import annotations

import hashlib
import os
import re
import sys
import tempfile
import unicodedata
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Iterable

try:
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit(
        "Missing dependency. Install pandas and openpyxl first:\n"
        "    python -m pip install pandas openpyxl"
    ) from exc


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".csv", ".tsv"}
REQUIRED_FIELDS = ("v1", "v2", "v3")
STUDENT_ID_FIELD = "student_id"


def canonical_header(value: object) -> str:
    """Return a case-insensitive header name with separators standardized."""
    text = unicodedata.normalize("NFKC", str(value)).strip().casefold()
    return re.sub(r"[\s-]+", "_", text)


def normalize_value(value: object) -> str | None:
    """Normalize formatting only; do not perform any fuzzy transformation."""
    if pd.isna(value):
        return None

    if isinstance(value, datetime):
        text = value.isoformat(sep=" ")
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)

    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"\s+", " ", text).strip().casefold()
    return text or None


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def prompt_for_source_path() -> Path:
    """Ask for and validate the full path of the ERC full-list file."""
    while True:
        raw_path = input("Enter the full path of the ERC full list: ").strip()
        raw_path = raw_path.strip('"').strip("'")
        candidate = Path(raw_path)

        if not candidate.is_absolute():
            print("Please enter an absolute/full path.")
            continue
        if not candidate.is_file():
            print("That file does not exist. Please try again.")
            continue
        if candidate.suffix.casefold() not in SUPPORTED_EXTENSIONS:
            supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            print(f"Unsupported input type. Supported types: {supported}")
            continue

        return candidate.resolve(strict=True)


def locate_required_columns(columns: Iterable[object]) -> dict[str, object]:
    """Locate v1/v2/v3 without changing the source column labels."""
    columns = list(columns)
    located: dict[str, object] = {}

    for required in REQUIRED_FIELDS:
        matches = [col for col in columns if canonical_header(col) == required]
        if len(matches) != 1:
            available = ", ".join(map(str, columns))
            if not matches:
                raise ValueError(
                    f"Required column '{required}' was not found. "
                    f"Available columns: {available}"
                )
            raise ValueError(
                f"More than one column resolves to '{required}'. "
                "Please make the source headers unique."
            )
        located[required] = matches[0]

    return located


def choose_excel_sheet(excel_file: pd.ExcelFile) -> str:
    """Automatically choose the only sheet containing v1, v2, and v3."""
    candidates: list[str] = []

    for sheet_name in excel_file.sheet_names:
        headers = pd.read_excel(excel_file, sheet_name=sheet_name, nrows=0).columns
        canonical = {canonical_header(col) for col in headers}
        if set(REQUIRED_FIELDS).issubset(canonical):
            candidates.append(sheet_name)

    if not candidates:
        raise ValueError(
            "No worksheet contains all three required columns: v1, v2, and v3."
        )
    if len(candidates) == 1:
        return candidates[0]

    print("More than one worksheet contains v1, v2, and v3:")
    for number, sheet_name in enumerate(candidates, start=1):
        print(f"  {number}. {sheet_name}")

    while True:
        response = input("Select the worksheet number: ").strip()
        if response.isdigit() and 1 <= int(response) <= len(candidates):
            return candidates[int(response) - 1]
        print("Please enter one of the listed worksheet numbers.")


def read_source_from_memory(source_path: Path, source_bytes: bytes) -> pd.DataFrame:
    """Parse the source bytes; the source path is never passed to a writer."""
    suffix = source_path.suffix.casefold()

    if suffix in {".xlsx", ".xlsm"}:
        with pd.ExcelFile(BytesIO(source_bytes), engine="openpyxl") as excel_file:
            sheet_name = choose_excel_sheet(excel_file)
            print(f"Using worksheet: {sheet_name}")
            return pd.read_excel(
                excel_file,
                sheet_name=sheet_name,
                dtype=object,
            )

    separator = "\t" if suffix == ".tsv" else ","
    last_error: UnicodeDecodeError | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            text = source_bytes.decode(encoding)
            return pd.read_csv(BytesIO(text.encode("utf-8")), sep=separator, dtype=object)
        except UnicodeDecodeError as exc:
            last_error = exc

    raise ValueError("The text input encoding could not be detected.") from last_error


def unused_column_name(columns: Iterable[object], preferred: str) -> str:
    existing = {canonical_header(col) for col in columns}
    candidate = preferred
    counter = 2
    while canonical_header(candidate) in existing:
        candidate = f"{preferred}_{counter}"
        counter += 1
    return candidate


def prepare_working_data(raw_data: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, object], str]:
    """Remove student_id and add a source-row audit column in memory only."""
    student_columns = [
        col for col in raw_data.columns if canonical_header(col) == STUDENT_ID_FIELD
    ]
    data = raw_data.drop(columns=student_columns).copy()

    required_columns = locate_required_columns(data.columns)
    source_row_column = unused_column_name(data.columns, "Source_Row")
    data.insert(0, source_row_column, range(2, len(data) + 2))

    return data, required_columns, source_row_column


def build_duplicate_report(
    data: pd.DataFrame,
    required_columns: dict[str, object],
    source_row_column: str,
    matching_fields: tuple[str, ...],
    group_prefix: str,
) -> pd.DataFrame:
    """Return every row belonging to a normalized exact-match group of size 2+."""
    working = data.copy()
    helper_columns: list[str] = []

    for position, field in enumerate(matching_fields, start=1):
        helper = f"__erc_normalized_key_{position}"
        while helper in working.columns:
            helper = f"_{helper}"
        working[helper] = working[required_columns[field]].map(normalize_value)
        helper_columns.append(helper)

    # A row with a blank required value cannot form a duplicate group.
    eligible = working.loc[working[helper_columns].notna().all(axis=1)].copy()

    duplicate_count_column = unused_column_name(data.columns, "Duplicate_Count")
    duplicate_group_column = unused_column_name(
        [*data.columns, duplicate_count_column], "Duplicate_Group_ID"
    )

    counts = (
        eligible.groupby(helper_columns, sort=True, dropna=False)
        .size()
        .reset_index(name=duplicate_count_column)
    )
    counts = counts.loc[counts[duplicate_count_column] >= 2].copy()
    counts = counts.sort_values(helper_columns, kind="stable").reset_index(drop=True)
    counts[duplicate_group_column] = [
        f"{group_prefix}-{number:04d}" for number in range(1, len(counts) + 1)
    ]

    report = eligible.merge(
        counts,
        on=helper_columns,
        how="inner",
        validate="many_to_one",
    )
    report = report.sort_values(
        [duplicate_group_column, source_row_column], kind="stable"
    )

    original_columns = list(data.columns)
    output_columns = [
        duplicate_group_column,
        duplicate_count_column,
        *original_columns,
    ]
    report = report.loc[:, output_columns].reset_index(drop=True)

    # student_id must never reach either report sheet.
    assert all(canonical_header(col) != STUDENT_ID_FIELD for col in report.columns)
    return report


def make_output_path(source_path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = f"{source_path.stem}_duplicate_report_{timestamp}"
    output_path = source_path.with_name(f"{base_name}.xlsx")
    counter = 2

    while output_path.exists():
        output_path = source_path.with_name(f"{base_name}_{counter}.xlsx")
        counter += 1

    if os.path.normcase(str(output_path.resolve())) == os.path.normcase(
        str(source_path.resolve())
    ):
        raise RuntimeError("Safety check failed: output path equals input path.")

    return output_path


def excel_safe_copy(frame: pd.DataFrame) -> pd.DataFrame:
    """Prevent text beginning with '=' from becoming a formula in the report."""
    safe = frame.copy()
    for column in safe.columns:
        safe[column] = safe[column].map(
            lambda value: f"'{value}"
            if isinstance(value, str) and value.startswith("=")
            else value
        )
    return safe


def format_worksheet(worksheet) -> None:
    """Apply compact, readable report formatting to a newly created sheet."""
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 26

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sample_end = min(worksheet.max_row, 501)
    for column_number in range(1, worksheet.max_column + 1):
        letter = get_column_letter(column_number)
        lengths = []
        for row_number in range(1, sample_end + 1):
            value = worksheet.cell(row=row_number, column=column_number).value
            lengths.append(len(str(value)) if value is not None else 0)
        worksheet.column_dimensions[letter].width = min(max(max(lengths, default=0) + 2, 12), 45)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_new_report(
    output_path: Path,
    section_report: pd.DataFrame,
    relaxed_report: pd.DataFrame,
) -> None:
    """Write only to a new temporary/output file; never to the source path."""
    temporary_file = tempfile.NamedTemporaryFile(
        mode="wb",
        suffix=".xlsx",
        prefix=".erc_duplicate_report_",
        dir=output_path.parent,
        delete=False,
    )
    temporary_path = Path(temporary_file.name)
    temporary_file.close()

    try:
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            excel_safe_copy(section_report).to_excel(
                writer,
                sheet_name="Section Constraint",
                index=False,
            )
            excel_safe_copy(relaxed_report).to_excel(
                writer,
                sheet_name="Relaxed Section Constraint",
                index=False,
            )

            format_worksheet(writer.book["Section Constraint"])
            format_worksheet(writer.book["Relaxed Section Constraint"])

        if output_path.exists():
            raise FileExistsError(
                f"Refusing to overwrite an existing report: {output_path}"
            )
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def count_groups(report: pd.DataFrame) -> int:
    if report.empty:
        return 0
    group_column = next(
        col for col in report.columns if canonical_header(col).startswith("duplicate_group_id")
    )
    return int(report[group_column].nunique())


def main() -> None:
    source_path = prompt_for_source_path()

    # The only source-file operation is a binary read.
    source_stat_before = source_path.stat()
    source_bytes = source_path.read_bytes()
    source_hash_before = sha256_bytes(source_bytes)

    raw_data = read_source_from_memory(source_path, source_bytes)
    data, required_columns, source_row_column = prepare_working_data(raw_data)

    section_report = build_duplicate_report(
        data=data,
        required_columns=required_columns,
        source_row_column=source_row_column,
        matching_fields=("v1", "v2", "v3"),
        group_prefix="SC",
    )
    relaxed_report = build_duplicate_report(
        data=data,
        required_columns=required_columns,
        source_row_column=source_row_column,
        matching_fields=("v1", "v3"),
        group_prefix="RC",
    )

    output_path = make_output_path(source_path)
    write_new_report(output_path, section_report, relaxed_report)

    # Verify byte-for-byte that the source remains unchanged.
    source_stat_after = source_path.stat()
    source_hash_after = sha256_bytes(source_path.read_bytes())
    source_unchanged = (
        source_hash_before == source_hash_after
        and source_stat_before.st_size == source_stat_after.st_size
        and source_stat_before.st_mtime_ns == source_stat_after.st_mtime_ns
    )
    if not source_unchanged:
        raise RuntimeError(
            "The source file changed while the report was running. "
            "The script did not write to it, but the report should be rerun."
        )

    print("\nDuplicate report created successfully.")
    print("Source verification: unchanged (SHA-256, size, and timestamp match).")
    print(
        "Section Constraint: "
        f"{count_groups(section_report)} duplicate groups, "
        f"{len(section_report)} matching rows"
    )
    print(
        "Relaxed Section Constraint: "
        f"{count_groups(relaxed_report)} duplicate groups, "
        f"{len(relaxed_report)} matching rows"
    )
    print(f"Report saved to: {output_path}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nCancelled. No report was written.")
        raise SystemExit(130)
    except Exception as exc:
        print(f"\nError: {exc}", file=sys.stderr)
        raise SystemExit(1)
