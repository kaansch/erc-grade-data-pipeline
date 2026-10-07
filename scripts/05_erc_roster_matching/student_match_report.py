#!/usr/bin/env python3
"""Create a five-sheet ERC student matching report without editing either input.

The workbook to be checked must contain at least three worksheets:

* Sheet 1: district=A, school=B, grade=C, section=D, student name=F.
* Sheets 2 and 3: district=A, school=B, grade=C, section=D,
  student name=E.

The ERC full-list workbook's first worksheet must have headers named ``v1``,
``v2``, and ``v3``. ``v1`` is ``district / school``, ``v2`` contains a
grade-prefixed section such as ``4-A``, and ``v3`` is the student name.

Matching rules:

* Direct, order-sensitive ``SequenceMatcher`` is used for every comparison.
* District and school must each score at least the location threshold.
* Every ERC student in the resolved district-school is scored across all
  sections. A student candidate qualifies only when its score is strictly
  greater than the student threshold.
* One qualifying candidate is green when its section matches and yellow when
  its section differs. Two or more candidates are orange and are detailed on
  the fifth sheet. No candidates, or a failed location, is red and is detailed
  on the fourth sheet.

Both inputs are opened with ``read_only=True``. Only a separately named output
workbook is written.

Interactive use (the ERC full list is requested first):

    python student_match_report.py

Paths may also be supplied on the command line:

    python student_match_report.py erc_full_list.xlsx source.xlsx report.xlsx
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import tempfile
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet


SUPPORTED_INPUT_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
REPORT_SHEET_COUNT = 3

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Aptos", size=11, color="000000")

GREEN_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
GREEN_FONT = Font(name="Aptos", size=11, color="006100")
YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFEB9C")
YELLOW_FONT = Font(name="Aptos", size=11, color="9C6500")
ORANGE_FILL = PatternFill(fill_type="solid", fgColor="F4B183")
ORANGE_FONT = Font(name="Aptos", size=11, color="9C5700")
RED_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
RED_FONT = Font(name="Aptos", size=11, color="9C0006")

ADVISORY_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
ADVISORY_FONT = Font(name="Aptos", size=11, color="7F6000")
EMPTY_MESSAGE_FILL = PatternFill(fill_type="solid", fgColor="E2F0D9")
EMPTY_MESSAGE_FONT = Font(name="Aptos", size=11, bold=True, color="375623")

OUTPUT_HEADERS = (
    "District",
    "School Name",
    "Section",
    "Student Name",
    "Matched ERC Student Name",
    "Name Similarity",
    "ERC Section",
    "ERC School",
    "ERC District",
)

MISMATCH_HEADERS = (
    "Source Sheet",
    "Source District",
    "Source School",
    "Source Section",
    "Source Student Name",
    "ERC District",
    "District Similarity",
    "ERC School",
    "School Similarity",
    "Best Candidate in Source Section",
    "Source-Section Similarity",
    "Source-Section ERC Section",
    "Best Candidate Across All Sections",
    "All-Sections Similarity",
    "All-Sections ERC Section",
    "Reason Not Matched",
)

MULTIPLE_HEADERS = (
    "Source Sheet",
    "Source District",
    "Source School",
    "Source Section",
    "Source Student Name",
    "Matched ERC Student Name",
    "Name Similarity",
    "ERC Section",
    "ERC School",
    "ERC District",
)

TURKISH_CASE_TRANSLATION = str.maketrans(
    {
        "I": "ı",
        "İ": "i",
        "Ş": "ş",
        "Ğ": "ğ",
        "Ü": "ü",
        "Ö": "ö",
        "Ç": "ç",
    }
)

PUNCTUATION_TRANSLATION = str.maketrans(
    {
        "’": "'",
        "‘": "'",
        "`": "'",
        "´": "'",
        "‐": "-",
        "‑": "-",
        "‒": "-",
        "–": "-",
        "—": "-",
        "−": "-",
    }
)

# Repair common Greek/Cyrillic homoglyphs that sometimes appear inside names
# copied from visually Latin text. For example, Greek ``ΜΑΗΜ`` becomes
# Latin ``MAHM`` before non-Latin characters are filtered.
CONFUSABLE_TRANSLATION = str.maketrans(
    {
        # Greek
        "Α": "A",
        "α": "a",
        "Β": "B",
        "β": "b",
        "Ε": "E",
        "ε": "e",
        "Ζ": "Z",
        "ζ": "z",
        "Η": "H",
        "η": "h",
        "Ι": "I",
        "ι": "i",
        "Κ": "K",
        "κ": "k",
        "Μ": "M",
        "μ": "m",
        "Ν": "N",
        "ν": "n",
        "Ο": "O",
        "ο": "o",
        "Ρ": "P",
        "ρ": "p",
        "Τ": "T",
        "τ": "t",
        "Υ": "Y",
        "υ": "y",
        "Χ": "X",
        "χ": "x",
        # Cyrillic
        "А": "A",
        "а": "a",
        "В": "B",
        "в": "b",
        "Е": "E",
        "е": "e",
        "К": "K",
        "к": "k",
        "М": "M",
        "м": "m",
        "Н": "H",
        "н": "h",
        "О": "O",
        "о": "o",
        "Р": "P",
        "р": "p",
        "С": "C",
        "с": "c",
        "Т": "T",
        "т": "t",
        "У": "Y",
        "у": "y",
        "Х": "X",
        "х": "x",
        "І": "I",
        "і": "i",
        "Ј": "J",
        "ј": "j",
    }
)


@dataclass(frozen=True)
class BaseRecord:
    row_number: int
    district: str
    school: str
    section: str
    student: str
    display_district: str
    display_school: str
    display_section: str
    display_student: str


@dataclass(frozen=True)
class BaseLocation:
    district: str
    school: str
    display_district: str
    display_school: str
    first_row_number: int


@dataclass(frozen=True)
class LocationResolution:
    accepted: bool
    location: BaseLocation | None
    district_score: float
    school_score: float
    reason: str = ""


@dataclass(frozen=True)
class ScoredCandidate:
    record: BaseRecord
    score: float


@dataclass(frozen=True)
class MatchOutcome:
    status: str
    location: BaseLocation | None
    district_score: float
    school_score: float
    single_match: ScoredCandidate | None = None
    qualifying_candidates: tuple[ScoredCandidate, ...] = ()
    best_source_section: ScoredCandidate | None = None
    best_all_sections: ScoredCandidate | None = None
    reason: str = ""


@dataclass(frozen=True)
class SourceAuditRecord:
    source_sheet: str
    district: str
    school: str
    section: str
    student: str
    outcome: MatchOutcome


@dataclass
class BaseLoadStats:
    usable_rows: int = 0
    skipped_blank_rows: int = 0
    skipped_invalid_rows: int = 0


@dataclass
class SheetStats:
    output_rows: int = 0
    green: int = 0
    yellow: int = 0
    orange: int = 0
    red_location: int = 0
    red_name: int = 0
    incomplete_source_fields: int = 0

    @property
    def red(self) -> int:
        return self.red_location + self.red_name


def cell_text(value: object) -> str:
    """Convert an Excel value to clean display text without changing its meaning."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_for_match(value: object) -> str:
    """Normalize case, Turkish letters, homoglyphs, punctuation, and whitespace."""
    text = unicodedata.normalize("NFKC", cell_text(value))
    text = text.translate(PUNCTUATION_TRANSLATION)
    text = text.translate(CONFUSABLE_TRANSLATION)

    # Python lowercasing is not Turkish-locale aware. Apply the special Turkish
    # uppercase mappings first, then fold accents and diacritics.
    text = text.translate(TURKISH_CASE_TRANSLATION).lower().replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))

    # Punctuation differences are separators; word order remains unchanged.
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def section_display(value: object) -> str:
    """Return only the section portion, removing prefixes such as ``4-``."""
    text = cell_text(value).translate(PUNCTUATION_TRANSLATION).strip()
    match = re.match(r"^\s*\d+\s*[-/]\s*(.*?)\s*$", text)
    return match.group(1).strip() if match else text


def normalize_section(value: object) -> str:
    """Normalize a source or ERC section while ignoring grade-like prefixes."""
    normalized = normalize_for_match(section_display(value))
    normalized = re.sub(r"\s+(?:sube|subesi)$", "", normalized).strip()
    return normalized


def split_base_location(value: object) -> tuple[str, str] | None:
    """Split ERC ``v1`` at the first slash into district and school."""
    text = cell_text(value)
    parts = re.split(r"\s*/\s*", text, maxsplit=1)
    if len(parts) != 2 or not parts[0].strip() or not parts[1].strip():
        return None
    return parts[0].strip(), parts[1].strip()


def string_similarity(left: str, right: str) -> float:
    """Return the direct, order-sensitive SequenceMatcher score from 0 to 100."""
    if not left or not right:
        return 0.0
    if left == right:
        return 100.0
    return (
        SequenceMatcher(None, left, right, autojunk=False).ratio() * 100.0
    )


class BaseMatcher:
    """Resolve one ERC district-school and classify every source student."""

    def __init__(
        self,
        records: Iterable[BaseRecord],
        *,
        student_threshold: float,
        location_threshold: float,
    ) -> None:
        self.student_threshold = student_threshold
        self.location_threshold = location_threshold

        records_by_location: dict[
            tuple[str, str], list[BaseRecord]
        ] = defaultdict(list)
        location_by_key: dict[tuple[str, str], BaseLocation] = {}

        for record in records:
            key = (record.district, record.school)
            records_by_location[key].append(record)
            if key not in location_by_key:
                location_by_key[key] = BaseLocation(
                    district=record.district,
                    school=record.school,
                    display_district=record.display_district,
                    display_school=record.display_school,
                    first_row_number=record.row_number,
                )

        self.records_by_location = {
            key: tuple(sorted(location_records, key=lambda item: item.row_number))
            for key, location_records in records_by_location.items()
        }
        self.locations = tuple(
            sorted(
                location_by_key.values(),
                key=lambda item: item.first_row_number,
            )
        )

    @staticmethod
    def _location_rank(
        source_district: str,
        source_school: str,
        location: BaseLocation,
        district_score: float,
        school_score: float,
    ) -> tuple[float, ...]:
        exact_pair = float(
            source_district == location.district
            and source_school == location.school
        )
        return (
            exact_pair,
            min(district_score, school_score),
            (district_score + school_score) / 2.0,
            school_score,
            district_score,
            -float(location.first_row_number),
        )

    @lru_cache(maxsize=None)
    def resolve_location(
        self,
        district: str,
        school: str,
    ) -> LocationResolution:
        """Resolve an actual ERC district-school pair using separate 95% gates."""
        best_any: tuple[
            tuple[float, ...], BaseLocation, float, float
        ] | None = None
        best_eligible: tuple[
            tuple[float, ...], BaseLocation, float, float
        ] | None = None

        for location in self.locations:
            district_score = string_similarity(district, location.district)
            school_score = string_similarity(school, location.school)
            rank = self._location_rank(
                district,
                school,
                location,
                district_score,
                school_score,
            )
            candidate = (rank, location, district_score, school_score)

            if best_any is None or rank > best_any[0]:
                best_any = candidate

            if (
                district
                and school
                and district_score >= self.location_threshold
                and school_score >= self.location_threshold
                and (best_eligible is None or rank > best_eligible[0])
            ):
                best_eligible = candidate

        if best_eligible is not None:
            _, location, district_score, school_score = best_eligible
            return LocationResolution(
                accepted=True,
                location=location,
                district_score=district_score,
                school_score=school_score,
            )

        if best_any is None:
            return LocationResolution(
                accepted=False,
                location=None,
                district_score=0.0,
                school_score=0.0,
                reason="The ERC list contains no usable district-school locations.",
            )

        _, location, district_score, school_score = best_any
        if not district or not school:
            missing = [
                label
                for label, value in (("district", district), ("school", school))
                if not value
            ]
            reason = f"Missing source field(s): {', '.join(missing)}."
        else:
            failures = []
            if district_score < self.location_threshold:
                failures.append(
                    "district similarity "
                    f"{district_score:.2f}% is below {self.location_threshold:.2f}%"
                )
            if school_score < self.location_threshold:
                failures.append(
                    "school similarity "
                    f"{school_score:.2f}% is below {self.location_threshold:.2f}%"
                )
            reason = "Location not accepted: " + "; ".join(failures) + "."

        return LocationResolution(
            accepted=False,
            location=location,
            district_score=district_score,
            school_score=school_score,
            reason=reason,
        )

    @lru_cache(maxsize=None)
    def classify(
        self,
        district: str,
        school: str,
        section: str,
        student: str,
    ) -> MatchOutcome:
        """Classify one source student using the final four-color algorithm."""
        location_result = self.resolve_location(district, school)
        location = location_result.location

        if not location_result.accepted or location is None:
            return MatchOutcome(
                status="red_location",
                location=location,
                district_score=location_result.district_score,
                school_score=location_result.school_score,
                reason=location_result.reason,
            )

        location_key = (location.district, location.school)
        records = self.records_by_location.get(location_key, ())
        scored = [
            ScoredCandidate(
                record=record,
                score=string_similarity(student, record.student),
            )
            for record in records
        ]
        scored.sort(
            key=lambda item: (
                -item.score,
                0 if item.record.section == section else 1,
                item.record.row_number,
            )
        )
        scored_tuple = tuple(scored)
        qualifying = tuple(
            item for item in scored_tuple if item.score > self.student_threshold
        )

        if len(qualifying) == 1:
            single = qualifying[0]
            status = (
                "green" if single.record.section == section else "yellow"
            )
            return MatchOutcome(
                status=status,
                location=location,
                district_score=location_result.district_score,
                school_score=location_result.school_score,
                single_match=single,
                qualifying_candidates=qualifying,
            )

        if len(qualifying) >= 2:
            return MatchOutcome(
                status="orange",
                location=location,
                district_score=location_result.district_score,
                school_score=location_result.school_score,
                qualifying_candidates=qualifying,
            )

        best_all = scored_tuple[0] if scored_tuple else None
        best_source_section = next(
            (
                item
                for item in scored_tuple
                if section and item.record.section == section
            ),
            None,
        )

        if not student:
            reason = "The source student name is empty after normalization."
        elif best_all is None:
            reason = (
                "The resolved ERC district-school contains no usable students."
            )
        else:
            reason = (
                "No ERC student in the resolved district-school scored higher "
                f"than {self.student_threshold:.2f}%. The best score across all "
                f"sections was {best_all.score:.2f}%."
            )

        return MatchOutcome(
            status="red_name",
            location=location,
            district_score=location_result.district_score,
            school_score=location_result.school_score,
            best_source_section=best_source_section,
            best_all_sections=best_all,
            reason=reason,
        )


def normalized_header(value: object) -> str:
    return normalize_for_match(value).replace(" ", "")


def locate_base_columns(
    worksheet: Worksheet,
    header_row: int,
) -> tuple[int, int, int]:
    """Find zero-based ``v1``/``v2``/``v3`` indexes."""
    header_values = next(
        worksheet.iter_rows(
            min_row=header_row,
            max_row=header_row,
            values_only=True,
        ),
        (),
    )
    indexes = {
        normalized_header(value): index
        for index, value in enumerate(header_values)
        if cell_text(value)
    }

    missing = [name for name in ("v1", "v2", "v3") if name not in indexes]
    if missing:
        raise ValueError(
            f"ERC sheet '{worksheet.title}' is missing header(s) "
            f"{', '.join(missing)} on row {header_row}."
        )
    return indexes["v1"], indexes["v2"], indexes["v3"]


def row_value(row: Sequence[object], index: int) -> object:
    return row[index] if index < len(row) else None


def read_base_records(
    base_path: Path,
    *,
    header_row: int,
) -> tuple[list[BaseRecord], BaseLoadStats, str]:
    """Read ERC records from the first worksheet without modifying the file."""
    workbook = load_workbook(
        filename=base_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        if not workbook.worksheets:
            raise ValueError("The ERC workbook has no worksheets.")

        worksheet = workbook.worksheets[0]
        v1_index, v2_index, v3_index = locate_base_columns(
            worksheet,
            header_row,
        )
        stats = BaseLoadStats()
        records: list[BaseRecord] = []

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=header_row + 1,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            v1 = row_value(row, v1_index)
            v2 = row_value(row, v2_index)
            v3 = row_value(row, v3_index)

            if not any((cell_text(v1), cell_text(v2), cell_text(v3))):
                stats.skipped_blank_rows += 1
                continue

            location_parts = split_base_location(v1)
            display_section = section_display(v2)
            display_student = cell_text(v3)
            normalized_section = normalize_section(v2)
            normalized_student = normalize_for_match(display_student)

            if (
                location_parts is None
                or not normalized_section
                or not normalized_student
            ):
                stats.skipped_invalid_rows += 1
                continue

            display_district, display_school = location_parts
            normalized_district = normalize_for_match(display_district)
            normalized_school = normalize_for_match(display_school)
            if not normalized_district or not normalized_school:
                stats.skipped_invalid_rows += 1
                continue

            records.append(
                BaseRecord(
                    row_number=row_number,
                    district=normalized_district,
                    school=normalized_school,
                    section=normalized_section,
                    student=normalized_student,
                    display_district=display_district,
                    display_school=display_school,
                    display_section=display_section,
                    display_student=display_student,
                )
            )
            stats.usable_rows += 1

        if not records:
            raise ValueError("No usable records were found in the ERC workbook.")
        return records, stats, worksheet.title
    finally:
        workbook.close()


def style_header_row(worksheet: Worksheet) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 30
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def style_report_sheet(worksheet: Worksheet) -> None:
    style_header_row(worksheet)
    widths = {
        "A": 24,
        "B": 46,
        "C": 12,
        "D": 32,
        "E": 32,
        "F": 17,
        "G": 13,
        "H": 46,
        "I": 24,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def style_student_result(cell: Cell, status: str) -> None:
    if status == "green":
        cell.fill = GREEN_FILL
        cell.font = GREEN_FONT
    elif status == "yellow":
        cell.fill = YELLOW_FILL
        cell.font = YELLOW_FONT
    elif status == "orange":
        cell.fill = ORANGE_FILL
        cell.font = ORANGE_FONT
    else:
        cell.fill = RED_FILL
        cell.font = RED_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")


def write_report_sheet(
    source_sheet: Worksheet,
    report_sheet: Worksheet,
    *,
    source_sheet_index: int,
    source_header_row: int,
    matcher: BaseMatcher,
) -> tuple[
    SheetStats,
    list[SourceAuditRecord],
    list[SourceAuditRecord],
]:
    """Write one of the three independently matched source report sheets."""
    report_sheet.append(OUTPUT_HEADERS)
    style_report_sheet(report_sheet)

    student_column_index = 5 if source_sheet_index == 0 else 4
    context = ["", "", "", ""]  # district, school, grade, section
    stats = SheetStats()
    mismatches: list[SourceAuditRecord] = []
    multiple_matches: list[SourceAuditRecord] = []

    for row in source_sheet.iter_rows(
        min_row=source_header_row + 1,
        max_col=6,
        values_only=True,
    ):
        values = list(row)
        if len(values) < 6:
            values.extend([None] * (6 - len(values)))

        # Forward-fill the class context for merged/grouped source layouts.
        for index in range(4):
            current = cell_text(values[index])
            if current:
                context[index] = current

        student_name = cell_text(values[student_column_index])
        if not student_name:
            continue

        district, school, _grade, raw_section = context
        output_section = section_display(raw_section)
        normalized_district = normalize_for_match(district)
        normalized_school = normalize_for_match(school)
        normalized_section = normalize_section(raw_section)
        normalized_student = normalize_for_match(student_name)

        outcome = matcher.classify(
            normalized_district,
            normalized_school,
            normalized_section,
            normalized_student,
        )

        match = outcome.single_match
        if match is not None:
            record = match.record
            output_values = (
                district,
                school,
                output_section,
                student_name,
                record.display_student,
                match.score / 100.0,
                record.display_section,
                record.display_school,
                record.display_district,
            )
        else:
            output_values = (
                district,
                school,
                output_section,
                student_name,
                "",
                None,
                "",
                "",
                "",
            )

        report_sheet.append(output_values)
        output_row = report_sheet.max_row
        report_sheet.row_dimensions[output_row].height = 22
        for cell in report_sheet[output_row]:
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")

        style_student_result(report_sheet.cell(output_row, 4), outcome.status)
        similarity_cell = report_sheet.cell(output_row, 6)
        similarity_cell.number_format = "0.00%"
        similarity_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        audit_record = SourceAuditRecord(
            source_sheet=source_sheet.title,
            district=district,
            school=school,
            section=output_section,
            student=student_name,
            outcome=outcome,
        )

        stats.output_rows += 1
        if outcome.status == "green":
            stats.green += 1
        elif outcome.status == "yellow":
            stats.yellow += 1
        elif outcome.status == "orange":
            stats.orange += 1
            multiple_matches.append(audit_record)
        elif outcome.status == "red_location":
            stats.red_location += 1
            mismatches.append(audit_record)
        else:
            stats.red_name += 1
            mismatches.append(audit_record)

        if (
            not normalized_district
            or not normalized_school
            or not normalized_section
        ):
            stats.incomplete_source_fields += 1

    report_sheet.auto_filter.ref = f"A1:I{report_sheet.max_row}"
    return stats, mismatches, multiple_matches


def write_empty_message(
    worksheet: Worksheet,
    *,
    last_column: str,
    message: str,
) -> None:
    worksheet.merge_cells(f"A2:{last_column}2")
    cell = worksheet["A2"]
    cell.value = message
    cell.fill = EMPTY_MESSAGE_FILL
    cell.font = EMPTY_MESSAGE_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[2].height = 24


def write_mismatch_sheet(
    worksheet: Worksheet,
    mismatches: Sequence[SourceAuditRecord],
) -> None:
    """Write location failures and zero-candidate students to Sheet 4."""
    worksheet.append(MISMATCH_HEADERS)
    style_header_row(worksheet)

    widths = {
        "A": 26,
        "B": 23,
        "C": 42,
        "D": 14,
        "E": 31,
        "F": 23,
        "G": 17,
        "H": 42,
        "I": 17,
        "J": 31,
        "K": 20,
        "L": 18,
        "M": 31,
        "N": 20,
        "O": 18,
        "P": 58,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    if not mismatches:
        write_empty_message(
            worksheet,
            last_column="P",
            message="No mismatched students were found.",
        )
        return

    for mismatch in mismatches:
        outcome = mismatch.outcome
        location = outcome.location
        same = outcome.best_source_section
        overall = outcome.best_all_sections

        worksheet.append(
            (
                mismatch.source_sheet,
                mismatch.district,
                mismatch.school,
                mismatch.section,
                mismatch.student,
                location.display_district if location else "",
                outcome.district_score / 100.0 if location else None,
                location.display_school if location else "",
                outcome.school_score / 100.0 if location else None,
                same.record.display_student if same else "",
                same.score / 100.0 if same else None,
                same.record.display_section if same else "",
                overall.record.display_student if overall else "",
                overall.score / 100.0 if overall else None,
                overall.record.display_section if overall else "",
                outcome.reason,
            )
        )
        output_row = worksheet.max_row
        for cell in worksheet[output_row]:
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")

        source_name_cell = worksheet.cell(output_row, 5)
        source_name_cell.fill = RED_FILL
        source_name_cell.font = RED_FONT

        for candidate_column in (10, 13):
            candidate_cell = worksheet.cell(output_row, candidate_column)
            if candidate_cell.value:
                candidate_cell.fill = ADVISORY_FILL
                candidate_cell.font = ADVISORY_FONT

        for score_column in (7, 9, 11, 14):
            score_cell = worksheet.cell(output_row, score_column)
            score_cell.number_format = "0.00%"
            score_cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        reason_cell = worksheet.cell(output_row, 16)
        reason_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )
        worksheet.row_dimensions[output_row].height = 34

    worksheet.auto_filter.ref = f"A1:P{worksheet.max_row}"


def write_multiple_candidates_sheet(
    worksheet: Worksheet,
    multiple_matches: Sequence[SourceAuditRecord],
) -> None:
    """Write every qualifying candidate for each orange source student."""
    worksheet.append(MULTIPLE_HEADERS)
    style_header_row(worksheet)

    widths = {
        "A": 26,
        "B": 23,
        "C": 42,
        "D": 14,
        "E": 31,
        "F": 31,
        "G": 17,
        "H": 14,
        "I": 42,
        "J": 23,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    if not multiple_matches:
        write_empty_message(
            worksheet,
            last_column="J",
            message="No students with multiple qualifying candidates were found.",
        )
        return

    for source_record in multiple_matches:
        for candidate in source_record.outcome.qualifying_candidates:
            record = candidate.record
            worksheet.append(
                (
                    source_record.source_sheet,
                    source_record.district,
                    source_record.school,
                    source_record.section,
                    source_record.student,
                    record.display_student,
                    candidate.score / 100.0,
                    record.display_section,
                    record.display_school,
                    record.display_district,
                )
            )
            output_row = worksheet.max_row
            for cell in worksheet[output_row]:
                cell.font = BODY_FONT
                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                )

            source_name_cell = worksheet.cell(output_row, 5)
            source_name_cell.fill = ORANGE_FILL
            source_name_cell.font = ORANGE_FONT

            candidate_name_cell = worksheet.cell(output_row, 6)
            candidate_name_cell.fill = ADVISORY_FILL
            candidate_name_cell.font = ADVISORY_FONT

            score_cell = worksheet.cell(output_row, 7)
            score_cell.number_format = "0.00%"
            score_cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            worksheet.row_dimensions[output_row].height = 22

    worksheet.auto_filter.ref = f"A1:J{worksheet.max_row}"


def validate_threshold(value: float, option_name: str) -> None:
    if not 0.0 <= value <= 100.0:
        raise ValueError(f"{option_name} must be between 0 and 100.")


def paths_refer_to_same_file(first: Path, second: Path) -> bool:
    """Protect inputs even when an existing output is a link to an input."""
    if first == second:
        return True
    try:
        return first.exists() and second.exists() and first.samefile(second)
    except OSError:
        return False


def resolved_input_path(value: str, label: str) -> Path:
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"{label} does not exist or is not a file: {path}")
    if path.suffix.lower() not in SUPPORTED_INPUT_EXTENSIONS:
        raise ValueError(
            f"{label} must be an .xlsx/.xlsm-compatible file, "
            f"not '{path.suffix}'."
        )
    return path


def resolved_output_path(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if path.suffix.lower() != ".xlsx":
        raise ValueError("The output workbook must use the .xlsx extension.")
    return path


def clean_user_path(value: str) -> str:
    """Trim whitespace and paired quotes from typed or dropped paths."""
    cleaned = value.strip()
    if (
        len(cleaned) >= 2
        and cleaned[0] == cleaned[-1]
        and cleaned[0] in {'"', "'"}
    ):
        cleaned = cleaned[1:-1].strip()
    return cleaned


def request_input_workbook(
    supplied_value: str | None,
    *,
    label: str,
    prompt: str,
) -> Path:
    """Resolve a CLI path, or repeatedly prompt until a valid path is entered."""
    if supplied_value is not None:
        return resolved_input_path(clean_user_path(supplied_value), label)

    while True:
        try:
            entered_value = input(prompt)
        except EOFError as error:
            raise ValueError(f"No path was supplied for {label.lower()}.") from error

        try:
            return resolved_input_path(clean_user_path(entered_value), label)
        except (FileNotFoundError, ValueError) as error:
            print(f"Error: {error}", file=sys.stderr)
            print("Please enter the full path again.", file=sys.stderr)


def default_output_path(source_path: Path) -> Path:
    """Choose a unique report filename beside the workbook being checked."""
    base_name = f"{source_path.stem}_student_match_report"
    candidate = source_path.with_name(f"{base_name}.xlsx")
    counter = 2
    while candidate.exists():
        candidate = source_path.with_name(f"{base_name}_{counter}.xlsx")
        counter += 1
    return candidate.resolve()


def save_workbook_atomically(workbook: Workbook, output_path: Path) -> None:
    """Save through a temporary file so failures leave no partial report."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}_",
            suffix=".xlsx",
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
        temporary_path = None
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def build_report(
    source_path: Path,
    base_path: Path,
    output_path: Path,
    *,
    student_threshold: float,
    location_threshold: float,
    source_header_row: int,
    base_header_row: int,
    overwrite: bool,
) -> tuple[BaseLoadStats, str, list[tuple[str, SheetStats]]]:
    """Build and save the report, returning compact processing statistics."""
    validate_threshold(student_threshold, "--student-threshold")
    validate_threshold(location_threshold, "--location-threshold")
    if source_header_row < 1 or base_header_row < 1:
        raise ValueError("Header row numbers must be 1 or greater.")

    if paths_refer_to_same_file(output_path, source_path) or paths_refer_to_same_file(
        output_path,
        base_path,
    ):
        raise ValueError("The output path must differ from both input paths.")
    if output_path.exists() and not overwrite:
        raise FileExistsError(
            f"Output already exists: {output_path}. "
            "Use --overwrite to replace it."
        )

    base_records, base_stats, base_sheet_name = read_base_records(
        base_path,
        header_row=base_header_row,
    )
    matcher = BaseMatcher(
        base_records,
        student_threshold=student_threshold,
        location_threshold=location_threshold,
    )

    source_workbook = load_workbook(
        filename=source_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    report_workbook = Workbook()
    report_workbook.remove(report_workbook.active)
    report_workbook.properties.title = "ERC Student Match Report"
    report_workbook.properties.subject = (
        "Source students independently checked against the ERC full list"
    )

    sheet_results: list[tuple[str, SheetStats]] = []
    all_mismatches: list[SourceAuditRecord] = []
    all_multiple_matches: list[SourceAuditRecord] = []

    try:
        if len(source_workbook.worksheets) < REPORT_SHEET_COUNT:
            raise ValueError(
                f"The source workbook has {len(source_workbook.worksheets)} "
                f"sheet(s); at least {REPORT_SHEET_COUNT} are required."
            )

        for sheet_index, source_sheet in enumerate(
            source_workbook.worksheets[:REPORT_SHEET_COUNT]
        ):
            report_sheet = report_workbook.create_sheet(
                title=source_sheet.title
            )
            stats, mismatches, multiple_matches = write_report_sheet(
                source_sheet,
                report_sheet,
                source_sheet_index=sheet_index,
                source_header_row=source_header_row,
                matcher=matcher,
            )
            sheet_results.append((source_sheet.title, stats))
            all_mismatches.extend(mismatches)
            all_multiple_matches.extend(multiple_matches)

        mismatch_sheet = report_workbook.create_sheet(
            title="Mismatched Students"
        )
        write_mismatch_sheet(mismatch_sheet, all_mismatches)

        multiple_sheet = report_workbook.create_sheet(
            title="Multiple Candidates"
        )
        write_multiple_candidates_sheet(
            multiple_sheet,
            all_multiple_matches,
        )

        save_workbook_atomically(report_workbook, output_path)
    finally:
        source_workbook.close()
        report_workbook.close()

    return base_stats, base_sheet_name, sheet_results


def parse_arguments(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create a five-sheet ERC student report using direct "
            "SequenceMatcher comparisons while opening both inputs read-only."
        ),
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "base_workbook",
        nargs="?",
        help="ERC full-list workbook containing v1/v2/v3; prompted when omitted",
    )
    parser.add_argument(
        "source_workbook",
        nargs="?",
        help="Three-sheet workbook to check; prompted when omitted",
    )
    parser.add_argument(
        "output_workbook",
        nargs="?",
        help=(
            "Optional new .xlsx report path; when omitted, a unique report "
            "name is created beside the workbook being checked"
        ),
    )
    parser.add_argument(
        "--student-threshold",
        type=float,
        default=90.0,
        help=(
            "Exclusive student-name threshold; a candidate must score higher "
            "than this value"
        ),
    )
    parser.add_argument(
        "--location-threshold",
        type=float,
        default=95.0,
        help=(
            "Inclusive threshold required separately for district and school"
        ),
    )
    parser.add_argument(
        "--source-header-row",
        type=int,
        default=1,
        help="Header row in every source sheet",
    )
    parser.add_argument(
        "--base-header-row",
        type=int,
        default=1,
        help="Header row in the ERC sheet",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow replacement of an existing output report (never an input)",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_arguments(argv)
    try:
        base_path = request_input_workbook(
            args.base_workbook,
            label="ERC full-list workbook",
            prompt="Enter the full path of the base Excel file (full list): ",
        )
        source_path = request_input_workbook(
            args.source_workbook,
            label="Workbook to be checked",
            prompt="Enter the full path of the Excel file to be checked: ",
        )
        output_path = (
            resolved_output_path(clean_user_path(args.output_workbook))
            if args.output_workbook is not None
            else default_output_path(source_path)
        )

        if args.output_workbook is None:
            print(f"The report will be created at: {output_path}")

        base_stats, base_sheet_name, sheet_results = build_report(
            source_path,
            base_path,
            output_path,
            student_threshold=args.student_threshold,
            location_threshold=args.location_threshold,
            source_header_row=args.source_header_row,
            base_header_row=args.base_header_row,
            overwrite=args.overwrite,
        )
    except (FileNotFoundError, FileExistsError, PermissionError, ValueError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 2

    print(
        f"ERC sheet '{base_sheet_name}': {base_stats.usable_rows} usable row(s), "
        f"{base_stats.skipped_invalid_rows} invalid row(s) skipped, "
        f"{base_stats.skipped_blank_rows} blank row(s) skipped."
    )
    for sheet_name, stats in sheet_results:
        print(
            f"{sheet_name}: {stats.output_rows} student(s) | "
            f"green={stats.green} | yellow={stats.yellow} | "
            f"orange={stats.orange} | red={stats.red} "
            f"(location={stats.red_location}, name={stats.red_name}) | "
            f"incomplete source fields={stats.incomplete_source_fields}"
        )
    print(f"Report written to: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
