#!/usr/bin/env python3
"""Create a four-sheet student matching report without modifying either input.

The source workbook is expected to contain at least three worksheets:

* Sheet 1: district=A, school=B, grade=C, section=D, student name=F.
* Sheets 2 and 3: district=A, school=B, grade=C, section=D,
  student name=E.

The base workbook's first worksheet must have headers named ``v1``, ``v2``,
and ``v3``. ``v1`` is ``district / school``, ``v2`` is ``4-section`` (the
constant ``4-`` prefix is ignored), and ``v3`` is the student name.

Both inputs are opened by openpyxl with ``read_only=True``. Only the separately
named output workbook is written. The first three output sheets mirror the
source sheets, and a fourth ``Mismatched Students`` sheet reports every red row
with its closest base name and similarity score.

Interactive use (the base file is requested first):

    python overview2student_match_report.py

The same paths can optionally be supplied on the command line in that order:

    python overview2student_match_report.py base.xlsx source.xlsx report.xlsx

The default fuzzy thresholds can be changed if a stricter or looser report is
needed:

    python overview2student_match_report.py base.xlsx source.xlsx report.xlsx --student-threshold 90 --location-threshold 94
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import tempfile
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet


SUPPORTED_INPUT_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
REPORT_SHEET_COUNT = 3

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Aptos", size=11, color="000000")

MATCH_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
MATCH_FONT = Font(name="Aptos", size=11, color="006100")
NO_MATCH_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
NO_MATCH_FONT = Font(name="Aptos", size=11, color="9C0006")

OUTPUT_HEADERS = ("District", "School Name", "Section", "Student Name")
MISMATCH_HEADERS = (
    "Source Sheet",
    "District",
    "School Name",
    "Section",
    "Student Name",
    "Best Base Student Match",
    "Name Similarity",
    "Best Base District",
    "Best Base School",
    "Best Base Section",
    "Reason Not Matched",
)

BEST_MATCH_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
BEST_MATCH_FONT = Font(name="Aptos", size=11, color="7F6000")
NO_MISMATCH_FILL = PatternFill(fill_type="solid", fgColor="E2F0D9")

TURKISH_CASE_TRANSLATION = str.maketrans(
    {
        "I": "ı",
        "İ": "i",
        "Ş": "ş",
        "Ğ": "ğ",
        "Ü": "ü",
        "Ö": "ö",
        "Ç": "ç",
    }
)

PUNCTUATION_TRANSLATION = str.maketrans(
    {
        "’": "'",
        "‘": "'",
        "`": "'",
        "´": "'",
        "‐": "-",
        "‑": "-",
        "‒": "-",
        "–": "-",
        "—": "-",
        "−": "-",
    }
)


@dataclass(frozen=True)
class BaseRecord:
    district: str
    school: str
    section: str
    student: str
    display_district: str
    display_school: str
    display_section: str
    display_student: str


@dataclass(frozen=True)
class MatchOutcome:
    status: str
    best_record: BaseRecord | None
    student_score: float
    reason: str = ""


@dataclass(frozen=True)
class MismatchRecord:
    source_sheet: str
    district: str
    school: str
    section: str
    student: str
    outcome: MatchOutcome


@dataclass
class BaseLoadStats:
    usable_rows: int = 0
    skipped_blank_rows: int = 0
    skipped_invalid_rows: int = 0


@dataclass
class SheetStats:
    output_rows: int = 0
    exact_matches: int = 0
    fuzzy_matches: int = 0
    unmatched: int = 0
    incomplete_location_rows: int = 0


def cell_text(value: object) -> str:
    """Convert an Excel value to clean display text without changing its meaning."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_for_match(value: object) -> str:
    """Normalize casing, Turkish characters, punctuation, and whitespace.

    Turkish letters are folded to their Latin equivalents for matching, so
    ``Şule`` and ``Sule`` normalize identically. The original source text is
    still written unchanged to the report.
    """
    text = unicodedata.normalize("NFKC", cell_text(value))
    text = text.translate(PUNCTUATION_TRANSLATION)

    # Python's default lowercasing is not Turkish-locale aware. Apply the two
    # special uppercase mappings first, then fold accents/diacritics.
    text = text.translate(TURKISH_CASE_TRANSLATION).lower().replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))

    # Treat punctuation and repeated whitespace as separators, not differences.
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def section_display(value: object) -> str:
    """Return only the section portion, removing prefixes such as ``4-``."""
    text = cell_text(value).translate(PUNCTUATION_TRANSLATION).strip()
    match = re.match(r"^\s*\d+\s*[-/]\s*(.*?)\s*$", text)
    return match.group(1).strip() if match else text


def normalize_section(value: object) -> str:
    """Normalize a source or base section while ignoring grade-like prefixes."""
    normalized = normalize_for_match(section_display(value))
    # Accept labels such as "A Şubesi" as section "A".
    normalized = re.sub(r"\s+(?:sube|subesi)$", "", normalized).strip()
    return normalized


def split_base_location(value: object) -> tuple[str, str] | None:
    """Split base v1 at the first slash into district and school."""
    text = cell_text(value)
    parts = re.split(r"\s*/\s*", text, maxsplit=1)
    if len(parts) != 2 or not parts[0].strip() or not parts[1].strip():
        return None
    return parts[0].strip(), parts[1].strip()


def string_similarity(left: str, right: str) -> float:
    """Return a 0-100 fuzzy score, allowing token-order differences."""
    if not left or not right:
        return 0.0
    if left == right:
        return 100.0

    direct = SequenceMatcher(None, left, right, autojunk=False).ratio()
    left_tokens = " ".join(sorted(left.split()))
    right_tokens = " ".join(sorted(right.split()))
    token_sorted = SequenceMatcher(
        None, left_tokens, right_tokens, autojunk=False
    ).ratio()
    return max(direct, token_sorted) * 100.0


class BaseMatcher:
    """Match within a section and resolved district/school location."""

    def __init__(
        self,
        records: Iterable[BaseRecord],
        *,
        student_threshold: float,
        location_threshold: float,
    ) -> None:
        self.student_threshold = student_threshold
        self.location_threshold = location_threshold

        records = tuple(records)
        exact_records: dict[tuple[str, str, str, str], BaseRecord] = {}
        records_by_location: dict[
            tuple[str, str, str], set[BaseRecord]
        ] = defaultdict(set)
        records_by_section: dict[str, set[BaseRecord]] = defaultdict(set)
        locations_by_section: dict[str, set[tuple[str, str, str]]] = defaultdict(set)

        for record in records:
            location = (record.district, record.school, record.section)
            exact_records.setdefault((*location, record.student), record)
            records_by_location[location].add(record)
            records_by_section[record.section].add(record)
            locations_by_section[record.section].add(location)

        record_sort_key = lambda record: (
            record.district,
            record.school,
            record.section,
            record.student,
            record.display_student,
        )
        self.all_records = tuple(sorted(records, key=record_sort_key))
        self.exact_records = exact_records
        self.records_by_location = {
            key: tuple(sorted(location_records, key=record_sort_key))
            for key, location_records in records_by_location.items()
        }
        self.records_by_section = {
            section: tuple(sorted(section_records, key=record_sort_key))
            for section, section_records in records_by_section.items()
        }
        self.locations_by_section = {
            section: tuple(sorted(locations))
            for section, locations in locations_by_section.items()
        }

    @lru_cache(maxsize=None)
    def _resolve_location(
        self, district: str, school: str, section: str
    ) -> tuple[str, str, str] | None:
        """Find an exact location first, then a high-confidence fuzzy location."""
        requested = (district, school, section)
        if requested in self.records_by_location:
            return requested

        if not district or not school or not section:
            return None

        best_location: tuple[str, str, str] | None = None
        best_score = -1.0

        # Section must always be exact. District and school may be fuzzy, but
        # each must independently clear the location threshold.
        for candidate in self.locations_by_section.get(section, ()):
            candidate_district, candidate_school, _ = candidate
            district_score = string_similarity(district, candidate_district)
            if district_score < self.location_threshold:
                continue

            school_score = string_similarity(school, candidate_school)
            if school_score < self.location_threshold:
                continue

            combined_score = (district_score * 0.35) + (school_score * 0.65)
            if combined_score > best_score:
                best_score = combined_score
                best_location = candidate

        return best_location

    def _best_student_at_location(
        self,
        location: tuple[str, str, str],
        student: str,
    ) -> tuple[BaseRecord | None, float]:
        """Return the closest student name within one resolved class location."""
        best_record: BaseRecord | None = None
        best_score = -1.0
        for record in self.records_by_location.get(location, ()):
            score = string_similarity(student, record.student)
            if score > best_score:
                best_record = record
                best_score = score
        return best_record, max(best_score, 0.0)

    @lru_cache(maxsize=None)
    def _advisory_best_name(
        self,
        district: str,
        school: str,
        section: str,
        student: str,
    ) -> tuple[BaseRecord | None, float]:
        """Find a report-only suggestion without changing match acceptance.

        Candidates from the same section are preferred. If the base contains no
        such section, every base record is considered. Name similarity is the
        primary ranking and district/school similarity breaks ties.
        """
        if not student:
            return None, 0.0

        candidates = self.records_by_section.get(section) or self.all_records
        best_record: BaseRecord | None = None
        best_name_score = -1.0
        best_location_score = -1.0

        for record in candidates:
            name_score = string_similarity(student, record.student)
            district_score = string_similarity(district, record.district)
            school_score = string_similarity(school, record.school)
            location_score = (district_score * 0.35) + (school_score * 0.65)
            if (name_score, location_score) > (
                best_name_score,
                best_location_score,
            ):
                best_record = record
                best_name_score = name_score
                best_location_score = location_score

        return best_record, max(best_name_score, 0.0)

    @lru_cache(maxsize=None)
    def classify(
        self, district: str, school: str, section: str, student: str
    ) -> MatchOutcome:
        """Return match status plus auditable best-match details."""
        if not district or not school or not section or not student:
            missing = [
                label
                for label, value in (
                    ("district", district),
                    ("school", school),
                    ("section", section),
                    ("student name", student),
                )
                if not value
            ]
            best_record, best_score = self._advisory_best_name(
                district, school, section, student
            )
            return MatchOutcome(
                status="unmatched",
                best_record=best_record,
                student_score=best_score,
                reason=f"Missing source field(s): {', '.join(missing)}.",
            )

        full_key = (district, school, section, student)
        exact_record = self.exact_records.get(full_key)
        if exact_record is not None:
            return MatchOutcome(
                status="exact",
                best_record=exact_record,
                student_score=100.0,
            )

        location = self._resolve_location(district, school, section)
        if location is None:
            best_record, best_score = self._advisory_best_name(
                district, school, section, student
            )
            if section not in self.records_by_section:
                reason = (
                    "No base records exist for this section; the displayed best "
                    "name was searched across all base sections."
                )
            else:
                reason = (
                    "District and/or school did not meet the "
                    f"{self.location_threshold:.1f}% location threshold."
                )
            return MatchOutcome(
                status="unmatched",
                best_record=best_record,
                student_score=best_score,
                reason=reason,
            )

        best_record, best_student_score = self._best_student_at_location(
            location, student
        )
        if best_student_score >= self.student_threshold:
            return MatchOutcome(
                status="fuzzy",
                best_record=best_record,
                student_score=best_student_score,
            )

        return MatchOutcome(
            status="unmatched",
            best_record=best_record,
            student_score=best_student_score,
            reason=(
                f"Best name similarity ({best_student_score:.1f}%) is below the "
                f"{self.student_threshold:.1f}% student threshold."
            ),
        )


def normalized_header(value: object) -> str:
    return normalize_for_match(value).replace(" ", "")


def locate_base_columns(
    worksheet: Worksheet, header_row: int
) -> tuple[int, int, int]:
    """Find zero-based v1/v2/v3 indexes from the requested header row."""
    header_values = next(
        worksheet.iter_rows(
            min_row=header_row,
            max_row=header_row,
            values_only=True,
        ),
        (),
    )
    indexes = {
        normalized_header(value): index
        for index, value in enumerate(header_values)
        if cell_text(value)
    }

    missing = [name for name in ("v1", "v2", "v3") if name not in indexes]
    if missing:
        missing_text = ", ".join(missing)
        raise ValueError(
            f"Base sheet '{worksheet.title}' is missing header(s) {missing_text} "
            f"on row {header_row}."
        )
    return indexes["v1"], indexes["v2"], indexes["v3"]


def row_value(row: Sequence[object], index: int) -> object:
    return row[index] if index < len(row) else None


def read_base_records(
    base_path: Path, *, header_row: int
) -> tuple[list[BaseRecord], BaseLoadStats, str]:
    """Read normalized base records from the first worksheet."""
    workbook = load_workbook(
        filename=base_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        if not workbook.worksheets:
            raise ValueError("The base workbook has no worksheets.")

        worksheet = workbook.worksheets[0]
        v1_index, v2_index, v3_index = locate_base_columns(worksheet, header_row)
        stats = BaseLoadStats()
        records: list[BaseRecord] = []

        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            v1 = row_value(row, v1_index)
            v2 = row_value(row, v2_index)
            v3 = row_value(row, v3_index)

            if not any((cell_text(v1), cell_text(v2), cell_text(v3))):
                stats.skipped_blank_rows += 1
                continue

            location = split_base_location(v1)
            display_section = section_display(v2)
            display_student = cell_text(v3)
            section = normalize_section(v2)
            student = normalize_for_match(display_student)
            if location is None or not section or not student:
                stats.skipped_invalid_rows += 1
                continue

            district, school = location
            normalized_district = normalize_for_match(district)
            normalized_school = normalize_for_match(school)
            if not normalized_district or not normalized_school:
                stats.skipped_invalid_rows += 1
                continue

            records.append(
                BaseRecord(
                    district=normalized_district,
                    school=normalized_school,
                    section=section,
                    student=student,
                    display_district=district,
                    display_school=school,
                    display_section=display_section,
                    display_student=display_student,
                )
            )
            stats.usable_rows += 1

        if not records:
            raise ValueError("No usable records were found in the base workbook.")
        return records, stats, worksheet.title
    finally:
        workbook.close()


def style_report_sheet(worksheet: Worksheet) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 24

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 52
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 34


def style_student_result(cell: Cell, matched: bool) -> None:
    cell.fill = MATCH_FILL if matched else NO_MATCH_FILL
    cell.font = MATCH_FONT if matched else NO_MATCH_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")


def write_report_sheet(
    source_sheet: Worksheet,
    report_sheet: Worksheet,
    *,
    source_sheet_index: int,
    source_header_row: int,
    matcher: BaseMatcher,
) -> tuple[SheetStats, list[MismatchRecord]]:
    """Copy source rows and color student cells by the match result."""
    report_sheet.append(OUTPUT_HEADERS)
    style_report_sheet(report_sheet)

    student_column_index = 5 if source_sheet_index == 0 else 4  # F, then E/E
    context = ["", "", "", ""]  # district, school, grade, section
    stats = SheetStats()
    mismatches: list[MismatchRecord] = []

    for row in source_sheet.iter_rows(
        min_row=source_header_row + 1,
        max_col=6,
        values_only=True,
    ):
        values = list(row)
        if len(values) < 6:
            values.extend([None] * (6 - len(values)))

        # Forward-fill the first four columns. This supports source sheets where
        # district, school, grade, or section cells are merged or shown once per
        # group of student rows.
        for index in range(4):
            current = cell_text(values[index])
            if current:
                context[index] = current

        student_name = cell_text(values[student_column_index])
        if not student_name:
            continue

        district, school, _grade, raw_section = context
        output_section = section_display(raw_section)

        normalized_district = normalize_for_match(district)
        normalized_school = normalize_for_match(school)
        normalized_section = normalize_section(raw_section)
        normalized_student = normalize_for_match(student_name)

        result = matcher.classify(
            normalized_district,
            normalized_school,
            normalized_section,
            normalized_student,
        )

        report_sheet.append((district, school, output_section, student_name))
        output_row = report_sheet.max_row

        for cell in report_sheet[output_row]:
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")

        matched = result.status in {"exact", "fuzzy"}
        style_student_result(report_sheet.cell(output_row, 4), matched)

        stats.output_rows += 1
        if result.status == "exact":
            stats.exact_matches += 1
        elif result.status == "fuzzy":
            stats.fuzzy_matches += 1
        else:
            stats.unmatched += 1
            mismatches.append(
                MismatchRecord(
                    source_sheet=source_sheet.title,
                    district=district,
                    school=school,
                    section=output_section,
                    student=student_name,
                    outcome=result,
                )
            )

        if not normalized_district or not normalized_school or not normalized_section:
            stats.incomplete_location_rows += 1

    if report_sheet.max_row >= 1:
        report_sheet.auto_filter.ref = f"A1:D{report_sheet.max_row}"
    return stats, mismatches


def write_mismatch_sheet(
    worksheet: Worksheet,
    mismatches: Sequence[MismatchRecord],
) -> None:
    """Write the combined fourth-sheet audit of all below-threshold rows."""
    worksheet.append(MISMATCH_HEADERS)
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 30

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    widths = {
        "A": 28,
        "B": 24,
        "C": 48,
        "D": 12,
        "E": 32,
        "F": 32,
        "G": 18,
        "H": 24,
        "I": 48,
        "J": 18,
        "K": 58,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    if not mismatches:
        worksheet.merge_cells("A2:K2")
        message_cell = worksheet["A2"]
        message_cell.value = "No mismatched students were found."
        message_cell.fill = NO_MISMATCH_FILL
        message_cell.font = Font(name="Aptos", size=11, bold=True, color="375623")
        message_cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[2].height = 24
        return

    for mismatch in mismatches:
        best = mismatch.outcome.best_record
        worksheet.append(
            (
                mismatch.source_sheet,
                mismatch.district,
                mismatch.school,
                mismatch.section,
                mismatch.student,
                best.display_student if best else "",
                mismatch.outcome.student_score / 100.0 if best else None,
                best.display_district if best else "",
                best.display_school if best else "",
                best.display_section if best else "",
                mismatch.outcome.reason,
            )
        )
        output_row = worksheet.max_row
        for cell in worksheet[output_row]:
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")

        source_name_cell = worksheet.cell(output_row, 5)
        source_name_cell.fill = NO_MATCH_FILL
        source_name_cell.font = NO_MATCH_FONT

        best_name_cell = worksheet.cell(output_row, 6)
        if best:
            best_name_cell.fill = BEST_MATCH_FILL
            best_name_cell.font = BEST_MATCH_FONT

        score_cell = worksheet.cell(output_row, 7)
        score_cell.number_format = "0.0%"
        score_cell.alignment = Alignment(horizontal="center", vertical="center")

        reason_cell = worksheet.cell(output_row, 11)
        reason_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )

    worksheet.auto_filter.ref = f"A1:K{worksheet.max_row}"


def validate_threshold(value: float, option_name: str) -> None:
    if not 0.0 <= value <= 100.0:
        raise ValueError(f"{option_name} must be between 0 and 100.")


def resolved_input_path(value: str, label: str) -> Path:
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"{label} does not exist or is not a file: {path}")
    if path.suffix.lower() not in SUPPORTED_INPUT_EXTENSIONS:
        raise ValueError(
            f"{label} must be an .xlsx/.xlsm-compatible file, not '{path.suffix}'."
        )
    return path


def resolved_output_path(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if path.suffix.lower() != ".xlsx":
        raise ValueError("The output workbook must use the .xlsx extension.")
    return path


def clean_user_path(value: str) -> str:
    """Trim whitespace and paired quotes from typed or drag-and-dropped paths."""
    cleaned = value.strip()
    if (
        len(cleaned) >= 2
        and cleaned[0] == cleaned[-1]
        and cleaned[0] in {'"', "'"}
    ):
        cleaned = cleaned[1:-1].strip()
    return cleaned


def request_input_workbook(
    supplied_value: str | None,
    *,
    label: str,
    prompt: str,
) -> Path:
    """Resolve a CLI path, or keep prompting until an existing workbook is given."""
    if supplied_value is not None:
        return resolved_input_path(clean_user_path(supplied_value), label)

    while True:
        try:
            entered_value = input(prompt)
        except EOFError as error:
            raise ValueError(f"No path was supplied for {label.lower()}.") from error

        try:
            return resolved_input_path(clean_user_path(entered_value), label)
        except (FileNotFoundError, ValueError) as error:
            print(f"Error: {error}", file=sys.stderr)
            print("Please enter the full path again.", file=sys.stderr)


def default_output_path(source_path: Path) -> Path:
    """Choose a new report filename beside the workbook being checked."""
    base_name = f"{source_path.stem}_student_match_report"
    candidate = source_path.with_name(f"{base_name}.xlsx")
    counter = 2
    while candidate.exists():
        candidate = source_path.with_name(f"{base_name}_{counter}.xlsx")
        counter += 1
    return candidate.resolve()


def save_workbook_atomically(workbook: Workbook, output_path: Path) -> None:
    """Save through a temporary file so a failed export leaves no partial report."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}_",
            suffix=".xlsx",
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
        temporary_path = None
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def build_report(
    source_path: Path,
    base_path: Path,
    output_path: Path,
    *,
    student_threshold: float,
    location_threshold: float,
    source_header_row: int,
    base_header_row: int,
    overwrite: bool,
) -> tuple[BaseLoadStats, str, list[tuple[str, SheetStats]]]:
    """Build and save the report, returning compact processing statistics."""
    validate_threshold(student_threshold, "--student-threshold")
    validate_threshold(location_threshold, "--location-threshold")
    if source_header_row < 1 or base_header_row < 1:
        raise ValueError("Header row numbers must be 1 or greater.")

    if output_path in {source_path, base_path}:
        raise ValueError("The output path must be different from both input paths.")
    if output_path.exists() and not overwrite:
        raise FileExistsError(
            f"Output already exists: {output_path}. Use --overwrite to replace it."
        )

    base_records, base_stats, base_sheet_name = read_base_records(
        base_path, header_row=base_header_row
    )
    matcher = BaseMatcher(
        base_records,
        student_threshold=student_threshold,
        location_threshold=location_threshold,
    )

    source_workbook = load_workbook(
        filename=source_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    report_workbook = Workbook()
    report_workbook.remove(report_workbook.active)
    report_workbook.properties.title = "Student Match Report"
    report_workbook.properties.subject = "Source students checked against base data"

    sheet_results: list[tuple[str, SheetStats]] = []
    all_mismatches: list[MismatchRecord] = []
    try:
        if len(source_workbook.worksheets) < REPORT_SHEET_COUNT:
            raise ValueError(
                f"The source workbook has {len(source_workbook.worksheets)} sheet(s); "
                f"at least {REPORT_SHEET_COUNT} are required."
            )

        for sheet_index, source_sheet in enumerate(
            source_workbook.worksheets[:REPORT_SHEET_COUNT]
        ):
            report_sheet = report_workbook.create_sheet(title=source_sheet.title)
            stats, sheet_mismatches = write_report_sheet(
                source_sheet,
                report_sheet,
                source_sheet_index=sheet_index,
                source_header_row=source_header_row,
                matcher=matcher,
            )
            sheet_results.append((source_sheet.title, stats))
            all_mismatches.extend(sheet_mismatches)

        mismatch_sheet = report_workbook.create_sheet(title="Mismatched Students")
        write_mismatch_sheet(mismatch_sheet, all_mismatches)

        save_workbook_atomically(report_workbook, output_path)
    finally:
        source_workbook.close()
        report_workbook.close()

    return base_stats, base_sheet_name, sheet_results


def parse_arguments(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create a green/red student matching report plus a fourth mismatch "
            "audit sheet while opening both input workbooks read-only."
        ),
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "base_workbook",
        nargs="?",
        help="Base workbook containing v1/v2/v3; prompted for when omitted",
    )
    parser.add_argument(
        "source_workbook",
        nargs="?",
        help="Three-sheet workbook to check; prompted for when omitted",
    )
    parser.add_argument(
        "output_workbook",
        nargs="?",
        help=(
            "Optional new .xlsx report path; when omitted, a unique report name "
            "is created beside the workbook being checked"
        ),
    )
    parser.add_argument(
        "--student-threshold",
        type=float,
        default=88.0,
        help="Minimum fuzzy score for the student name",
    )
    parser.add_argument(
        "--location-threshold",
        type=float,
        default=90.0,
        help="Minimum fuzzy score required separately for district and school",
    )
    parser.add_argument(
        "--source-header-row",
        type=int,
        default=1,
        help="Header row in every source sheet",
    )
    parser.add_argument(
        "--base-header-row",
        type=int,
        default=1,
        help="Header row in the base sheet",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow replacement of an existing output report (never an input)",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_arguments(argv)
    try:
        base_path = request_input_workbook(
            args.base_workbook,
            label="Base workbook",
            prompt="Enter the full path of the base Excel file (full list): ",
        )
        source_path = request_input_workbook(
            args.source_workbook,
            label="Workbook to be checked",
            prompt="Enter the full path of the Excel file to be checked: ",
        )
        output_path = (
            resolved_output_path(clean_user_path(args.output_workbook))
            if args.output_workbook is not None
            else default_output_path(source_path)
        )

        if args.output_workbook is None:
            print(f"The report will be created at: {output_path}")

        base_stats, base_sheet_name, sheet_results = build_report(
            source_path,
            base_path,
            output_path,
            student_threshold=args.student_threshold,
            location_threshold=args.location_threshold,
            source_header_row=args.source_header_row,
            base_header_row=args.base_header_row,
            overwrite=args.overwrite,
        )
    except (FileNotFoundError, FileExistsError, PermissionError, ValueError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 2

    print(
        f"Base sheet '{base_sheet_name}': {base_stats.usable_rows} usable row(s), "
        f"{base_stats.skipped_invalid_rows} invalid row(s) skipped, "
        f"{base_stats.skipped_blank_rows} blank row(s) skipped."
    )
    for sheet_name, stats in sheet_results:
        print(
            f"{sheet_name}: {stats.output_rows} student(s) | "
            f"green={stats.exact_matches + stats.fuzzy_matches} "
            f"(normalized exact={stats.exact_matches}, fuzzy={stats.fuzzy_matches}) | "
            f"red={stats.unmatched} | "
            f"incomplete location={stats.incomplete_location_rows}"
        )
    print(f"Report written to: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
