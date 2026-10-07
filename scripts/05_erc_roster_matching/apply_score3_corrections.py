#!/usr/bin/env python3
"""Apply manually approved Sheet-4 corrections to a new workbook copy.

The input is the five-sheet report produced by ``student_match_report.py``
after a ``My score`` column has been inserted into ``Mismatched Students``.
For every Sheet-4 row whose score is 3, this script:

* resolves the target among the first three sheets using Source Sheet plus
  source district, school, section, and student name;
* requires that the target resolve to exactly one row;
* copies the best source-section candidate into the target row's ERC fields;
* changes that source student's name cell to a distinct manual-correction
  light-green style, different from the automatic-match green.

Ambiguous, missing, or unsafe corrections are reported and skipped. The input
workbook is never saved. A separately named workbook is written atomically.

Interactive use:

    python apply_score3_corrections.py

The input and optional output can also be supplied on the command line:

    python apply_score3_corrections.py scored_report.xlsx corrected_report.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import math
import os
import re
import sys
import tempfile
import unicodedata
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


SUPPORTED_EXTENSION = ".xlsx"
SOURCE_SHEET_COUNT = 3
MISMATCH_SHEET_TITLE = "Mismatched Students"

# Automatic matches in the original report use C6EFCE. Keep manual approvals
# visibly separate with a stronger light green.
MANUAL_CORRECTION_FILL = PatternFill(fill_type="solid", fgColor="A9D18E")
MANUAL_CORRECTION_FONT_COLOR = "375623"

SOURCE_HEADERS = (
    "District",
    "School Name",
    "Section",
    "Student Name",
    "Matched ERC Student Name",
    "Name Similarity",
    "ERC Section",
    "ERC School",
    "ERC District",
)

AUDIT_HEADERS = (
    "Source Sheet",
    "Source District",
    "Source School",
    "Source Section",
    "Source Student Name",
    "ERC District",
    "My score",
    "ERC School",
    "Best Candidate in Source Section",
    "Source-Section Similarity",
    "Source-Section ERC Section",
)

SOURCE_IDENTITY_HEADERS = (
    "District",
    "School Name",
    "Section",
    "Student Name",
)

AUDIT_IDENTITY_HEADERS = (
    "Source District",
    "Source School",
    "Source Section",
    "Source Student Name",
)

CORRECTION_MAPPING = (
    ("Matched ERC Student Name", "Best Candidate in Source Section"),
    ("Name Similarity", "Source-Section Similarity"),
    ("ERC Section", "Source-Section ERC Section"),
    ("ERC School", "ERC School"),
    ("ERC District", "ERC District"),
)


@dataclass
class CorrectionStats:
    """Processing totals and warnings returned to callers and the CLI."""

    score_three_rows: int = 0
    corrected_rows: int = 0
    ambiguous_rows: int = 0
    missing_rows: int = 0
    unsafe_rows: int = 0
    formula_score_rows: int = 0
    warnings: list[str] = field(default_factory=list)
    input_sha256: str = ""

    @property
    def skipped_rows(self) -> int:
        return (
            self.ambiguous_rows
            + self.missing_rows
            + self.unsafe_rows
            + self.formula_score_rows
        )


@dataclass(frozen=True)
class SourceSheetIndex:
    """Header positions and identity indexes for one source report sheet."""

    worksheet: Worksheet
    headers: dict[str, int]
    exact_rows: dict[tuple[str, str, str, str], tuple[int, ...]]
    normalized_rows: dict[tuple[str, str, str, str], tuple[int, ...]]


@dataclass(frozen=True)
class PendingCorrection:
    """A fully validated correction that has not yet changed the workbook."""

    audit_row: int
    source_index: SourceSheetIndex
    source_row: int
    candidate_values: tuple[object, object, object, object, object]


def cell_text(value: object) -> str:
    """Return a stable trimmed text representation for identity lookups."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_text(value: object) -> str:
    """Normalize harmless text differences without performing fuzzy matching."""
    text = unicodedata.normalize("NFKC", cell_text(value))
    # Apply Turkish uppercase I rules before case folding. Removing combining
    # marks then makes visually equivalent forms such as ``İ`` and ``i`` equal.
    # This is deterministic normalization, not a similarity/fuzzy comparison.
    text = text.replace("I", "ı").replace("İ", "i").casefold().replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.split())


def normalized_header(value: object) -> str:
    """Normalize an English report header independently of its column letter."""
    return re.sub(r"[^a-z0-9]+", "", normalized_text(value))


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.lstrip().startswith("=")


def is_score_three(value: object) -> bool:
    """Accept numeric 3 and ordinary text forms such as 3, 3.0, or 3,0."""
    if isinstance(value, bool) or value is None or is_formula(value):
        return False

    if isinstance(value, float) and not math.isfinite(value):
        return False

    text = cell_text(value).replace(",", ".")
    if not text:
        return False

    try:
        return Decimal(text) == Decimal(3)
    except InvalidOperation:
        return False


def sha256_file(path: Path) -> str:
    """Hash a file without loading it all into memory."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def paths_refer_to_same_file(first: Path, second: Path) -> bool:
    """Reject the input path itself and existing links to it."""
    if first == second:
        return True
    try:
        return first.exists() and second.exists() and first.samefile(second)
    except OSError:
        return False


def clean_user_path(value: str) -> str:
    """Trim whitespace and paired quotes from typed or dropped paths."""
    cleaned = value.strip()
    if (
        len(cleaned) >= 2
        and cleaned[0] == cleaned[-1]
        and cleaned[0] in {'"', "'"}
    ):
        cleaned = cleaned[1:-1].strip()
    return cleaned


def resolved_input_path(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(
            f"The manually scored report does not exist or is not a file: {path}"
        )
    if path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The manually scored report must be an .xlsx file.")
    return path


def resolved_output_path(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The output workbook must use the .xlsx extension.")
    return path


def request_input_workbook(supplied_value: str | None) -> Path:
    """Resolve a CLI path, or keep prompting until a valid path is entered."""
    if supplied_value is not None:
        return resolved_input_path(clean_user_path(supplied_value))

    while True:
        try:
            entered_value = input(
                "Enter the full path of the manually scored Excel report: "
            )
        except EOFError as error:
            raise ValueError("No input workbook path was supplied.") from error

        try:
            return resolved_input_path(clean_user_path(entered_value))
        except (FileNotFoundError, ValueError) as error:
            print(f"Error: {error}", file=sys.stderr)
            print("Please enter the full path again.", file=sys.stderr)


def default_output_path(input_path: Path) -> Path:
    """Choose a unique sibling path without overwriting an earlier result."""
    base_name = f"{input_path.stem}_score3_corrected"
    candidate = input_path.with_name(f"{base_name}.xlsx")
    counter = 2
    while candidate.exists():
        candidate = input_path.with_name(f"{base_name}_{counter}.xlsx")
        counter += 1
    return candidate.resolve()


def save_workbook_atomically(workbook: Workbook, output_path: Path) -> None:
    """Save through a temporary file so failures leave no partial workbook."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}_",
            suffix=".xlsx",
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
        temporary_path = None
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def locate_headers(
    worksheet: Worksheet,
    required_headers: Sequence[str],
    *,
    header_row: int = 1,
) -> dict[str, int]:
    """Return one-based columns for required headers, rejecting duplicates."""
    columns_by_normalized_header: dict[str, list[int]] = defaultdict(list)
    for cell in worksheet[header_row]:
        key = normalized_header(cell.value)
        if key:
            columns_by_normalized_header[key].append(cell.column)

    result: dict[str, int] = {}
    missing: list[str] = []
    duplicates: list[str] = []
    for header in required_headers:
        columns = columns_by_normalized_header.get(normalized_header(header), [])
        if not columns:
            missing.append(header)
        elif len(columns) > 1:
            duplicates.append(header)
        else:
            result[header] = columns[0]

    problems: list[str] = []
    if missing:
        problems.append(f"missing header(s): {', '.join(missing)}")
    if duplicates:
        problems.append(f"duplicate header(s): {', '.join(duplicates)}")
    if problems:
        raise ValueError(
            f"Sheet '{worksheet.title}' has " + "; ".join(problems) + "."
        )
    return result


def locate_mismatch_sheet(workbook: Workbook) -> tuple[Worksheet, dict[str, int]]:
    """Find Sheet 4 by title, with a validated fourth-sheet fallback."""
    if MISMATCH_SHEET_TITLE in workbook.sheetnames:
        worksheet = workbook[MISMATCH_SHEET_TITLE]
        return worksheet, locate_headers(worksheet, AUDIT_HEADERS)

    if len(workbook.worksheets) < 4:
        raise ValueError(
            "The workbook has fewer than four sheets and no "
            f"'{MISMATCH_SHEET_TITLE}' sheet."
        )

    worksheet = workbook.worksheets[3]
    try:
        headers = locate_headers(worksheet, AUDIT_HEADERS)
    except ValueError as error:
        raise ValueError(
            f"Could not identify '{MISMATCH_SHEET_TITLE}'. "
            "The fourth sheet does not have the required headers."
        ) from error
    return worksheet, headers


def source_identity_from_row(
    worksheet: Worksheet,
    row_number: int,
    headers: dict[str, int],
) -> tuple[str, str, str, str]:
    return tuple(
        cell_text(worksheet.cell(row_number, headers[header]).value)
        for header in SOURCE_IDENTITY_HEADERS
    )  # type: ignore[return-value]


def audit_identity_from_row(
    worksheet: Worksheet,
    row_number: int,
    headers: dict[str, int],
) -> tuple[str, str, str, str]:
    return tuple(
        cell_text(worksheet.cell(row_number, headers[header]).value)
        for header in AUDIT_IDENTITY_HEADERS
    )  # type: ignore[return-value]


def normalized_identity(
    identity: tuple[str, str, str, str],
) -> tuple[str, str, str, str]:
    return tuple(normalized_text(value) for value in identity)  # type: ignore[return-value]


def build_source_index(worksheet: Worksheet) -> SourceSheetIndex:
    """Index every data row by the four source identity fields."""
    headers = locate_headers(worksheet, SOURCE_HEADERS)
    exact: dict[tuple[str, str, str, str], list[int]] = defaultdict(list)
    normalized: dict[tuple[str, str, str, str], list[int]] = defaultdict(list)

    for row_number in range(2, worksheet.max_row + 1):
        identity = source_identity_from_row(worksheet, row_number, headers)
        if not any(identity):
            continue
        exact[identity].append(row_number)
        normalized[normalized_identity(identity)].append(row_number)

    return SourceSheetIndex(
        worksheet=worksheet,
        headers=headers,
        exact_rows={key: tuple(rows) for key, rows in exact.items()},
        normalized_rows={key: tuple(rows) for key, rows in normalized.items()},
    )


def resolve_source_rows(
    source_index: SourceSheetIndex,
    identity: tuple[str, str, str, str],
) -> tuple[int, ...]:
    """Prefer exact cleaned identity; use only a unique-safe text fallback."""
    exact_rows = source_index.exact_rows.get(identity, ())
    if exact_rows:
        return exact_rows
    return source_index.normalized_rows.get(normalized_identity(identity), ())


def values_equivalent(existing: object, proposed: object) -> bool:
    """Compare existing ERC values conservatively for overwrite protection."""
    if is_blank(existing) and is_blank(proposed):
        return True
    if isinstance(existing, bool) or isinstance(proposed, bool):
        return existing == proposed
    if isinstance(existing, (int, float, Decimal)) and isinstance(
        proposed, (int, float, Decimal)
    ):
        try:
            return Decimal(str(existing)) == Decimal(str(proposed))
        except InvalidOperation:
            return False
    return cell_text(existing) == cell_text(proposed)


def candidate_values_from_audit(
    worksheet: Worksheet,
    row_number: int,
    headers: dict[str, int],
) -> tuple[object, object, object, object, object]:
    return tuple(
        worksheet.cell(row_number, headers[audit_header]).value
        for _, audit_header in CORRECTION_MAPPING
    )  # type: ignore[return-value]


def warning_prefix(
    audit_sheet: Worksheet,
    audit_row: int,
    audit_headers: dict[str, int],
) -> str:
    source_sheet = cell_text(
        audit_sheet.cell(audit_row, audit_headers["Source Sheet"]).value
    )
    student = cell_text(
        audit_sheet.cell(audit_row, audit_headers["Source Student Name"]).value
    )
    return (
        f"Sheet 4 row {audit_row} (source sheet '{source_sheet}', "
        f"student '{student or '[blank]'}')"
    )


def add_warning(
    stats: CorrectionStats,
    audit_sheet: Worksheet,
    audit_row: int,
    audit_headers: dict[str, int],
    message: str,
) -> None:
    stats.warnings.append(
        f"{warning_prefix(audit_sheet, audit_row, audit_headers)}: "
        f"{message} Skipped."
    )


def apply_score3_corrections(
    input_path: Path,
    output_path: Path,
) -> CorrectionStats:
    """Create a corrected copy and return processing statistics."""
    input_path = input_path.expanduser().resolve()
    output_path = output_path.expanduser().resolve()

    if input_path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The input workbook must be an .xlsx file.")
    if output_path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The output workbook must use the .xlsx extension.")
    if not input_path.is_file():
        raise FileNotFoundError(f"Input workbook does not exist: {input_path}")
    if paths_refer_to_same_file(input_path, output_path):
        raise ValueError("The output path must be different from the input path.")
    if output_path.exists():
        raise FileExistsError(
            f"Output already exists and will not be overwritten: {output_path}"
        )

    stats = CorrectionStats(input_sha256=sha256_file(input_path))
    workbook = load_workbook(input_path, data_only=False)
    try:
        if len(workbook.worksheets) < SOURCE_SHEET_COUNT + 1:
            raise ValueError("The workbook must contain at least four sheets.")

        source_indexes = [
            build_source_index(worksheet)
            for worksheet in workbook.worksheets[:SOURCE_SHEET_COUNT]
        ]
        exact_sheet_names = {
            cell_text(index.worksheet.title): index for index in source_indexes
        }
        normalized_sheet_names: dict[str, list[SourceSheetIndex]] = defaultdict(list)
        for index in source_indexes:
            normalized_sheet_names[normalized_text(index.worksheet.title)].append(index)

        audit_sheet, audit_headers = locate_mismatch_sheet(workbook)
        score_column = audit_headers["My score"]
        pending: list[PendingCorrection] = []

        for audit_row in range(2, audit_sheet.max_row + 1):
            score_value = audit_sheet.cell(audit_row, score_column).value
            if is_formula(score_value):
                stats.formula_score_rows += 1
                add_warning(
                    stats,
                    audit_sheet,
                    audit_row,
                    audit_headers,
                    "the My score cell is a formula, which openpyxl cannot calculate",
                )
                continue
            if not is_score_three(score_value):
                continue

            stats.score_three_rows += 1
            source_sheet_value = cell_text(
                audit_sheet.cell(
                    audit_row, audit_headers["Source Sheet"]
                ).value
            )
            source_index = exact_sheet_names.get(source_sheet_value)
            if source_index is None:
                sheet_candidates = normalized_sheet_names.get(
                    normalized_text(source_sheet_value), []
                )
                if len(sheet_candidates) == 1:
                    source_index = sheet_candidates[0]

            if source_index is None:
                stats.missing_rows += 1
                add_warning(
                    stats,
                    audit_sheet,
                    audit_row,
                    audit_headers,
                    "the named source sheet was not found among the first three sheets",
                )
                continue

            identity = audit_identity_from_row(
                audit_sheet, audit_row, audit_headers
            )
            if not all(identity):
                stats.unsafe_rows += 1
                add_warning(
                    stats,
                    audit_sheet,
                    audit_row,
                    audit_headers,
                    "one or more source identity fields are blank",
                )
                continue

            source_rows = resolve_source_rows(source_index, identity)
            if not source_rows:
                stats.missing_rows += 1
                add_warning(
                    stats,
                    audit_sheet,
                    audit_row,
                    audit_headers,
                    "no source row matched district, school, section, and student",
                )
                continue
            if len(source_rows) > 1:
                stats.ambiguous_rows += 1
                add_warning(
                    stats,
                    audit_sheet,
                    audit_row,
                    audit_headers,
                    f"{len(source_rows)} source rows matched the full identity",
                )
                continue

            candidate_values = candidate_values_from_audit(
                audit_sheet, audit_row, audit_headers
            )
            missing_candidate_fields = [
                audit_header
                for (_, audit_header), value in zip(
                    CORRECTION_MAPPING, candidate_values
                )
                if is_blank(value)
            ]
            if missing_candidate_fields:
                stats.unsafe_rows += 1
                add_warning(
                    stats,
                    audit_sheet,
                    audit_row,
                    audit_headers,
                    "required candidate field(s) are blank: "
                    + ", ".join(missing_candidate_fields),
                )
                continue

            source_row = source_rows[0]
            conflicting_fields: list[str] = []
            for (source_header, _), proposed_value in zip(
                CORRECTION_MAPPING, candidate_values
            ):
                existing_value = source_index.worksheet.cell(
                    source_row, source_index.headers[source_header]
                ).value
                if not is_blank(existing_value) and not values_equivalent(
                    existing_value, proposed_value
                ):
                    conflicting_fields.append(source_header)

            if conflicting_fields:
                stats.unsafe_rows += 1
                add_warning(
                    stats,
                    audit_sheet,
                    audit_row,
                    audit_headers,
                    "existing ERC data conflict in: "
                    + ", ".join(conflicting_fields),
                )
                continue

            pending.append(
                PendingCorrection(
                    audit_row=audit_row,
                    source_index=source_index,
                    source_row=source_row,
                    candidate_values=candidate_values,
                )
            )

        pending_by_target: dict[tuple[int, int], list[PendingCorrection]] = (
            defaultdict(list)
        )
        for correction in pending:
            pending_by_target[
                (id(correction.source_index.worksheet), correction.source_row)
            ].append(correction)

        for corrections in pending_by_target.values():
            if len(corrections) > 1:
                for correction in corrections:
                    stats.ambiguous_rows += 1
                    add_warning(
                        stats,
                        audit_sheet,
                        correction.audit_row,
                        audit_headers,
                        "more than one score-3 instruction points to the same source row",
                    )
                continue

            correction = corrections[0]
            worksheet = correction.source_index.worksheet
            headers = correction.source_index.headers
            for (source_header, _), value in zip(
                CORRECTION_MAPPING, correction.candidate_values
            ):
                worksheet.cell(
                    correction.source_row, headers[source_header]
                ).value = value

            student_cell = worksheet.cell(
                correction.source_row, headers["Student Name"]
            )
            student_cell.fill = MANUAL_CORRECTION_FILL
            green_font = copy(student_cell.font)
            green_font.color = MANUAL_CORRECTION_FONT_COLOR
            student_cell.font = green_font
            stats.corrected_rows += 1

        if sha256_file(input_path) != stats.input_sha256:
            raise RuntimeError(
                "The input workbook changed while it was being processed; "
                "the output was not written."
            )

        save_workbook_atomically(workbook, output_path)
    finally:
        workbook.close()

    if sha256_file(input_path) != stats.input_sha256:
        raise RuntimeError(
            "The input workbook changed while the output was being saved."
        )
    return stats


def parse_arguments(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Copy a manually scored five-sheet report and apply unambiguous "
            "My score = 3 source-section corrections."
        )
    )
    parser.add_argument(
        "input_workbook",
        nargs="?",
        help="Full path of the manually scored .xlsx report",
    )
    parser.add_argument(
        "output_workbook",
        nargs="?",
        help="Optional full path for the new corrected .xlsx workbook",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_arguments(argv)
    try:
        input_path = request_input_workbook(args.input_workbook)
        output_path = (
            resolved_output_path(clean_user_path(args.output_workbook))
            if args.output_workbook is not None
            else default_output_path(input_path)
        )
        if args.output_workbook is None:
            print(f"The corrected copy will be created at: {output_path}")

        stats = apply_score3_corrections(
            input_path,
            output_path,
        )
    except (
        FileNotFoundError,
        FileExistsError,
        PermissionError,
        RuntimeError,
        ValueError,
    ) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 2

    for warning in stats.warnings:
        print(f"WARNING: {warning}", file=sys.stderr)

    print(
        f"My score = 3 rows: {stats.score_three_rows} | "
        f"corrected: {stats.corrected_rows} | "
        f"ambiguous skipped: {stats.ambiguous_rows} | "
        f"other skipped: {stats.missing_rows + stats.unsafe_rows}"
    )
    if stats.formula_score_rows:
        print(
            f"Formula score cells skipped: {stats.formula_score_rows}",
            file=sys.stderr,
        )
    print(f"Input workbook unchanged (SHA-256 verified): {input_path}")
    print(f"Corrected workbook written to: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
