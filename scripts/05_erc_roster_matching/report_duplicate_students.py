#!/usr/bin/env python3
"""Add a sixth-sheet duplicate audit to a copy of a student report workbook.

Only the first three worksheets are analyzed. They are treated independently,
and a possible duplicate is two or more rows with the same normalized:

    District + School Name + Section + Student Name

Normalization is deterministic only: Unicode normalization, Turkish-aware
case normalization, diacritic normalization, and whitespace cleanup. No fuzzy
matching, SequenceMatcher, edit-distance scoring, or similarity threshold is
used.

The input workbook is never overwritten. A new workbook copy is created, and
the only intentional workbook change is a new sixth sheet named
``Duplicate Students``. Existing worksheet cells are never written to.

Interactive use:

    python report_duplicate_students.py

Command-line use:

    python report_duplicate_students.py input.xlsx [output.xlsx]
"""

from __future__ import annotations

import argparse
import hashlib
import math
import os
import re
import sys
import tempfile
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


SUPPORTED_EXTENSION = ".xlsx"
SOURCE_SHEET_COUNT = 3
MINIMUM_EXISTING_SHEETS = 5
REPORT_SHEET_TITLE = "Duplicate Students"

REQUIRED_SOURCE_HEADERS = (
    "District",
    "School Name",
    "Section",
    "Student Name",
)

REPORT_HEADERS = (
    "Source Sheet",
    "Source Row",
    "District",
    "School Name",
    "Section",
    "Student Name",
    "Duplicate Type",
    "Duplicate Group",
    "Occurrence Count",
)

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Aptos", size=11, color="000000")
EXACT_FILL = PatternFill(fill_type="solid", fgColor="FCE4D6")
NORMALIZED_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
GROUP_FILL_A = PatternFill(fill_type="solid", fgColor="F7F9FC")
GROUP_FILL_B = PatternFill(fill_type="solid", fgColor="EAF2F8")
THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9E1F2"),
)


ExactCellKey = tuple[str, object]
NormalizedKey = tuple[str, str, str, str]


@dataclass(frozen=True)
class Occurrence:
    sheet_number: int
    sheet_name: str
    row_number: int
    values: tuple[object, object, object, object]
    exact_key: tuple[ExactCellKey, ExactCellKey, ExactCellKey, ExactCellKey]


@dataclass(frozen=True)
class DuplicateReportRow:
    source_sheet: str
    source_row: int
    district: object
    school: object
    section: object
    student: object
    duplicate_type: str
    group_id: str
    occurrence_count: int

    def as_tuple(self) -> tuple[object, ...]:
        return (
            self.source_sheet,
            self.source_row,
            self.district,
            self.school,
            self.section,
            self.student,
            self.duplicate_type,
            self.group_id,
            self.occurrence_count,
        )


@dataclass
class DuplicateStats:
    scanned_rows: int = 0
    complete_rows: int = 0
    incomplete_rows: int = 0
    duplicate_groups: int = 0
    duplicate_occurrences: int = 0
    exact_groups: int = 0
    normalized_groups: int = 0
    per_sheet: dict[str, tuple[int, int]] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    input_sha256: str = ""


def cell_text(value: object) -> str:
    """Return a stable, trimmed representation suitable for normalization."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_text(value: object) -> str:
    """Normalize text deterministically without any fuzzy comparison."""
    text = unicodedata.normalize("NFKC", cell_text(value))
    text = text.replace("I", "ı").replace("İ", "i").casefold().replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.split())


def normalized_header(value: object) -> str:
    """Normalize an English report header independently of its column."""
    return re.sub(r"[^a-z0-9]+", "", normalized_text(value))


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def exact_cell_key(value: object) -> ExactCellKey:
    """Return a type-aware key for classifying an exact duplicate group."""
    if value is None:
        return "blank", None
    if isinstance(value, bool):
        return "boolean", value
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and not math.isfinite(value):
            return "float", repr(value)
        try:
            return "number", Decimal(str(value))
        except InvalidOperation:
            return "other", repr(value)
    if isinstance(value, str):
        return "text", value
    return type(value).__name__, repr(value)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_user_path(value: str) -> str:
    cleaned = value.strip()
    if (
        len(cleaned) >= 2
        and cleaned[0] == cleaned[-1]
        and cleaned[0] in {'"', "'"}
    ):
        cleaned = cleaned[1:-1].strip()
    return cleaned


def resolved_input_path(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Input workbook does not exist: {path}")
    if path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The input workbook must be an .xlsx file.")
    return path


def resolved_output_path(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The output workbook must use the .xlsx extension.")
    return path


def request_input_path(supplied_value: str | None) -> Path:
    if supplied_value is not None:
        return resolved_input_path(clean_user_path(supplied_value))

    while True:
        try:
            entered = input("Enter the full path of the latest Excel workbook: ")
        except EOFError as error:
            raise ValueError("No input workbook path was supplied.") from error

        try:
            return resolved_input_path(clean_user_path(entered))
        except (FileNotFoundError, ValueError) as error:
            print(f"Error: {error}", file=sys.stderr)
            print("Please enter the full path again.", file=sys.stderr)


def default_output_path(input_path: Path) -> Path:
    base_name = f"{input_path.stem}_with_duplicate_report"
    candidate = input_path.with_name(f"{base_name}.xlsx")
    counter = 2
    while candidate.exists():
        candidate = input_path.with_name(f"{base_name}_{counter}.xlsx")
        counter += 1
    return candidate.resolve()


def paths_refer_to_same_file(first: Path, second: Path) -> bool:
    if first == second:
        return True
    try:
        return first.exists() and second.exists() and first.samefile(second)
    except OSError:
        return False


def save_workbook_atomically(workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}_",
            suffix=SUPPORTED_EXTENSION,
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
        temporary_path = None
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def locate_headers(
    worksheet: Worksheet,
    required_headers: Sequence[str],
) -> dict[str, int]:
    """Locate required headers in row 1 and reject missing/duplicate labels."""
    columns_by_header: dict[str, list[int]] = defaultdict(list)
    for cell in worksheet[1]:
        key = normalized_header(cell.value)
        if key:
            columns_by_header[key].append(cell.column)

    result: dict[str, int] = {}
    missing: list[str] = []
    duplicates: list[str] = []
    for header in required_headers:
        columns = columns_by_header.get(normalized_header(header), [])
        if not columns:
            missing.append(header)
        elif len(columns) > 1:
            duplicates.append(header)
        else:
            result[header] = columns[0]

    problems: list[str] = []
    if missing:
        problems.append(f"missing header(s): {', '.join(missing)}")
    if duplicates:
        problems.append(f"duplicate header(s): {', '.join(duplicates)}")
    if problems:
        raise ValueError(
            f"Sheet '{worksheet.title}' has " + "; ".join(problems) + "."
        )
    return result


def collect_duplicate_rows(
    workbook,
    stats: DuplicateStats,
) -> list[DuplicateReportRow]:
    report_rows: list[DuplicateReportRow] = []

    for sheet_number, worksheet in enumerate(
        workbook.worksheets[:SOURCE_SHEET_COUNT],
        start=1,
    ):
        headers = locate_headers(worksheet, REQUIRED_SOURCE_HEADERS)
        groups: dict[NormalizedKey, list[Occurrence]] = defaultdict(list)

        for row_number in range(2, worksheet.max_row + 1):
            values = tuple(
                worksheet.cell(row_number, headers[header]).value
                for header in REQUIRED_SOURCE_HEADERS
            )
            if not any(not is_blank(value) for value in values):
                continue

            stats.scanned_rows += 1
            if any(is_blank(value) for value in values):
                stats.incomplete_rows += 1
                stats.warnings.append(
                    f"Sheet '{worksheet.title}' row {row_number} has a blank "
                    "district, school, section, or student value and was skipped."
                )
                continue

            stats.complete_rows += 1
            normalized_key = tuple(normalized_text(value) for value in values)
            exact_key = tuple(exact_cell_key(value) for value in values)
            groups[normalized_key].append(
                Occurrence(
                    sheet_number=sheet_number,
                    sheet_name=worksheet.title,
                    row_number=row_number,
                    values=values,  # type: ignore[arg-type]
                    exact_key=exact_key,  # type: ignore[arg-type]
                )
            )

        duplicate_groups = [
            occurrences
            for occurrences in groups.values()
            if len(occurrences) >= 2
        ]
        duplicate_groups.sort(key=lambda items: min(item.row_number for item in items))

        sheet_group_count = 0
        sheet_occurrence_count = 0
        for group_number, occurrences in enumerate(duplicate_groups, start=1):
            occurrences.sort(key=lambda item: item.row_number)
            group_id = f"S{sheet_number}-D{group_number:04d}"
            duplicate_type = (
                "Exact"
                if len({item.exact_key for item in occurrences}) == 1
                else "Normalized"
            )

            stats.duplicate_groups += 1
            stats.duplicate_occurrences += len(occurrences)
            sheet_group_count += 1
            sheet_occurrence_count += len(occurrences)
            if duplicate_type == "Exact":
                stats.exact_groups += 1
            else:
                stats.normalized_groups += 1

            for occurrence in occurrences:
                district, school, section, student = occurrence.values
                report_rows.append(
                    DuplicateReportRow(
                        source_sheet=occurrence.sheet_name,
                        source_row=occurrence.row_number,
                        district=district,
                        school=school,
                        section=section,
                        student=student,
                        duplicate_type=duplicate_type,
                        group_id=group_id,
                        occurrence_count=len(occurrences),
                    )
                )

        stats.per_sheet[worksheet.title] = (
            sheet_group_count,
            sheet_occurrence_count,
        )

    return report_rows


def style_report_sheet(worksheet: Worksheet) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 34

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    group_order: dict[str, int] = {}
    for row_number in range(2, worksheet.max_row + 1):
        group_id = worksheet.cell(row_number, 8).value
        if not group_id:
            continue
        if group_id not in group_order:
            group_order[group_id] = len(group_order)
        group_fill = (
            GROUP_FILL_A if group_order[group_id] % 2 == 0 else GROUP_FILL_B
        )
        for cell in worksheet[row_number]:
            cell.font = BODY_FONT
            cell.fill = group_fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

        duplicate_type_cell = worksheet.cell(row_number, 7)
        duplicate_type_cell.fill = (
            EXACT_FILL
            if duplicate_type_cell.value == "Exact"
            else NORMALIZED_FILL
        )
        worksheet.cell(row_number, 2).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        worksheet.cell(row_number, 5).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        worksheet.cell(row_number, 7).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        worksheet.cell(row_number, 8).alignment = Alignment(
            horizontal="center", vertical="center"
        )
        worksheet.cell(row_number, 9).alignment = Alignment(
            horizontal="center", vertical="center"
        )

    widths = {
        "A": 28,
        "B": 12,
        "C": 22,
        "D": 42,
        "E": 12,
        "F": 32,
        "G": 18,
        "H": 18,
        "I": 18,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    last_column = get_column_letter(len(REPORT_HEADERS))
    worksheet.auto_filter.ref = f"A1:{last_column}{worksheet.max_row}"


def add_duplicate_report(
    input_path: Path,
    output_path: Path,
) -> DuplicateStats:
    """Create a workbook copy containing only one new duplicate-report sheet."""
    input_path = input_path.expanduser().resolve()
    output_path = output_path.expanduser().resolve()

    if input_path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The input workbook must be an .xlsx file.")
    if output_path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The output workbook must use the .xlsx extension.")
    if not input_path.is_file():
        raise FileNotFoundError(f"Input workbook does not exist: {input_path}")
    if paths_refer_to_same_file(input_path, output_path):
        raise ValueError(
            "The output path must differ from the input path so the input "
            "workbook remains untouched."
        )
    if output_path.exists():
        raise FileExistsError(
            f"Output already exists and will not be overwritten: {output_path}"
        )

    stats = DuplicateStats(input_sha256=sha256_file(input_path))
    workbook = load_workbook(input_path, data_only=False, keep_links=True)
    try:
        if len(workbook.worksheets) < MINIMUM_EXISTING_SHEETS:
            raise ValueError(
                "The latest workbook must contain at least five existing sheets "
                "before the duplicate report is added as Sheet 6."
            )
        if REPORT_SHEET_TITLE in workbook.sheetnames:
            raise ValueError(
                f"The workbook already contains a '{REPORT_SHEET_TITLE}' sheet; "
                "it was not replaced."
            )

        report_rows = collect_duplicate_rows(workbook, stats)
        report_sheet = workbook.create_sheet(REPORT_SHEET_TITLE, index=5)
        report_sheet.append(REPORT_HEADERS)
        if report_rows:
            for report_row in report_rows:
                report_sheet.append(report_row.as_tuple())
        else:
            report_sheet.append(
                (
                    "No possible duplicate students were found in the first "
                    "three sheets.",
                )
            )
            report_sheet.merge_cells(
                start_row=2,
                start_column=1,
                end_row=2,
                end_column=len(REPORT_HEADERS),
            )
            report_sheet["A2"].alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            report_sheet["A2"].font = Font(
                name="Aptos",
                size=11,
                italic=True,
                color="666666",
            )

        style_report_sheet(report_sheet)

        if sha256_file(input_path) != stats.input_sha256:
            raise RuntimeError(
                "The input workbook changed while it was being processed; "
                "the output was not written."
            )

        save_workbook_atomically(workbook, output_path)
    finally:
        workbook.close()

    if sha256_file(input_path) != stats.input_sha256:
        raise RuntimeError("The input workbook changed while saving the output.")
    return stats


def parse_arguments(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create a copy of a five-sheet student report and add a sixth sheet "
            "listing normalized, non-fuzzy duplicate students from the first "
            "three sheets."
        )
    )
    parser.add_argument(
        "input_workbook",
        nargs="?",
        help="Full path of the latest .xlsx student report",
    )
    parser.add_argument(
        "output_workbook",
        nargs="?",
        help="Optional full path for the new workbook copy",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_arguments(argv)
    try:
        input_path = request_input_path(args.input_workbook)
        output_path = (
            resolved_output_path(clean_user_path(args.output_workbook))
            if args.output_workbook is not None
            else default_output_path(input_path)
        )
        if args.output_workbook is None:
            print(f"The duplicate-report copy will be created at: {output_path}")

        stats = add_duplicate_report(input_path, output_path)
    except (
        FileNotFoundError,
        FileExistsError,
        PermissionError,
        RuntimeError,
        ValueError,
    ) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 2

    for warning in stats.warnings:
        print(f"WARNING: {warning}", file=sys.stderr)

    for sheet_name, (group_count, occurrence_count) in stats.per_sheet.items():
        print(
            f"{sheet_name}: duplicate groups={group_count}, "
            f"reported rows={occurrence_count}"
        )
    print(
        f"Total duplicate groups: {stats.duplicate_groups} "
        f"(exact={stats.exact_groups}, normalized={stats.normalized_groups}) | "
        f"reported occurrences: {stats.duplicate_occurrences} | "
        f"incomplete rows skipped: {stats.incomplete_rows}"
    )
    print("The input workbook is unchanged (SHA-256 verified).")
    print(f"Workbook with Sheet 6 written to: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
