#!/usr/bin/env python3
"""Measure structured-name improvements in an existing Mismatched Students sheet.

The mismatch sheet defines the historical red cohort. Count every nonblank data
row, including unresolved cases, and exclude its header. Search the complete ERC
school for each original name; never require a preselected accepted name pair.
Both workbooks remain read-only. Output contains local, private record details.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from student_match_report import (
    BaseMatcher, cell_text, normalize_for_match as norm, normalize_section,
    read_base_records,
)
from structured_name_match import assess_unmatched_student


REQUIRED_HEADERS = (
    'Source Sheet', 'Source District', 'Source School', 'Source Section',
    'Source Student Name',
)


def read_mismatch_cohort(report_path, sheet_name='Mismatched Students'):
    """Read membership from row contents, not max_row, fills, or accepted pairs."""
    book = load_workbook(report_path, read_only=True, data_only=True)
    try:
        if sheet_name not in book.sheetnames:
            raise ValueError(
                f"{report_path.name} has no {sheet_name!r} sheet. Supply the original "
                'student matching report, not all_grades_final.xlsx.'
            )
        sheet = book[sheet_name]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        indexes = {norm(value): i for i, value in enumerate(header) if cell_text(value)}
        missing = [label for label in REQUIRED_HEADERS if norm(label) not in indexes]
        if missing:
            raise ValueError(f'Mismatch sheet is missing headers: {missing}')

        def value(row, label):
            index = indexes.get(norm(label))
            return row[index] if index is not None and index < len(row) else None

        records = []
        for sheet_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(cell_text(item) for item in row):
                continue
            # The original report writer uses this merged-cell message when
            # the cohort is empty. It is not a student record.
            if (cell_text(row[0]) == 'No mismatched students were found.'
                    and not any(cell_text(item) for item in row[1:])):
                continue
            source_sheet = cell_text(value(row, 'Source Sheet'))
            if not source_sheet:
                raise ValueError(f'Missing Source Sheet on mismatch row {sheet_row}; row was not discarded')
            records.append({
                'sheet': source_sheet, 'mismatch_row': sheet_row,
                'source_district': cell_text(value(row, 'Source District')),
                'source_school': cell_text(value(row, 'Source School')),
                'source_section': cell_text(value(row, 'Source Section')),
                'source_name': cell_text(value(row, 'Source Student Name')),
                'recorded_reason': cell_text(value(row, 'Reason Not Matched')),
                'recorded_best_name': cell_text(value(row, 'Best Candidate Across All Sections')),
                'recorded_best_score': value(row, 'All-Sections Similarity'),
                # Kept as source annotations only, never used to pick a match.
                'manual_score': value(row, 'My score'),
                'manual_score2': value(row, 'My Score2'),
            })
        # Preserve original grade-sheet ordering where present in the workbook.
        sheet_order = [name for name in book.sheetnames if any(r['sheet'] == name for r in records)]
        sheet_order.extend(dict.fromkeys(r['sheet'] for r in records if r['sheet'] not in sheet_order))
        return records, sheet_order, sheet.max_row
    finally:
        book.close()


def summarize_cohort(rows, sheet_order=()):
    keys = ('original_red_records', 'red_to_green', 'red_to_yellow', 'flagged',
            'still_red', 'baseline_now_nonred', 'red_location_now')
    counts = {name: dict.fromkeys(keys, 0) for name in sheet_order}
    for row in rows:
        group = counts.setdefault(row['sheet'], dict.fromkeys(keys, 0))
        group['original_red_records'] += 1
        promoted = row['decision'] in {'Green + review', 'Yellow + review'}
        if promoted and row['baseline_status'] != 'red_name':
            raise ValueError('An unchanged nonred/location result cannot be credited to a name rule')
        group['red_to_green'] += row['decision'] == 'Green + review'
        group['red_to_yellow'] += row['decision'] == 'Yellow + review'
        group['flagged'] += row['review_required']
        group['baseline_now_nonred'] += not row['baseline_status'].startswith('red_')
        group['red_location_now'] += row['baseline_status'] == 'red_location'
        group['still_red'] += row['baseline_status'].startswith('red_') and not promoted

    def rates(group):
        total = group['original_red_records']
        return {
            **group,
            'not_promoted_by_rules': total - group['red_to_green'] - group['red_to_yellow'],
            'red_to_green_percent': 100 * group['red_to_green'] / total if total else None,
            'red_reduction_percent': 100 * (group['red_to_green'] + group['red_to_yellow']) / total if total else None,
        }

    total = {key: sum(group[key] for group in counts.values()) for key in keys}
    return {'per_sheet': {name: rates(group) for name, group in counts.items()}, 'overall': rates(total)}


def print_summary(summary):
    width = max([len('Worksheet'), len('OVERALL'), *map(len, summary['per_sheet'])])
    print('\nHistorical mismatch cohort: new green / original red student records')
    print(f"{'Worksheet':<{width}}  {'Original red':>12}  {'To green':>9}  {'To yellow':>9}  "
          f"{'Not promoted':>12}  {'Improvement':>12}")
    for name, group in [*summary['per_sheet'].items(), ('OVERALL', summary['overall'])]:
        value = group['red_to_green_percent']
        rate = 'n.a.' if value is None else f'{value:.2f}%'
        print(f"{name:<{width}}  {group['original_red_records']:>12}  {group['red_to_green']:>9}  "
              f"{group['red_to_yellow']:>9}  {group['not_promoted_by_rules']:>12}  {rate:>12}")
    print('Header excluded; all historical red data rows retained. Every new match requires review.')
    if summary['overall']['baseline_now_nonred']:
        print(f"{summary['overall']['baseline_now_nonred']} historical rows now qualify under the unchanged "
              'baseline; these are disclosed separately and receive no credit as rule improvements.')


def audit_mismatches(report_path, roster_path, output_path, *, expected_red_records=None,
                     mismatch_sheet='Mismatched Students'):
    if output_path.resolve() in {report_path.resolve(), roster_path.resolve()}:
        raise ValueError('Output must differ from both inputs')
    if output_path.exists():
        raise FileExistsError(f'Output already exists: {output_path}')
    hashes = {str(p.resolve()): hashlib.sha256(p.read_bytes()).hexdigest()
              for p in (report_path, roster_path)}
    cohort, sheet_order, physical_rows = read_mismatch_cohort(report_path, mismatch_sheet)
    if expected_red_records is not None and len(cohort) != expected_red_records:
        raise ValueError(f'Expected {expected_red_records} red records, found {len(cohort)}; header excluded')
    print(f'Read {len(cohort)} historical red student records from {mismatch_sheet}', flush=True)
    records, stats, _ = read_base_records(roster_path, header_row=1)
    matcher = BaseMatcher(records, student_threshold=90, location_threshold=95)
    results = []
    candidates = []
    for i, row in enumerate(cohort, 1):
        result = assess_unmatched_student(
            matcher, norm(row['source_district']), norm(row['source_school']),
            normalize_section(row['source_section']), norm(row['source_name']),
        )
        baseline = result.baseline
        selected = result.promoted or (result.candidates[0] if result.candidates else None)
        if result.promoted:
            decision = result.promoted[1].decision
            reason = result.promoted[1].reason
        elif not baseline.status.startswith('red_'):
            decision = 'Baseline already nonred'
            reason = 'Not credited: unchanged matcher now qualifies this historical red record'
        elif result.candidates:
            decision = 'Review only'
            reason = selected[1].reason
        else:
            decision = 'Unchanged red'
            reason = baseline.reason
        if selected:
            record, pair = selected
            target_fields = {
                'erc_name':record.display_student, 'erc_section':record.display_section,
                'erc_school':record.display_school, 'erc_district':record.display_district,
                'erc_row':record.row_number, 'baseline_pair_score':pair.baseline / 100,
                'adjusted_pair_score':pair.adjusted / 100,
                'rule_score':pair.evidence.rule_score / 100 if pair.evidence.rule_score is not None else None,
                'pattern':pair.evidence.pattern, 'fuzzy_parts':pair.evidence.fuzzy_parts,
                'full_parts_similarity':(pair.evidence.full_parts_similarity / 100
                                         if pair.evidence.full_parts_similarity is not None else None),
            }
        else:
            target_fields = dict.fromkeys(('erc_name','erc_section','erc_school','erc_district','erc_row',
                                           'baseline_pair_score','adjusted_pair_score','rule_score','pattern',
                                           'fuzzy_parts','full_parts_similarity'))
        best = baseline.best_all_sections or baseline.single_match
        if best is None and baseline.qualifying_candidates:
            best = baseline.qualifying_candidates[0]
        results.append({
            **row, **target_fields, 'baseline_status':baseline.status,
            'baseline_best_name':best.record.display_student if best else None,
            'baseline_best_score':best.score / 100 if best else None,
            'decision':decision, 'reason':reason,
            'review_required':bool(result.candidates), 'pattern_candidates':len(result.candidates),
        })
        for record, pair in result.candidates:
            candidates.append({
                'sheet':row['sheet'], 'mismatch_row':row['mismatch_row'],
                'source_name':row['source_name'], 'erc_name':record.display_student,
                'erc_row':record.row_number, 'erc_section':record.display_section,
                'baseline':pair.baseline/100, 'adjusted':pair.adjusted/100,
                'pattern':pair.evidence.pattern, 'fuzzy_parts':pair.evidence.fuzzy_parts,
                'full_parts_similarity':(pair.evidence.full_parts_similarity / 100
                                         if pair.evidence.full_parts_similarity is not None else None),
                'decision':pair.decision, 'reason':pair.reason,
                'rivals':[{'erc_row':r.row_number,'erc_name':r.display_student,
                           'erc_section':r.display_section,'baseline':old/100,'proposed':new/100}
                          for r,old,new,_ in pair.rivals],
            })
        if i % 500 == 0 or i == len(cohort):
            print(f'Evaluated {i}/{len(cohort)} original red records', flush=True)
    after = {str(p.resolve()): hashlib.sha256(p.read_bytes()).hexdigest()
             for p in (report_path, roster_path)}
    if after != hashes:
        raise AssertionError('An input file changed during the audit')
    if len(results) != len(cohort):
        raise AssertionError('Historical cohort membership changed')
    summary = summarize_cohort(results, sheet_order)
    payload = {
        'population':'historical_mismatch_sheet', 'report':str(report_path.resolve()),
        'roster':str(roster_path.resolve()), 'mismatch_sheet':mismatch_sheet,
        'worksheet_rows_including_header':physical_rows, 'roster_rows':stats.usable_rows,
        'hashes':hashes,
        'definition':'100 * original mismatch records newly Green + review / all original mismatch data rows',
        'summary':summary, 'rows':results, 'review_candidates':candidates,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    print_summary(summary)
    print(f'Saved {output_path}')
    return payload


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('matching_report', type=Path)
    parser.add_argument('erc_roster', type=Path)
    parser.add_argument('output_json', type=Path)
    parser.add_argument('--mismatch-sheet', default='Mismatched Students')
    parser.add_argument('--expected-red-records', type=int,
                        help='Optional check on data-row count; exclude the header')
    args = parser.parse_args()
    audit_mismatches(args.matching_report,args.erc_roster,args.output_json,
                     expected_red_records=args.expected_red_records,mismatch_sheet=args.mismatch_sheet)


if __name__ == '__main__':
    main()
