#!/usr/bin/env python3
"""Add ERC student availability and ambiguity reports to a workbook copy.

The program reads two input workbooks:

* the ERC full list, whose first sheet has ``v1`` (``District / School``),
  ``v2`` (a grade-prefixed section such as ``4-A``), and ``v3`` (student);
* the latest six-sheet student report, whose first three sheets contain the
  populated ERC match fields created by the earlier matching workflow.

For every nonblank ERC full-list row, Sheet 7 reports whether the ERC student
appears in each of the first three report sheets. Matching is deterministic
and non-fuzzy. The key is normalized ERC district + ERC school + matched ERC
student name. Section is deliberately excluded from the key. A report row is
eligible only when its ``ERC Section`` is nonblank.

The original source ``Student Name`` is never used as identity data. Only after
one unambiguous report row has been selected through the ERC fields is that
row's source name written to Sheet 7 and its fill/font copied to the result
cell.

Sheet 8 lists duplicate ERC keys and cases where one ERC key resolves to more
than one row in one of the first three report sheets. The program never edits
either input and refuses to overwrite an existing output file.

Interactive use (ERC path is requested first):

    python report_erc_availability.py

Command-line use:

    python report_erc_availability.py erc.xlsx report.xlsx output.xlsx
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
import tempfile
import unicodedata
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


SUPPORTED_EXTENSION = ".xlsx"
SOURCE_SHEET_COUNT = 3
MINIMUM_EXISTING_SHEETS = 6
AVAILABILITY_SHEET_TITLE = "ERC Student Availability"
DUPLICATE_SHEET_TITLE = "ERC Availability Duplicates"

SOURCE_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "source_student": ("Student Name",),
    "erc_student": ("Matched ERC Student Name", "ERC Student Name"),
    "erc_section": ("ERC Section",),
    "erc_school": ("ERC School",),
    "erc_district": ("ERC District",),
}

AVAILABILITY_FIXED_HEADERS = (
    "ERC District",
    "ERC School",
    "ERC Section",
    "ERC Student Name",
)

DUPLICATE_HEADERS = (
    "Duplicate Type",
    "Duplicate Group",
    "Occurrence Number",
    "Occurrence Count",
    "ERC District",
    "ERC School",
    "ERC Student Name",
    "ERC Section",
    "Report Sheet",
    "Report Row",
    "Report ERC Section",
    "Details",
)

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Aptos", size=11, color="000000")
NOT_PRESENT_FILL = PatternFill(fill_type="solid", fgColor="E7E6E6")
NOT_PRESENT_FONT = Font(name="Aptos", size=11, color="595959")
DUPLICATE_FILL = PatternFill(fill_type="solid", fgColor="F4B183")
DUPLICATE_FONT = Font(name="Aptos", size=11, color="9C5700")
INVALID_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
INVALID_FONT = Font(name="Aptos", size=11, color="9C0006")
GROUP_FILL_A = PatternFill(fill_type="solid", fgColor="F7F9FC")
GROUP_FILL_B = PatternFill(fill_type="solid", fgColor="EAF2F8")
THIN_BORDER = Border(bottom=Side(style="thin", color="D9E1F2"))


MatchKey = tuple[str, str, str]


@dataclass(frozen=True)
class ErcRecord:
    """One nonblank data row from the ERC full list."""

    source_order: int
    district: str
    school: str
    section: str
    student: str
    key: MatchKey | None
    problem: str = ""


@dataclass(frozen=True)
class ReportOccurrence:
    """One eligible ERC-filled row in a source report sheet."""

    sheet_number: int
    sheet_name: str
    row_number: int
    source_student: str
    erc_district: str
    erc_school: str
    erc_student: str
    erc_section: str


@dataclass
class SourceSheetIndex:
    """Header positions and match-key index for one report sheet."""

    worksheet: Worksheet
    sheet_number: int
    headers: dict[str, int]
    rows_by_key: dict[MatchKey, tuple[ReportOccurrence, ...]]


@dataclass
class AvailabilityStats:
    erc_rows: int = 0
    usable_erc_rows: int = 0
    invalid_erc_rows: int = 0
    eligible_report_rows: int = 0
    incomplete_report_rows: int = 0
    present_cells: int = 0
    absent_cells: int = 0
    duplicate_cells: int = 0
    erc_duplicate_groups: int = 0
    report_duplicate_groups: int = 0
    duplicate_report_rows: int = 0
    warnings: list[str] = field(default_factory=list)
    erc_input_sha256: str = ""
    report_input_sha256: str = ""


def cell_text(value: object) -> str:
    """Return a stable, trimmed display string for an Excel value."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_text(value: object) -> str:
    """Normalize harmless text differences without fuzzy matching."""
    text = unicodedata.normalize("NFKC", cell_text(value))

    # Turkish uppercase I needs locale-aware handling before case folding.
    text = text.replace("I", "ı").replace("İ", "i")
    text = text.casefold().replace("ı", "i")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))

    # Treat punctuation as separators and collapse whitespace. This is an
    # equality normalization only; no similarity score is calculated.
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def normalized_header(value: object) -> str:
    """Normalize an English header independently of its column position."""
    return normalized_text(value).replace(" ", "")


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def section_display(value: object) -> str:
    """Remove a leading grade prefix such as ``4-`` from an ERC section."""
    text = cell_text(value)
    match = re.match(r"^\s*\d+\s*[-/]\s*(.*?)\s*$", text)
    return match.group(1).strip() if match else text


def split_erc_location(value: object) -> tuple[str, str] | None:
    """Split ERC ``v1`` at its first slash into district and school."""
    text = cell_text(value)
    parts = re.split(r"\s*/\s*", text, maxsplit=1)
    if len(parts) != 2 or not parts[0].strip() or not parts[1].strip():
        return None
    return parts[0].strip(), parts[1].strip()


def make_match_key(
    district: object,
    school: object,
    student: object,
) -> MatchKey | None:
    """Build the section-free exact-normalized identity key."""
    key = (
        normalized_text(district),
        normalized_text(school),
        normalized_text(student),
    )
    return key if all(key) else None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_user_path(value: str) -> str:
    cleaned = value.strip()
    if (
        len(cleaned) >= 2
        and cleaned[0] == cleaned[-1]
        and cleaned[0] in {'"', "'"}
    ):
        cleaned = cleaned[1:-1].strip()
    return cleaned


def resolved_input_path(value: str, label: str) -> Path:
    path = Path(value).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"{label} does not exist: {path}")
    if path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError(f"{label} must be an .xlsx file.")
    return path


def resolved_output_path(value: str) -> Path:
    path = Path(value).expanduser().resolve()
    if path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The output workbook must use the .xlsx extension.")
    return path


def request_input_path(
    supplied_value: str | None,
    *,
    prompt: str,
    label: str,
) -> Path:
    if supplied_value is not None:
        return resolved_input_path(clean_user_path(supplied_value), label)

    while True:
        try:
            entered = input(prompt)
        except EOFError as error:
            raise ValueError(f"No path was supplied for {label.lower()}.") from error
        try:
            return resolved_input_path(clean_user_path(entered), label)
        except (FileNotFoundError, ValueError) as error:
            print(f"Error: {error}", file=sys.stderr)
            print("Please enter the full path again.", file=sys.stderr)


def default_output_path(report_path: Path) -> Path:
    base_name = f"{report_path.stem}_with_erc_availability"
    candidate = report_path.with_name(f"{base_name}.xlsx")
    counter = 2
    while candidate.exists():
        candidate = report_path.with_name(f"{base_name}_{counter}.xlsx")
        counter += 1
    return candidate.resolve()


def paths_refer_to_same_file(first: Path, second: Path) -> bool:
    if first == second:
        return True
    try:
        return first.exists() and second.exists() and first.samefile(second)
    except OSError:
        return False


def save_workbook_atomically(workbook, output_path: Path) -> None:
    """Save to a temporary sibling, then atomically move it into place."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}_",
            suffix=SUPPORTED_EXTENSION,
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
        temporary_path = None
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def locate_exact_headers(
    worksheet: Worksheet,
    required_headers: Iterable[str],
    *,
    header_row: int = 1,
) -> dict[str, int]:
    """Locate required row-1 headers and reject missing or duplicate labels."""
    columns: dict[str, list[int]] = defaultdict(list)
    for cell in worksheet[header_row]:
        key = normalized_header(cell.value)
        if key:
            columns[key].append(cell.column)

    result: dict[str, int] = {}
    missing: list[str] = []
    duplicates: list[str] = []
    for header in required_headers:
        matches = columns.get(normalized_header(header), [])
        if not matches:
            missing.append(header)
        elif len(matches) > 1:
            duplicates.append(header)
        else:
            result[header] = matches[0]

    problems: list[str] = []
    if missing:
        problems.append(f"missing header(s): {', '.join(missing)}")
    if duplicates:
        problems.append(f"duplicate header(s): {', '.join(duplicates)}")
    if problems:
        raise ValueError(
            f"Sheet '{worksheet.title}' has " + "; ".join(problems) + "."
        )
    return result


def locate_source_headers(worksheet: Worksheet) -> dict[str, int]:
    """Locate the logical ERC/source columns using accepted header aliases."""
    columns: dict[str, list[int]] = defaultdict(list)
    for cell in worksheet[1]:
        key = normalized_header(cell.value)
        if key:
            columns[key].append(cell.column)

    result: dict[str, int] = {}
    problems: list[str] = []
    for logical_name, aliases in SOURCE_HEADER_ALIASES.items():
        found: list[int] = []
        for alias in aliases:
            found.extend(columns.get(normalized_header(alias), []))
        found = sorted(set(found))
        if not found:
            problems.append(f"missing {' or '.join(aliases)}")
        elif len(found) > 1:
            problems.append(f"duplicate/ambiguous {' or '.join(aliases)}")
        else:
            result[logical_name] = found[0]

    if problems:
        raise ValueError(
            f"Sheet '{worksheet.title}' has " + "; ".join(problems) + "."
        )
    return result


def load_erc_records(
    erc_path: Path,
    stats: AvailabilityStats,
) -> list[ErcRecord]:
    """Read every nonblank ERC data row from the first worksheet."""
    workbook = load_workbook(erc_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        headers = locate_exact_headers(worksheet, ("v1", "v2", "v3"))
        value_indexes = {
            header: column_number - 1
            for header, column_number in headers.items()
        }
        last_required_column = max(headers.values())
        records: list[ErcRecord] = []
        source_order = 0

        # A read-only worksheet is a forward-only XML stream. Repeated
        # worksheet.cell(row, column) calls can rescan that stream and become
        # extremely slow as the row count grows. Read each ERC row once instead.
        for row_number, row_values in enumerate(
            worksheet.iter_rows(
                min_row=2,
                max_col=last_required_column,
                values_only=True,
            ),
            start=2,
        ):
            v1 = row_values[value_indexes["v1"]]
            v2 = row_values[value_indexes["v2"]]
            v3 = row_values[value_indexes["v3"]]
            if all(is_blank(value) for value in (v1, v2, v3)):
                continue

            source_order += 1
            stats.erc_rows += 1
            location = split_erc_location(v1)
            district = location[0] if location else ""
            school = location[1] if location else ""
            section = section_display(v2)
            student = cell_text(v3)

            problems: list[str] = []
            if location is None:
                problems.append("v1 is not in 'District / School' format")
            if not student:
                problems.append("v3 student name is blank")
            if not section:
                problems.append("v2 section is blank")

            key = make_match_key(district, school, student)
            if key is None and not problems:
                problems.append("district, school, or student is empty after normalization")

            problem = "; ".join(problems)
            if problem:
                stats.invalid_erc_rows += 1
                stats.warnings.append(
                    f"ERC sheet '{worksheet.title}' row {row_number}: {problem}."
                )
            else:
                stats.usable_erc_rows += 1

            records.append(
                ErcRecord(
                    source_order=source_order,
                    district=district,
                    school=school,
                    section=section,
                    student=student,
                    key=None if problem else key,
                    problem=problem,
                )
            )
        return records
    finally:
        workbook.close()


def build_source_index(
    worksheet: Worksheet,
    sheet_number: int,
    stats: AvailabilityStats,
) -> SourceSheetIndex:
    """Index eligible rows using only populated ERC match fields."""
    headers = locate_source_headers(worksheet)
    grouped: dict[MatchKey, list[ReportOccurrence]] = defaultdict(list)

    for row_number in range(2, worksheet.max_row + 1):
        erc_section_value = worksheet.cell(
            row_number, headers["erc_section"]
        ).value

        # The agreed rule is explicit: blank ERC Section means not matched.
        if is_blank(erc_section_value):
            continue

        erc_district = cell_text(
            worksheet.cell(row_number, headers["erc_district"]).value
        )
        erc_school = cell_text(
            worksheet.cell(row_number, headers["erc_school"]).value
        )
        erc_student = cell_text(
            worksheet.cell(row_number, headers["erc_student"]).value
        )
        erc_section = section_display(erc_section_value)
        source_student = cell_text(
            worksheet.cell(row_number, headers["source_student"]).value
        )
        key = make_match_key(erc_district, erc_school, erc_student)

        if key is None:
            stats.incomplete_report_rows += 1
            stats.warnings.append(
                f"Report sheet '{worksheet.title}' row {row_number} has a "
                "nonblank ERC Section but an empty ERC district, school, or "
                "matched student name; the row was ignored."
            )
            continue

        stats.eligible_report_rows += 1
        grouped[key].append(
            ReportOccurrence(
                sheet_number=sheet_number,
                sheet_name=worksheet.title,
                row_number=row_number,
                source_student=source_student,
                erc_district=erc_district,
                erc_school=erc_school,
                erc_student=erc_student,
                erc_section=erc_section,
            )
        )

    return SourceSheetIndex(
        worksheet=worksheet,
        sheet_number=sheet_number,
        headers=headers,
        rows_by_key={
            key: tuple(sorted(rows, key=lambda item: item.row_number))
            for key, rows in grouped.items()
        },
    )


def style_header(worksheet: Worksheet, column_count: int) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 34
    for cell in worksheet.iter_cols(
        min_row=1,
        max_row=1,
        min_col=1,
        max_col=column_count,
    ):
        header_cell = cell[0]
        header_cell.fill = HEADER_FILL
        header_cell.font = HEADER_FONT
        header_cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )


def set_status_style(cell, fill: PatternFill, font: Font) -> None:
    cell.fill = copy(fill)
    cell.font = copy(font)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def copy_source_color(
    source_index: SourceSheetIndex,
    source_row: int,
    target_cell,
) -> None:
    """Copy display color only; never read the source-name value for matching."""
    source_cell = source_index.worksheet.cell(
        source_row, source_index.headers["source_student"]
    )
    target_cell.fill = copy(source_cell.fill)
    target_cell.font = copy(source_cell.font)
    target_cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )


def write_availability_sheet(
    worksheet: Worksheet,
    records: Sequence[ErcRecord],
    source_indexes: Sequence[SourceSheetIndex],
    erc_groups: dict[MatchKey, tuple[ErcRecord, ...]],
    stats: AvailabilityStats,
) -> None:
    dynamic_headers = tuple(
        f"Sheet {index.sheet_number}: {index.worksheet.title}"
        for index in source_indexes
    )
    worksheet.append(AVAILABILITY_FIXED_HEADERS + dynamic_headers)

    for record in records:
        worksheet.append(
            (
                record.district,
                record.school,
                record.section,
                record.student,
            )
            + ("",) * len(source_indexes)
        )
        output_row = worksheet.max_row

        for offset, source_index in enumerate(source_indexes, start=5):
            result_cell = worksheet.cell(output_row, offset)

            if record.key is None:
                result_cell.value = "Invalid ERC data"
                set_status_style(result_cell, INVALID_FILL, INVALID_FONT)
                continue

            if len(erc_groups[record.key]) > 1:
                result_cell.value = "Duplicate ERC key — see Sheet 8"
                set_status_style(result_cell, DUPLICATE_FILL, DUPLICATE_FONT)
                stats.duplicate_cells += 1
                continue

            occurrences = source_index.rows_by_key.get(record.key, ())
            if not occurrences:
                result_cell.value = "Not Present"
                set_status_style(result_cell, NOT_PRESENT_FILL, NOT_PRESENT_FONT)
                stats.absent_cells += 1
            elif len(occurrences) == 1:
                # Matching has already completed using only the exact-normalized
                # ERC key. The original source name is read now solely so Sheet
                # 7 shows the spelling that can be searched in that source sheet.
                result_cell.value = occurrences[0].source_student
                copy_source_color(
                    source_index, occurrences[0].row_number, result_cell
                )
                stats.present_cells += 1
            else:
                result_cell.value = "Duplicate — see Sheet 8"
                set_status_style(result_cell, DUPLICATE_FILL, DUPLICATE_FONT)
                stats.duplicate_cells += 1

    style_header(worksheet, len(AVAILABILITY_FIXED_HEADERS) + len(source_indexes))
    for row_number in range(2, worksheet.max_row + 1):
        for column in range(1, 5):
            cell = worksheet.cell(row_number, column)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.cell(row_number, 3).alignment = Alignment(
            horizontal="center", vertical="center"
        )

    widths = {"A": 22, "B": 42, "C": 13, "D": 31}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    for column_number in range(5, 5 + len(source_indexes)):
        worksheet.column_dimensions[get_column_letter(column_number)].width = 32

    last_column = get_column_letter(
        len(AVAILABILITY_FIXED_HEADERS) + len(source_indexes)
    )
    worksheet.auto_filter.ref = f"A1:{last_column}{worksheet.max_row}"


def write_duplicate_sheet(
    worksheet: Worksheet,
    erc_groups: dict[MatchKey, tuple[ErcRecord, ...]],
    source_indexes: Sequence[SourceSheetIndex],
    stats: AvailabilityStats,
) -> None:
    worksheet.append(DUPLICATE_HEADERS)
    group_number = 0

    duplicate_erc_groups = [
        records for records in erc_groups.values() if len(records) > 1
    ]
    duplicate_erc_groups.sort(key=lambda records: records[0].source_order)

    for records in duplicate_erc_groups:
        group_number += 1
        group_id = f"ERC-{group_number:04d}"
        stats.erc_duplicate_groups += 1
        for occurrence_number, record in enumerate(records, start=1):
            worksheet.append(
                (
                    "Duplicate in ERC Full List",
                    group_id,
                    occurrence_number,
                    len(records),
                    record.district,
                    record.school,
                    record.student,
                    record.section,
                    "",
                    "",
                    "",
                    "Section is ignored, so these ERC records have the same key.",
                )
            )
            stats.duplicate_report_rows += 1

    erc_record_by_key = {
        key: records[0] for key, records in erc_groups.items()
    }
    for source_index in source_indexes:
        duplicate_items = [
            (key, occurrences)
            for key, occurrences in source_index.rows_by_key.items()
            if len(occurrences) > 1
        ]
        duplicate_items.sort(
            key=lambda item: min(row.row_number for row in item[1])
        )

        for key, occurrences in duplicate_items:
            group_number += 1
            group_id = f"REPORT-{group_number:04d}"
            stats.report_duplicate_groups += 1
            erc_record = erc_record_by_key.get(key)
            for occurrence_number, occurrence in enumerate(
                occurrences, start=1
            ):
                worksheet.append(
                    (
                        "Multiple rows in report sheet",
                        group_id,
                        occurrence_number,
                        len(occurrences),
                        erc_record.district if erc_record else occurrence.erc_district,
                        erc_record.school if erc_record else occurrence.erc_school,
                        erc_record.student if erc_record else occurrence.erc_student,
                        erc_record.section if erc_record else "",
                        occurrence.sheet_name,
                        occurrence.row_number,
                        occurrence.erc_section,
                        (
                            "More than one eligible row has this section-free ERC key."
                            if erc_record
                            else "The duplicated report key is not present in the supplied ERC list."
                        ),
                    )
                )
                stats.duplicate_report_rows += 1

    if stats.duplicate_report_rows == 0:
        worksheet.append(
            (
                "No duplicates found",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "No duplicate ERC keys or multiple report-sheet matches were found.",
            )
        )

    style_header(worksheet, len(DUPLICATE_HEADERS))
    group_colors: dict[str, PatternFill] = {}
    for row_number in range(2, worksheet.max_row + 1):
        group_id = cell_text(worksheet.cell(row_number, 2).value)
        if group_id and group_id not in group_colors:
            group_colors[group_id] = (
                GROUP_FILL_A
                if len(group_colors) % 2 == 0
                else GROUP_FILL_B
            )
        group_fill = group_colors.get(group_id)
        for cell in worksheet[row_number]:
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if group_fill is not None:
                cell.fill = copy(group_fill)

        for column in (2, 3, 4, 8, 10, 11):
            worksheet.cell(row_number, column).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

    widths = {
        "A": 30,
        "B": 18,
        "C": 18,
        "D": 18,
        "E": 22,
        "F": 42,
        "G": 31,
        "H": 13,
        "I": 28,
        "J": 13,
        "K": 18,
        "L": 58,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    worksheet.auto_filter.ref = f"A1:L{worksheet.max_row}"


def verify_output(
    output_path: Path,
    expected_availability_rows: int,
    expected_duplicate_rows: int,
) -> None:
    """Reopen the saved file and verify the two newly written sheets."""
    workbook = load_workbook(output_path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) < 8:
            raise RuntimeError("The saved output has fewer than eight sheets.")
        if workbook.worksheets[6].title != AVAILABILITY_SHEET_TITLE:
            raise RuntimeError("Sheet 7 was not saved with the expected title.")
        if workbook.worksheets[7].title != DUPLICATE_SHEET_TITLE:
            raise RuntimeError("Sheet 8 was not saved with the expected title.")
        if workbook.worksheets[6].max_row != expected_availability_rows + 1:
            raise RuntimeError("The saved Sheet 7 row count is incorrect.")
        if workbook.worksheets[7].max_row != max(1, expected_duplicate_rows) + 1:
            raise RuntimeError("The saved Sheet 8 row count is incorrect.")
    finally:
        workbook.close()


def create_erc_availability_report(
    erc_path: Path,
    report_path: Path,
    output_path: Path,
) -> AvailabilityStats:
    """Create a new report copy containing Sheets 7 and 8."""
    erc_path = erc_path.expanduser().resolve()
    report_path = report_path.expanduser().resolve()
    output_path = output_path.expanduser().resolve()

    for path, label in (
        (erc_path, "ERC full-list workbook"),
        (report_path, "six-sheet report workbook"),
    ):
        if path.suffix.lower() != SUPPORTED_EXTENSION:
            raise ValueError(f"The {label} must be an .xlsx file.")
        if not path.is_file():
            raise FileNotFoundError(f"The {label} does not exist: {path}")

    if output_path.suffix.lower() != SUPPORTED_EXTENSION:
        raise ValueError("The output workbook must use the .xlsx extension.")
    if paths_refer_to_same_file(output_path, erc_path) or paths_refer_to_same_file(
        output_path, report_path
    ):
        raise ValueError("The output path must differ from both input paths.")
    if output_path.exists():
        raise FileExistsError(
            f"Output already exists and will not be overwritten: {output_path}"
        )

    stats = AvailabilityStats(
        erc_input_sha256=sha256_file(erc_path),
        report_input_sha256=sha256_file(report_path),
    )
    records = load_erc_records(erc_path, stats)
    erc_groups_list: dict[MatchKey, list[ErcRecord]] = defaultdict(list)
    for record in records:
        if record.key is not None:
            erc_groups_list[record.key].append(record)
    erc_groups = {
        key: tuple(sorted(group, key=lambda item: item.source_order))
        for key, group in erc_groups_list.items()
    }

    workbook = load_workbook(report_path, data_only=False, keep_links=True)
    try:
        if len(workbook.worksheets) < MINIMUM_EXISTING_SHEETS:
            raise ValueError(
                "The report workbook must contain at least six existing sheets."
            )
        for title in (AVAILABILITY_SHEET_TITLE, DUPLICATE_SHEET_TITLE):
            if title in workbook.sheetnames:
                raise ValueError(
                    f"The report already contains '{title}'; it was not replaced."
                )

        source_indexes = [
            build_source_index(worksheet, sheet_number, stats)
            for sheet_number, worksheet in enumerate(
                workbook.worksheets[:SOURCE_SHEET_COUNT], start=1
            )
        ]

        # Insert at positions 7 and 8 so any later optional sheets are retained
        # and shifted, while the first six existing sheets remain unchanged.
        availability_sheet = workbook.create_sheet(
            AVAILABILITY_SHEET_TITLE, index=6
        )
        duplicate_sheet = workbook.create_sheet(DUPLICATE_SHEET_TITLE, index=7)

        write_availability_sheet(
            availability_sheet,
            records,
            source_indexes,
            erc_groups,
            stats,
        )
        write_duplicate_sheet(
            duplicate_sheet,
            erc_groups,
            source_indexes,
            stats,
        )

        if sha256_file(erc_path) != stats.erc_input_sha256:
            raise RuntimeError(
                "The ERC input changed during processing; no output was written."
            )
        if sha256_file(report_path) != stats.report_input_sha256:
            raise RuntimeError(
                "The report input changed during processing; no output was written."
            )

        save_workbook_atomically(workbook, output_path)
    finally:
        workbook.close()

    verify_output(
        output_path,
        expected_availability_rows=len(records),
        expected_duplicate_rows=stats.duplicate_report_rows,
    )

    if sha256_file(erc_path) != stats.erc_input_sha256:
        raise RuntimeError("The ERC input changed while the output was saved.")
    if sha256_file(report_path) != stats.report_input_sha256:
        raise RuntimeError("The report input changed while the output was saved.")
    return stats


def parse_arguments(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Copy a six-sheet student report and add ERC availability and "
            "duplicate-detail sheets without fuzzy matching."
        )
    )
    parser.add_argument(
        "erc_workbook",
        nargs="?",
        help="Full path of the ERC full-list .xlsx workbook",
    )
    parser.add_argument(
        "report_workbook",
        nargs="?",
        help="Full path of the latest six-sheet .xlsx report",
    )
    parser.add_argument(
        "output_workbook",
        nargs="?",
        help="Optional full path for the new .xlsx workbook",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_arguments(argv)
    try:
        erc_path = request_input_path(
            args.erc_workbook,
            prompt="Enter the full path of the ERC full-list Excel file: ",
            label="ERC full-list workbook",
        )
        report_path = request_input_path(
            args.report_workbook,
            prompt="Enter the full path of the latest six-sheet report: ",
            label="Six-sheet report workbook",
        )
        output_path = (
            resolved_output_path(clean_user_path(args.output_workbook))
            if args.output_workbook is not None
            else default_output_path(report_path)
        )
        if args.output_workbook is None:
            print(f"The new workbook will be created at: {output_path}")

        stats = create_erc_availability_report(
            erc_path, report_path, output_path
        )
    except (
        FileNotFoundError,
        FileExistsError,
        PermissionError,
        RuntimeError,
        ValueError,
    ) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 2

    for warning in stats.warnings:
        print(f"WARNING: {warning}", file=sys.stderr)

    print(
        f"ERC rows reported: {stats.erc_rows} | "
        f"usable={stats.usable_erc_rows} | invalid={stats.invalid_erc_rows}"
    )
    print(
        f"Sheet-7 results: present={stats.present_cells} | "
        f"not present={stats.absent_cells} | "
        f"duplicate/ambiguous={stats.duplicate_cells}"
    )
    print(
        f"Sheet-8 duplicate groups: ERC={stats.erc_duplicate_groups} | "
        f"report sheets={stats.report_duplicate_groups} | "
        f"reported rows={stats.duplicate_report_rows}"
    )
    print(
        "Both input workbooks are unchanged (SHA-256 verified).\n"
        f"New workbook written to: {output_path}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
