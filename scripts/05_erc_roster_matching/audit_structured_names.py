#!/usr/bin/env python3
"""Audit historical mismatches, or explicitly choose a paired calibration audit.

Both XLSX inputs are read-only. Prints per-sheet red-to-green improvement rates
and writes local JSON for the companion spreadsheet renderer; never edits
grades, changes mappings, or exports private data online.
The optional paired mode is a calibration subset, never a replacement for the
historical red cohort in the original report.
"""
from __future__ import annotations
import argparse
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from student_match_report import BaseMatcher, normalize_for_match as norm, normalize_section, read_base_records
from structured_name_match import assess_existing_pair


def file_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def baseline_red_status(matcher, district, school, section, student, target, pair_score):
    """Return the original matcher's red category, or an empty string.

    An existing pair above the threshold proves a non-red result only if its
    ERC record belongs to the resolved school. Otherwise run the original full
    classifier: a low pair score alone does not establish a red source row.
    The shortcut avoids rescoring the whole school for most accepted pairs.
    """
    location = matcher.resolve_location(district, school)
    if not location.accepted or location.location is None:
        return 'red_location'
    if (target is not None and pair_score > matcher.student_threshold
            and (target.district, target.school) ==
                (location.location.district, location.location.school)):
        return ''
    outcome = matcher.classify(district, school, section, student)
    return outcome.status if outcome.status in {'red_name', 'red_location'} else ''


def summarize_red_improvement(rows, sheet_names=()):
    """Count records, using every baseline red row as the denominator.

    Rates are percentages on a 0-100 scale; zero denominators produce None.
    Overall rates use summed counts, never an average of sheet percentages.
    Yellow promotions reduce red workload but are not red-to-green successes.
    """
    keys = ('records', 'baseline_red', 'red_name', 'red_location',
            'red_to_green', 'red_to_yellow', 'red_score_increases',
            'name_red_to_green', 'location_red_to_green')
    by_sheet = {name: dict.fromkeys(keys, 0) for name in sheet_names}
    for row in rows:
        counts = by_sheet.setdefault(row['sheet'], dict.fromkeys(keys, 0))
        counts['records'] += 1
        if not row['baseline_is_red']:
            continue
        reason = row['baseline_red_reason']
        if reason not in {'red_name', 'red_location'}:
            raise ValueError(f"Missing baseline red reason on {row['sheet']} row {row['row']}")
        counts['baseline_red'] += 1
        counts[reason] += 1
        counts['red_score_increases'] += row['adjusted'] > row['baseline']
        # Stored scores use a 0-1 scale. Crossing the exclusive 90% cutoff is
        # necessary, while the guarded decision also checks context/ambiguity.
        promoted = row['adjusted'] > row['baseline'] and row['adjusted'] > 0.9
        counts['red_to_green'] += promoted and row['decision'] == 'Green + review'
        counts['red_to_yellow'] += promoted and row['decision'] == 'Yellow + review'
        counts['name_red_to_green'] += promoted and row['decision'] == 'Green + review' and reason == 'red_name'
        counts['location_red_to_green'] += promoted and row['decision'] == 'Green + review' and reason == 'red_location'

    def with_rates(counts):
        before = counts['baseline_red']
        resolved = counts['red_to_green'] + counts['red_to_yellow']
        return {
            **counts,
            'red_remaining': before - resolved,
            'red_to_green_percent': 100.0 * counts['red_to_green'] / before if before else None,
            'red_reduction_percent': 100.0 * resolved / before if before else None,
            'name_red_to_green_percent': (100.0 * counts['name_red_to_green'] / counts['red_name']
                                          if counts['red_name'] else None),
        }

    overall = {key: sum(counts[key] for counts in by_sheet.values()) for key in keys}
    return {
        'definition': '100 * baseline-red records becoming Green + review / all baseline-red records',
        'denominator': 'Original automatic matcher: red_name or red_location; not the existing manual fill color',
        'scope': 'Records in the supplied paired workbook; repeated grade periods count separately',
        'per_sheet': {name: with_rates(counts) for name, counts in by_sheet.items()},
        'overall': with_rates(overall),
    }


def print_red_improvement(report):
    print('\nImprovement among baseline red records (red -> green / baseline red):')
    width = max([len('Worksheet'), len('OVERALL'), *map(len, report['per_sheet'])])
    print(f"{'Worksheet':<{width}}  {'Red before':>10}  {'To green':>9}  {'Red left':>9}  "
          f"{'All-red rate':>12}  {'Name-only rate':>14}")
    for name, counts in [*report['per_sheet'].items(), ('OVERALL', report['overall'])]:
        rate = counts['red_to_green_percent']
        display = 'n.a.' if rate is None else f'{rate:.2f}%'
        name_rate = counts['name_red_to_green_percent']
        name_display = 'n.a.' if name_rate is None else f'{name_rate:.2f}%'
        print(f"{name:<{width}}  {counts['baseline_red']:>10}  {counts['red_to_green']:>9}  "
              f"{counts['red_remaining']:>9}  {display:>12}  {name_display:>14}")
    print('All-red rate includes location failures; name-only rate uses red_name records only.')
    print('All newly green records still require manual review. Counts are records, not unique students.')


def audit(source: Path, roster: Path, output: Path):
    if output.resolve() in (source.resolve(),roster.resolve()):
        raise ValueError("Output must not overwrite an input")
    if output.exists():
        raise FileExistsError(f"Output already exists: {output}")
    before_hashes = {str(p.resolve()): file_hash(p) for p in (source,roster)}
    records, load_stats, _ = read_base_records(roster,header_row=1)
    matcher = BaseMatcher(records,student_threshold=90,location_threshold=95)
    by_key = defaultdict(list)
    for r in records:
        by_key[r.district,r.school,r.section,r.student].append(r)
    result_rows, review_rows, rival_rows = [],[],[]
    counts, patterns, per_sheet = Counter(),Counter(),defaultdict(Counter)
    identities = defaultdict(set)
    book = load_workbook(source,read_only=True,data_only=True)
    source_sheets = list(book.sheetnames)
    try:
        for sheet in book:
            headers = [norm(c) for c in next(sheet.iter_rows(max_row=1,values_only=True))]
            expected = {0:"erc ilce",1:"erc okul",3:"erc sube",4:"erc ad soyad",
                        5:"raw data ilce",6:"raw data okul",7:"raw data sube",8:"raw data ad soyad"}
            if any(len(headers)<=i or headers[i]!=v for i,v in expected.items()):
                raise ValueError(f"Unexpected paired-workbook headers on {sheet.title}")
            for source_row,cells in enumerate(sheet.iter_rows(min_row=2,max_col=9),2):
                raw = [c.value for c in cells]
                if not any(v is not None for v in raw):
                    continue
                target_key = norm(raw[0]),norm(raw[1]),normalize_section(raw[3]),norm(raw[4])
                targets = by_key.get(target_key,[])
                target = targets[0] if len(targets)==1 else None
                assessment = assess_existing_pair(matcher,norm(raw[5]),norm(raw[6]),
                                                  normalize_section(raw[7]),norm(raw[8]),target,norm(raw[4]))
                red_reason = baseline_red_status(
                    matcher, norm(raw[5]), norm(raw[6]), normalize_section(raw[7]),
                    norm(raw[8]), target, assessment.baseline,
                )
                ev = assessment.evidence
                increased = assessment.adjusted > assessment.baseline
                crossed = increased and assessment.baseline <= 90 < assessment.adjusted
                green_crossed = crossed and assessment.decision == "Green + review"
                counts['records'] += 1
                counts['baseline_at_or_below_90'] += assessment.baseline<=90
                counts['score_increases'] += increased
                counts['score_decreases'] += assessment.adjusted<assessment.baseline
                counts['crossed_90'] += crossed
                counts['promoted_green'] += green_crossed
                counts['promoted_yellow'] += crossed and assessment.decision == "Yellow + review"
                counts['unrelated_records'] += not bool(ev.pattern)
                counts['unrelated_changed'] += not ev.pattern and increased
                counts['unchanged_scores'] += not increased
                counts['flagged'] += bool(ev.pattern)
                counts['review_only'] += assessment.decision == "Review only"
                per_sheet[sheet.title]['records'] += 1
                per_sheet[sheet.title]['flagged'] += bool(ev.pattern)
                per_sheet[sheet.title]['increased'] += increased
                per_sheet[sheet.title]['promoted_green'] += green_crossed
                pair_id = (norm(raw[0]),norm(raw[1]),normalize_section(raw[3]),norm(raw[4]),norm(raw[8]))
                if ev.pattern:
                    patterns[ev.pattern] += 1
                    identities['flagged_pairs'].add(pair_id)
                if green_crossed:
                    identities['promoted_green_pairs'].add(pair_id)
                if increased:
                    identities['increased_pairs'].add(pair_id)
                item = {
                    'sheet': sheet.title, 'row':source_row,
                    'source_name':raw[8], 'erc_name':raw[4],
                    'baseline':assessment.baseline/100,'adjusted':assessment.adjusted/100,
                    'rule_score': ev.rule_score/100 if ev.rule_score is not None else None,
                    'pattern':ev.pattern,'decision':assessment.decision,
                    'reason':assessment.reason,'baseline_status':assessment.baseline_status,
                    'baseline_is_red':bool(red_reason),'baseline_red_reason':red_reason,
                    'erc_district':raw[0],'erc_school':raw[1],'erc_section':raw[3],
                    'source_district':raw[5],'source_school':raw[6],'source_section':raw[7],
                    'erc_row':target.row_number if target else None,
                    'original_fill':cells[8].fill.fgColor.rgb if cells[8].fill.fgColor.type=='rgb' else '',
                    'original_candidates':assessment.baseline_candidates,
                    'augmented_candidates':assessment.augmented_candidates,
                    'rival_score':assessment.rival_score/100 if assessment.rival_score is not None else None,
                    'initials':ev.initials,'omitted':ev.omitted,'anchors':ev.anchors,
                    'fuzzy_parts':ev.fuzzy_parts,
                    'full_parts_similarity':(ev.full_parts_similarity/100
                                             if ev.full_parts_similarity is not None else None),
                    'crossed':crossed, 'green_crossed':green_crossed,
                }
                result_rows.append(item)
                if ev.pattern:
                    review_rows.append(len(result_rows))
                    for r,old,new,pattern in assessment.rivals:
                        rival_rows.append({'sheet':sheet.title,'row':source_row,'source_name':raw[8],
                                           'mapped_name':raw[4],'rival_row':r.row_number,
                                           'rival_name':r.display_student,'rival_section':r.display_section,
                                           'baseline':old/100,'proposed':new/100,'pattern':pattern})
            print(f"Audited {sheet.title}: {per_sheet[sheet.title]['records']} rows",flush=True)
    finally:
        book.close()
    after_hashes = {str(p.resolve()):file_hash(p) for p in (source,roster)}
    if before_hashes!=after_hashes or counts['score_decreases'] or counts['unrelated_changed']:
        raise AssertionError("Preservation check failed")
    summary = dict(counts)
    summary.update({k:len(v) for k,v in identities.items()})
    red_improvement = summarize_red_improvement(result_rows, source_sheets)
    for name, metrics in red_improvement['per_sheet'].items():
        per_sheet[name] = {**per_sheet[name], **metrics}
    payload = {'population':'paired_calibration_subset',
               'source':str(source.resolve()),'roster':str(roster.resolve()),
               'hashes':before_hashes,'roster_rows':load_stats.usable_rows,
               'summary':summary,'patterns':patterns,'per_sheet':dict(per_sheet),
               'red_improvement':red_improvement,
               'rows':result_rows,'review_rows':review_rows,'rivals':rival_rows}
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(json.dumps(payload,ensure_ascii=False,indent=1),encoding='utf-8')
    print_red_improvement(red_improvement)
    print(f'Full audit saved to {output}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input_workbook',type=Path)
    parser.add_argument('erc_roster',type=Path)
    parser.add_argument('output_json',type=Path)
    parser.add_argument('--population',choices=('mismatches','paired'),default='mismatches',
                        help='Historical mismatch sheet by default; paired is only a calibration subset')
    args = parser.parse_args()
    if args.population == 'paired':
        audit(args.input_workbook,args.erc_roster,args.output_json)
    else:
        from audit_mismatched_students import audit_mismatches
        audit_mismatches(args.input_workbook,args.erc_roster,args.output_json)
