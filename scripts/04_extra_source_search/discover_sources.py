"""Find possible grade sources in PDF and spreadsheet folders.

This script was added after the original folder search. It lists candidates
for review and does not change the source files. Run without arguments to
enter paths interactively, or run --help for command-line options.
"""
from __future__ import annotations

import argparse
import collections
import hashlib
import itertools
import json
import logging
import os
from pathlib import Path
import re
import sys
import unicodedata
from zipfile import ZipFile
import xml.etree.ElementTree as ET

SUPPORTED = {'.pdf', '.xlsx', '.xlsm', '.xls', '.csv', '.tsv', '.docx', '.udf', '.rtf', '.doc'}
SKIP_DIRS = {'.git', '__pycache__', 'node_modules', '.venv'}


def normalize(value):
    """Loose keyword comparison only. Do not use this for section identity."""
    text = unicodedata.normalize('NFKD', str(value or '')).casefold().replace('ı', 'i')
    text = ''.join(c for c in text if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def section_letters(text):
    # Turkish case mapping deliberately keeps I and İ distinct.
    text = unicodedata.normalize('NFC', str(text)).replace('i', 'İ').replace('ı', 'I').upper()
    return sorted(set(re.findall(r'(?<!\d)[234]\s*[/\-]\s*([A-ZÇĞİÖŞÜ])(?![A-ZÇĞİÖŞÜ])', text)))


def evidence(text):
    normalized = normalize(text)
    grades = sorted(set(re.findall(r'\b([234])\s*sinif\b', normalized)))
    terms = sorted(set(re.findall(r'\b([12])\s*donem\b', normalized)))
    for grade in re.findall(r'(?<!\d)([234])\s*[/\-]\s*[A-Za-zÇĞİÖŞÜçğıöşü]', text):
        if grade not in grades: grades.append(grade)
    headings = [s for s in ('turkce', 'matematik', 'fen bilimleri', 'hayat bilgisi', 'ders', 'not', 'puan', 'davranis', 'karne', 'ortalama')
                if re.search(r'\b'+re.escape(s)+r'\b', normalized)]
    return {'grades': sorted(grades), 'terms': terms, 'sections': section_letters(text), 'grade_headings': headings}


def sha256(path):
    result = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024*1024), b''): result.update(chunk)
    return result.hexdigest()


def school_keyword(value):
    # The generic school suffix is not distinctive. No fuzzy identity match is
    # invented here; the user can pass an alias separately when needed.
    return re.sub(r'\s+ilkokulu$', '', normalize(value))


def missing_entries(report):
    """Read the three ERC missing-section sheets with their source row IDs."""
    if not report:
        return []
    from openpyxl import load_workbook
    book = load_workbook(report, read_only=True, data_only=True)
    entries = []
    try:
        for sheet in book.worksheets:
            title = normalize(sheet.title)
            grade_match = re.search(r'\b([234])\s*sinif\b', title)
            grade = grade_match.group(1) if grade_match else ''
            term_match = re.search(r'\b([12])\s*d', title)
            term = term_match.group(1) if term_match else ''
            if grade == '3' and not term:
                # Some source workbook titles contain a replacement character
                # in DÖNEM; the standalone term digit still survives.
                term = next((x for x in re.findall(r'\b[12]\b', title)), '')
            for rowno, row in enumerate(sheet.iter_rows(min_row=2, max_col=3, values_only=True), 2):
                if not all(row):
                    continue
                section = str(row[2]).strip().translate(str.maketrans({'i':'İ','ı':'I'})).upper()
                section = section.split('/')[-1].strip()
                if len(section) != 1:
                    continue
                entries.append({'report_sheet':sheet.title, 'report_row':rowno,
                                'district':str(row[0]).strip(), 'school':str(row[1]).strip(),
                                'section':section, 'grade':grade, 'term':term,
                                'keyword':school_keyword(row[1])})
    finally:
        book.close()
    return entries


def school_keywords(entries, extra):
    return sorted({school_keyword(t) for t in extra if school_keyword(t)} |
                  {entry['keyword'] for entry in entries if entry['keyword']})


def units(path, args):
    """Yield (kind, sheet, position, text, multiple_sheets)."""
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        from pypdf import PdfReader
        with path.open('rb') as stream:
            reader = PdfReader(stream)
            for index, page in enumerate(reader.pages, 1):
                if args.max_pdf_pages and index > args.max_pdf_pages: break
                yield 'pdf', '', index, page.extract_text() or '', False
    elif suffix in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            multiple_sheets = len(book.worksheets) > 1
            for sheet in book.worksheets:
                # Some workbooks claim that row 1,048,576 is used although
                # their actual XML ends much earlier. Read the real rows.
                sheet.reset_dimensions()
                rows = sheet.iter_rows(values_only=True)
                if args.max_excel_rows: rows = itertools.islice(rows, args.max_excel_rows)
                for rowno, row in enumerate(rows, 1):
                    line=' | '.join(str(v) for v in row if v is not None and str(v).strip())
                    if line:
                        yield 'excel', sheet.title, rowno, line, multiple_sheets
        finally: book.close()
    elif suffix=='.xls':
        import xlrd
        book=xlrd.open_workbook(path, on_demand=True)
        try:
            multiple_sheets = book.nsheets > 1
            for sheet in book.sheets():
                limit=min(sheet.nrows,args.max_excel_rows or sheet.nrows)
                for rowno in range(1,limit+1):
                    line=' | '.join(str(v) for v in sheet.row_values(rowno-1) if v is not None and str(v).strip())
                    if line:
                        yield 'excel', sheet.name, rowno, line, multiple_sheets
        finally: book.release_resources()
    elif suffix in {'.csv','.tsv'}:
        content=path.read_text(encoding=args.text_encoding)
        yield 'file', '', 0, content, False
    elif suffix in {'.docx','.udf'}:
        with ZipFile(path) as archive:
            for name in archive.namelist():
                if (suffix=='.docx' and name=='word/document.xml') or (suffix=='.udf' and name.lower().endswith('.xml')):
                    root=ET.fromstring(archive.read(name))
                    yield 'file', '', 0, ' '.join(root.itertext()), False
    elif suffix=='.rtf':
        # RTF needs a real decoder for dependable Turkish characters.
        raise ValueError('RTF requires conversion to text/ DOCX; not silently treated as no match')
    else:
        raise ValueError('Legacy DOC requires conversion to DOCX or PDF')


def consecutive_ranges(numbers):
    """Return exact contiguous row or page ranges, without fixed-size blocks."""
    ordered = sorted(set(numbers))
    if not ordered:
        return
    start = end = ordered[0]
    for number in ordered[1:]:
        if number == end + 1:
            end = number
        else:
            yield start, end
            start = end = number
    yield start, end


def candidate_location(kind, sheet, multiple_sheets, positions):
    if kind == 'excel':
        ranges = [f'row {start}' if start == end else f'rows {start}-{end}'
                  for start,end in consecutive_ranges(positions)]
        rows = '; '.join(ranges)
        return f'{sheet}, {rows}' if multiple_sheets else rows
    if kind == 'pdf':
        return '; '.join(f'page {start}' if start == end else f'pages {start}-{end}'
                         for start,end in consecutive_ranges(positions))
    return ''


def write_candidates(path, hits):
    """Create the short, human-readable candidate list."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = 'Candidates'
    sheet.append(('File', 'Location', 'Why candidate'))
    for hit in hits:
        sheet.append((hit['path'], hit['location'], hit['reason']))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(fill_type='solid', fgColor='23405A')
        cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 24
    for column, width in {'A':80, 'B':38, 'C':76}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    book.save(path)
    book.close()


def scan(args):
    output=args.output_dir.resolve()
    roots=[p.resolve() for p in args.root]
    if any(not p.exists() for p in roots): raise ValueError('Every --root must exist')
    # The output may be beside source directories, never inside one.
    if any(output==p or output.is_relative_to(p) for p in roots):
        raise ValueError('--output-dir must be outside every scanned root')
    if output.exists() and any(output.iterdir()): raise ValueError('Use an empty output directory to preserve earlier runs')
    output.mkdir(parents=True,exist_ok=True)
    entries=missing_entries(args.missing_report)
    keywords=school_keywords(entries,args.school)
    if not args.inventory_only and not keywords: raise ValueError('Content scan needs --missing-report or --school')
    patterns={keyword:re.compile(r'\b'+re.escape(keyword)+r'\b') for keyword in keywords}
    labels={entry['keyword']:entry['school'] for entry in entries if entry['keyword']}
    labels.update({school_keyword(name):name for name in args.school if school_keyword(name)})
    inventory=[];hits=[];errors=[];seen_paths=set();first_by_hash={};counts=collections.Counter()
    for root in roots:
        if root.is_file(): paths=[root]
        else:
            paths=[]
            for directory,dirs,names in os.walk(root,onerror=lambda e: errors.append({'path':str(e.filename),'error':str(e)})):
                dirs[:]=sorted(d for d in dirs if d not in SKIP_DIRS and not Path(directory,d).is_symlink())
                paths.extend(Path(directory,n) for n in sorted(names))
        for path in paths:
            if path.suffix.lower() not in SUPPORTED or path.name.startswith(('~$','._')):continue
            identity=str(path).casefold()
            if identity in seen_paths:continue
            seen_paths.add(identity)
            record={'path':str(path),'root':str(root),'extension':path.suffix.lower(),'status':'pending','sha256':'','duplicate_of':''}
            try:
                before=path.stat(); digest=sha256(path);record['sha256']=digest;record['bytes']=before.st_size
                record['duplicate_of']=first_by_hash.get(digest,'')
                if args.inventory_only:
                    record['status']='inventoried'
                else:
                    file_hits=[];count=0;empty=0
                    path_text=normalize(str(path))
                    path_terms={k for k,pattern in patterns.items() if pattern.search(path_text)}
                    content_locations=collections.defaultdict(set)
                    matched_in_content=set()
                    for kind,sheet,index,text,multiple_sheets in units(path,args):
                        count+=1
                        if not text.strip():empty+=1;continue
                        normalized=normalize(text)
                        for keyword,pattern in patterns.items():
                            if pattern.search(normalized):
                                matched_in_content.add(keyword)
                                content_locations[(keyword,kind,sheet,multiple_sheets)].add(index)
                    for (keyword,kind,sheet,multiple_sheets),positions in sorted(content_locations.items()):
                        name=labels.get(keyword,keyword)
                        reason=f'{name}: "{keyword}" found in file text.'
                        file_hits.append({'path':str(path),
                                          'location':candidate_location(kind,sheet,multiple_sheets,positions),
                                          'reason':reason})
                    for keyword in sorted(path_terms-matched_in_content):
                        name=labels.get(keyword,keyword)
                        file_hits.append({'path':str(path),'location':'',
                                          'reason':f'{name}: "{keyword}" found in file path.'})
                    record['units']=count;record['empty_units']=empty
                    record['status']='needs_ocr' if path.suffix.lower()=='.pdf' and (count==0 or count==empty) else ('text_scanned_with_empty_pages' if empty else 'text_scanned')
                    if args.max_pdf_pages and path.suffix.lower()=='.pdf' or args.max_excel_rows and path.suffix.lower() in {'.xlsx','.xlsm','.xls'}:
                        record['status']+='__limited'
                after=path.stat()
                if after.st_size!=before.st_size or after.st_mtime_ns!=before.st_mtime_ns:
                    raise RuntimeError('Source changed during scan; review this file again')
                if digest not in first_by_hash:
                    first_by_hash[digest]=str(path)
                if not args.inventory_only:
                    hits.extend(file_hits)
            except Exception as exc:
                record['status']='error';record['error']=f'{type(exc).__name__}: {exc}'
                errors.append({'path':str(path),'error':record['error']})
            inventory.append(record);counts[record['status']]+=1
    manifest={'method':'deterministic keyword screening','roots':[str(p) for p in roots],
              'inventory_only':args.inventory_only,'limits':{'pdf_pages':args.max_pdf_pages,'excel_rows':args.max_excel_rows},
              'file_instances':len(inventory),'unique_hashes':len({r['sha256'] for r in inventory if r['sha256']}),
              'candidate_rows':len(hits),'candidate_files':len({hit['path'] for hit in hits}),
              'missing_report_entries':len(entries),
              'statuses':dict(counts),'errors':len(errors),
              'caveats':['Keywords are candidates, not confirmed recovered sections.','Textless PDF pages need OCR or visual review.','Blank formula caches are not proof of absent grades.','I and İ sections are distinct.']}
    write_candidates(output/'candidates.xlsx',hits)
    for name,data in [('inventory.json',inventory),('errors.json',errors),('manifest.json',manifest)]:
        (output/name).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False,indent=2))
    return 2 if errors else 0


def prompted_path(prompt):
    """Accept a path pasted from File Explorer, with or without quotes."""
    value = input(prompt).strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in ('"', "'"):
        value = value[1:-1].strip()
    if not value:
        raise ValueError('A path is required.')
    return Path(value).expanduser()


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    interactive = not argv
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root',type=Path,action='append',required=not interactive,help='File or folder; repeat for multiple roots')
    parser.add_argument('--missing-report',type=Path,help='ERC report: district, school, section in the first three columns')
    parser.add_argument('--school',action='append',default=[],help='Explicit school keyword or alias; may repeat')
    parser.add_argument('--output-dir',type=Path,required=not interactive)
    parser.add_argument('--inventory-only',action='store_true',help='Hash and list files without reading document contents')
    parser.add_argument('--max-pdf-pages',type=int,default=0,help='0 = all pages; limits are recorded')
    parser.add_argument('--max-excel-rows',type=int,default=0,help='0 = all rows per worksheet; limits are recorded')
    parser.add_argument('--text-encoding',default='utf-8-sig')
    args=parser.parse_args(argv)
    if interactive:
        try:
            root = prompted_path('Extra folder containing PDF and Excel files: ')
            report = prompted_path('Missing-sections Excel file (from overview1): ')
            output = prompted_path('New results folder outside the extra folder: ')
        except EOFError:
            parser.error('Run this script in an interactive terminal so you can enter the paths.')
        except ValueError as exc:
            parser.error(str(exc))
        if not root.exists():
            parser.error(f'Extra folder was not found: {root}')
        if not report.is_file():
            parser.error(f'Missing-sections Excel file was not found: {report}')
        args.root = [root]
        args.missing_report = report
        args.output_dir = output
    if min(args.max_pdf_pages,args.max_excel_rows)<0:parser.error('Limits must be nonnegative')
    logging.getLogger('pypdf').setLevel(logging.ERROR)
    if interactive:
        print('Scanning source files. This can take a while for large folders.', flush=True)
    result = scan(args)
    print(f'Results saved in: {args.output_dir.resolve()}')
    return result


if __name__=='__main__':raise SystemExit(main())
