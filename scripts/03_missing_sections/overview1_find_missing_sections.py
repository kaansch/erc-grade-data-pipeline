#!/usr/bin/env python3
"""Find district-school sections missing from the allgrades workbook.

The two source workbooks are opened through binary read-only file handles and
with openpyxl's read-only mode. This program never calls save() on either input.

Default interactive usage:

    python find_missing_sections.py

The program then asks for the allgrades path first and the
TREATMENT_CLASSROOMS path second. The output is created beside allgrades as
missing_sections.xlsx (or a numbered new filename if that name already exists).

Explicit paths:

    python find_missing_sections.py ^
        --allgrades "C:\\path\\allgrades.xlsx" ^
        --treatment "C:\\path\\TREATMENT_CLASSROOMS.xlsx" ^
        --output "C:\\path\\missing_sections.xlsx"

Use --overwrite only when an existing output workbook may be replaced.
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
import unicodedata
from collections import OrderedDict
from contextlib import contextmanager
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import BinaryIO, Iterator

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - depends on the user's environment
    raise SystemExit(
        "This program requires openpyxl. Install it with: pip install openpyxl"
    ) from exc


ALLGRADES_SHEETS = (
    "3. SINIF 1. DÖNEM",
    "3. SINIF 2. DÖNEM",
    "4. SINIF",
)

OUTPUT_HEADERS = ("District", "School", "Missing Section")
SUPPORTED_INPUT_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
DEFAULT_DISTRICT_FUZZY_THRESHOLD = 90.0
DEFAULT_SCHOOL_FUZZY_THRESHOLD = 92.0
DEFAULT_FUZZY_MARGIN = 3.0

_WHITESPACE_RE = re.compile(r"\s+")
_NON_ALPHANUMERIC_RE = re.compile(r"[^a-z0-9]+")
_TREATMENT_SECTION_RE = re.compile(
    r"^4\s*[-‐‑‒–—−]\s*(.+)$",
    flags=re.IGNORECASE,
)
_TURKISH_LOWER_TRANSLATION = str.maketrans({"I": "ı", "İ": "i", "Ş": "ş", "Ğ": "ğ", "Ü": "ü", "Ö": "ö", "Ç": "ç"})
_TURKISH_UPPER_TRANSLATION = str.maketrans({"i": "İ", "ı": "I", "ş": "Ş", "ğ": "Ğ", "ü": "Ü", "ö": "Ö", "ç": "Ç"})
_FUZZY_ASCII_TRANSLATION = str.maketrans(
    {"ç": "c", "ğ": "g", "ı": "i", "ö": "o", "ş": "s", "ü": "u"}
)


class WorkbookDataError(ValueError):
    """Raised when workbook content does not match the required structure."""


@dataclass
class SchoolRequirement:
    """Required sections for one canonical district-school pair."""

    district: str
    school: str
    sections: list[str] = field(default_factory=list)
    source_rows: list[int] = field(default_factory=list)
    _section_set: set[str] = field(default_factory=set, repr=False)

    def add_section(self, section: str, row_number: int) -> None:
        self.source_rows.append(row_number)
        if section not in self._section_set:
            self._section_set.add(section)
            self.sections.append(section)


@dataclass(frozen=True)
class MissingCombination:
    district: str
    school: str
    section: str


@dataclass(frozen=True)
class NameMatchNotice:
    source_district: str
    source_school: str
    target_district: str
    target_school: str
    district_score: float
    school_score: float


@dataclass(frozen=True)
class NameMatchIssue:
    district: str
    school: str
    reason: str


@dataclass(frozen=True)
class PairResolution:
    reference_key: tuple[str, str] | None
    status: str
    district_score: float = 0.0
    school_score: float = 0.0
    reason: str = ""


@dataclass
class SheetScan:
    missing: list[MissingCombination]
    rows_examined: int
    incomplete_rows: int
    invalid_section_rows: int
    fuzzy_matches: list[NameMatchNotice]
    ambiguous_names: list[NameMatchIssue]
    unmatched_names: list[NameMatchIssue]
    fuzzy_matched_rows: int


def clean_text(value: object) -> str:
    """Normalize Unicode and whitespace without changing the source cell."""

    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    return _WHITESPACE_RE.sub(" ", text).strip()


def comparison_key(value: object) -> str:
    """Create a whitespace-tolerant, Turkish-aware, case-insensitive key."""

    text = clean_text(value).translate(_TURKISH_LOWER_TRANSLATION).casefold()
    return unicodedata.normalize("NFC", text)


def fuzzy_key(value: object) -> str:
    """Create an accent-, punctuation-, whitespace-, and case-tolerant key."""

    text = comparison_key(value).translate(_FUZZY_ASCII_TRANSLATION)
    return _NON_ALPHANUMERIC_RE.sub(" ", text).strip()


def similarity_score(left: object, right: object) -> float:
    """Return a 0-100 character similarity score for two normalized names."""

    left_key = fuzzy_key(left)
    right_key = fuzzy_key(right)
    if not left_key or not right_key:
        return 0.0
    return SequenceMatcher(
        None,
        left_key,
        right_key,
        autojunk=False,
    ).ratio() * 100.0


def best_unique_fuzzy_match(
    source: str,
    candidates: list[tuple[str, str]],
    threshold: float,
    margin: float,
) -> tuple[str | None, float, str]:
    """Return a confident candidate key, its score, and a status message."""

    scored = sorted(
        (
            (similarity_score(source, display_name), key, display_name)
            for key, display_name in candidates
        ),
        reverse=True,
    )
    if not scored:
        return None, 0.0, "no reference candidates"

    best_score, best_key, best_display = scored[0]
    if best_score < threshold:
        return (
            None,
            best_score,
            f"best candidate {best_display!r} scored {best_score:.1f}, "
            f"below threshold {threshold:.1f}",
        )

    if len(scored) > 1:
        runner_up_score, _runner_up_key, runner_up_display = scored[1]
        if best_score - runner_up_score < margin:
            return (
                None,
                best_score,
                f"ambiguous between {best_display!r} ({best_score:.1f}) and "
                f"{runner_up_display!r} ({runner_up_score:.1f}); required "
                f"margin is {margin:.1f}",
            )

    return best_key, best_score, "matched"


class ReferenceNameMatcher:
    """Resolve allgrades names to reference pairs without cross-district mixing."""

    def __init__(
        self,
        requirements: OrderedDict[tuple[str, str], SchoolRequirement],
        district_threshold: float,
        school_threshold: float,
        margin: float,
    ) -> None:
        self.requirements = requirements
        self.district_threshold = district_threshold
        self.school_threshold = school_threshold
        self.margin = margin
        self.districts: OrderedDict[
            str, tuple[str, OrderedDict[str, SchoolRequirement]]
        ] = OrderedDict()

        for (district_key, school_key), requirement in requirements.items():
            if district_key not in self.districts:
                self.districts[district_key] = (
                    requirement.district,
                    OrderedDict(),
                )
            self.districts[district_key][1][school_key] = requirement

    def resolve(self, district: str, school: str) -> PairResolution:
        district_key = comparison_key(district)
        school_key = comparison_key(school)
        exact_pair = (district_key, school_key)
        if exact_pair in self.requirements:
            return PairResolution(exact_pair, "exact", 100.0, 100.0)

        district_was_fuzzy = False
        district_score = 100.0
        if district_key not in self.districts:
            matched_district, district_score, reason = best_unique_fuzzy_match(
                district,
                [
                    (candidate_key, display_name)
                    for candidate_key, (display_name, _schools) in self.districts.items()
                ],
                self.district_threshold,
                self.margin,
            )
            if matched_district is None:
                status = "ambiguous" if reason.startswith("ambiguous") else "unmatched"
                return PairResolution(
                    None,
                    status,
                    district_score=district_score,
                    reason=f"district: {reason}",
                )
            district_key = matched_district
            district_was_fuzzy = True

        _district_display, schools = self.districts[district_key]
        if school_key in schools:
            resolved_key = (district_key, school_key)
            status = "fuzzy" if district_was_fuzzy else "exact"
            return PairResolution(
                resolved_key,
                status,
                district_score=district_score,
                school_score=100.0,
            )

        matched_school, school_score, reason = best_unique_fuzzy_match(
            school,
            [
                (candidate_key, requirement.school)
                for candidate_key, requirement in schools.items()
            ],
            self.school_threshold,
            self.margin,
        )
        if matched_school is None:
            status = "ambiguous" if reason.startswith("ambiguous") else "unmatched"
            return PairResolution(
                None,
                status,
                district_score=district_score,
                school_score=school_score,
                reason=f"school: {reason}",
            )

        return PairResolution(
            (district_key, matched_school),
            "fuzzy",
            district_score=district_score,
            school_score=school_score,
        )


def normalize_section_letter(value: object) -> str:
    """Return one uppercase section letter, or raise for an invalid value."""

    section = clean_text(value)
    if len(section) != 1 or not section.isalpha():
        raise WorkbookDataError(
            f"expected one section letter, but found {value!r}"
        )
    return section.translate(_TURKISH_UPPER_TRANSLATION).upper()


def normalize_treatment_section(value: object) -> str:
    """Convert a TREATMENT_CLASSROOMS value such as '4-E' to 'E'."""

    section = clean_text(value)
    match = _TREATMENT_SECTION_RE.fullmatch(section)
    if match is None:
        raise WorkbookDataError(
            f"expected a section in the form '4-E', but found {value!r}"
        )
    return normalize_section_letter(match.group(1))


def file_sha256(path: Path) -> str:
    """Hash a file using read-only access."""

    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@contextmanager
def open_excel_read_only(path: Path) -> Iterator[object]:
    """Open an Excel workbook using only an rb handle and read-only mode."""

    source: BinaryIO
    with path.open("rb") as source:
        workbook = load_workbook(
            source,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            yield workbook
        finally:
            workbook.close()


def read_required_sections(
    treatment_path: Path,
) -> OrderedDict[tuple[str, str], SchoolRequirement]:
    """Read every reference row and build the required section mapping."""

    requirements: OrderedDict[tuple[str, str], SchoolRequirement] = OrderedDict()

    with open_excel_read_only(treatment_path) as workbook:
        if len(workbook.sheetnames) != 1:
            raise WorkbookDataError(
                "TREATMENT_CLASSROOMS must contain exactly one worksheet; "
                f"found {len(workbook.sheetnames)}: {workbook.sheetnames}"
            )

        worksheet = workbook[workbook.sheetnames[0]]

        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=2,
                min_col=1,
                max_col=3,
                values_only=True,
            ),
            start=2,
        ):
            v1, v2, _trstatus = values
            combined_name = clean_text(v1)
            section_text = clean_text(v2)

            if not combined_name and not section_text:
                continue
            if not combined_name or not section_text:
                raise WorkbookDataError(
                    f"TREATMENT_CLASSROOMS row {row_number}: Columns A and B "
                    "must both contain values."
                )
            if "/" not in combined_name:
                raise WorkbookDataError(
                    f"TREATMENT_CLASSROOMS row {row_number}: Column A must be "
                    f"in 'District / School' form; found {v1!r}."
                )

            district_text, school_text = combined_name.split("/", maxsplit=1)
            district = clean_text(district_text)
            school = clean_text(school_text)
            if not district or not school:
                raise WorkbookDataError(
                    f"TREATMENT_CLASSROOMS row {row_number}: District and "
                    f"school must both be non-empty; found {v1!r}."
                )

            try:
                section = normalize_treatment_section(v2)
            except WorkbookDataError as exc:
                raise WorkbookDataError(
                    f"TREATMENT_CLASSROOMS row {row_number}: {exc}"
                ) from exc

            key = (comparison_key(district), comparison_key(school))
            if key not in requirements:
                requirements[key] = SchoolRequirement(
                    district=district,
                    school=school,
                )
            requirements[key].add_section(section, row_number)

    if not requirements:
        raise WorkbookDataError(
            "TREATMENT_CLASSROOMS contains no district-school-section data rows."
        )

    anomalies: list[str] = []
    for requirement in requirements.values():
        if len(requirement.source_rows) != 3 or len(requirement.sections) != 3:
            anomalies.append(
                f"{requirement.district} / {requirement.school}: "
                f"rows {requirement.source_rows}, distinct sections "
                f"{requirement.sections}"
            )

    if anomalies:
        shown = "\n  - ".join(anomalies[:20])
        remainder = len(anomalies) - 20
        suffix = f"\n  ... and {remainder} more" if remainder > 0 else ""
        raise WorkbookDataError(
            "Every reference district-school must have exactly three rows "
            "and three distinct sections. Problems found:\n  - "
            f"{shown}{suffix}"
        )

    return requirements


def scan_allgrades_sheet(
    worksheet: object,
    requirements: OrderedDict[tuple[str, str], SchoolRequirement],
    matcher: ReferenceNameMatcher,
) -> SheetScan:
    """Examine one allgrades worksheet row by row."""

    present: set[tuple[str, str, str]] = set()
    rows_examined = 0
    incomplete_rows = 0
    invalid_section_rows = 0
    fuzzy_matched_rows = 0
    resolution_cache: dict[tuple[str, str], PairResolution] = {}
    fuzzy_matches: OrderedDict[tuple[str, str], NameMatchNotice] = OrderedDict()
    ambiguous_names: OrderedDict[tuple[str, str], NameMatchIssue] = OrderedDict()
    unmatched_names: OrderedDict[tuple[str, str], NameMatchIssue] = OrderedDict()

    for values in worksheet.iter_rows(
        min_row=2,
        min_col=1,
        max_col=4,
        values_only=True,
    ):
        rows_examined += 1
        district_value, school_value, _ignored_column_c, section_value = values

        district = clean_text(district_value)
        school = clean_text(school_value)
        section_text = clean_text(section_value)

        if not district and not school and not section_text:
            continue
        if not district or not school or not section_text:
            incomplete_rows += 1
            continue

        try:
            section = normalize_section_letter(section_value)
        except WorkbookDataError:
            invalid_section_rows += 1
            continue

        source_key = (comparison_key(district), comparison_key(school))
        if source_key not in resolution_cache:
            resolution = matcher.resolve(district, school)
            resolution_cache[source_key] = resolution

            if resolution.status == "fuzzy" and resolution.reference_key is not None:
                target = requirements[resolution.reference_key]
                fuzzy_matches[source_key] = NameMatchNotice(
                    source_district=district,
                    source_school=school,
                    target_district=target.district,
                    target_school=target.school,
                    district_score=resolution.district_score,
                    school_score=resolution.school_score,
                )
            elif resolution.status == "ambiguous":
                ambiguous_names[source_key] = NameMatchIssue(
                    district=district,
                    school=school,
                    reason=resolution.reason,
                )
            elif resolution.status == "unmatched":
                unmatched_names[source_key] = NameMatchIssue(
                    district=district,
                    school=school,
                    reason=resolution.reason,
                )
        else:
            resolution = resolution_cache[source_key]

        if resolution.reference_key is None:
            continue
        if resolution.status == "fuzzy":
            fuzzy_matched_rows += 1

        reference_district, reference_school = resolution.reference_key
        present.add(
            (
                reference_district,
                reference_school,
                section,
            )
        )

    missing: list[MissingCombination] = []
    for (district_key, school_key), requirement in requirements.items():
        for required_section in requirement.sections:
            if (district_key, school_key, required_section) not in present:
                missing.append(
                    MissingCombination(
                        district=requirement.district,
                        school=requirement.school,
                        section=required_section,
                    )
                )

    return SheetScan(
        missing=missing,
        rows_examined=rows_examined,
        incomplete_rows=incomplete_rows,
        invalid_section_rows=invalid_section_rows,
        fuzzy_matches=list(fuzzy_matches.values()),
        ambiguous_names=list(ambiguous_names.values()),
        unmatched_names=list(unmatched_names.values()),
        fuzzy_matched_rows=fuzzy_matched_rows,
    )


def scan_allgrades(
    allgrades_path: Path,
    requirements: OrderedDict[tuple[str, str], SchoolRequirement],
    district_threshold: float,
    school_threshold: float,
    fuzzy_margin: float,
) -> OrderedDict[str, SheetScan]:
    """Scan each required allgrades sheet independently."""

    results: OrderedDict[str, SheetScan] = OrderedDict()
    matcher = ReferenceNameMatcher(
        requirements,
        district_threshold=district_threshold,
        school_threshold=school_threshold,
        margin=fuzzy_margin,
    )

    with open_excel_read_only(allgrades_path) as workbook:
        missing_sheets = [
            sheet_name
            for sheet_name in ALLGRADES_SHEETS
            if sheet_name not in workbook.sheetnames
        ]
        if missing_sheets:
            raise WorkbookDataError(
                "allgrades is missing required worksheet(s): "
                + ", ".join(missing_sheets)
                + f". Found: {workbook.sheetnames}"
            )

        for sheet_name in ALLGRADES_SHEETS:
            results[sheet_name] = scan_allgrades_sheet(
                workbook[sheet_name],
                requirements,
                matcher,
            )

    return results


def format_output_sheet(worksheet: object, data_row_count: int) -> None:
    """Apply compact, readable formatting to one output worksheet."""

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Aptos", size=11, color="000000")
    header_border = Border(
        bottom=Side(style="medium", color="17365D")
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = header_border

    if data_row_count:
        for row in worksheet.iter_rows(
            min_row=2,
            max_row=data_row_count + 1,
            min_col=1,
            max_col=3,
        ):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(horizontal="left", vertical="center")

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 24

    for column_number, width in enumerate((25, 52, 19), start=1):
        worksheet.column_dimensions[get_column_letter(column_number)].width = width


def write_output_workbook(
    output_path: Path,
    results: OrderedDict[str, SheetScan],
) -> None:
    """Create and save only the new output workbook."""

    output_workbook = Workbook()
    output_workbook.properties.title = "Missing District-School Sections"
    output_workbook.properties.subject = (
        "Required classroom sections absent from each allgrades worksheet"
    )

    try:
        for index, sheet_name in enumerate(ALLGRADES_SHEETS):
            if index == 0:
                worksheet = output_workbook.active
                worksheet.title = sheet_name
            else:
                worksheet = output_workbook.create_sheet(title=sheet_name)

            worksheet.append(OUTPUT_HEADERS)
            for missing in results[sheet_name].missing:
                worksheet.append(
                    (missing.district, missing.school, missing.section)
                )

            format_output_sheet(
                worksheet,
                data_row_count=len(results[sheet_name].missing),
            )

        output_workbook.save(output_path)
    finally:
        output_workbook.close()


def verify_output_workbook(
    output_path: Path,
    results: OrderedDict[str, SheetScan],
) -> None:
    """Reopen the new workbook read-only and verify its basic contents."""

    with open_excel_read_only(output_path) as workbook:
        if tuple(workbook.sheetnames) != ALLGRADES_SHEETS:
            raise RuntimeError(
                "Output verification failed: worksheet names or order changed."
            )

        for sheet_name in ALLGRADES_SHEETS:
            worksheet = workbook[sheet_name]
            header = tuple(
                worksheet.cell(row=1, column=column).value
                for column in range(1, 4)
            )
            if header != OUTPUT_HEADERS:
                raise RuntimeError(
                    f"Output verification failed in {sheet_name!r}: "
                    f"unexpected header {header!r}."
                )

            actual_rows = sum(
                1
                for row in worksheet.iter_rows(
                    min_row=2,
                    min_col=1,
                    max_col=3,
                    values_only=True,
                )
                if any(value is not None for value in row)
            )
            expected_rows = len(results[sheet_name].missing)
            if actual_rows != expected_rows:
                raise RuntimeError(
                    f"Output verification failed in {sheet_name!r}: expected "
                    f"{expected_rows} data rows, found {actual_rows}."
                )


def normalized_path_identity(path: Path) -> str:
    """Return a case-normalized absolute path for collision checks."""

    return os.path.normcase(str(path.resolve()))


def prompted_path(prompt: str) -> Path:
    """Ask for one path and accept quoted paths pasted from Windows Explorer."""

    try:
        entered = input(prompt).strip()
    except EOFError as exc:
        raise ValueError("No file path was entered.") from exc

    if len(entered) >= 2 and entered[0] == entered[-1] and entered[0] in "\"'":
        entered = entered[1:-1].strip()
    if not entered:
        raise ValueError("No file path was entered.")
    return Path(entered)


def next_default_output_path(allgrades_path: Path) -> Path:
    """Choose a new output name beside allgrades without overwriting a file."""

    output_directory = allgrades_path.parent
    candidate = output_directory / "missing_sections.xlsx"
    if not candidate.exists():
        return candidate

    number = 2
    while True:
        candidate = output_directory / f"missing_sections_{number}.xlsx"
        if not candidate.exists():
            return candidate
        number += 1


def validate_paths(
    allgrades_path: Path,
    treatment_path: Path,
    output_path: Path,
    overwrite: bool,
) -> None:
    """Validate paths before opening any workbook."""

    for label, path in (
        ("allgrades", allgrades_path),
        ("TREATMENT_CLASSROOMS", treatment_path),
    ):
        if not path.is_file():
            raise FileNotFoundError(f"{label} workbook was not found: {path}")
        if path.suffix.lower() not in SUPPORTED_INPUT_SUFFIXES:
            raise ValueError(
                f"{label} must be an Open XML Excel workbook "
                f"({', '.join(sorted(SUPPORTED_INPUT_SUFFIXES))}); found {path.name}."
            )

    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("The output workbook must use the .xlsx extension.")

    identities = {
        "allgrades": normalized_path_identity(allgrades_path),
        "treatment": normalized_path_identity(treatment_path),
        "output": normalized_path_identity(output_path),
    }
    if len(set(identities.values())) != 3:
        raise ValueError(
            "The two inputs and the output must be three different files. "
            "Refusing to risk overwriting an input workbook."
        )

    if output_path.exists() and not overwrite:
        raise FileExistsError(
            f"Output already exists: {output_path}. Choose another path or "
            "run with --overwrite."
        )


def validate_fuzzy_settings(
    district_threshold: float,
    school_threshold: float,
    fuzzy_margin: float,
) -> None:
    """Reject unsafe or nonsensical fuzzy-matching settings."""

    for label, value in (
        ("district fuzzy threshold", district_threshold),
        ("school fuzzy threshold", school_threshold),
    ):
        if not 0.0 <= value <= 100.0:
            raise ValueError(f"The {label} must be between 0 and 100.")
    if not 0.0 <= fuzzy_margin <= 100.0:
        raise ValueError("The fuzzy-match margin must be between 0 and 100.")


def parse_arguments(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Compare required TREATMENT_CLASSROOMS sections with each "
            "allgrades sheet and create a new missing-sections workbook."
        )
    )
    parser.add_argument(
        "--allgrades",
        type=Path,
        default=None,
        help="Path to allgrades.xlsx; if omitted, the program asks for it.",
    )
    parser.add_argument(
        "--treatment",
        type=Path,
        default=None,
        help=(
            "Path to TREATMENT_CLASSROOMS.xlsx; if omitted, the program "
            "asks for it after the allgrades path."
        ),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help=(
            "Path for the new output workbook; if omitted, a new "
            "missing_sections workbook is created beside allgrades."
        ),
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow replacement of an existing output workbook only.",
    )
    parser.add_argument(
        "--district-threshold",
        type=float,
        default=DEFAULT_DISTRICT_FUZZY_THRESHOLD,
        help=(
            "Minimum 0-100 district similarity score "
            f"(default: {DEFAULT_DISTRICT_FUZZY_THRESHOLD:.0f})."
        ),
    )
    parser.add_argument(
        "--school-threshold",
        type=float,
        default=DEFAULT_SCHOOL_FUZZY_THRESHOLD,
        help=(
            "Minimum 0-100 school similarity score "
            f"(default: {DEFAULT_SCHOOL_FUZZY_THRESHOLD:.0f})."
        ),
    )
    parser.add_argument(
        "--fuzzy-margin",
        type=float,
        default=DEFAULT_FUZZY_MARGIN,
        help=(
            "Required lead over the second-best fuzzy candidate "
            f"(default: {DEFAULT_FUZZY_MARGIN:.0f})."
        ),
    )
    return parser.parse_args(argv)


def run(args: argparse.Namespace) -> int:
    allgrades_entered = args.allgrades
    if allgrades_entered is None:
        allgrades_entered = prompted_path(
            "Enter the path to the allgrades Excel file: "
        )

    treatment_entered = args.treatment
    if treatment_entered is None:
        treatment_entered = prompted_path(
            "Enter the path to the TREATMENT_CLASSROOMS Excel file: "
        )

    allgrades_path = allgrades_entered.expanduser().resolve()
    treatment_path = treatment_entered.expanduser().resolve()
    if args.output is None:
        output_path = next_default_output_path(allgrades_path).resolve()
    else:
        output_path = args.output.expanduser().resolve()

    validate_fuzzy_settings(
        args.district_threshold,
        args.school_threshold,
        args.fuzzy_margin,
    )
    validate_paths(
        allgrades_path,
        treatment_path,
        output_path,
        overwrite=args.overwrite,
    )

    input_hashes_before = {
        allgrades_path: file_sha256(allgrades_path),
        treatment_path: file_sha256(treatment_path),
    }

    requirements = read_required_sections(treatment_path)
    results = scan_allgrades(
        allgrades_path,
        requirements,
        district_threshold=args.district_threshold,
        school_threshold=args.school_threshold,
        fuzzy_margin=args.fuzzy_margin,
    )

    input_hashes_after_reading = {
        allgrades_path: file_sha256(allgrades_path),
        treatment_path: file_sha256(treatment_path),
    }
    if input_hashes_before != input_hashes_after_reading:
        raise RuntimeError(
            "An input workbook changed while it was being examined. "
            "No output workbook was written."
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    write_output_workbook(output_path, results)
    verify_output_workbook(output_path, results)

    input_hashes_after_output = {
        allgrades_path: file_sha256(allgrades_path),
        treatment_path: file_sha256(treatment_path),
    }
    if input_hashes_before != input_hashes_after_output:
        raise RuntimeError(
            "Input-integrity verification failed: an input workbook changed "
            "during the run."
        )

    print(f"Created: {output_path}")
    print(f"Reference district-school pairs: {len(requirements)}")
    for sheet_name in ALLGRADES_SHEETS:
        result = results[sheet_name]
        print(
            f"{sheet_name}: {len(result.missing)} missing combination(s); "
            f"{result.rows_examined} data row(s) examined; "
            f"{result.incomplete_rows} incomplete row(s); "
            f"{result.invalid_section_rows} invalid-section row(s); "
            f"{len(result.fuzzy_matches)} distinct fuzzy name match(es) "
            f"covering {result.fuzzy_matched_rows} row(s); "
            f"{len(result.ambiguous_names)} ambiguous name pair(s); "
            f"{len(result.unmatched_names)} unmatched name pair(s)."
        )
        for notice in result.fuzzy_matches[:10]:
            print(
                "  Fuzzy: "
                f"{notice.source_district} / {notice.source_school} -> "
                f"{notice.target_district} / {notice.target_school} "
                f"(district {notice.district_score:.1f}, "
                f"school {notice.school_score:.1f})"
            )
        if len(result.fuzzy_matches) > 10:
            print(f"  ... {len(result.fuzzy_matches) - 10} more fuzzy match(es)")

        for issue_label, issues in (
            ("Ambiguous", result.ambiguous_names),
            ("Unmatched", result.unmatched_names),
        ):
            for issue in issues[:5]:
                print(
                    f"  {issue_label}: {issue.district} / {issue.school} "
                    f"({issue.reason})"
                )
            if len(issues) > 5:
                print(f"  ... {len(issues) - 5} more {issue_label.lower()} pair(s)")
    print("Input workbook SHA-256 hashes were unchanged.")
    return 0


def main(argv: list[str] | None = None) -> int:
    try:
        args = parse_arguments(argv)
        return run(args)
    except (FileNotFoundError, FileExistsError, ValueError, RuntimeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
